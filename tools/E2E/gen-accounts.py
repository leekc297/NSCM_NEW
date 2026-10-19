# -*- coding: utf-8 -*-
r"""
gen-accounts — 통합테스트_데이터셋(04_TESTDATA_SET.xlsx) 의 테스트계정·권한매트릭스·시나리오 → accounts.json / scenarios.csv
  python gen-accounts.py [--xlsx <path>]   (기본: e2e.config.json testDataXlsx)
  accounts.json : [{id, name, groups[], ws, divisions[], screens{DMF0082:'DM Planner',...}, usable, note}]
  authMatrix    : [{group, screen, div, editMeasure, editDay}]
  scenarios.csv : TESTDATA_SET 시트 그대로(시나리오ID, DIV, 화면, 테스트계정, ITEM_ID, SCM_SITE_ID, PROJECT_CODE …)
비밀번호는 없다 — dev-login(id) 만 사용.
"""
import sys as _sys
try: _sys.stdout.reconfigure(encoding='utf-8', errors='replace'); _sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception: pass
import sys, os, json, csv
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
cfg = json.load(open(os.path.join(HERE, 'e2e.config.json'), encoding='utf-8-sig'))
xl = sys.argv[sys.argv.index('--xlsx') + 1] if '--xlsx' in sys.argv else cfg['testDataXlsx']
wb = openpyxl.load_workbook(xl, read_only=True, data_only=True)
s = lambda v: '' if v is None else str(v).strip()

acc = {}
ws = wb['테스트계정_권한해석']; rows = [list(map(s, r)) for r in ws.iter_rows(values_only=True)]
hdr = rows[0]
for r in rows[1:]:
    if not r[0]: continue
    screens = {hdr[i].replace(' 해석', ''): r[i] for i in range(4, len(hdr)) if hdr[i].endswith('해석') and r[i] and r[i] != 'None'}
    acc[r[0]] = {'id': r[0], 'name': r[1], 'groups': [g.strip() for g in r[2].split(',') if g.strip()], 'ws': r[3] == 'Y', 'adminAlso': r[7] == 'Y',
                 'screens': screens, 'divisions': [], 'usable': r[8].startswith('사용'), 'note': r[8]}
# 시나리오 → 계정별 DIV
ws = wb['TESTDATA_SET']; rows = [list(map(s, r)) for r in ws.iter_rows(values_only=True)]
h = rows[0]; idx = {k: i for i, k in enumerate(h)}
scen = []
for r in rows[1:]:
    if not r[0]: continue
    d = {k: r[i] for k, i in idx.items()}
    scen.append(d)
    for a in [d.get('테스트계정', '')] + [x.strip() for x in d.get('대체계정', '').split(',')]:
        if a and a in acc and d['DIV'] and d['DIV'] not in acc[a]['divisions']: acc[a]['divisions'].append(d['DIV'])
# 권한매트릭스
ws = wb['권한매트릭스_실측']; rows = [list(map(s, r)) for r in ws.iter_rows(values_only=True)]
matrix = [{'group': r[0], 'screen': r[1], 'div': r[2], 'editMeasure': r[3], 'editDay': r[4].strip()} for r in rows[1:] if r[0]]
divs = sorted({m['div'].split('_')[0] for m in matrix if m['div']})
for a in acc.values():
    if not a['divisions']: a['divisions'] = ['40'] if a['adminAlso'] else divs[:1]
out = {'_source': os.path.basename(xl), 'divisions': divs, 'accounts': sorted(acc.values(), key=lambda a: (not a['adminAlso'], a['id'])), 'authMatrix': matrix}
json.dump(out, open(os.path.join(HERE, 'accounts.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=2)
with open(os.path.join(HERE, 'scenarios.csv'), 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=h); w.writeheader(); w.writerows(scen)
print(f"accounts.json 계정 {len(acc)} (ADMIN {sum(1 for a in acc.values() if a['adminAlso'])}) · 권한매트릭스 {len(matrix)} · DIV {divs} · scenarios.csv {len(scen)}행")
