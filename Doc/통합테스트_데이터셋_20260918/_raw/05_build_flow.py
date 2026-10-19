# -*- coding: utf-8 -*-
"""Phase 1 산출 — 01_PROCESS_FLOW.md / 02_PROCESS_FLOW.xlsx / 03_PROCESS_FLOW.html"""
import json, pathlib, sys, datetime, html
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from importlib import import_module
fd = import_module("04_flow_def") if False else None
import importlib.util
spec = importlib.util.spec_from_file_location("flow_def",
        pathlib.Path(__file__).resolve().parent / "04_flow_def.py")
FD = importlib.util.module_from_spec(spec); spec.loader.exec_module(FD)

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

D = pathlib.Path(__file__).resolve().parent
OUT = D.parent
META = json.loads((D / "01_meta.json").read_text(encoding="utf-8"))
NOW = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
M = META["meta"]
HDRNOTE = ("ZIOUI1 실데이터 기준 · host {} · user {} · 조회 {}".format(M["host"], M["user"], M["fetched"]))

DIVNAME = {"40": "40 (소형)", "81": "81 (자동차)", "82": "82 (ESS)"}
# mindmap 은 노드 텍스트의 () 를 shape 문법으로 해석하므로 괄호 없는 표기를 쓴다
MINDNAME = {"40": "40 소형", "81": "81 자동차", "82": "82 ESS"}

# ── 01 MD (Mermaid) ────────────────────────────────────────────────
def mermaid_div(div):
    ls = ["```mermaid", "flowchart TD"]
    prev = None
    for s in FD.STEPS:
        if s[1] != div:
            continue
        sid, _, _, kind, scr, auth, meas, day, ledger, job, sp, ver, _no = s
        nid = sid.replace("-", "_")
        # ⚠ Mermaid: 라벨에 () 가 들어가면 flowchart 노드 문법과 충돌한다
        #   ('DM Planner (W/S)', 'LOC(단가)', ' (MON) ' 등) → 라벨 전체를 "..." 로 감싼다.
        #   따옴표 라벨 안에서는 " 만 금지이므로 ' 로 치환한다.
        def q(text):
            return '"{}"'.format(text.replace('"', "'"))

        if kind == "배치":
            label = "{}<br/>배치 수동실행<br/>{}".format(sid, sp.split("(")[0])
            shape = "[[{}]]".format(q(label))
        elif kind == "검증":
            label = "{}<br/>{} 검증<br/>{}".format(sid, FD.scr_label(scr), meas)
            shape = "{{{{{}}}}}".format(q(label))
        else:
            label = "{}<br/>{} / {}<br/>{}{}".format(
                sid, FD.scr_label(scr), auth, meas,
                "" if day == "-" else " ({})".format(day))
            shape = "[{}]".format(q(label))
        ls.append("  {}{}".format(nid, shape))
        if prev:
            ls.append("  {} --> {}".format(prev, nid))
        prev = nid
    ls.append("```")
    return "\n".join(ls)


md = ["# 통합테스트 PROCESS FLOW (배치 JOB 연결)", "",
      "> 생성 {} · {}".format(NOW, HDRNOTE),
      "> **배치는 수동 실행 전제** (ZIOUI1 의 관련 JOB 은 전부 `ENABLED=FALSE`)", "",
      "## 전체 개요 (마인드맵)", "",
      "```mermaid", "mindmap",
      "  root((통합테스트 데이터 SET))"]
for div in ("40", "81", "82"):
    md.append("    {}".format(MINDNAME[div]))
    seen = []
    for s in FD.STEPS:
        if s[1] == div and s[4] not in ("-",) and s[4] not in seen:
            seen.append(s[4])
            md.append("      {}".format(FD.scr_label_plain(s[4])))
    md.append("      배치 Cutoff")
md += ["    DMF0035 검증"]
for t, sp, led in FD.WATERFALL:
    md.append("      {}".format(t))
md += ["```", ""]
# 마인드맵 블록(mindmap ~ 마지막 노드)을 HTML 에서도 재사용한다
MIND = chr(10).join(md[md.index("mindmap"):-2])

for div in ("40", "81", "82"):
    md += ["", "## DIVISION {}".format(DIVNAME[div]), "", mermaid_div(div), ""]
    md += ["| " + " | ".join(FD.COLS_X) + " |",
           "|" + "---|" * len(FD.COLS_X)]
    for s in FD.steps_x():
        if s[1] == div:
            md.append("| " + " | ".join("`{}`".format(x) if i == 11 and x else str(x)
                                        for i, x in enumerate(s)) + " |")
    md.append("")

md += ["", "## DMF0035 Waterfall 3 TYPE ↔ 조회 SP / 원장", "",
       "| TYPE | PROGRAM_ID | 화면명 | 조회 SP | 집계 원장 |", "|---|---|---|---|---|"]
for t, sp, led in FD.WATERFALL:
    pid = FD.WATERFALL_PROGRAM[t]
    md.append("| {} | {} | {} | `{}` | `{}` |".format(t, pid, FD.scr_name(pid), sp, led))
md += ["", "## 배치 JOB 원 스케줄 (참고 — 현재 전부 ENABLED=FALSE)", "",
       "| JOB | 주기 |", "|---|---|"]
for j in META["jobs"]:
    md.append("| `{}` | {} |".format(j["JOB_NAME"], (j["REPEAT_INTERVAL"] or "")[:60]))
(OUT / "01_PROCESS_FLOW.md").write_text("\n".join(md), encoding="utf-8")
print("MD 저장")

# ── 02 XLSX ────────────────────────────────────────────────────────
HF, HFONT = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True, size=9)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "PROCESS_FLOW"
ws.append(FD.COLS_X)
for c in ws[1]:
    c.fill, c.font = HF, HFONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
FILL = {"입력": "E2EFDA", "배치": "FFF2CC", "검증": "FCE4D6", "선행/배치": "FFF2CC", "선행/입력": "E2EFDA"}
for s in FD.steps_x():
    ws.append(s)
    f = FILL.get(s[3])
    if f:
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=f)
for i, w in enumerate([10, 6, 6, 10, 10, 22, 17, 26, 8, 30, 46, 62, 40, 24], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.font = Font(size=9)

ws2 = wb.create_sheet("DMF0035_3TYPE")
ws2.append(["TYPE", "PROGRAM_ID", "화면명", "조회 SP", "집계 원장"])
for c in ws2[1]:
    c.fill, c.font = HF, HFONT
for t, sp, led in FD.WATERFALL:
    pid = FD.WATERFALL_PROGRAM[t]
    ws2.append([t, pid, FD.scr_name(pid), sp, led])
for i, w in enumerate([16, 12, 24, 58, 34], 1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws3 = wb.create_sheet("배치JOB_원스케줄")
ws3.append(["JOB", "ENABLED", "주기", "수동실행 SP(추출)"])
for c in ws3[1]:
    c.fill, c.font = HF, HFONT
import re as _re
for j in META["jobs"]:
    act = _re.sub(r"\s+", " ", j["JOB_ACTION"] or "")
    calls = [m for m in _re.findall(r"(PR_[A-Z0-9_]+\s*\([^)]*\))", act)]
    ws3.append([j["JOB_NAME"], j["ENABLED"], j["REPEAT_INTERVAL"], "\n".join(calls[:6])])
for i, w in enumerate([56, 10, 46, 80], 1):
    ws3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
for row in ws3.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.font = Font(size=9)
try:
    wb.save(OUT / "02_PROCESS_FLOW.xlsx")
    print("XLSX 저장")
except PermissionError:
    # Excel 에서 열려 있으면 잠긴다. 나머지 산출물 생성은 계속한다.
    print("XLSX 저장 생략 — 파일이 열려 있음 (닫은 뒤 재실행 필요)")

# ── 03 HTML ────────────────────────────────────────────────────────
def esc(x):
    return html.escape(str(x))


rows_html = ""
for s in FD.steps_x():
    cls = {"입력": "inp", "배치": "bat", "검증": "ver"}.get(s[3], "inp")
    rows_html += "<tr class='{}'>".format(cls) + "".join(
        "<td>{}</td>".format(esc(x)) for x in s) + "</tr>\n"

mer_mind = "<h3>전체 개요</h3>" + chr(10) + "<pre class='mermaid'>" + esc(MIND) + "</pre>"
mer = "\n\n".join("<h3>DIVISION {}</h3>\n<pre class='mermaid'>{}</pre>".format(
    esc(DIVNAME[d]), esc(mermaid_div(d).replace("```mermaid\n", "").replace("\n```", "")))
    for d in ("40", "81", "82"))

htmlt = """<!DOCTYPE html><html lang="ko"><head><meta charset="utf-8">
<title>통합테스트 PROCESS FLOW</title>
<script src="https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.min.js"></script>
<script>mermaid.initialize({{startOnLoad:true,theme:'neutral'}});</script>
<style>
 body{{font-family:'Malgun Gothic',system-ui,sans-serif;margin:24px;color:#222;background:#fff}}
 h1{{border-bottom:3px solid #1F4E78;padding-bottom:8px;color:#1F4E78}}
 h2{{margin-top:34px;color:#1F4E78}}
 .note{{background:#F2F7FB;border-left:4px solid #1F4E78;padding:10px 14px;margin:14px 0;font-size:14px}}
 table{{border-collapse:collapse;width:100%;font-size:12px;margin:12px 0}}
 th{{background:#1F4E78;color:#fff;padding:7px;position:sticky;top:0}}
 td{{border:1px solid #D0D7DE;padding:6px;vertical-align:top}}
 tr.inp td{{background:#E2EFDA}} tr.bat td{{background:#FFF2CC}} tr.ver td{{background:#FCE4D6}}
 code{{background:#F3F4F6;padding:1px 4px;border-radius:3px;font-size:11px}}
 pre.mermaid{{background:#fff;text-align:center}}
</style></head><body>
<h1>통합테스트 PROCESS FLOW</h1>
<div class="note"><b>생성</b> {now} &nbsp;|&nbsp; {hdr}<br>
<b>배치는 수동 실행 전제</b> — ZIOUI1 의 관련 JOB 은 전부 <code>ENABLED=FALSE</code></div>
<h2>프로세스 도식</h2>
{mer}
<h2>단계 전체</h2>
<table><thead><tr>{th}</tr></thead><tbody>
{rows}
</tbody></table>
<h2>DMF0035 Waterfall 3 TYPE</h2>
<table><thead><tr><th>TYPE</th><th>PROGRAM_ID</th><th>화면명</th><th>조회 SP</th><th>집계 원장</th></tr></thead><tbody>
{wf}
</tbody></table>
</body></html>"""
htmlt = htmlt.format(
    now=NOW, hdr=esc(HDRNOTE), mer=mer_mind + mer,
    th="".join("<th>{}</th>".format(esc(c)) for c in FD.COLS_X),
    rows=rows_html,
    wf="".join("<tr><td>{}</td><td>{}</td><td>{}</td>"
               "<td><code>{}</code></td><td><code>{}</code></td></tr>".format(
        esc(a), esc(FD.WATERFALL_PROGRAM[a]), esc(FD.scr_name(FD.WATERFALL_PROGRAM[a])),
        esc(b), esc(c)) for a, b, c in FD.WATERFALL))
(OUT / "03_PROCESS_FLOW.html").write_text(htmlt, encoding="utf-8")
print("HTML 저장")
