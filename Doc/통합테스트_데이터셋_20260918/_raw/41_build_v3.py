# -*- coding: utf-8 -*-
"""v3 — 08_SAMPLEDATA_STEPNO_v3.xlsx 생성

사용자 요구 (2026-09-20)
  · 40TestData / 81TestData / 82TestData 3시트는 **이전(06) 양식 36컬럼을 그대로 유지**하고
    v2(07) 에서 신규 생성된 데이터까지 그 양식 안으로 합류시킨다.
  · '기존' 으로 분류되는 행은 반드시 **DB 에 실존하는 데이터**여야 한다.
  · 그 다음 Step 별 Sheet 를 추가 생성한다.

행 원천 (모두 2026-09-18 ZIOUI1 실측 산출물 — 본 스크립트는 값을 만들어내지 않는다)
  v1(06)  : 06_SAMPLEDATA_STEPNO.xlsx 3시트 299행
            └ 04_TESTDATA_SET.xlsx(115행 실측) + 11_sspm_gap.json + 12_master_if.json
  v2(07)  : 07_SAMPLEDATA_STEPNO_v2.xlsx STEP 시트 30개
            └ 31_pick.py 가 ZIOUI1 에서 직접 선별한 행 (INCOTERMS EDIT_FLAG='Y' 정규화 포함)

⚠ 2026-09-20 현재 ZIOUI1 은 서비스 미등록(인스턴스 다운)이라 재측정이 불가하다.
  v2 행의 `기간(실측)` `현재값(실측)` 2컬럼은 09-18 회차에 수집되지 않았으므로
  `미측정(재조회 필요)` 로 명시한다 — 임의값 금지(anti-simplification).
  DB 복구 후 `42_refill_measure.py` 로 채운다.

⚠ 32_build.py 결함 정정 — `item` 변수가 `item = rec.get("ITEM_ID")` 로 덮여
  07 의 `항목` 컬럼에 ITEM_ID 가 들어갔다(40TestData 1-1 '385665' 등).
  v3 는 항목을 DIVISION 시트에서 다시 읽는다.
"""
import json
import pathlib
import importlib.util
from collections import OrderedDict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
OUT_DIR = HERE.parent
SRC_V1 = OUT_DIR / "06_SAMPLEDATA_STEPNO.xlsx"
SRC_V2 = OUT_DIR / "07_SAMPLEDATA_STEPNO_v2.xlsx"
OUT = OUT_DIR / "08_SAMPLEDATA_STEPNO_v3.xlsx"
DELIVER = pathlib.Path(r"C:\TRANFILE\RECEIVE\08_SAMPLEDATA_STEPNO_v3.xlsx")

_spec = importlib.util.spec_from_file_location("stepdef", HERE / "40_stepdef_v3.py")
SD = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(SD)

META = json.loads((HERE / "01_meta.json").read_text(encoding="utf-8"))
ACCR = json.loads((HERE / "08_accounts.json").read_text(encoding="utf-8"))
PNAME = json.loads((HERE / "09_program_names.json").read_text(encoding="utf-8"))["names"]
LAYOUT = json.loads((HERE / "30_layouts.json").read_text(encoding="utf-8"))

PEOPLE = {p["USER_ID"]: p for p in ACCR["people"]}
RECO = ACCR["recommend"]

# 31_pick.editable_incoterms() 실측 (SCM_ADM_LOOKUP_MASTER_DETAIL.ATTRIBUTE2)
INCO_Y = {"CFR", "CIF", "CIP", "CPT", "DAP", "DDP", "EXW", "FAS", "FCA", "FOB"}
INCO_N = {"A01", "DAT", "DDU", "Z01", "Z02", "Z04"}

# 권한 × 화면 × DIV -> 편집 MEASURE / 편집요일 (GET_ADMIN_GROUP_MEASURE 실측)
AUTH_MEAS, AUTH_DAY = {}, {}
for _m in META["auth_measure"]:
    AUTH_MEAS.setdefault((_m["AUTH"], _m["SCR"], _m["DIV"]), set()).add(_m["MEASURE"])
    _d = (_m.get("EDIT_DAY") or "").strip()
    if _d:
        AUTH_DAY.setdefault((_m["AUTH"], _m["SCR"], _m["DIV"]), set()).add(_d)

V1_COLS = ["StepNo", "Step", "항목", "샘플지정", "데이터출처", "분류기준", "DIV", "구분",
           "엑셀항목", "화면", "실제화면", "화면명", "권한그룹", "테스트계정", "대체계정",
           "해석authGroup(실측)", "W/S판정(실측)",
           "ITEM_ID", "SCM_SITE_ID", "PROJECT_CODE", "FCST_TYPE", "INCOTERMS", "DOM_EXP",
           "PLANT_SITE_ID", "입력MEASURE", "편집요일", "권한메타근거", "대상원장", "등록값명세",
           "데이터보유(실측)", "기간(실측)", "현재값(실측)", "테스트행위", "기대결과",
           "검증화면", "대응결함No"]
NEW_COLS = ["END_CUSTOMER", "INCOTERMS_EDIT_FLAG", "CELL_PROD_SITE_ID", "CURRENCY",
            "화면LAYOUT원천", "LAYOUT컬럼수", "G열요구사항", "STEP시트", "행출처"]
COLS = V1_COLS + NEW_COLS

V1_WIDTH = [8, 12, 26, 8, 20, 46, 5, 9, 26, 9, 9, 22, 17, 13, 30, 20, 8,
            22, 26, 18, 11, 10, 8, 14, 24, 8, 30, 40, 60, 18, 22, 30, 34, 34, 10, 22]
NEW_WIDTH = [20, 9, 18, 9, 40, 7, 40, 10, 8]

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(color="FFFFFF", bold=True, size=9)
FILL = {"기존(데이터보유 상위)": "E2EFDA", "기존(IF 실수신)": "E2EFDA",
        "신규(SSPM 미생성)": "FFF2CC", "신규(기준정보·IF)": "DDEBF7",
        "신규(VMI 가능 SITE)": "FCE4D6"}
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NOT_MEASURED = "미측정(재조회 필요)"
KEY7 = ["ITEM_ID", "SCM_SITE_ID", "PROJECT_CODE", "FCST_TYPE", "INCOTERMS",
        "DOM_EXP", "PLANT_SITE_ID"]

# ── DIVISION 시트 파싱 — Step 대분류 / 항목 / G열 ────────────────────────────
SPECIAL_82 = {("UI : Sales Price", "DP Planner"): "5-1",
              ("UI : Sales Price", "admin"): "5-2"}


def parse_division(ws, div):
    scan, cur_step, last, sub = [], "", None, 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        c = [("" if x is None else str(x).strip()) for x in row] + [""] * 8
        if not any(c[:7]):
            continue
        if c[1]:
            cur_step = c[1]
        sno, item = c[0], c[2]
        if not item:                       # 섹션 헤더 행
            last = None
            continue
        if div == "82" and (item, c[4]) in SPECIAL_82:
            key = SPECIAL_82[(item, c[4])]
            last, sub = key, 0
        elif sno:
            key, last, sub = sno, sno, 0
        else:                              # 하위 행 (81 3-2 의 IF 4표)
            if last is None:
                continue
            key = last + chr(ord("a") + sub)
            sub += 1
        scan.append((0 if c[5] == "샘플" else 1, key,
                     {"Step": cur_step, "항목": item, "시간": c[3],
                      "권한비고": c[4], "샘플Link": c[5], "G열": c[6]}))
    out = {}
    for _pri, key, info in sorted(scan, key=lambda x: x[0]):
        out.setdefault(key, info)
    return out


# ── 계정 해석 (13_build_sampledata.pick_acct 와 동일 규칙) ─────────────────────
def pick_acct(auth, scr):
    if auth in ("-", "기준정보 관리"):
        cand = sorted(RECO.get("DM Planner", {}).get("사용가능", []))
        u = cand[0] if cand else "(없음)"
        return (u, ", ".join(cand[1:])[:60], "메타 미등재(제한 없음 전제)",
                PEOPLE.get(u, {}).get("wsPlannerYn", "-"))
    if auth in ("DM ADMIN/ADMIN", "ADMIN"):
        cand = sorted(u for u, p in PEOPLE.items() if p["admin_masked"])
        u = cand[0] if cand else "(없음)"
        return (u, ", ".join(cand[1:])[:60], "ADMIN",
                PEOPLE.get(u, {}).get("wsPlannerYn", "-"))
    r = RECO.get(auth)
    if not r or not r["사용가능"]:
        return "(없음)", "", "미해석", "-"
    ok = sorted(r["사용가능"], key=lambda u: (0 if PEOPLE[u]["groups"] == [auth] else 1,
                                             len(PEOPLE[u]["groups"]), u))
    u = ok[0]
    p = PEOPLE[u]
    resolved = (p["resolved"].get(scr) or "메타 미등재(제한 없음 전제)"
                if scr in ("DMF0082", "DMF0021", "DMF0042", "DMF0003")
                else "메타 미등재(제한 없음 전제)")
    alt = ", ".join(ok[1:])[:60]
    if r["사용불가(ADMIN 해석)"]:
        alt += "  |  ⚠사용불가(ADMIN 해석): " + ", ".join(r["사용불가(ADMIN 해석)"])
    return u, alt[:160], str(resolved), p["wsPlannerYn"]


def acct_detail(user, scr):
    p = PEOPLE.get(user)
    if not p:
        return "메타 미등재(제한 없음 전제)", "-"
    resolved = (p["resolved"].get(scr) or "메타 미등재(제한 없음 전제)"
                if scr in ("DMF0082", "DMF0021", "DMF0042", "DMF0003")
                else "메타 미등재(제한 없음 전제)")
    return str(resolved), p["wsPlannerYn"]


def auth_key(auth, scr, div):
    if (auth, scr, div) in AUTH_MEAS:
        return (auth, scr, div)
    for k in sorted(AUTH_MEAS):
        if k[0] == auth and k[1] == scr and k[2].startswith(div):
            return k
    return None


def measures_for(auth, scr, div, declared):
    k = auth_key(auth, scr, div)
    if not k:
        return declared
    allowed = AUTH_MEAS[k]
    toks = [t.strip() for t in str(declared).split("/") if t.strip()]
    hit = [t for t in toks if t.split("(")[0].strip() in allowed]
    return " / ".join(hit) if hit else declared


def edit_day(auth, scr, div):
    """편집 가능 요일 — GET_ADMIN_GROUP_MEASURE.ATTRIBUTE1 실측 (MEASURE 단위)."""
    k = auth_key(auth, scr, div)
    if not k or k not in AUTH_DAY:
        return "-"
    pairs = sorted({(m["MEASURE"], (m.get("EDIT_DAY") or "").strip())
                    for m in META["auth_measure"]
                    if (m["AUTH"], m["SCR"], m["DIV"]) == k
                    and (m.get("EDIT_DAY") or "").strip()})
    return ", ".join("{}={}".format(a, b) for a, b in pairs) or "-"


def meta_basis(auth, scr, div):
    k = auth_key(auth, scr, div)
    if not k:
        return "메타 미등재(제한 없음 전제)"
    return "GET_ADMIN_GROUP_MEASURE[{}/{}/{}] = {}".format(
        auth, scr, k[2], ", ".join(sorted(AUTH_MEAS[k])))


def inco_flag(v):
    v = str(v or "").strip()
    if not v or v == "-":
        return "-"
    if v in INCO_Y:
        return "Y"
    if v in INCO_N:
        return "N"
    return "?"


def spec_text(cols, rec, src_desc, limit=900):
    parts = []
    for c in cols:
        v = rec.get(c)
        if v in (None, "", "-"):
            continue
        parts.append("{}={}".format(c, v))
    s = " / ".join(parts)
    if not s:
        return "-"
    if len(s) > limit:
        s = s[:limit] + " …(STEP 시트 전체 참조)"
    return "LAYOUT[{}] ← {}".format(src_desc, s)


def sheet_rows(ws, header_row=1, first_data=2):
    hdr = [c.value for c in next(ws.iter_rows(min_row=header_row, max_row=header_row))]
    out = []
    for r in ws.iter_rows(min_row=first_data, values_only=True):
        if not any(x not in (None, "") for x in r):
            continue
        out.append(OrderedDict((h, r[i] if i < len(r) else None)
                               for i, h in enumerate(hdr) if h is not None))
    return hdr, out


def main():
    wb1 = openpyxl.load_workbook(SRC_V1, data_only=True)
    wb2 = openpyxl.load_workbook(SRC_V2, data_only=True)

    DIVTEXT = {d: parse_division(wb2[SD.SHEET_OF[d]], d) for d in ("40", "81", "82")}
    STEPDEF = {(s[0], s[1]): s for s in SD.STEPS}

    # ── STEP 별 대응결함No (v1 에서 승계) ──
    defect = {}
    for d in ("40", "81", "82"):
        _, rows = sheet_rows(wb1[SD.OUT_SHEET[d]])
        for r in rows:
            v = str(r.get("대응결함No") or "-").strip()
            if v and v != "-":
                defect.setdefault((d, str(r["StepNo"])), set()).add(v)

    rows_by_div = {"40": [], "81": [], "82": []}
    v2_by_step = {}
    step_sheets = []

    # ── ① v1(06) 299행 — 이전 양식 그대로 승계 ────────────────────────────
    for d in ("40", "81", "82"):
        _, rows = sheet_rows(wb1[SD.OUT_SHEET[d]])
        for r in rows:
            sno = str(r["StepNo"])
            sd = STEPDEF.get((d, sno))
            t = DIVTEXT[d].get(sno, {})
            rec = {c: ("" if r.get(c) is None else r.get(c)) for c in V1_COLS}
            rec["항목"] = t.get("항목") or rec["항목"]
            rec["엑셀항목"] = rec["항목"]
            rec["END_CUSTOMER"] = "-"
            rec["INCOTERMS_EDIT_FLAG"] = inco_flag(rec.get("INCOTERMS"))
            rec["CELL_PROD_SITE_ID"] = {
                "40": NOT_MEASURED,
                "81": "(해당없음 — 81 축 = PLANT_SITE_ID)",
                "82": "(해당없음 — 82 화면 LAYOUT 미보유)"}[d]
            rec["CURRENCY"] = "-"
            rec["화면LAYOUT원천"] = "-"
            rec["LAYOUT컬럼수"] = ""
            rec["G열요구사항"] = t.get("G열") or "-"
            rec["STEP시트"] = ("{}_{}".format(d, sno)
                             if sd and sd[9] else "-")
            rec["행출처"] = "v1(06)"
            rows_by_div[d].append(rec)

    # ── ② v2(07) STEP 시트 — 이전 양식으로 변환 합류 ───────────────────────
    td2 = {}
    for d in ("40", "81", "82"):
        _, rows = sheet_rows(wb2[SD.OUT_SHEET[d]])
        for r in rows:
            td2.setdefault((d, str(r["StepNo"])), []).append(r)

    for (d, sno, smp, scr, real, ledger, groups, meas, kind, lkey) in SD.STEPS:
        if not lkey:
            continue
        name = "{}_{}".format(d, sno)
        if name not in wb2.sheetnames:
            print("  !! STEP 시트 없음 {}".format(name))
            continue
        L = LAYOUT.get(lkey, {})
        hdr, srows = sheet_rows(wb2[name], header_row=1, first_data=3)
        meta_n = 8
        # 30_layouts 가 division 토큰 중복 등록분을 그대로 실어 같은 COLUMN_NAME 이
        # 여러 번 나온다(07 기준 109개). 입력 LAYOUT 으로는 1회면 되므로 선순위 1개만 남긴다.
        lay_cols, _seen = [], set()
        for h in hdr[meta_n:]:
            if h and h not in _seen:
                _seen.add(h)
                lay_cols.append(h)
        t = DIVTEXT[d].get(sno, {})
        tds = list(td2.get((d, sno), []))
        if len(tds) != len(srows):
            print("  !! 행수 불일치 {} STEP {} / TestData {}".format(
                name, len(srows), len(tds)))
        for i, sr in enumerate(srows):
            td = tds[i] if i < len(tds) else {}
            grp = str(sr.get("권한그룹") or td.get("권한그룹") or "-")
            acct = str(sr.get("테스트계정") or td.get("테스트계정") or "-")
            src_kind = str(sr.get("데이터출처") or "-")
            basis = str(sr.get("분류기준") or "-")
            # LAYOUT 실값 + TestData 보정값 (IF LAYOUT 은 7키가 컬럼에 없다)
            val = {c: sr.get(c) for c in lay_cols}
            get = lambda c: (val.get(c) if val.get(c) not in (None, "")
                             else (td.get(c) if td.get(c) not in (None, "") else "-"))
            if acct != "-" and acct in PEOPLE:
                resolved, wsyn = acct_detail(acct, real)
                _, alt, _, _ = pick_acct(grp, real)
            else:
                acct2, alt, resolved, wsyn = pick_acct(grp, real)
                acct = acct if acct != "-" else acct2
            m_real = val.get("MEASURE") or td.get("입력MEASURE")
            rec = {
                "StepNo": sno, "Step": t.get("Step", ""), "항목": t.get("항목", ""),
                "샘플지정": smp, "데이터출처": src_kind, "분류기준": basis,
                "DIV": d, "구분": kind, "엑셀항목": t.get("항목", ""),
                "화면": scr, "실제화면": real, "화면명": PNAME.get(real, "-"),
                "권한그룹": grp, "테스트계정": acct, "대체계정": alt,
                "해석authGroup(실측)": resolved, "W/S판정(실측)": wsyn,
                "ITEM_ID": get("ITEM_ID"), "SCM_SITE_ID": get("SCM_SITE_ID"),
                "PROJECT_CODE": get("PROJECT_CODE"), "FCST_TYPE": get("FCST_TYPE"),
                "INCOTERMS": get("INCOTERMS"), "DOM_EXP": get("DOM_EXP"),
                "PLANT_SITE_ID": get("PLANT_SITE_ID"),
                "입력MEASURE": (m_real if m_real not in (None, "", "-")
                              else measures_for(grp, real, d, meas)),
                "편집요일": edit_day(grp, real, d),
                "권한메타근거": meta_basis(grp, real, d),
                "대상원장": ledger,
                "등록값명세": spec_text(lay_cols, val, "{} / {}컬럼".format(
                    L.get("kind", "?"), len(lay_cols))),
                "데이터보유(실측)": ("없음 → 신규입력" if src_kind.startswith("신규")
                               else "있음"),
                # 신규 행은 원장에 없으므로 실측 대상이 아니다(v1 규약 = '-').
                # 기존 행만 재측정 대기 표기 → 42_refill_measure.py 가 채운다.
                "기간(실측)": ("-" if src_kind.startswith("신규") else NOT_MEASURED),
                "현재값(실측)": ("-" if src_kind.startswith("신규") else NOT_MEASURED),
                "테스트행위": "{} 에서 {} 권한으로 {} 입력/검증".format(
                    real if real != "(IF)" else ledger, grp,
                    (measures_for(grp, real, d, meas)
                     if meas not in ("-", "", None) else t.get("항목", "대상"))),
                "기대결과": ("신규 키가 정상 저장되고 후속 STEP 에 반영"
                        if src_kind.startswith("신규")
                        else "정상 저장 및 후속 STEP 반영"),
                "검증화면": ("-" if real == "(IF)" and kind == "IF" else "DMF0035"),
                "대응결함No": ("STEP 승계: " + ", ".join(sorted(defect[(d, sno)]))
                           if (d, sno) in defect else "-"),
                "END_CUSTOMER": get("END_CUSTOMER"),
                "INCOTERMS_EDIT_FLAG": inco_flag(get("INCOTERMS")),
                "CELL_PROD_SITE_ID": get("CELL_PROD_SITE_ID"),
                "CURRENCY": get("CURRENCY"),
                "화면LAYOUT원천": str(td.get("화면LAYOUT원천") or "-"),
                "LAYOUT컬럼수": len(lay_cols),
                "G열요구사항": t.get("G열") or "-",
                "STEP시트": name, "행출처": "v2(07)",
            }
            rows_by_div[d].append(rec)
            v2_by_step.setdefault((d, sno), []).append(rec)
        step_sheets.append((name, d, sno, hdr, srows, lay_cols, meta_n))

    # ── ③ 동일 키 병합 ────────────────────────────────────────────────────
    # v1 행(기간/현재값 실측 보유)을 남기고, 같은 키의 v2 행이 가진
    # 신규 컬럼(LAYOUT·CELL_PROD_SITE·CURRENCY·END_CUSTOMER·등록값명세)을 채워 넣는다.
    MERGE_IN = ["END_CUSTOMER", "CELL_PROD_SITE_ID", "CURRENCY", "화면LAYOUT원천",
                "LAYOUT컬럼수", "STEP시트", "등록값명세", "권한메타근거", "편집요일"]
    DEDUP = ["StepNo", "데이터출처", "권한그룹", "테스트계정"] + KEY7 + ["입력MEASURE"]
    # ⚠ 병합은 v1↔v2 **교차**만 허용한다. 같은 회차 안에서는 합치지 않는다
    #   — IF LAYOUT 행은 7키 투영이 전부 '-' 라 서로 달라도 키가 같아지기 때문.
    merged_log = {}
    for d in ("40", "81", "82"):
        v1rows = [r for r in rows_by_div[d] if r["행출처"] == "v1(06)"]
        v2rows = [r for r in rows_by_div[d] if r["행출처"] == "v2(07)"]
        idx = {}
        for r in v1rows:
            idx.setdefault(tuple(str(r.get(c, "")) for c in DEDUP), []).append(r)
        keep = list(v1rows)
        for r in v2rows:
            k = tuple(str(r.get(c, "")) for c in DEDUP)
            pool = idx.get(k)
            if pool:
                base = pool.pop(0)
                for c in MERGE_IN:
                    v = r.get(c, "")
                    if v not in ("", "-", None) and base.get(c, "") in ("", "-", None,
                                                                        NOT_MEASURED):
                        base[c] = v
                base["행출처"] = "v1(06)+v2(07)"
                key = (d, r["StepNo"], str(r["데이터출처"]))
                merged_log[key] = merged_log.get(key, 0) + 1
                continue
            keep.append(r)
        if len(keep) != len(rows_by_div[d]):
            print("  DIV {} 동일키 {}행 병합".format(d, len(rows_by_div[d]) - len(keep)))
        rows_by_div[d] = keep
    for k, v in sorted(merged_log.items()):
        print("     병합 {} {} {} — {}행".format(*k, v))

    # ── ④ 저장 ──────────────────────────────────────────────────────────
    # 07 을 그대로 복사해 DIVISION 3시트 서식을 무손실 보존하고, 나머지는 재생성한다.
    import shutil
    shutil.copyfile(SRC_V2, OUT)
    out = openpyxl.load_workbook(OUT)
    for n in list(out.sheetnames):
        if n not in [SD.SHEET_OF[d] for d in ("40", "81", "82")]:
            del out[n]

    width = V1_WIDTH + NEW_WIDTH
    for d in ("40", "81", "82"):
        ws = out.create_sheet(SD.OUT_SHEET[d])
        ws.append(COLS)
        for c in ws[1]:
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
        ordered = sorted(rows_by_div[d],
                         key=lambda x: (SD.snum(x["StepNo"]), str(x["데이터출처"]),
                                        str(x["권한그룹"]), str(x["ITEM_ID"])))
        for r in ordered:
            ws.append([r.get(c, "") for c in COLS])
            f = FILL.get(str(r["데이터출처"]))
            for c in ws[ws.max_row]:
                if f:
                    c.fill = PatternFill("solid", fgColor=f)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(size=9)
                c.border = BORDER
        for i, w in enumerate(width, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "D2"
        ws.row_dimensions[1].height = 30
        cnt = {}
        for r in ordered:
            cnt[str(r["데이터출처"])] = cnt.get(str(r["데이터출처"]), 0) + 1
        print("{} — {:,}행  {}".format(SD.OUT_SHEET[d], len(ordered), cnt))

    # ── ⑤ STEP 별 시트 ─────────────────────────────────────────────────
    META_COLS = ["StepNo", "Step", "항목", "DIV", "구분", "화면", "실제화면", "화면명",
                 "권한그룹", "테스트계정", "대체계정", "해석authGroup(실측)",
                 "W/S판정(실측)", "편집요일", "데이터출처", "분류기준", "대상원장"]
    td_index = v2_by_step

    for (name, d, sno, hdr, srows, lay_cols, meta_n) in step_sheets:
        sd = STEPDEF[(d, sno)]
        ws = out.create_sheet(name)
        head = META_COLS + lay_cols
        ws.append(head)
        lbls = [c.value for c in next(wb2[name].iter_rows(min_row=2, max_row=2))]
        ws.append(["-"] * len(META_COLS) +
                  [lbls[meta_n + i] if meta_n + i < len(lbls) else c
                   for i, c in enumerate(lay_cols)])
        tds = td_index.get((d, sno), [])
        for i, sr in enumerate(srows):
            td = tds[i] if i < len(tds) else {}
            ws.append([td.get(c, "") for c in META_COLS] +
                      [("" if sr.get(c) is None else sr.get(c)) for c in lay_cols])
            f = FILL.get(str(td.get("데이터출처", "")))
            for c in ws[ws.max_row]:
                if f:
                    c.fill = PatternFill("solid", fgColor=f)
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(size=9)
                c.border = BORDER
        for c in ws[1]:
            c.fill, c.font = HDR_FILL, HDR_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDER
        for c in ws[2]:
            c.fill, c.font = SUB_FILL, Font(size=8, italic=True)
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for i in range(1, len(head) + 1):
            w = 8
            for r in range(1, min(ws.max_row, 200) + 1):
                v = ws.cell(r, i).value
                if v is not None:
                    w = max(w, min(38, len(str(v)) + 2))
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "D3"
        ws.row_dimensions[1].height = 30
        print("  {:10s} {:3d}행 x {:3d}컬럼".format(name, len(srows), len(head)))

    order = ([SD.SHEET_OF[d] for d in ("40", "81", "82")] +
             [SD.OUT_SHEET[d] for d in ("40", "81", "82")] +
             sorted([n for n, *_ in step_sheets],
                    key=lambda n: (n.split("_")[0], SD.snum(n.split("_")[1]))))
    out._sheets = [out[n] for n in order if n in out.sheetnames]
    out.save(OUT)
    print("\n저장 → {}".format(OUT))
    try:
        import shutil
        shutil.copyfile(OUT, DELIVER)
        print("전달본 → {}".format(DELIVER))
    except Exception as e:
        print("전달본 복사 실패: {}".format(e))


if __name__ == "__main__":
    main()
