# -*- coding: utf-8 -*-
"""v3 STEP 시트 LAYOUT 공백 채움 — ZIOUI1 복구 후 실행 (읽기 전용 조회)

문제 (2026-09-20 census)
  08_SAMPLEDATA_STEPNO_v3.xlsx 의 STEP 시트 30개에 **전(全)행 공백인 LAYOUT 컬럼 411개**
  가 남아 있다. 원인은 31_pick.existing_rows() 가 화면 LAYOUT 전체가 아니라
  고정 화이트리스트(DIVISION_ID/SCM_SITE_ID/ITEM_ID/PROJECT_CODE/INCOTERMS/…)만
  SELECT 하기 때문이다. IF 행도 `공통컬럼` 교집합만 투영해 나머지가 비었다.

본 스크립트는 각 STEP 시트 행의 **키로 원장을 다시 읽어 SELECT \\* 한 값**으로
공백 LAYOUT 셀만 채운다. 값 생성은 하지 않는다 — 조회되지 않으면 공백을 유지하고
`_공백사유` 를 리포트한다(anti-simplification).

  python 44_refill_layout.py            # 리포트만
  python 44_refill_layout.py --write    # 08 xlsx 에 반영

⚠ UPLOAD LAYOUT(81_4-3)은 컬럼명이 엑셀 헤더 라벨('Customer Model*')이라
  원장 컬럼과 직접 매칭되지 않는다. SCM_ADM_EXCEL_COLUMNS ↔ DM_PLN_CUSTOMER_PSI
  매핑이 별도로 필요하므로 본 스크립트는 대상에서 제외하고 리포트만 한다.
⚠ 최초 실행 시 반드시 --write 없이 리포트를 먼저 확인할 것.
"""
import importlib.util
import pathlib
import sys

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
XLSX = HERE.parent / "08_SAMPLEDATA_STEPNO_v3.xlsx"
META_N = 17            # STEP 시트 메타 컬럼 수 (41_build_v3.META_COLS)
LABEL_ROW = 2          # 2행 = 컬럼 라벨

_p = importlib.util.spec_from_file_location("pick", HERE / "31_pick.py")
PICK = importlib.util.module_from_spec(_p)
_p.loader.exec_module(PICK)

_s = importlib.util.spec_from_file_location("stepdef", HERE / "40_stepdef_v3.py")
SD = importlib.util.module_from_spec(_s)
_s.loader.exec_module(SD)

# 실제화면 -> 원장 (UI LAYOUT 행의 재조회 대상)
UI_TABLE = {
    "DMF0082": "DM_PLN_MKT_FCST", "DMF0003": "DM_PLN_MKT_FCST",
    "DMF0021": "DM_PLN_SALES_PLAN", "DMF0020": "SCM_MST_FUTURE_PRICE",
    "DMF0005": "SCM_MST_SITE_ITEM", "DMF0007": "DM_MST_CUSTOMER_FCST",
    "DMF0072": "DM_CUST_MODEL_MAPPING_M",
}
# IF STEP 의 값 출처 표 — 단일 원천은 42_refill_measure.IF_SRC
_r = importlib.util.spec_from_file_location("rf42", HERE / "42_refill_measure.py")
RF42 = importlib.util.module_from_spec(_r)
_r.loader.exec_module(RF42)
IF_TABLE = RF42.IF_SRC
SKIP = {"81_4-3"}      # UPLOAD LAYOUT — 라벨↔컬럼 매핑 미확보
SENTINEL = ("미측정", "해당없음", "재조회", "(")   # 값이 아니라 표기인 셀 — 필터로 쓰면 0건
KEYCOLS = ["ITEM_ID", "SCM_SITE_ID", "PROJECT_CODE", "FCST_TYPE", "INCOTERMS",
           "DOM_EXP", "PLANT_SITE_ID", "CELL_PROD_SITE_ID", "CUSTOMER_CODE",
           "SOLDTO_CODE", "IFSEQNO", "IF_SEQ", "SEQ"]

_cols = {}


def cols_of(cur, owner, table):
    k = (owner, table)
    if k not in _cols:
        cur.execute("""SELECT COLUMN_NAME FROM ALL_TAB_COLUMNS
                        WHERE OWNER=:o AND TABLE_NAME=:t""", {"o": owner, "t": table})
        _cols[k] = {r[0] for r in cur}
    return _cols[k]


def fetch_one(cur, owner, table, div, key):
    have = cols_of(cur, owner, table)
    if not have:
        return None, "표 없음 {}.{}".format(owner, table)
    where, bind = [], {}
    if "DIVISION_ID" in have:
        where.append(PICK.div_pred(div))
    n = 0
    for c, v in key.items():
        if c not in have or v in (None, "", "-") or any(t in str(v) for t in SENTINEL):
            continue
        where.append("{} = :b{}".format(c, n))
        bind["b{}".format(n)] = v
        n += 1
    if n == 0:
        return None, "조회키 0"
    cur.execute("SELECT * FROM {}.{} WHERE {} FETCH FIRST 1 ROWS ONLY".format(
        owner, table, " AND ".join(where)), bind)
    names = [d[0] for d in cur.description]
    r = cur.fetchone()
    if not r:
        return None, "키 미존재"
    return dict(zip(names, r)), "OK"


def main(write=False):
    wb = openpyxl.load_workbook(XLSX)
    db = PICK.connect()
    cur = db.cursor()
    filled = blanks = 0
    report = []
    for name in wb.sheetnames:
        if "_" not in name or name.split("_")[0] not in ("40", "81", "82"):
            continue
        div, sno = name.split("_", 1)
        ws = wb[name]
        hdr = [c.value for c in ws[1]]
        lay = [(i + 1, h) for i, h in enumerate(hdr) if i >= META_N and h]
        if name in SKIP:
            report.append((name, 0, 0, "SKIP — UPLOAD 라벨↔컬럼 매핑 미확보"))
            continue
        owner, table = IF_TABLE.get(name, (None, None))
        if not table:
            scr = str(ws.cell(3, 7).value or "")     # 실제화면
            table, owner = UI_TABLE.get(scr), "SCM_MGR"
        if not table:
            report.append((name, 0, 0, "원장 미지정 — 건너뜀"))
            continue
        f = b = 0
        reasons = {}
        for row in ws.iter_rows(min_row=3):
            if not any(c.value not in (None, "") for c in row):
                continue
            cur_vals = {hdr[i]: row[i].value for i in range(len(hdr)) if hdr[i]}
            key = {c: cur_vals.get(c) for c in KEYCOLS if cur_vals.get(c)}
            empties = [(ci, h) for ci, h in lay if row[ci - 1].value in (None, "")]
            if not empties:
                continue
            try:
                rec, why = fetch_one(cur, owner, table, div, key)
            except Exception as e:
                rec, why = None, "조회오류: " + str(e).split("\n")[0][:60]
            if not rec:
                b += len(empties)
                reasons[why] = reasons.get(why, 0) + 1
                continue
            for ci, h in empties:
                v = rec.get(h)
                if v in (None, ""):
                    b += 1
                    reasons["원장에도 NULL/컬럼없음"] = reasons.get(
                        "원장에도 NULL/컬럼없음", 0) + 1
                    continue
                if write:
                    row[ci - 1].value = v
                f += 1
        filled += f
        blanks += b
        report.append((name, f, b, "{}.{} · ".format(owner, table) +
                       ", ".join("{}×{}".format(k, v) for k, v in
                                 sorted(reasons.items(), key=lambda x: -x[1])[:3])))
    print("채움 {:,}셀 / 잔여공백 {:,}셀".format(filled, blanks))
    for n, f, b, note in report:
        print("  {:10s} 채움 {:5d}  공백 {:5d}  {}".format(n, f, b, note[:110]))
    if write:
        wb.save(XLSX)
        print("\n반영 저장 → {}".format(XLSX))
        print("→ 반드시 43_verify_v3.py 재실행할 것")
    return 0


if __name__ == "__main__":
    sys.exit(main("--write" in sys.argv))
