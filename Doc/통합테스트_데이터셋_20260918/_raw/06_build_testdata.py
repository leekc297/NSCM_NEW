# -*- coding: utf-8 -*-
"""
Phase 2 — 04_TESTDATA_SET.xlsx (테스트 대상 지정서)

프로세스 단계(04_flow_def) × 대상 ITEM 5종(03_targets_v2) 을 교차해
**ZIOUI1 실존 행**만으로 시나리오 행을 만든다. 예측값 생성 금지.
"""
import json, re, pathlib, sys, datetime, importlib.util
sys.path.insert(0, r"C:\Projects\NSCM\_artifacts\_diagnostics")
import oracledb
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

D = pathlib.Path(__file__).resolve().parent
OUT = D.parent
spec = importlib.util.spec_from_file_location("flow_def", D / "04_flow_def.py")
FD = importlib.util.module_from_spec(spec); spec.loader.exec_module(FD)
META = json.loads((D / "01_meta.json").read_text(encoding="utf-8"))
TG = json.loads((D / "03_targets_v2.json").read_text(encoding="utf-8"))

# ── 권한 -> 계정 ────────────────────────────────────────────────────────────
# ⚠ 정정(2026-09-18): 권한그룹 보유만으로 계정을 고르면 안 된다.
#   FE 가 권한을 해석하는 CMBO 3종(FcstPrice/SalesPlan/BizPlan UserAuthGroup)은
#   ORDER BY DECODE(USER_GR_ID,'ADMIN',1,2) 로 ADMIN 을 1행째로 올리고 FE 는 1행만 채택한다.
#   → ADMIN 겸직 계정(minyonglee1, jaehun_cho)은 화면에서 항상 ADMIN 으로 해석되어
#     DM Planner / SCP Planner 권한 경계를 검증할 수 없다.
#   08_acct_resolve.py 가 계정별 CMBO 실행 결과를 실측해 08_accounts.json 에 넣는다.
ACCR = json.loads((D / "08_accounts.json").read_text(encoding="utf-8"))
PEOPLE = {p["USER_ID"]: p for p in ACCR["people"]}
RECO = ACCR["recommend"]


def _rank(uid, auth):
    """순수 단일권한 계정 우선 → 보유그룹 수 적은 순 → 사전순."""
    p = PEOPLE[uid]
    pure = 0 if p["groups"] == [auth] else 1
    return (pure, len(p["groups"]), uid)


def pick_acct(auth, scr):
    """(계정, 대체계정, 해석authGroup 실측, W/S판정 실측) — 실측 해석 기준."""
    if auth == "-":
        return "(권한무관)", "", "-", "-"
    if auth == "DM ADMIN/ADMIN":
        # ADMIN 겸직 계정이 오히려 정답인 유일한 단계 (DMF0049 G0 게이트 우회)
        cand = [u for u, p in PEOPLE.items() if p["admin_masked"]]
        cand.sort()
        if not cand:
            return "(없음)", "", "미해석", "-"
        u = cand[0]
        return u, ", ".join(cand[1:])[:60], "ADMIN", PEOPLE[u]["wsPlannerYn"]

    r = RECO.get(auth)
    if not r:
        return "(없음)", "", "미해석", "-"
    ok = sorted(r["사용가능"], key=lambda u: _rank(u, auth))
    if not ok:
        return "(없음)", "", "미해석", "-"
    u = ok[0]
    p = PEOPLE[u]
    # 해석 CMBO 가 있는 화면만 실측값, 나머지는 메타 미등재로 명시 (DMF0082 값 전용 금지)
    if scr in ("DMF0082", "DMF0021", "DMF0042", "DMF0003"):
        resolved = p["resolved"].get(scr) or "메타 미등재(제한 없음 전제)"
    else:
        resolved = "메타 미등재(제한 없음 전제)"
    alt = ", ".join(ok[1:])[:60]
    if r["사용불가(ADMIN 해석)"]:
        alt += "  |  ⚠사용불가(ADMIN 해석): " + ", ".join(r["사용불가(ADMIN 해석)"])
    return u, alt[:160], str(resolved), p["wsPlannerYn"]

# 권한 x 화면 x DIV -> 편집 메저·요일 (실측)
AM = {}
for r in META["auth_measure"]:
    AM.setdefault((r["AUTH"], r["SCR"], r["DIV"]), []).append(
        (r["MEASURE"], r["EDIT_DAY"]))

VER = {}
for v in META["versions"]:
    VER.setdefault(v["DIVISION_ID"], {})[v["CUTOFF_TYPE"]] = v["LATEST_VER"]


def conn():
    sh = pathlib.Path(r"C:\Projects\NSCM\_artifacts\_apply_corporate\apply.sh")
    m = re.search(r"CONN='([^/]+)/\"([^\"]+)\"@([^']+)'", sh.read_text(encoding="utf-8", errors="ignore"))
    return oracledb.connect(user=m.group(1), password=m.group(2), dsn=m.group(3))


COLS = ["시나리오ID", "프로세스단계", "DIV", "구분", "화면", "화면명", "권한그룹", "테스트계정", "대체계정",
        "해석authGroup(실측)", "W/S판정(실측)",
        "ITEM_ID", "SCM_SITE_ID", "PROJECT_CODE", "FCST_TYPE", "INCOTERMS", "DOM_EXP",
        "PLANT_SITE_ID", "입력MEASURE", "편집요일", "권한메타근거",
        "대상원장", "데이터보유(실측)", "기간(실측)", "현재값(실측)",
        "테스트행위", "기대결과", "검증화면", "대응결함No"]

rows = []
G1_FAIL = []
with conn() as db:
    cur = db.cursor()
    for s in FD.STEPS:
        sid, div, order, kind, scr, auth, meas, day, ledger, job, sp, ver, defno = s
        if kind in ("배치",):                      # 배치는 지정서 대상 아님
            continue
        targets = TG["divisions"].get(div, [])
        for idx, t in enumerate(targets, 1):
            acct, alt, resolved, wsyn = pick_acct(auth, scr)
            key = (auth, scr, div)
            meta_ev = ""
            if key in AM:
                meta_ev = ", ".join("{}{}".format(m, "(%s)" % d if d and d.strip() else "")
                                    for m, d in AM[key])
            else:
                meta_ev = "메타 미등재(제한 없음 전제)"

            # 실측: 이 조합의 해당 measure 현황
            have, period, curval = "", "", ""
            mname = meas.split("(")[0].strip()
            if ledger == "DM_PLN_MKT_FCST" and mname in (
                    "MKT_FCST", "SCM_FCST", "CUST_FCST", "CUSTOMER FCST"):
                mm = mname if mname != "CUSTOMER FCST" else "CUST_FCST"
                hit = [x for x in t["measures"] if x["MEASURE"] == mm]
                if hit:
                    h = hit[0]
                    have = "있음 {}행".format(h["CNT"])
                    period = "{} ~ {}".format(h["MIN_YMD"], h["MAX_YMD"])
                    curval = "SUM(QTY)={}".format(h["SUM_QTY"])
                else:
                    have = "없음 → 신규입력"
            elif ledger == "DM_PLN_CFM_SELLING_PRICE":
                ps = t.get("price_selling") or []
                have = "있음 {}종".format(len(ps)) if ps else "없음"
                if ps:
                    period = "{} ~ {}".format(ps[0]["MIN_YMD"], ps[0]["MAX_YMD"])
                    curval = "; ".join("{} {} avg={}".format(p["MEASURE"], p["CURRENCY"],
                                                            p["AVG_PRICE"]) for p in ps[:3])
            elif ledger == "SCM_MST_FUTURE_PRICE":
                pf = t.get("price_future") or []
                have = "있음 {}종".format(len(pf)) if pf else "없음"
                if pf:
                    period = "{} ~ {}".format(pf[0]["MIN_YMD"], pf[0]["MAX_YMD"])
                    curval = "; ".join("{} {} avg={}".format(p["MEASURE"], p["CURRENCY"],
                                                            p["AVG_PRICE"]) for p in pf[:3])
            elif ledger in ("DM_PLN_SALES_PLAN", "DM_PLN_BIZPLAN"):
                # BIZPLAN 은 월 단위 컬럼(YYYYMM), SALES_PLAN 은 일 단위(YYYYMMDD)
                dcol = "YYYYMM" if ledger == "DM_PLN_BIZPLAN" else "YYYYMMDD"
                cur.execute(
                    "SELECT COUNT(*), MIN({c}), MAX({c}) FROM SCM_MGR.{t} "
                    "WHERE DIVISION_ID=:d AND ITEM_ID=:i AND SCM_SITE_ID=:s".format(
                        c=dcol, t=ledger),
                    d=div, i=t["ITEM_ID"], s=t["SCM_SITE_ID"])
                c0, mn, mx = cur.fetchone()
                have = "있음 {}행".format(c0) if c0 else "없음 → 신규입력"
                period = "{} ~ {}".format(mn, mx) if c0 else ""

            if have.startswith("없음"):
                G1_FAIL.append((sid, div, t["ITEM_ID"], mname, ledger))

            act = {"입력": "{} 에서 {} 입력/수정".format(scr, meas),
                   "선행/입력": "{} 로 {} 유입 확인".format(scr, meas),
                   "선행/배치": "EDI 배치 실행 후 {} 유입 확인".format(meas),
                   "검증": "{} 3 TYPE 조회".format(scr)}.get(kind, kind)

            rows.append([
                "{}-I{}".format(sid, idx), sid, div, kind, scr, FD.scr_name(scr), auth, acct, alt,
                resolved, wsyn,
                t["ITEM_ID"], t["SCM_SITE_ID"], t["PROJECT_CODE"], t["FCST_TYPE"],
                t["INCOTERMS"], t["DOM_EXP"], t["PLANT_SITE_ID"],
                meas, day, meta_ev, ledger, have, period, curval,
                act, ver, "DMF0035" if kind != "검증" else "-", defno])

# ── XLSX ──
HF, HFONT = PatternFill("solid", fgColor="1F4E78"), Font(color="FFFFFF", bold=True, size=9)
wb = openpyxl.Workbook()
ws = wb.active; ws.title = "TESTDATA_SET"
ws.append(COLS)
for c in ws[1]:
    c.fill, c.font = HF, HFONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
FILL = {"40": "E2EFDA", "81": "DDEBF7", "82": "FCE4D6"}
for r in rows:
    ws.append(r)
    for c in ws[ws.max_row]:
        c.fill = PatternFill("solid", fgColor=FILL.get(r[2], "FFFFFF"))
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.font = Font(size=9)
for i, w in enumerate([13, 10, 5, 10, 9, 22, 17, 13, 30, 18, 8, 20, 24, 16, 14, 11, 8, 14,
                       22, 8, 34, 28, 16, 22, 34, 34, 34, 10, 22], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "C2"

ws2 = wb.create_sheet("대상ITEM_실측")
ws2.append(["DIV", "ITEM_ID", "SCM_SITE_ID", "PROJECT_CODE", "FCST_TYPE", "INCOTERMS",
            "DOM_EXP", "PLANT_SITE_ID", "FOCUS보유", "FOCUS결손", "기간", "총행수",
            "보유MEASURE(실측)"])
for c in ws2[1]:
    c.fill, c.font = HF, HFONT
for div, ts in TG["divisions"].items():
    for t in ts:
        ws2.append([div, t["ITEM_ID"], t["SCM_SITE_ID"], t["PROJECT_CODE"], t["FCST_TYPE"],
                    t["INCOTERMS"], t["DOM_EXP"], t["PLANT_SITE_ID"],
                    ", ".join(t["focus_have"]), ", ".join(t["focus_miss"]) or "-",
                    "{} ~ {}".format(t["MIN_YMD"], t["MAX_YMD"]), t["ROW_CNT"],
                    ", ".join("{}({})".format(m["MEASURE"], m["CNT"]) for m in t["measures"])])
for i, w in enumerate([6, 20, 26, 16, 14, 11, 8, 14, 34, 20, 22, 10, 90], 1):
    ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
for row in ws2.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True); c.font = Font(size=9)

ws3 = wb.create_sheet("권한매트릭스_실측")
ws3.append(["권한그룹", "화면", "DIV", "편집MEASURE", "편집요일"])
for c in ws3[1]:
    c.fill, c.font = HF, HFONT
for r in META["auth_measure"]:
    ws3.append([r["AUTH"], r["SCR"], r["DIV"], r["MEASURE"], r["EDIT_DAY"] or "-"])
for i, w in enumerate([20, 11, 9, 20, 10], 1):
    ws3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws4 = wb.create_sheet("테스트계정_권한해석")
ws4.append(["USER_ID", "USER_NM", "보유그룹(전체)", "W/S판정(MktFcstWsPlannerYn)",
            "DMF0082 해석", "DMF0021 해석", "DMF0042 해석", "ADMIN겸직", "판정"])
for c in ws4[1]:
    c.fill, c.font = HF, HFONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for p in ACCR["people"]:
    r = p["resolved"]
    verdict = ("⚠ 권한경계 검증 불가 — 항상 ADMIN 으로 해석됨 (DM ADMIN/ADMIN 단계 전용)"
               if p["admin_masked"] else "사용 가능")
    ws4.append([p["USER_ID"], p["USER_NM"], ", ".join(p["groups"]), p["wsPlannerYn"],
                str(r["DMF0082"]), str(r["DMF0021"]), str(r["DMF0042"]),
                "Y" if p["admin_masked"] else "N", verdict])
    if p["admin_masked"]:
        for c in ws4[ws4.max_row]:
            c.fill = PatternFill("solid", fgColor="FCE4D6")
for i, w in enumerate([16, 12, 52, 24, 20, 20, 20, 10, 54], 1):
    ws4.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
for row in ws4.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True); c.font = Font(size=9)

ws6 = wb.create_sheet("CMBO_사전확인")
ws6.append(["CMBO SQL_ID", "USE_YN", "최종수정", "역할", "미등록/미사용 시 영향"])
for c in ws6[1]:
    c.fill, c.font = HF, HFONT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ROLE = {
 "MktFcstWsPlannerYn": ("DMF0082 W/S 권한 판정 단일 원천",
   "⚠ 0행/미등록이면 wsPlannerYn=false 로 조용히 degrade — AS_In_Req·W/S FCST Type·AS FROM SITE 검증이 오류 없이 전부 꺼진다"),
 "FcstPriceUserAuthGroup": ("DMF0082 authGroup 해석", "권한별 편집 measure 판정 불가"),
 "SalesPlanUserAuthGroup": ("DMF0021 authGroup 해석", "권한별 편집 measure 판정 불가"),
 "BizPlanUserAuthGroup": ("DMF0042 authGroup 해석 (SP IN 14 P_USERINFO 필수)", "조회 자체가 권한 미해석 상태로 나감"),
 "FcstTypeWs": ("W/S FCST Type 집합(셀 개방 ∩ 저장검증)", "컬럼은 열리는데 MSG00476 미발화 (2026-08-06 사고)"),
 "FcstTypeEditMode": ("W/S 전용 FCST Type 등록본 6종", "W/S FCST Type 스왑 불가"),
 "SalesFuturePriceEditableDivision": ("DMF0020 편집가능 division 화이트리스트", "82 단가 입력 불가"),
 "MyDivision": ("DMF0035 division 선택", "조회 불가 (MSG00097)"),
 "DMF0042_BP_Measure": ("DMF0042 Insert 행 measure 후보", "삽입 차단 MSG00113"),
 "Dmf0003ScpPlannerEditFlag": ("DMF0003 SCP Planner 시간창", "0행 = 편집 차단 MSG00477"),
}
for cid, v in ACCR["combos"].items():
    role, impact = ROLE.get(cid, ("", ""))
    ws6.append([cid, v["USE_YN"], v["LAST_UPDATE"], role, impact])
    if v["USE_YN"] != "Y":
        for c in ws6[ws6.max_row]:
            c.fill = PatternFill("solid", fgColor="FCE4D6")
for i, w in enumerate([34, 9, 18, 46, 76], 1):
    ws6.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
for row in ws6.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True); c.font = Font(size=9)

ws5 = wb.create_sheet("G1_데이터결손")
ws5.append(["단계ID", "DIV", "ITEM_ID", "MEASURE", "원장", "조치"])
for c in ws5[1]:
    c.fill, c.font = HF, HFONT
for g in G1_FAIL:
    ws5.append(list(g) + ["신규 입력 시나리오로 사용 또는 SEED 보완"])
for i, w in enumerate([12, 6, 20, 14, 30, 40], 1):
    ws5.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

wb.save(OUT / "04_TESTDATA_SET.xlsx")
print("지정서 {:,}행 / 결손 {:,}건 저장".format(len(rows), len(G1_FAIL)))
(D / "06_gaps.json").write_text(json.dumps(G1_FAIL, ensure_ascii=False, indent=1),
                                encoding="utf-8")
