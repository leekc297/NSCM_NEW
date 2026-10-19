# -*- coding: utf-8 -*-
"""
Phase 5 — 06_SAMPLEDATA_STEPNO.xlsx 생성

원본 3시트를 그대로 복사하고 40TestData / 81TestData / 82TestData 를 덧붙인다.
원본 파일(C:\\TRANFILE\\RECEIVE\\...)은 수정하지 않는다.

행 구성 (데이터출처)
  기존(데이터보유 상위)  오전 04_TESTDATA_SET.xlsx 115행 — 프로세스단계 → StepNo 치환
  신규(SSPM 미생성)      SSPM 키는 있는데 MKT_FCST 0건인 키 5종 × 샘플STEP × 권한그룹
  신규(기준정보·IF)      기준정보·IF STEP 의 실존 레코드 5건 (등록값 명세)

`분류기준` 컬럼에 **왜 이 행이 그 구분에 속하는지** 실측 근거를 싣는다 (사용자 지시 2026-09-18).
"""
import json
import pathlib
import importlib.util
import shutil

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

D = pathlib.Path(__file__).resolve().parent
OUT = D.parent
SRC = pathlib.Path(r"C:\TRANFILE\RECEIVE\DM_통합테스트_SampleData_20260918.xlsx")

spec = importlib.util.spec_from_file_location("stepmap", D / "10_stepmap.py")
SM = importlib.util.module_from_spec(spec)
spec.loader.exec_module(SM)

GAP = json.loads((D / "11_sspm_gap.json").read_text(encoding="utf-8"))
MIF = json.loads((D / "12_master_if.json").read_text(encoding="utf-8"))
TG = json.loads((D / "03_targets_v2.json").read_text(encoding="utf-8"))
ACCR = json.loads((D / "08_accounts.json").read_text(encoding="utf-8"))
PNAME = json.loads((D / "09_program_names.json").read_text(encoding="utf-8"))["names"]

PEOPLE = {p["USER_ID"]: p for p in ACCR["people"]}
RECO = ACCR["recommend"]

# 권한 × 화면 × DIV -> 편집 가능 MEASURE (GET_ADMIN_GROUP_MEASURE 실측)
META = json.loads((D / "01_meta.json").read_text(encoding="utf-8"))
AUTH_MEAS = {}
for _m in META["auth_measure"]:
    AUTH_MEAS.setdefault((_m["AUTH"], _m["SCR"], _m["DIV"]), set()).add(_m["MEASURE"])


def measures_for(auth, scr, div, declared):
    """STEP 의 입력MEASURE 를 권한별로 좁힌다.

    STEP 정의는 'SalesAmt_LOC / RTF_SYS_RP' 처럼 권한 2개분을 한 칸에 적어 둔다.
    권한 메타가 있으면 그 권한이 실제 편집 가능한 것만 남긴다(임의 부여 금지).
    메타 미등재면 STEP 표기를 그대로 쓴다."""
    allowed = set()
    for k, v in AUTH_MEAS.items():
        if k[0] == auth and k[1] == scr and (k[2] == div or k[2].startswith(div)):
            allowed |= v
    toks = [t.strip() for t in str(declared).split("/") if t.strip()]
    if not allowed:
        return declared
    hit = [t for t in toks if t.split("(")[0].strip() in allowed]
    return " / ".join(hit) if hit else declared

# (DIV, StepNo) -> 12_master_if 의 원장 키
MIF_KEY = {
    ("40", "1-1"): "SSPM", ("81", "1-1"): "SSPM", ("82", "1-1"): "SSPM",
    ("81", "3-1"): "NN",   ("81", "3-2"): "EDI",
    ("81", "4-1"): "CFM",  ("81", "4-2"): "CFM",   # 4-2 도 Cust FCST Master (사용자 정정)
    ("40", "3-1"): "RSLT", ("81", "5-1"): "RSLT", ("82", "4-1"): "RSLT",
    ("82", "3-1"): "VMI",
}

COLS = ["StepNo", "Step", "항목", "샘플지정", "데이터출처", "분류기준", "DIV", "구분",
        "엑셀항목", "화면", "실제화면", "화면명", "권한그룹", "테스트계정", "대체계정",
        "해석authGroup(실측)", "W/S판정(실측)",
        "ITEM_ID", "SCM_SITE_ID", "PROJECT_CODE", "FCST_TYPE", "INCOTERMS", "DOM_EXP",
        "PLANT_SITE_ID", "입력MEASURE", "편집요일", "권한메타근거", "대상원장", "등록값명세",
        "데이터보유(실측)", "기간(실측)", "현재값(실측)", "테스트행위", "기대결과",
        "검증화면", "대응결함No"]


def pick_acct(auth, scr):
    """계정 선정 — ADMIN 겸직 계정은 권한경계 검증 불가라 제외 (08_accounts.json 실측)"""
    if auth in ("-", "기준정보 관리"):
        cand = sorted(RECO.get("DM Planner", {}).get("사용가능", []))
        u = cand[0] if cand else "(없음)"
        return (u, ", ".join(cand[1:])[:60], "메타 미등재(제한 없음 전제)",
                PEOPLE.get(u, {}).get("wsPlannerYn", "-"))
    if auth in ("DM ADMIN/ADMIN", "ADMIN"):
        cand = sorted(u for u, p in PEOPLE.items() if p["admin_masked"])
        u = cand[0] if cand else "(없음)"
        return (u, ", ".join(cand[1:])[:60], "ADMIN",
                PEOPLE.get(u, {}).get("wsPlannerYn", "-"))
    r = RECO.get(auth)
    if not r or not r["사용가능"]:
        return "(없음)", "", "미해석", "-"
    ok = sorted(r["사용가능"], key=lambda u: (0 if PEOPLE[u]["groups"] == [auth] else 1,
                                             len(PEOPLE[u]["groups"]), u))
    u = ok[0]
    p = PEOPLE[u]
    resolved = (p["resolved"].get(scr) or "메타 미등재(제한 없음 전제)"
                if scr in ("DMF0082", "DMF0021", "DMF0042", "DMF0003")
                else "메타 미등재(제한 없음 전제)")
    alt = ", ".join(ok[1:])[:60]
    if r["사용불가(ADMIN 해석)"]:
        alt += "  |  ⚠사용불가(ADMIN 해석): " + ", ".join(r["사용불가(ADMIN 해석)"])
    return u, alt[:160], str(resolved), p["wsPlannerYn"]


# ── 원본 엑셀에서 StepNo 별 Step/항목/시간/권한비고 (82 5-1 중복 → 5-1/5-2) ──
src_wb = openpyxl.load_workbook(SRC)
STEP_TEXT = {}
for div, sheet in SM.SHEET_OF.items():
    ws = src_wb[sheet]
    cur_step = ""
    for row in ws.iter_rows(min_row=2, values_only=True):
        c = [("" if x is None else str(x).strip()) for x in row]
        if not any(c):
            continue
        if c[1]:
            cur_step = c[1]
        sno, item = c[0], c[2]
        if not sno or not item:
            continue
        key = sno
        if div == "82" and sno == "5-1" and "Sales Price" in item:
            # 판가입력 2행이 원본에서 둘 다 '5-1' → 5-1 / 5-2 로 분리 (사용자 확정)
            key = "5-1" if "DP Planner" in c[4] else "5-2"
        STEP_TEXT[(div, key)] = {"Step": cur_step, "항목": item,
                                 "시간": c[3], "권한비고": c[4]}

# ── 오전 115행 ──
mws = openpyxl.load_workbook(OUT / "04_TESTDATA_SET.xlsx")["TESTDATA_SET"]
mi = {n: k for k, n in enumerate(c.value for c in mws[1])}
MORN_ROWS = list(mws.iter_rows(min_row=2, values_only=True))

# ITEM 별 '기존' 분류기준 — FOCUS 메저 보유 실측
CRIT = {}
for div, ts in TG["divisions"].items():
    for t in ts:
        CRIT[(div, t["ITEM_ID"])] = (
            "메저보유 상위 — FOCUS 보유 {} / 결손 {} · 총 {}행 ({}~{})".format(
                ",".join(t["focus_have"]) or "-", ",".join(t["focus_miss"]) or "없음",
                t["ROW_CNT"], t["MIN_YMD"], t["MAX_YMD"]))

rows_by_div = {"40": [], "81": [], "82": []}

for div, sno, smp, scr, real, led, auths, meas, morn, kind in SM.STEPS:
    st = STEP_TEXT.get((div, sno), {"Step": "", "항목": "", "시간": "", "권한비고": ""})
    base = dict(StepNo=sno, Step=st["Step"], 항목=st["항목"], 샘플지정=smp, DIV=div,
                구분=kind, 엑셀항목=st["항목"], 화면=scr, 실제화면=real,
                화면명=PNAME.get(real, "-"), 대상원장=led, 입력MEASURE=meas)

    # A. 기존(데이터보유 상위)
    for r in MORN_ROWS:
        if r[mi["프로세스단계"]] not in morn:
            continue
        item = r[mi["ITEM_ID"]]
        rows_by_div[div].append(dict(base, **{
            "데이터출처": "기존(데이터보유 상위)",
            "분류기준": CRIT.get((div, item), "오전 선별 ITEM"),
            "권한그룹": r[mi["권한그룹"]], "테스트계정": r[mi["테스트계정"]],
            "대체계정": r[mi["대체계정"]],
            "해석authGroup(실측)": r[mi["해석authGroup(실측)"]],
            "W/S판정(실측)": r[mi["W/S판정(실측)"]],
            "ITEM_ID": item, "SCM_SITE_ID": r[mi["SCM_SITE_ID"]],
            "PROJECT_CODE": r[mi["PROJECT_CODE"]], "FCST_TYPE": r[mi["FCST_TYPE"]],
            "INCOTERMS": r[mi["INCOTERMS"]], "DOM_EXP": r[mi["DOM_EXP"]],
            "PLANT_SITE_ID": r[mi["PLANT_SITE_ID"]],
            "입력MEASURE": r[mi["입력MEASURE"]], "편집요일": r[mi["편집요일"]],
            "권한메타근거": r[mi["권한메타근거"]], "등록값명세": "-",
            "데이터보유(실측)": r[mi["데이터보유(실측)"]],
            "기간(실측)": r[mi["기간(실측)"]], "현재값(실측)": r[mi["현재값(실측)"]],
            "테스트행위": r[mi["테스트행위"]], "기대결과": r[mi["기대결과"]],
            "검증화면": r[mi["검증화면"]], "대응결함No": r[mi["대응결함No"]]}))

    if smp != "Y":
        continue

    # B. 신규(SSPM 미생성) — 권한그룹 × 5키
    for auth in auths:
        acct, alt, resolved, wsyn = pick_acct(auth, real)
        meas_a = measures_for(auth, real, div, meas)
        for g in GAP.get(div, []):
            rows_by_div[div].append(dict(base, **{
                "데이터출처": "신규(SSPM 미생성)", "분류기준": g["분류기준"],
                "입력MEASURE": meas_a,
                "권한그룹": auth, "테스트계정": acct, "대체계정": alt,
                "해석authGroup(실측)": resolved, "W/S판정(실측)": wsyn,
                "ITEM_ID": g["ITEM_ID"], "SCM_SITE_ID": g["SCM_SITE_ID"],
                "PROJECT_CODE": g["PROJECT_CODE"], "FCST_TYPE": g["FCST_TYPE"],
                "INCOTERMS": g["INCOTERMS"], "DOM_EXP": g["DOM_EXP"],
                "PLANT_SITE_ID": g["PLANT_SITE_ID"], "편집요일": "-",
                "권한메타근거": "부속키 출처 = " + g["부속키출처"],
                "등록값명세": ("SSPM 실값 — CURRENCY={} / DEPT_ID={} / OWNER={} /"
                              " CUST_GROUP={}").format(
                    g.get("CURRENCY"), g.get("DEPT_ID"), g.get("OWNER"), g.get("CUST_GROUP")),
                "데이터보유(실측)": "없음 → 신규입력", "기간(실측)": "-", "현재값(실측)": "-",
                "테스트행위": "{} 에서 {} 신규 입력".format(real, meas_a),
                "기대결과": "신규 키가 정상 저장되고 후속 STEP 에 반영",
                "검증화면": "DMF0035", "대응결함No": "-"}))

    # C. 신규(기준정보·IF) — 등록값 명세
    mk = MIF_KEY.get((div, sno))
    if mk and mk in MIF:
        sp = MIF[mk]
        acct, alt, resolved, wsyn = pick_acct(auths[0], real)
        for rec in sp["by_div"].get(div, []):
            kv = " / ".join("{}={}".format(c, rec.get(c)) for c in sp["key_cols"])
            vv = " / ".join("{}={}".format(c, rec.get(c)) for c in sp["val_cols"])
            rows_by_div[div].append(dict(base, **{
                "데이터출처": "신규(기준정보·IF)",
                "분류기준": "{} 실존 레코드 (데이터 보유 상위)".format(sp["table"]),
                "권한그룹": auths[0], "테스트계정": acct, "대체계정": alt,
                "해석authGroup(실측)": resolved, "W/S판정(실측)": wsyn,
                "ITEM_ID": rec.get("ITEM_ID", "-"), "SCM_SITE_ID": rec.get("SCM_SITE_ID", "-"),
                "PROJECT_CODE": rec.get("PROJECT_CODE", "-"),
                "FCST_TYPE": rec.get("FCST_TYPE", "-"),
                "INCOTERMS": rec.get("INCOTERMS", "-"), "DOM_EXP": rec.get("DOM_EXP", "-"),
                "PLANT_SITE_ID": "-", "편집요일": "-", "권한메타근거": "-",
                "등록값명세": "{} ← 키[{}] 값[{}]".format(sp["table"], kv, vv),
                "데이터보유(실측)": "실존", "기간(실측)": "-", "현재값(실측)": "-",
                "테스트행위": "{} 에 위 키/값으로 등록·수신".format(sp["table"]),
                "기대결과": "등록/수신분이 후속 STEP 에 반영",
                "검증화면": real if real != "(IF)" else "-", "대응결함No": "-"}))

# ── 엑셀 생성 (원본 3시트 보존 + 3시트 추가) ──
DEST = OUT / "06_SAMPLEDATA_STEPNO.xlsx"
shutil.copy(SRC, DEST)
wb = openpyxl.load_workbook(DEST)
HF = PatternFill("solid", fgColor="1F4E78")
HFONT = Font(color="FFFFFF", bold=True, size=9)
FILL = {"기존(데이터보유 상위)": "E2EFDA", "신규(SSPM 미생성)": "FFF2CC",
        "신규(기준정보·IF)": "DDEBF7"}
WIDTH = [8, 12, 26, 8, 20, 46, 5, 9, 26, 9, 9, 22, 17, 13, 30, 20, 8,
         22, 26, 18, 11, 10, 8, 14, 24, 8, 30, 40, 60, 18, 22, 30, 34, 34, 10, 22]


def snum(s):
    a, _, b = str(s).partition("-")
    try:
        return (int(a), int(b or 0))
    except ValueError:
        return (999, 999)


for div in ("40", "81", "82"):
    ws = wb.create_sheet(SM.OUT_SHEET[div])
    ws.append(COLS)
    for c in ws[1]:
        c.fill, c.font = HF, HFONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in sorted(rows_by_div[div],
                    key=lambda x: (snum(x["StepNo"]), x["데이터출처"],
                                   str(x["권한그룹"]), str(x["ITEM_ID"]))):
        ws.append([r.get(c, "") for c in COLS])
        f = FILL.get(r["데이터출처"])
        for c in ws[ws.max_row]:
            if f:
                c.fill = PatternFill("solid", fgColor=f)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(size=9)
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    src_cnt = {}
    for r in rows_by_div[div]:
        src_cnt[r["데이터출처"]] = src_cnt.get(r["데이터출처"], 0) + 1
    print("{} — {:,}행  {}".format(SM.OUT_SHEET[div], len(rows_by_div[div]), src_cnt))

wb.save(DEST)
print("\n06_SAMPLEDATA_STEPNO.xlsx 저장 — 시트 {}".format(wb.sheetnames))
