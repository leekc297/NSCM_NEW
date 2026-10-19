# -*- coding: utf-8 -*-
"""Phase 4 — 06_SAMPLEDATA_STEPNO(CHO) 보완본 생성

사용자 요구 (2026-09-18)
  · G열 요구사항 반영 — 화면/IF/UPLOAD/N:N LAYOUT 기준 샘플 작성, RTF 점검 STEP·EDI IF 4표 신설
  · 40 INCOTERMS 미사용 값(EDIT_FLAG='N') 제시 금지
  · Cell Prod Site 필수 컬럼 공백 해소
  · CRUD UI = division 별 USER_SELECTABLE_FLAG='N' 컬럼 전체 데이터 생성
  · DMF0082 VMI 메저를 쓸 수 있는 SITE 선정해 VMI 샘플 추가
  · 기존 40_1-1 ~ 81_12-1 STEP 시트는 삭제하고 전체 STEP 시트를 재생성
  · TestData 3시트에 전량 생성

원본 C:\\TRANFILE\\RECEIVE\\06_SAMPLEDATA_STEPNO(CHO).xlsx 는 수정하지 않는다.
"""
import json
import pathlib
import shutil
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import importlib

pick = importlib.import_module("31_pick")
lay = importlib.import_module("30_layouts")

HERE = pathlib.Path(__file__).resolve().parent
OUTDIR = HERE.parent
SRC = pathlib.Path(r"C:\TRANFILE\RECEIVE\06_SAMPLEDATA_STEPNO(CHO).xlsx")
OUT = OUTDIR / "07_SAMPLEDATA_STEPNO_v2.xlsx"
DELIVER = pathlib.Path(r"C:\TRANFILE\RECEIVE\07_SAMPLEDATA_STEPNO_v2.xlsx")
N_PER = 5

DIV_SHEETS = ["40 DIVISION", "81 DIVISION", "82 DIVISION"]
TESTDATA = {"40": "40TestData", "81": "81TestData", "82": "82TestData"}

# ── STEP 정의 (G열 요구 반영) ─────────────────────────────────────────────
# (DIV, StepNo, 항목, 구분, 화면/원천, 권한그룹들, 대상원장, G열요구, 성격)
STEPS = [
 ("40", "1-1",  "SSPM",                  "기준정보", "DMF0005", ["기준정보 관리"],
  "SCM_MST_SITE_ITEM", "기존 SPMS 데이터 필요, 해당 화면 LAYOUT 기준으로 샘플 작성", "UI"),
 ("40", "3-1",  "IF 수신 [M to O]",       "IF",      "(IF)",    ["-"],
  "INF_MGR.XIF_ERP_SALES_RESULT_R", "IF LAYOUT으로 샘플 작성", "IF"),
 ("40", "4-1",  "UI : FCST&PRICE MGMT",  "UI",      "DMF0082", ["DM Planner"],
  "DM_PLN_CFM_SELLING_PRICE, DM_PLN_MKT_FCST", "-", "UI"),
 ("40", "6-1",  "UI : FCST&PRICE MGMT",  "UI",      "DMF0082", ["SCP Planner"],
  "DM_PLN_MKT_FCST", "-", "UI"),
 ("40", "9-1",  "RTF 생성 점검",           "검증",     "DMF0082", ["DM Planner", "SCP Planner"],
  "DM_PLN_MKT_FCST(RTF)", "RTF 발행 이후 검증 샘플", "UI"),
 ("40", "10-1", "Sales Plan 입력",        "UI",      "DMF0021", ["DM Planner", "SCP Planner"],
  "DM_PLN_SALES_PLAN", "-", "UI"),

 ("81", "1-1",  "SSPM",                  "기준정보", "DMF0005", ["기준정보 관리"],
  "SCM_MST_SITE_ITEM", "기존 SPMS 데이터 필요, 해당 화면 LAYOUT 기준으로 샘플 작성", "UI"),
 ("81", "3-1",  "N:N Master Mapping",    "기준정보", "DMF0072", ["기준정보 관리"],
  "DM_CUST_MODEL_MAPPING_M", "기존 정보 추가, N:N MAPPING LAYOUT 기준 샘플작성", "UI"),
 ("81", "3-2a", "EDI Call-off IF — HEAD", "IF",     "(IF)",    ["-"],
  "INF_MGR.XIF_ERP_EDI_CALLOFF_HEAD_R", "신규 정보 기준추가, I/F LAYOUT으로 작성", "IF"),
 ("81", "3-2b", "EDI Call-off IF — ITEM", "IF",     "(IF)",    ["-"],
  "INF_MGR.XIF_ERP_EDI_CALLOFF_ITEM_R", "신규 정보 기준추가, I/F LAYOUT으로 작성", "IF"),
 ("81", "3-2c", "EDI Call-off IF — PARTNER", "IF",  "(IF)",    ["-"],
  "INF_MGR.XIF_ERP_EDI_CALLOFF_PARTNER_R", "신규 정보 기준추가, I/F LAYOUT으로 작성", "IF"),
 ("81", "3-2d", "CUST DEMAND MAPPING IF", "IF",     "(IF)",    ["-"],
  "INF_MGR.XIF_CUST_DEMAND_MAPPING_MST_R", "신규 정보 기준추가, I/F LAYOUT으로 작성", "IF"),
 ("81", "4-1",  "Customer FCST Master",  "기준정보", "DMF0007", ["기준정보 관리"],
  "DM_MST_CUSTOMER_FCST", "신규+기존 정보 기준으로 작성, 해당 화면 입력 기준 LAYOUT", "UI"),
 ("81", "4-2",  "Customer FCST Mapping", "기준정보", "DMF0007", ["기준정보 관리"],
  "DM_MST_CUSTOMER_FCST", "신규+기존 정보 기준으로 작성, 해당 화면 입력 기준 LAYOUT", "UI"),
 ("81", "4-3",  "Customer FCST Upload",  "UI",      "DMF0014", ["DM Planner"],
  "DM_PLN_CUSTOMER_PSI", "신규+기존 정보 기준으로 작성, UPLOAD LAYOUT 기준으로 작성", "UPLOAD"),
 ("81", "5-1",  "IF 수신 [M to O]",       "IF",      "(IF)",    ["-"],
  "INF_MGR.XIF_ERP_SALES_RESULT_R", "IF LAYOUT으로 샘플 작성", "IF"),
 ("81", "6-1",  "UI : FCST&PRICE MGMT",  "UI",      "DMF0082", ["DM Planner"],
  "SCM_MST_FUTURE_PRICE, DM_PLN_MKT_FCST", "-", "UI"),
 ("81", "6-2",  "UI : FCST&PRICE MGMT",  "UI",      "DMF0082", ["DM Planner (W/S)"],
  "DM_PLN_MKT_FCST", "-", "UI"),
 ("81", "8-1",  "UI : FCST&PRICE MGMT",  "UI",      "DMF0082", ["SCP Planner"],
  "DM_PLN_MKT_FCST", "-", "UI"),
 ("81", "11-1", "RTF 생성 점검",           "검증",     "DMF0082", ["DM Planner", "SCP Planner"],
  "DM_PLN_MKT_FCST(RTF)", "RTF 발행 이후 검증 샘플", "UI"),
 ("81", "12-1", "Sales Plan 입력",        "UI",      "DMF0021", ["DM Planner", "SCP Planner"],
  "DM_PLN_SALES_PLAN", "-", "UI"),

 ("82", "1-1",  "SSPM",                  "기준정보", "DMF0005", ["기준정보 관리"],
  "SCM_MST_SITE_ITEM", "기존 SPMS 데이터 필요, 해당 화면 LAYOUT 기준으로 샘플 작성", "UI"),
 ("82", "3-1",  "재고 Data IF",            "IF",      "(IF)",    ["-"],
  "FP_MGR.SCM_STOCK_H (= SCM_DNM_STOCK_SN 시노님)",
  "SCM_DNM_STOCK_SN 데이터 기준으로 샘플작성", "IF"),
 ("82", "4-1",  "IF 수신 [M to O]",       "IF",      "(IF)",    ["-"],
  "INF_MGR.XIF_ERP_SALES_RESULT_R", "IF LAYOUT으로 샘플 작성", "IF"),
 ("82", "5-1",  "UI : Sales Price (DP Planner)", "UI", "DMF0020", ["DM Planner"],
  "SCM_MST_FUTURE_PRICE", "-", "UI"),
 ("82", "5-2",  "UI : Sales Price (admin)",      "UI", "DMF0020", ["ADMIN"],
  "SCM_MST_FUTURE_PRICE", "-", "UI"),
 ("82", "6-1",  "UI : FCST&PRICE MGMT",  "UI",      "DMF0003", ["DM Planner"],
  "DM_PLN_MKT_FCST", "82 는 DMF0082 조회전용 → 실제화면 DMF0003", "UI"),
 ("82", "6-2",  "UI : FCST&PRICE MGMT",  "UI",      "DMF0003", ["DM Planner (W/S)"],
  "DM_PLN_MKT_FCST", "82 는 DMF0082 조회전용 → 실제화면 DMF0003", "UI"),
 ("82", "9-1",  "RTF 생성 점검",           "검증",     "DMF0003", ["DM Planner"],
  "DM_PLN_MKT_FCST(RTF)", "RTF 발행 이후 검증 샘플", "UI"),
 ("82", "10-1", "Sales Plan 입력",        "UI",      "DMF0021", ["DM Planner"],
  "DM_PLN_SALES_PLAN", "-", "UI"),
]

VMI_STEP = ("40", "4-1")   # VMI 샘플을 붙일 STEP (DMF0082 · 40)

TD_COLS = [
 "StepNo", "항목", "구분", "DIV", "샘플지정", "데이터출처", "분류기준",
 "화면/원천", "화면LAYOUT원천", "LAYOUT컬럼수", "G열요구사항",
 "권한그룹", "테스트계정", "대체계정",
 "ITEM_ID", "SCM_SITE_ID", "PROJECT_CODE", "END_CUSTOMER",
 "FCST_TYPE", "INCOTERMS", "INCOTERMS_EDIT_FLAG",
 "DOM_EXP", "CELL_PROD_SITE_ID", "PLANT_SITE_ID", "CURRENCY",
 "입력MEASURE", "대상원장", "데이터보유(실측)", "테스트행위", "기대결과", "비고",
]

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(color="FFFFFF", bold=True, size=9)
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
NEW_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def autosize(ws, ncol, maxw=42):
    for c in range(1, ncol + 1):
        w = 8
        for r in range(1, min(ws.max_row, 200) + 1):
            v = ws.cell(r, c).value
            if v is not None:
                w = max(w, min(maxw, len(str(v)) + 2))
        ws.column_dimensions[get_column_letter(c)].width = w


def accounts_for(meta8, group):
    rec = meta8.get("recommend", {}).get(group, {})
    ok = rec.get("사용가능", [])
    if ok:
        return ok[0], ", ".join(ok[1:5]) or "-"
    holders = meta8.get("holders", {}).get(group, [])
    return (holders[0] if holders else "-"), ", ".join(holders[1:5]) or "-"


def main():
    meta8 = json.loads((HERE / "08_accounts.json").read_text(encoding="utf-8"))
    sspm_gap = json.loads((HERE / "11_sspm_gap.json").read_text(encoding="utf-8"))
    layouts = json.loads((HERE / "30_layouts.json").read_text(encoding="utf-8"))

    db = pick.connect()
    cur = db.cursor()
    inco = pick.editable_incoterms(cur)
    inco_set = set(inco)
    cellvals = {d: pick.cell_site_values(cur, d) for d in ("40", "81", "82")}
    # division 별 최빈 입력가능 INCOTERMS (최종 폴백)
    inco_top = {}
    for d in ("40", "81", "82"):
        cur.execute("""SELECT INCOTERMS FROM (
                         SELECT INCOTERMS, COUNT(*) N FROM SCM_MGR.DM_PLN_MKT_FCST
                          WHERE {dd} AND INCOTERMS IN ({ii})
                          GROUP BY INCOTERMS ORDER BY 2 DESC) WHERE ROWNUM=1""".format(
            dd=pick.div_pred(d), ii=",".join("'"+c+"'" for c in inco)))
        r = cur.fetchone()
        inco_top[d] = r[0] if r else inco[0]
    print("division 최빈 입력가능 INCOTERMS:", inco_top)

    wb = openpyxl.load_workbook(SRC)
    # ── 기존 STEP 시트 전부 삭제 (사용자 지시) ──
    removed = []
    for n in list(wb.sheetnames):
        if n in DIV_SHEETS or n in TESTDATA.values():
            continue
        removed.append(n)
        del wb[n]
    print("삭제한 기존 STEP 시트 {}개: {}".format(len(removed), ", ".join(removed)))

    td_rows = {"40": [], "81": [], "82": []}
    step_sheets = []

    for (div, sno, item, kind_kr, screen, groups, ledger, greq, lkind) in STEPS:
        key = "{}|{}".format(div, sno)
        L = layouts.get(key)
        if not L:
            print("  !! LAYOUT 없음 {}".format(key))
            continue

        # ── LAYOUT 컬럼 목록 ──
        if L["kind"] == "ui":
            cols = [c["col"] for c in L["cols"]]
            labels = {c["col"]: c["label"] for c in L["cols"]}
            src_desc = "SCM_MST_SHEET_COLUMNS {} / USER_SELECTABLE_FLAG='N'".format(L["args"][0])
            if not cols:   # DMF0072 처럼 'N' 이 0건인 화면 — 전체 표시 컬럼으로 폴백
                cur.execute("""SELECT COLUMN_NAME, COLUMN_LABEL_NAME
                                 FROM SCM_MGR.SCM_MST_SHEET_COLUMNS
                                WHERE PROGRAM_ID=:p AND VALID_YN='Y' AND SHEET_NAME='main'
                                ORDER BY SORT_ORDER_NO""", {"p": L["args"][0]})
                rr = cur.fetchall()
                cols = [r[0] for r in rr]
                labels = {r[0]: r[1] for r in rr}
                src_desc = ("SCM_MST_SHEET_COLUMNS {} / 전체(USER_SELECTABLE_FLAG='N' 0건 폴백)"
                            .format(L["args"][0]))
        elif L["kind"] == "if":
            cols = [c["col"] for c in L["cols"]]
            labels = {c["col"]: c["col"] for c in L["cols"]}
            src_desc = "{}.{} 실제 컬럼".format(*L["args"])
        else:
            toks = []
            for s in L["cols"]:
                toks += s["tokens"]
            cols = toks
            labels = {t: t for t in toks}
            src_desc = "SCM_ADM_EXCEL_COLUMNS {} HEADER_CONTENT".format(L["args"][0])

        # ── 실데이터 선별 ──
        picked = []
        if L["kind"] == "if":
            owner, table = L["args"]
            rs, note = pick.if_rows(cur, owner, table, div, N_PER)
            src_desc += " · " + note
            for r in rs:
                picked.append(("기존(IF 실수신)", note, r))
            if not rs:
                print("   !! {} IF 행 0 — {}".format(key, note))
        elif L["kind"] == "upload":
            cur.execute("""SELECT * FROM SCM_MGR.DM_PLN_CUSTOMER_PSI
                            WHERE {d} FETCH FIRST :n ROWS ONLY""".format(d=pick.div_pred(div)),
                        {"n": N_PER})
            for r in pick.rows_as_dicts(cur):
                picked.append(("기존(데이터보유 상위)", "DM_PLN_CUSTOMER_PSI 실행", r))
        else:
            pid = L["args"][0]
            for r in pick.existing_rows(cur, pid, div, inco, N_PER):
                picked.append(("기존(데이터보유 상위)",
                               "ITEM 행수 상위 (해당 ITEM {:,}행) · INCOTERMS EDIT_FLAG='Y'"
                               .format(r.get("ITEM_CNT", 0)), r))
            for g in sspm_gap.get(div, [])[:N_PER]:
                picked.append(("신규(SSPM 미생성)",
                               "SSPM 등록 O / MKT_FCST 0건", dict(g)))

        # VMI 추가 (지정 STEP 한정)
        if (div, sno) == VMI_STEP:
            vms = [m["COLUMN_NAME"] for m in pick.vmi_measures(cur, div)]
            for r in pick.vmi_rows(cur, div, inco, N_PER):
                r = dict(r)
                r["_VMI"] = ", ".join(vms)
                picked.append(("신규(VMI 가능 SITE)",
                               "SCM_MST_SITE.VMI_YN='Y' · DMF0082 VMI 메저 {}".format(vms), r))

        # ── STEP 시트 ──
        name = "{}_{}".format(div, sno)
        ws = wb.create_sheet(name)
        meta_cols = ["StepNo", "항목", "DIV", "화면/원천", "권한그룹", "테스트계정",
                     "데이터출처", "분류기준"]
        head = meta_cols + cols
        for i, h in enumerate(head, 1):
            ws.cell(1, i, h)
        # 2행 = 컬럼 라벨(설명)
        for i, h in enumerate(head, 1):
            ws.cell(2, i, "-" if i <= len(meta_cols) else labels.get(h, h))
        for c in range(1, len(head) + 1):
            ws.cell(2, c).fill = SUB_FILL
            ws.cell(2, c).font = Font(size=8, italic=True)

        # ── 키 단위 정규화 (INCOTERMS 입력가능 · Cell Prod Site 필수값) ──
        norm = []
        for src_kind, basis, rec in picked:
            rec = dict(rec)
            note = []
            if L["kind"] == "ui":
                site = rec.get("SCM_SITE_ID") or ""
                item = rec.get("ITEM_ID") or ""
                if "INCOTERMS" in cols or rec.get("INCOTERMS") is not None:
                    v, s = pick.resolve_incoterms(cur, div, site, item,
                                                  rec.get("INCOTERMS"), inco, inco_top[div])
                    if v != rec.get("INCOTERMS"):
                        note.append("INCOTERMS {}→{} ({})".format(
                            rec.get("INCOTERMS") or "공백", v, s))
                    rec["INCOTERMS"] = v
                want = pick.CELL_COL[div]
                v, s = pick.resolve_cell_site(cur, div, site, item,
                                              rec.get(want) if want else None, cellvals[div])
                if want:
                    if rec.get(want) != v:
                        note.append("{} {}→{} ({})".format(
                            want, rec.get(want) or "공백", v, s))
                    rec[want] = v
                    other = "PLANT_SITE_ID" if want == "CELL_PROD_SITE_ID" else "CELL_PROD_SITE_ID"
                    rec.setdefault(other, "(해당없음 — {} 축)".format(div))
                else:
                    rec["CELL_PROD_SITE_ID"] = v
                    rec["PLANT_SITE_ID"] = v
            else:
                # IF / UPLOAD LAYOUT — 화면 축이 아니므로 Cell Prod Site 개념이 없다
                tag = "(해당없음 — {} LAYOUT)".format(
                    "IF" if L["kind"] == "if" else "UPLOAD")
                for c in ("CELL_PROD_SITE_ID", "PLANT_SITE_ID"):
                    if not rec.get(c):
                        rec[c] = tag
            norm.append((src_kind, basis + ((" · " + " / ".join(note)) if note else ""), rec))
        picked = norm

        rowno = 3
        for gi, grp in enumerate(groups):
            acc, alt = accounts_for(meta8, grp)
            for src_kind, basis, rec in picked:
                ws.cell(rowno, 1, sno)
                ws.cell(rowno, 2, item)
                ws.cell(rowno, 3, div)
                ws.cell(rowno, 4, screen if screen != "(IF)" else ledger)
                ws.cell(rowno, 5, grp)
                ws.cell(rowno, 6, acc)
                ws.cell(rowno, 7, src_kind)
                ws.cell(rowno, 8, basis)
                for i, c in enumerate(cols, len(meta_cols) + 1):
                    v = rec.get(c)
                    if v is None:
                        v = ""
                    ws.cell(rowno, i, str(v) if not isinstance(v, (int, float)) else v)
                if src_kind.startswith("신규"):
                    for c in range(1, len(head) + 1):
                        ws.cell(rowno, c).fill = NEW_FILL
                rowno += 1

                # TestData 통합행
                td_rows[div].append({
                    "StepNo": sno, "항목": item, "구분": kind_kr, "DIV": div,
                    "샘플지정": "Y", "데이터출처": src_kind, "분류기준": basis,
                    "화면/원천": screen if screen != "(IF)" else ledger,
                    "화면LAYOUT원천": src_desc, "LAYOUT컬럼수": len(cols),
                    "G열요구사항": greq,
                    "권한그룹": grp, "테스트계정": acc, "대체계정": alt,
                    "ITEM_ID": rec.get("ITEM_ID", ""),
                    "SCM_SITE_ID": rec.get("SCM_SITE_ID", ""),
                    "PROJECT_CODE": rec.get("PROJECT_CODE", ""),
                    "END_CUSTOMER": rec.get("END_CUSTOMER", ""),
                    "FCST_TYPE": rec.get("FCST_TYPE", ""),
                    "INCOTERMS": rec.get("INCOTERMS", ""),
                    "INCOTERMS_EDIT_FLAG": ("Y" if rec.get("INCOTERMS") in inco_set
                                            else ("-" if not rec.get("INCOTERMS") else "N")),
                    "DOM_EXP": rec.get("DOM_EXP", ""),
                    "CELL_PROD_SITE_ID": rec.get("CELL_PROD_SITE_ID", ""),
                    "PLANT_SITE_ID": rec.get("PLANT_SITE_ID", ""),
                    "CURRENCY": rec.get("CURRENCY", ""),
                    "입력MEASURE": rec.get("_VMI") or rec.get("MEASURE", ""),
                    "대상원장": ledger,
                    "데이터보유(실측)": ("없음 → 신규입력" if src_kind.startswith("신규(SSPM")
                                   else "있음"),
                    "테스트행위": "{} 에서 {} 로 입력/검증".format(screen, grp),
                    "기대결과": "정상 저장 및 후속 STEP 반영",
                    "비고": greq,
                })
        style_header(ws, len(head))
        autosize(ws, len(head))
        step_sheets.append((name, rowno - 3, len(head)))
        print("  {:8s} {:3d}행 x {:3d}컬럼  ({})".format(name, rowno - 3, len(head), src_desc))

    # ── TestData 3시트 ──
    for div, sname in TESTDATA.items():
        if sname in wb.sheetnames:
            del wb[sname]
        ws = wb.create_sheet(sname)
        for i, h in enumerate(TD_COLS, 1):
            ws.cell(1, i, h)
        for r, rec in enumerate(td_rows[div], 2):
            for i, h in enumerate(TD_COLS, 1):
                ws.cell(r, i, rec.get(h, ""))
            if str(rec.get("데이터출처", "")).startswith("신규"):
                for c in range(1, len(TD_COLS) + 1):
                    ws.cell(r, c).fill = NEW_FILL
        style_header(ws, len(TD_COLS))
        autosize(ws, len(TD_COLS))
        print("  {:12s} {:4d}행 x {}컬럼".format(sname, len(td_rows[div]), len(TD_COLS)))

    # 시트 순서 — DIVISION → TestData → STEP
    order = DIV_SHEETS + list(TESTDATA.values()) + [n for n, _, _ in step_sheets]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]

    wb.save(OUT)
    print("\n저장 → {}".format(OUT))
    try:
        shutil.copyfile(OUT, DELIVER)
        print("전달본 → {}".format(DELIVER))
    except Exception as e:
        print("전달본 복사 실패(열려 있을 수 있음): {}".format(e))


if __name__ == "__main__":
    main()
