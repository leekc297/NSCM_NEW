# -*- coding: utf-8 -*-
"""STEP 시트 LAYOUT 잔여 공백 사유 시트 `_LAYOUT_GAP` 생성 (2026-09-21)

44_refill_layout.py 로 채우고도 남는 공백 컬럼이 왜 비었는지 **행마다 사유를 실측으로
판정**해 한 시트에 싣는다. 공백을 임의값으로 메우지 않고 사유를 명시하는 것이 원칙이다.

사유 4분류
  A 원장에 컬럼 없음   화면 LAYOUT 에는 있으나 저장 원장에 해당 컬럼이 없다
                       (= 화면 표시/파생 컬럼이거나 다른 표에서 조인해 오는 값)
  B 원장 값 NULL       컬럼은 있고 매칭 레코드도 있는데 값이 NULL
  C 신규 행 — 입력대상  해당 행이 '신규' 분류라 원장에 아직 없음 (= 테스터가 입력할 칸)
  D 매핑 미확보        UPLOAD LAYOUT (81_4-3) — 엑셀 헤더 라벨 ↔ 원장 컬럼 매핑 별도 필요
"""
import importlib.util
import pathlib
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = pathlib.Path(__file__).resolve().parent
XLSX = HERE.parent / "08_SAMPLEDATA_STEPNO_v3.xlsx"
SHEET = "_LAYOUT_GAP"

_m = importlib.util.spec_from_file_location("rl", HERE / "44_refill_layout.py")
RL = importlib.util.module_from_spec(_m)
_m.loader.exec_module(RL)
PICK = RL.PICK

COLS = ["STEP시트", "DIV", "StepNo", "LAYOUT컬럼", "컬럼라벨", "원장", "공백행수",
        "전체행수", "사유코드", "사유", "신규행비율"]


def main():
    wb = openpyxl.load_workbook(XLSX)
    db = PICK.connect()
    cur = db.cursor()
    out = []
    for name in wb.sheetnames:
        if "_" not in name or name.split("_")[0] not in ("40", "81", "82"):
            continue
        div, sno = name.split("_", 1)
        ws = wb[name]
        hdr = [c.value for c in ws[1]]
        lbl = [c.value for c in ws[2]]
        data = [r for r in ws.iter_rows(min_row=3)
                if any(c.value not in (None, "") for c in r)]
        if not data:
            continue
        src_i = hdr.index("데이터출처")
        n_new = sum(1 for r in data if str(r[src_i].value or "").startswith("신규"))
        if name in RL.SKIP:
            owner = table = None
            have = set()
        else:
            owner, table = RL.IF_TABLE.get(name, (None, None))
            if not table:
                scr = str(ws.cell(3, hdr.index("실제화면") + 1).value or "")
                table, owner = RL.UI_TABLE.get(scr), "SCM_MGR"
            have = RL.cols_of(cur, owner, table) if table else set()
        for i in range(RL.META_N, len(hdr)):
            h = hdr[i]
            if not h:
                continue
            nb = sum(1 for r in data if i >= len(r) or r[i].value in (None, ""))
            if not nb:
                continue
            if name in RL.SKIP:
                code, why = "D", "UPLOAD LAYOUT — 엑셀 헤더 라벨 ↔ 원장 컬럼 매핑 미확보"
            elif not table:
                code, why = "D", "원장 미지정"
            elif h not in have:
                code, why = "A", "{}.{} 에 컬럼 없음 — 화면 표시/파생 컬럼".format(owner, table)
            elif nb == n_new and n_new:
                code, why = "C", "'신규' 분류 행 — 원장 미존재(테스터 입력 대상)"
            else:
                code, why = "B", "{}.{} 에 컬럼은 있으나 매칭 레코드 값이 NULL".format(
                    owner, table)
            out.append([name, div, sno, h, lbl[i] if i < len(lbl) else "",
                        "{}.{}".format(owner, table) if table else "-",
                        nb, len(data), code, why,
                        "{}/{}".format(n_new, len(data))])
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)
    ws.append(COLS)
    for r in out:
        ws.append(r)
    fill = {"A": "FCE4D6", "B": "FFF2CC", "C": "E2EFDA", "D": "D9D9D9"}
    for row in ws.iter_rows(min_row=2):
        f = fill.get(str(row[8].value))
        for c in row:
            if f:
                c.fill = PatternFill("solid", fgColor=f)
            c.font = Font(size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.font = Font(color="FFFFFF", bold=True, size=9)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate([11, 5, 8, 30, 26, 34, 9, 9, 9, 56, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(XLSX)
    cnt = {}
    for r in out:
        cnt[r[8]] = cnt.get(r[8], 0) + 1
    print("_LAYOUT_GAP {}행 — 사유 {}".format(len(out), cnt))
    print("저장 → {}".format(XLSX))
    return 0


if __name__ == "__main__":
    sys.exit(main())
