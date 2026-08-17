#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""DMF0074 Excel Upload 샘플 데이터 생성기 — division 40 기준 3건 (업로드 통과 보장).

양식 = `SCM_ADM_EXCEL_COLUMNS.PROGRAM_ID='DMF0074_40'` HEADER_CONTENT 실측 **53열**
       (division 40 은 Project Code / Project 2열이 없다 — 55열 'DMF0074' 양식을 쓰면 SP 의
        40번 2칸 시프트와 어긋나 전 행 'NULL Metal Link.' 로 탈락한다).
키 조합 = `P_DMF0074_GET_METAL_LINK_TEMPLATE`(Excel Form) 실조회분에서 선별 —
       C19 'Invalid Sales Plan Waterfall MODEL.' 를 통과하는 실재 조합만 사용.

샘플 3건 (검증 분기를 골고루 덮도록 구성)
  1행 Metal Link = 'N'            → 광물 검증 게이트 밖 (광물 42열 전부 공란)
  2행 Metal Link = 'Y' + Ni 1종   → Key_Ni='M035'(LME) + Period/Price/Index 5값 완전입력
  3행 Metal Link = 'Y' + Ni·Co 2종 + Al 미연동 → 연동 2종은 5값 완전입력,
      미연동 코드(M051)는 5값 **공란**이 정답(NEED EMPTY VALUE 규칙)

생성: python Doc/화면분석서/_gen_DMF0074_upload_sample.py
"""
import io
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = pathlib.Path(__file__).parent / "DMF0074_ExcelUpload_Sample_div40_3rows.xlsx"
TPL = pathlib.Path(
    r"C:\Users\ejene\AppData\Local\Temp\claude\c--Projects-NSCM"
    r"\55375f24-fdd4-4ea6-9d5a-270f246a74ff\scratchpad\template_rows.json")
PLAN_YYYYMM = "202609"          # 당월(202608) 이후 — C14 'Past Date.' 회피
METALS = ["Ni", "Co", "Li2CO3", "LIOH", "Al", "Mn", "Cu"]
BASE_COLS = [("Div", "DIVISION_ID"), ("Plan Date", "PLAN_YYYYMM"), ("Team", "DEPT_NAME"),
             ("End Customer", "END_CUSTOMER"), ("Model Code", "ITEM_ID"),
             ("Metal Link", "METAL_LINK_YN"), ("Sales Site", "SCM_SITE_ID"),
             ("Currency", "CURRENCY"), ("FCST type", "FCST_TYPE"),
             ("Exception", "METAL_EXCEPTION_YN"), ("Energy", "ENERGY")]
GROUPS = ["Key", "PeriodFr", "PeriodTo", "ContractPrice", "ChangePrice", "MetalIndex"]
HEADERS = [h for h, _ in BASE_COLS] + [f"{g}_{m}" for g in GROUPS for m in METALS]

REQ = ["CURRENCY", "FCST_TYPE", "END_CUSTOMER", "SCM_SITE_ID", "ITEM_ID"]
rows_src, seen = [], set()
for t in json.loads(TPL.read_text(encoding="utf-8"))["40"]:
    if any(str(t.get(k, "-")) in ("-", "", "None") for k in REQ):
        continue
    key = tuple(t.get(x) for x in REQ)
    if key in seen:
        continue
    seen.add(key)
    rows_src.append(t)
if len(rows_src) < 3:
    sys.exit("⛔ 템플릿에서 유효 키 조합 3건 확보 실패")


def base_row(src):
    r = {h: "" for h in HEADERS}
    for hdr, key in BASE_COLS:
        v = src.get(key, "")
        r[hdr] = "" if v in ("-", None) else v
    r["Div"], r["Plan Date"], r["Exception"] = "40", PLAN_YYYYMM, "N"
    return r


def metal(r, m, code, fr=None, to=None, cp=None, chg=None, idx=None):
    """연동 코드면 5값 전부, 미연동 코드면 5값 공란(NEED EMPTY VALUE 규칙).

    ⚠ Period(Fr/To) 는 **문자열**로 넣는다 — 숫자 셀로 넣으면 BE 엑셀 파서가 `6.0` 으로 변환하고
      SP C42 가 `INSTR(값,'.') > 0` 을 오류로 보아 'Invalid INTEGER 0 ~ 12.' 로 탈락한다
      (2026-08-18 실측 재현: 숫자 6 → 스테이징 '6.0' → ERROR).
    """
    r[f"Key_{m}"] = code
    if fr is None:
        return
    r[f"PeriodFr_{m}"], r[f"PeriodTo_{m}"] = str(fr), str(to)   # C43: Fr > To / C42: 0~12 정수
    r[f"ContractPrice_{m}"], r[f"ChangePrice_{m}"], r[f"MetalIndex_{m}"] = cp, chg, idx

samples, notes = [], []

s1 = base_row(rows_src[0])
s1["Metal Link"] = "N"
samples.append(s1)
notes.append("① Metal Link='N' — 메탈 미연동 행. 광물 42열 전부 공란(검증 게이트 밖)")

s2 = base_row(rows_src[1])
s2["Metal Link"] = "Y"
metal(s2, "Ni", "M035", 6, 0, 18.5, 0.75, 0.42)
samples.append(s2)
notes.append("② Metal Link='Y' + Ni 1종 — Key_Ni='M035'(LME) + Period(6>0)·Price·Index 5값 완전입력")

s3 = base_row(rows_src[2])
s3["Metal Link"] = "Y"
metal(s3, "Ni", "M035", 3, 1, 17.2, 0.6, 0.38)
metal(s3, "Co", "M037", 3, 1, 24.9, 1.1, 0.15)
metal(s3, "Al", "M051")          # 미연동 — 5값 공란이 정답
samples.append(s3)
notes.append("③ Metal Link='Y' + Ni·Co 2종 완전입력 + Al='M051'(미연동)은 5값 공란 "
             "— 연동/미연동 혼재 패턴")

wb = Workbook()
ws = wb.active
ws.title = "Metal Link Price"
ws.append(HEADERS)
for r in samples:
    ws.append([r[h] for h in HEADERS])

fill = PatternFill("solid", fgColor="1F4E79")
for c in ws[1]:
    c.fill, c.font = fill, Font(color="FFFFFF", bold=True, size=10)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for i, h in enumerate(HEADERS, start=1):
    ws.column_dimensions[get_column_letter(i)].width = max(11, min(22, len(h) + 3))
ws.freeze_panes = "A2"

# 안내 시트 (업로드 대상 시트는 첫 시트 하나 — 안내는 뒤에 둔다)
gd = wb.create_sheet("README")
gd.append(["DMF0074 Metal Link Price — Excel Upload 샘플 (division 40 · 3건)"])
gd.append([])
gd.append(["항목", "내용"])
for k, v in [
    ("양식", "DMF0074_40 카탈로그 53열 (Div/Plan Date/Team/End Customer/Model Code/Metal Link/"
             "Sales Site/Currency/FCST type/Exception/Energy + 광물 6그룹×7원소 42)"),
    ("업로드 경로", "화면 [Upload] → 검증(P_MODE=0) 결과 확인 → [OK] 로 적재(P_MODE=1). "
                    "Error 가 1건이라도 있으면 적재 자체가 차단된다"),
    ("Plan Date", f"{PLAN_YYYYMM} — 당월 이전이면 'Past Date.' 로 탈락하므로 실행 시점에 맞게 갱신할 것"),
    ("키 조합", "Excel Form(템플릿) 실조회분에서 선별 — 금주 Sales Plan Waterfall MRF 'Y' 버전에 "
                "존재하는 조합만 통과한다(Invalid Sales Plan Waterfall MODEL. 회피)"),
    ("헤더 문자열", "1바이트라도 다르면 그 열 전량이 조용히 탈락한다 — 헤더 행은 수정 금지"),
    ("광물 값 규칙", "Key 가 연동 코드면 Period(Fr)/Period(To)/Contract Price/Change Price/"
                     "Metal Index 5값을 모두 채워야 하고, '미연동' 코드(M036/M040/M044/M051/M055/"
                     "M057/M058)면 5값을 모두 비워야 한다"),
    ("Period 규칙", "0~12 정수이며 Period(Fr) > Period(To) 여야 한다(같거나 작으면 탈락)"),
    ("⚠ Period 입력 형식", "Period(Fr)/Period(To) 는 **텍스트 서식**으로 입력할 것. 엑셀에서 숫자로 "
                           "입력하면 서버 파서가 '6.0' 으로 변환하고 검증이 소수점을 오류로 보아 "
                           "'Invalid INTEGER 0 ~ 12.' 로 탈락한다(2026-08-18 실측). "
                           "본 샘플은 텍스트로 넣어 두었으니 셀 서식을 바꾸지 말 것"),
    ("40 division 주의", "Project Code/Project 열이 없는 전용 양식이다. 55열 'DMF0074' 양식을 쓰면 "
                         "전 행이 'NULL Metal Link.' 로 탈락한다"),
    ("검증 규칙 전문", "Doc/화면분석서/DMF0074_MetalLinkPrice_ExcelUpload_Validation_v1.1.xlsx (49판정)"),
    ("⚠ 환경 전제", "DM_BPS_METAL_LINK_PRICE_H.HISTORY_SEQ identity 시퀀스가 기존 이력 최대값보다 "
                    "뒤처져 있으면 [OK](P_MODE=1) 가 ORA-00001 로 실패한다(다건 업로드에서 항상 재현). "
                    "사내는 2026-08-18 재시드 완료(61→1565). 클라우드는 적용 전 "
                    "_artifacts/_apply_corporate/DM_BPS_METAL_LINK_PRICE_H_IDENTITY_RESEED_20260818.sql "
                    "①진단 쿼리로 확인할 것"),
    ("실검증 결과", "2026-08-18 사내 실측 — ①검증(P_MODE=0) Total 3 / Success 3 / Error 0, "
                    "②적재(P_MODE=1) CODE=0 EVENT=3 → 원장 3행 + 이력 3행 생성, "
                    "검증 후 해당 3행 삭제해 원장 원복(36,870 / 1,539 동일)"),
    ("⚠ 적재 행 수", "원장은 METAL_KIND_CODE 단위로 저장된다. 2행(Ni 1종)=1행, 3행(Ni·Co 2종)=2행이 "
                     "적재되고, **1행(Metal Link='N')은 검증은 통과하지만 적재 행이 생성되지 않는다** "
                     "(연동 광물이 없어 저장할 메탈 행이 없음). 3건 모두 원장에 남겨야 하면 "
                     "1행도 Metal Link='Y' + 광물 1종 이상으로 바꿀 것"),
]:
    gd.append([k, v])
gd.append([])
gd.append(["행", "구성"])
for i, n in enumerate(notes, start=1):
    gd.append([f"{i}행", n])
for c in gd[3]:
    c.fill, c.font = fill, Font(color="FFFFFF", bold=True, size=10)
gd.column_dimensions["A"].width = 18
gd.column_dimensions["B"].width = 118
for row in gd.iter_rows(min_row=1):
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=True)
gd["A1"].font = Font(bold=True, size=12, color="1F4E79")

wb.save(OUT)
print(f"생성 -> {OUT}")
for i, (s, n) in enumerate(zip(samples, notes), start=1):
    print(f"  {i}행 {s['Model Code']:<16} {s['Sales Site']:<24} {s['Currency']} "
          f"MetalLink={s['Metal Link']}  | {n[:60]}")

# 업로드 검증용 바이트도 반환 (동일 파일을 그대로 검증에 사용)
buf = io.BytesIO()
wb2 = Workbook()
w2 = wb2.active
w2.title = "Metal Link Price"
w2.append(HEADERS)
for r in samples:
    w2.append([r[h] for h in HEADERS])
wb2.save(buf)
(OUT.parent / "_dmf0074_sample_upload_only.xlsx").write_bytes(buf.getvalue())
print(f"검증용(단일시트) -> {OUT.parent / '_dmf0074_sample_upload_only.xlsx'}")
