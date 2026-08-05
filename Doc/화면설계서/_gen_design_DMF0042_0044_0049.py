# -*- coding: utf-8 -*-
"""DMF0042 / DMF0044 / DMF0049 화면설계서 8시트 생성 (hook #105 §1-A)
   원천: 사내 메타 실측(2026-07-29) + UI분석서/선행 MD 승계 + FE 소스 실측"""
import json, pathlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XImage

DOC = pathlib.Path(r"C:\Projects\NSCM\Doc\화면설계서")
CAP = pathlib.Path(r"C:\Projects\NSCM\Doc\_figma\CAP_SVG")
SCR = pathlib.Path(r"C:\Users\ejene\AppData\Local\Temp\claude\c--Projects-NSCM\adf13ccc-2d21-421e-8c5a-38922bd857e4\scratchpad")
UI = json.load(open(SCR / "ui3.json", encoding="utf-8"))
UIF = json.load(open(SCR / "ui3_full.json", encoding="utf-8"))

TITLE_F = Font(size=14, bold=True, color="1F4E79")
SEC_FILL = PatternFill("solid", fgColor="1F4E79")
SEC_F = Font(bold=True, color="FFFFFF", size=11)
HDR_FILL = PatternFill("solid", fgColor="4472A8")
HDR_F = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)


class SB:
    def __init__(self, ws, widths):
        self.ws, self.r, self.n = ws, 0, len(widths)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def title(self, t):
        self.r += 1
        c = self.ws.cell(self.r, 1, t); c.font = TITLE_F
        self.r += 1

    def sec(self, t):
        self.r += 1
        for i in range(1, self.n + 1):
            self.ws.cell(self.r, i).fill = SEC_FILL
        c = self.ws.cell(self.r, 1, t); c.font = SEC_F
        self.ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=self.n)
        c.alignment = Alignment(horizontal="left", vertical="center")

    def hdr(self, cols):
        self.r += 1
        for i, v in enumerate(cols, 1):
            c = self.ws.cell(self.r, i, v); c.fill = HDR_FILL; c.font = HDR_F
            c.alignment = Alignment(vertical="center", wrap_text=True); c.border = THIN

    def row(self, cols):
        """값이 있는 마지막 열 이후가 전부 비면 그 열부터 끝까지 가로 머지 + 좌측 정렬
        (사용자 요청 2026-07-29 — 1열만 채워진 행은 전 컬럼 머지)."""
        self.r += 1
        vals = list(cols) + [None] * (self.n - len(cols))
        idx = [i for i, v in enumerate(vals) if v is not None and str(v).strip()]
        for i, v in enumerate(vals, 1):
            c = self.ws.cell(self.r, i, v)
            c.alignment = WRAP
            c.border = THIN
        if idx and idx[-1] < self.n - 1:
            self._merge_left(self.r, idx[-1] + 1)

    def text(self, t):
        self.r += 1
        self.ws.cell(self.r, 1, t).alignment = WRAP
        self._merge_left(self.r, 1)

    def _merge_left(self, r, start_col=1):
        """start_col ~ 마지막 컬럼을 가로 머지 + 좌측 정렬(줄바꿈 유지)."""
        if self.n < 2 or start_col >= self.n:
            return
        self.ws.merge_cells(start_row=r, start_column=start_col, end_row=r, end_column=self.n)
        self.ws.cell(r, start_col).alignment = Alignment(
            horizontal="left", vertical="top", wrap_text=True
        )

    def gap(self):
        self.r += 1


def ov(mid, key, default=""):
    return UI[mid]["overview"].get(key, default)


# ── 선행 MD 본문 → 설계서 자연어 정리 (hook #105: 라인번호·원시코드 표기 제거, 식별자만 유지) ──
import re as _re

_LINE_REF = _re.compile(r"\s*\((?:정본\s*)?L\d[\d,\s~\-–L]*\)")      # (L123~L456) / (정본 L12, L34)
_LINE_REF2 = _re.compile(r"\s*[—-]\s*L\d[\d,\s~\-–L]*(?=[\s,)]|$)")  # — L123~L456
_BACKTICK = _re.compile(r"`([^`]*)`")
_BOLD = _re.compile(r"\*\*([^*]*)\*\*")


_LINE_TOKEN = _re.compile(r"(?:정본\s*)?L\d+(?:\s*[~\-–]\s*L?\d+)?")  # L123 / L123~L456


def clean(text: str) -> str:
    """MD 본문 1행을 설계서 서술체로 정리 — 내용은 보존하고 표기(백틱·굵게·라인번호)만 제거한다."""
    t = _BACKTICK.sub(r"\1", text)
    t = _BOLD.sub(r"\1", t)
    t = _LINE_REF.sub("", t)
    t = _LINE_REF2.sub("", t)
    t = _LINE_TOKEN.sub("", t)              # 괄호 안 다른 설명과 섞인 라인번호까지 제거
    t = _re.sub(r"[,\s]*\)", ")", t)        # "(설명, )" → "(설명)"
    t = _re.sub(r"\(\s*[,;]\s*", "(", t)
    t = _re.sub(r"\(\s*\)", "", t)
    t = _re.sub(r"\s+([,.])", r"\1", t)
    t = _re.sub(r"\s{2,}", " ", t)
    return t.strip(" ,;—-")


def md_detail_rows(md_path):
    """MD 전문을 (구분, 항목, 내용) 행으로 변환. 표·코드블록 포함 전 행 보존."""
    rows, cur_h2, cur_h3, in_code = [], "", "", False
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        s = line.strip()
        if s.startswith("```"):
            in_code = not in_code
            continue
        if not s or s.startswith("---"):
            continue
        if s.startswith("# "):
            continue
        if s.startswith("## "):
            cur_h2, cur_h3 = clean(s[3:]), ""
            rows.append(["■ 이벤트", cur_h2, ""])
            continue
        if s.startswith("### "):
            cur_h3 = clean(s[4:])
            rows.append(["· 단계", cur_h3, ""])
            continue
        if s.startswith("#### "):
            cur_h3 = clean(s[5:])
            rows.append(["· 단계", cur_h3, ""])
            continue
        # MD 표 — 구분선은 버리고 셀 단위로 분해해 읽기 좋게 배치
        if s.startswith("|") and s.endswith("|"):
            cells = [clean(c) for c in s.strip("|").split("|")]
            if all(_re.fullmatch(r":?-{2,}:?", c or "") for c in cells if c is not None):
                continue
            cells = [c for c in cells]
            rows.append(["  표", cells[0] if cells else "", " · ".join(cells[1:])])
            continue
        body = clean(s.lstrip("-*").strip() if s.startswith(("- ", "* ")) else s)
        if not body:
            continue
        kind = "  코드/식별자" if in_code else "  내용"
        rows.append([kind, cur_h3 or cur_h2, body])
    return rows


# ── 화면별 실측 데이터 (사내 DB / FE 소스 2026-07-29) ─────────────────────────
SPEC = {
 "DMF0042": {
  "file": "DMF0042_BizPlan_화면설계서_v1.1.xlsx",
  "name": "Business Plan", "page": "business-plan", "cap": "DMF0042.png",
  "route": "demand/demand-plan/business-plan", "url": "business-plan", "mnu_desc": "Business Plan",
  "sheets": ['Biz Plan', 'cutoffVersion', 'dayRange', 'exchangeRate', 'futurePrice',
             'futurePriceData', 'mdfVersion', 'spreadTitle', 'titleSet'],
  "btn_rows": 23, "sp_cnt": 15, "cols": [("Biz Plan", 63, 60, 63)], "measures": 9,
  "search": [
    ("Division *", "카탈로그 DivisionField", "SCM0001_GET_DIVISION", "변경 시 하위 ALL reset + Grid 초기화 + CRUD reset(재조회는 Search)"),
    ("Team", "카탈로그", "TeambyDivision 계열", "Division 종속"),
    ("Sales Site", "카탈로그", "SalesSite 계열", "Division 종속"),
    ("Project", "카탈로그", "-", "40 은 비활성(AS-IS 정합)"),
    ("Model", "카탈로그", "-", ""),
    ("Owner", "카탈로그", "-", ""),
    ("Period", "카탈로그 PeriodField", "-", "datePickerFix 내장(높이 36px)"),
    ("End Customer", "인라인 콤보", "EndCustomer(81·82) / EndCustomerByDivision40(40)", "Division 40 에서 원천 스왑"),
    ("Biz Type", "인라인 콤보", "BizTypeByDivision(81·82) / DmBizType(40)", "Division 40 에서 원천 스왑"),
    ("Plant Site", "인라인 콤보", "PlantSitebyEndCustomer", "Division+EndCustomer 서버 바인드(값 quoting 의무). 40 만 비활성"),
    ("Item Type", "인라인 콤보", "ItemType(ITEM_TYPE_CODE)", ""),
    ("Bucket", "인라인 콤보", "BucketType2(MONTHLY/QUARTERLY/YEARLY)", "기본 Monthly. 변경 시 변경분 존재하면 MSG00227 로 차단"),
    ("Sum(체크박스)", "체크박스", "-", "합계 열 표시 토글"),
  ],
  "buttons": [
    ("SEARCH", "공통", "조회 — 메인 데이터 + 동적 컬럼(spreadTitle) 재구성"),
    ("INSERT", "공통", "행 추가(권한 hasButton 게이트)"),
    ("DELETE", "공통", "행 삭제 마킹"),
    ("SAVE", "공통", "'^' bulk 저장 — 변경(dirty) 셀만 수집"),
    ("SORT / FILTER / EXCEL / COLUMN_SETUP", "공통", "정렬·필터·엑셀·컬럼설정(공용 모달)"),
    ("MEASURE_SELECT", "특화", "Measure 선택 모달(SHEET_MEASURE 9종)"),
    ("EXCHANGE_RATE", "특화", "환율 팝업 — 조회/저장(exchangeRate 시트)"),
    ("CUTOFF", "특화", "Cutoff 버전 반영(cutoffVersion 시트)"),
    ("REPLACE", "특화", "Cutoff 기준 대체 저장(SAVE_DM_PLN_BP_REPLACE_BY_CUTOFF)"),
    ("SET_PRICE", "특화", "미래단가 반영(futurePrice·futurePriceData 시트 → SAVE_..._PRICE_CONV)"),
  ],
  "grid": [
    ("Biz Plan (메인)", "행 = 키 × Measure(9), 열 = 연(Y_)/분기(Q_)/월(M_) 3계층 PIVOT", "63컬럼(가시 60) · 전 컬럼 셀 콤보 메타 보유 · MEASURE 컬럼 VISIBLE_FLAG=N(메타 그대로)"),
    ("동적 컬럼", "spreadTitle SEARCH 결과로 기간 컬럼 생성", "Bucket(M/Q/Y)에 따라 열 구성 변경. 그룹+leaf 전부 allowSorting/allowMerging false"),
    ("편집 규칙", "EDIT_FLAG(행 DEPT_ID=내 팀 OR ADMIN) × Measure ∈ {Business_Plan, LOC}", "편집 셀 Yellow / 잠금 Gray / 기본 White+zebra(3색 표준)"),
    ("입력 파생", "Business_Plan·LOC 입력 시 분기→월 자동 배분", "Y_/Q_ SUM 은 footer 가 아닌 셀 계산"),
  ],
  "popups": ["ExchangeRatePopup(환율)", "Cutoff Version", "Future Price(단가) 2단", "MDF Version", "TitleSet(공용)"],
 },
 "DMF0044": {
  "file": "DMF0044_SalesReports_화면설계서_v1.1.xlsx",
  "name": "Sales Reports", "page": "sales-reports", "cap": "DMF0044.png",
  "route": "demand/demand-plan/sales-reports", "url": "sales-reports", "mnu_desc": "Sales Reports",
  "sheets": ['MDF Reports', 'exchangeRate', 'measure', 'spreadTitle', 'textbox', 'titleSet', 'unitedVersion'],
  "btn_rows": 28, "sp_cnt": 18, "cols": [("MDF Reports", 46, 39, 46), ("MDFVerpopup", 12, 12, 12)], "measures": 18,
  "search": [
    ("Division *", "DivisionField(공용)", "SCM0001_GET_DIVISION", "변경 시 하위 reset + Grid 초기화"),
    ("Report Type", "드롭다운", "MDF_TYPE(REPORT_SP_W=SP Version / REPORT_BP_M=BP Version)", "버전 축 결정"),
    ("Year", "드롭다운", "CALENDAR_YYYY", "선택 연도의 월(M_yyyymm) PIVOT 기준"),
    ("Data Type", "드롭다운", "MDF_TYPE2 (ALL / 1=Cutoff / 2=ETC)", "⭐ 편집 게이트1 — ALL 이면 편집 전면 차단(MSG00321)"),
    ("Version", "버전 팝업 연동 입력", "unitedVersion SEARCH", "United Version 팝업에서 선택"),
  ],
  "buttons": [
    ("SEARCH", "공통", "조회 — 버전·연도 기준 데이터 + 동적 월 컬럼"),
    ("INSERT", "공통", "행 추가 — **Data Type=2(ETC) 일 때만 활성**"),
    ("DELETE", "공통", "행 삭제 마킹"),
    ("SAVE", "공통", "저장 — Data Type 2 → SAVE_SALES_PLAN_REPORT_C"),
    ("SAVE_CF", "특화", "Cutoff 저장 — Data Type 1 → SAVE_SALES_PLAN_REPORT_C_CF"),
    ("SORT / FILTER / EXCEL / COLUMN_SETUP", "공통", "정렬·필터·엑셀·컬럼설정"),
    ("MEASURE_SELECT", "특화", "Measure 선택(SHEET_MEASURE 18종)"),
    ("EXCHANGE_RATE", "특화", "환율 팝업 — 제목/조회/데이터/저장 4배선(exchangeRate 시트)"),
    ("TEXTBOX", "특화", "리포트 코멘트 조회·저장(textbox 시트)"),
    ("SP_VER", "특화", "버전 팝업 진입(unitedVersion)"),
    ("unitedVersion ACTION / ACTION_BP / TRANSFER", "팝업 내", "판매계획 버전 복사 / BP 버전 복사 / 모바일 리포트 생성"),
  ],
  "grid": [
    ("MDF Reports (메인)", "행 = 키 × Measure(18), 열 = 선택 연도의 월(M_yyyymm) 동적 PIVOT(연 span 2계층)", "46컬럼(가시 39) · 전 컬럼 콤보 메타"),
    ("동적 컬럼", "spreadTitle SEARCH 로 월 컬럼 생성 + 연 그룹 헤더", "2단 헤더 48px(24+24)"),
    ("⭐ 편집 게이트(이중)", "게이트1 = Data Type(ALL 차단 / 2=ETC 행추가+REPORT_C / 1=Cutoff REPORT_C_CF)", "게이트2 = SP 이중 게이트(Layer·버전 상태)"),
    ("MDFVerpopup", "버전 팝업 그리드", "12컬럼(전부 가시)"),
  ],
  "popups": ["United Version(정본 — 유일 호출)", "Exchange Rate", "Textbox(코멘트)", "TitleSet(공용)",
             "(참고) Sales_VersionPopUp·VersionCopy·MDFVersionPopUp = 현행 호출 경로 없음"],
 },
 "DMF0049": {
  "file": "DMF0049_EtcSalesReports_화면설계서_v1.1.xlsx",
  "name": "ETC Sales Reports", "page": "etc-sales-reports", "cap": "DMF0049.png",
  "route": "demand/demand-plan/etc-sales-reports", "url": "etc-sales-reports", "mnu_desc": "ETC Sales Reports",
  "sheets": ['MDF Reports', 'measure', 'spreadTitle', 'titleSet', 'unitedVersion'],
  "btn_rows": 19, "sp_cnt": 11, "cols": [("MDF Reports", 42, 33, 42)], "measures": 3,
  "search": [
    ("Division *", "DivisionField(공용)", "SCM0001_GET_DIVISION", "변경 시 하위 reset + Grid 초기화"),
    ("Report Type", "드롭다운", "MDF_TYPE", "버전 축"),
    ("Year", "드롭다운", "CALENDAR_YYYY", "월(M_) PIVOT 기준"),
    ("Plan/Result", "드롭다운", "계획/실적 구분", "AS-IS 기타매출 축 분리 기준"),
    ("Version", "버전 팝업 연동 입력", "unitedVersion SEARCH", "United Version 팝업 공유(DMF0044 SP 재사용)"),
  ],
  "buttons": [
    ("SEARCH", "공통", "조회 — 기타 매출 데이터 + 동적 월 컬럼"),
    ("INSERT / DELETE / SAVE", "공통", "행 추가·삭제·저장(SAVE_SALES_ETC_REPORT_C)"),
    ("SORT / FILTER / EXCEL / COLUMN_SETUP", "공통", "정렬·필터·엑셀·컬럼설정"),
    ("MEASURE_SELECT", "특화", "Measure 선택(SHEET_MEASURE 3종)"),
    ("SP_VER", "특화", "버전 팝업 진입"),
    ("unitedVersion ACTION / ACTION_BP / TRANSFER", "팝업 내", "판매계획·BP 버전 복사 / 모바일 리포트 생성 — **DMF0044 PKG 재사용**"),
  ],
  "grid": [
    ("MDF Reports (메인)", "행 = 키 × Measure(3: SalesAmt_LOC / SalesAmt_KRW / R…), 열 = 월(M_) 동적 PIVOT", "42컬럼(가시 33) · 40 전용 리포트 그룹 컬럼 다수"),
    ("동적 컬럼", "spreadTitle SEARCH 로 월 컬럼 생성", "DIVISION_ID 토큰으로 컬럼 노출 분기"),
    ("편집 규칙", "DMF0044 대비 단순 — Measure 3종 축", "편집 셀 Yellow / 잠금 Gray"),
  ],
  "popups": ["United Version(DMF0044 공유)", "TitleSet(공용)"],
 },
}

BTN_DETAIL = json.load(open(SCR / "btn3.json", encoding="utf-8")) if (SCR / "btn3.json").exists() else {}


def build(mid):
    s = SPEC[mid]
    wb = openpyxl.Workbook(); wb.remove(wb.active)

    # 1.개요
    b = SB(wb.create_sheet("1.개요"), [24, 122])
    b.title(f"{mid} {s['name']} — 화면설계서 v1.1")
    b.sec("화면 정의")
    rows = [
        ("화면 ID / 명", ov(mid, "화면 ID / 명")),
        ("도메인 / 메뉴경로", ov(mid, "도메인 / 메뉴")),
        ("메뉴 실측(SCM_ADM_MNU_M)", f"MNU_ID={mid} / UPPR=Demand_Plan / MNU_URL='{s['url']}' / ROUTE_PATH='{s['route']}' / MNU_DESC='{s['mnu_desc']}' / UNIT_TP_CD='D' (2026-07-29 실측)"),
        ("화면 목적", ov(mid, "화면 목적")),
        ("AS-IS 폼 (정본)", ov(mid, "AS-IS 폼 (정본)")),
        ("TO-BE 경로", f"nscm-fe/features/demand/src/pages/{s['page']} (구현 완료 — 2026-07-29 실측)"),
        ("사용 Division", ov(mid, "사용 Division")),
        ("주요 SP", ov(mid, "주요 SP")),
        ("시트(sheetName)", ov(mid, "시트(sheetName)")),
        ("편집 게이트", ov(mid, "편집 게이트") or ov(mid, "⭐ 편집 게이트 (이중)") or "선행 MD §3 참조"),
        ("제외 범위", "배치·인터페이스(발송·집계 배치)는 화면 범위 외. 공용 모달(TitleSet 등)은 공통 영역 링크."),
        ("선행 산출물", f"{mid} 주요이벤트_상세분석.md + 제약매트릭스_상세.md + UI분석서(Doc/화면분석서)"),
    ]
    for k, v in rows:
        b.row([k, v])

    # 2.전체흐름 (UI분석서 승계)
    b = SB(wb.create_sheet("2.전체흐름"), [30, 26, 40, 44, 56])
    b.title("전체흐름")
    for r in UIF[mid]["2.전체흐름"]:
        if len(r) == 1:
            b.sec(clean(r[0])) if len(r[0]) < 60 else b.text(clean(r[0]))
        else:
            b.row([clean(x) for x in r[:5]])

    # 3.LAYOUT
    b = SB(wb.create_sheet("3.LAYOUT"), [26, 110])
    b.title("LAYOUT")
    b.sec("영역 구성")
    b.hdr(["영역", "내용"])
    b.row(["조회조건(SearchArea)", f"{len(s['search'])}종 — 1단 4개(gridSize 3) · 전 객체 높이 36px(hook #94) · Division 변경 = 하위 reset + Grid 초기화"])
    b.row(["SubTitle 공통버튼", "hotkeyConfig 기반(권한 hasButton 게이트) + customButtons(화면 특화)"])
    b.row(["메인 그리드", s["grid"][0][1]])
    b.row(["팝업/모달", " · ".join(s["popups"])])
    b.gap()
    b.sec(f"AS-IS 캡쳐 (원본: Doc/_figma/CAP_SVG/{s['cap']})")
    try:
        im = XImage(str(CAP / s["cap"])); im.width, im.height = 950, 500
        b.ws.add_image(im, f"A{b.r + 2}")
    except Exception as e:
        b.text(f"(이미지 삽입 불가: {e})")

    # 4.조회조건·버튼
    b = SB(wb.create_sheet("4.조회조건·버튼"), [24, 22, 44, 56])
    b.title("조회조건 · 버튼")
    b.sec(f"A. 조회조건 {len(s['search'])}종")
    b.hdr(["항목", "형식", "원천(CMBO/SP)", "동작·종속"])
    for x in s["search"]:
        b.row(list(x))
    b.gap()
    b.sec(f"B. 버튼 (BTN_MPG {s['btn_rows']}행 / 시트 {len(s['sheets'])})")
    b.hdr(["버튼", "구분", "동작", ""])
    for x in s["buttons"]:
        b.row([x[0], x[1], x[2], ""])
    b.gap()
    b.sec("C. 권한")
    b.text("공통·특화 버튼 모두 useTabBtnPermissions + hasButton(BTN_CD) 게이트를 통과해야 활성된다. BTN_MPG 등록 시트·버튼 조합과 그룹 권한(USER_GR_MNU_BTN_MPG_M)이 일치해야 하며, 누락 시 '권한없음' 이 발생한다.")

    # 5.GRID
    b = SB(wb.create_sheet("5.GRID"), [26, 58, 62])
    b.title("GRID")
    b.sec("A. 그리드 구성")
    b.hdr(["구분", "내용", "메타·표준"])
    for x in s["grid"]:
        b.row(list(x))
    b.gap()
    b.sec("B. 컬럼 메타 실측 (SCM_MST_SHEET_COLUMNS / SHEET_MEASURE, 2026-07-29)")
    b.hdr(["시트", "컬럼 수", "가시(VISIBLE_FLAG=Y)"])
    for sheet, c, v, _ in s["cols"]:
        b.row([sheet, c, v])
    b.row(["SHEET_MEASURE", s["measures"], "Measure 선택 모달 원천"])
    b.gap()
    b.sec("C. 그리드 표준")
    b.text("공용 DataGrid 사용(인라인 FlexGrid 금지) / 컬럼 = SHEET_COLUMNS 동적 fetch(하드코딩 금지) / No. 코너 헤더 표준 / 셀 배경 3색(White+zebra · 편집 Yellow · 잠금 Gray) / 동적 컬럼은 NVL wrap + 그룹·leaf 모두 allowSorting·allowMerging false + 헤더 높이 명시(1단 48 / 2단 24 / 3단 16).")

    # 6.EVENT
    b = SB(wb.create_sheet("6.EVENT"), [14, 40, 110])
    b.title("EVENT (AS-IS 전수 — 선행 MD 승계)")
    b.hdr(["#", "언제", "무엇이 어떻게"])
    md = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서") / f"{mid}_{ {'DMF0042':'BizPlan','DMF0044':'SalesReports','DMF0049':'EtcSalesReports'}[mid] }_주요이벤트_상세분석.md"
    n = 0
    for line in md.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line.startswith("## ") and not line.startswith("## 원천"):
            n += 1
            head = clean(line[3:])
            b.row([f"E{n:02d}", head.split("—")[0].strip()[:60], head])
        elif line.startswith("### ") and n:
            sub = clean(line[4:])
            b.row(["", "   " + sub.split("(")[0].strip()[:56], sub])
    b.gap()
    b.sec(f"상세 — 이벤트 단계별 전수 (원천 {md.name} 전문 승계)")
    b.hdr(["구분", "항목", "내용"])
    for kind, head, body in md_detail_rows(md):
        b.row([kind, head, body])

    # 7.VALIDATION (UI분석서 승계)
    b = SB(wb.create_sheet("7.VALIDATION"), [14, 44, 58, 26])
    b.title("VALIDATION (화면 + SP 취합)")
    for r in UIF[mid]["3.VALIDATION취합"]:
        if len(r) == 1:
            b.sec(clean(r[0])) if len(r[0]) < 60 else b.text(clean(r[0]))
        else:
            b.row([clean(x) for x in r[:4]])

    # 8.API·SP
    b = SB(wb.create_sheet("8.API·SP"), [20, 34, 46, 40])
    b.title("API · SP")
    b.sec(f"A. 시트 × API × SP (BTN_MPG {s['btn_rows']}행 / SP {s['sp_cnt']}종 — 2026-07-29 실측)")
    b.hdr(["SHT_NM", "MNU_BTN_CD", "API_ID", "API_URL(SP)"])
    for r in BTN_DETAIL.get(mid, []):
        b.row(r)
    b.gap()
    b.sec("B. SP·LOOKUP·콤보 (UI분석서 승계)")
    for r in UIF[mid]["9.SP_LOOKUP_콤보"]:
        if len(r) == 1:
            b.text(clean(r[0]))
        else:
            b.row([clean(x) for x in r[:4]])

    out = DOC / s["file"]
    wb.save(out)
    return out


made = []
for mid in ("DMF0042", "DMF0044", "DMF0049"):
    made.append(build(mid))

for p in made:
    wb = openpyxl.load_workbook(p)
    names = wb.sheetnames
    assert len(names) == 8, f"{p.name}: {len(names)}시트"
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    assert "단위테스트" not in v, f"{p.name}/{ws.title} 금지어"
    wb.close()
    print(f"OK {p.name} — 8시트 assert PASS / {names}")
