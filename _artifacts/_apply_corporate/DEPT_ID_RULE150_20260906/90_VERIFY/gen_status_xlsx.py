# -*- coding: utf-8 -*-
"""hook #150 DEPT_ID 원천 변경 — 최초 대상 18 객체 처리현황 xlsx 생성 (4 시트)"""
import json, pathlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PKG = pathlib.Path(__file__).resolve().parents[1]
OUT = PKG / "DEPT_ID_RULE150_처리현황_20260906.xlsx"

rows = [
    ("PACKAGE BODY", "PKG_DMF0014", "변경대상없음", "6442/6499/6524 이미 ITM.DEPT_ID (2026-04-01 전환), AS-IS ①③ 은 주석", "—", "—"),
    ("PACKAGE BODY", "PKG_DMF0020", "9건 수정", "SALES_PRICE_TEMPLATE 2729·2805 / OWNER_DATA 3800 (3쿼리) / FUTURE_PRICE 5285·5633 / POPUP 8005·8083·8473",
     "GET_SALES_PRICE_TEMPLATE 40/81/82 = 0/121/107 동일 · GET_OWNER_DATA 0/0/0 동일 · GET_SCM_MST_FUTURE_PRICE 8,652/9,956/10,520 동일 · POPUP 40 2,152 동일 · 81 2,454 동일 · 82 38,659→38,672 (+13)",
     "POPUP DEPT_ID: 81 NULL 122→2 / 82 NULL 0→3,063 (SCM_MST_SITE_ITEM.DEPT_ID 결측 아이템) / 40 동일 · OWNER_DATA 부서후보(V_TMPDEPT) 81 108→109"),
    ("PACKAGE BODY", "PKG_DMF0022", "1건 수정", "1802 GT_DM22_SCM_MST_SITE_ITEM.DEPT_ID 스칼라 → A.DEPT_ID", "40/81/82 = 1,679/1,156/736 동일", "없음 (DEPT 분포 13/16/12 동일, 결과 해시 동일)"),
    ("PACKAGE BODY", "PKG_DMF0035", "1건 수정", "2247 GT_DM35 스칼라 → A.DEPT_ID", "370/517/199 동일", "없음 (해시 동일)"),
    ("PACKAGE BODY", "PKG_DMF0040", "1건 수정", "1307 E.DEPT_ID 추가 + 1335 TEAM=DEPT_INFO ON E.DEPT_ID + 1347 K ON E.DEPT_ID (동적 SQL)", "81/82 = 394/476 동일 · 40 은 전/후 모두 300s 초과(측정불가)", "없음 (해시 동일)"),
    ("PACKAGE BODY", "PKG_DMF0044", "5건 수정", "3084 스칼라 4개 / 8588 / 8718 / 8825 / 8847",
     "SAVE SP → 블록 SQL 대조: 8588 div별 동일 · 3084 23,199 전건 동일값 · 8718 0=0 (룩업 타입 부재 사각) · 8825 82분기 38,621→35,557",
     "8825: item DEPT_ID NULL 3,064행 분기1 이탈 (분기2 UNION 으로 유지)"),
    ("PACKAGE BODY", "PKG_DMF0049", "변경대상없음", "DEPT_ID 는 DM_PLN_SALES_ETC_REPORT_C 저장값 소비 + TB_FX_USER 본인 부서(권한)", "—", "—"),
    ("PACKAGE BODY", "PKG_DMF0051", "2건 수정", "1177 OW.DEPT_ID → SMSI.DEPT_ID / 1205·1215 GT_DM51_SCM_CITY DEPT = SITE_ITEM PK 4키", "40 = 13,076 동일 · 81/82 = 0 동일", "없음 (해시 동일)"),
    ("PACKAGE BODY", "PKG_DMF0074", "4건 수정", "96·201·632·711 TEAM=DEPT_INFO ON SITE_ITEM/SI.DEPT_ID", "METAL_LINK_PRICE 0/4,390/480 동일 · TEMPLATE 342/145/74 동일", "없음 (해시 동일)"),
    ("PACKAGE BODY", "PKG_DMF0082", "6건 수정", "844 TEAM+K / 5066 UPLOAD DI ON SI.DEPT_ID / 7407 TEAM+K / 10395 GT_PJT / 10488 / 11045",
     "POPUP 40 2,220 동일 · 81 2,454 동일 · 82 측정불가(전 0 / 후 300s 초과) · GET_MAIN_DATA(_CURRENT) 전/후 모두 ORA-00904 SORT_KEY (기존 결함, 측정불가)",
     "POPUP 81 NULL 122→2 · 5066 END_CUSTOMER 후보 81 333→340 · 82 1,297→1,375 (상위집합)"),
    ("PROCEDURE", "DMF0020_GET_SCM_MST_FUTURE_PRICE", "2건 수정", "697 PJT SI.DEPT_ID / 1044 TEAM=DEPT_INFO ON SITE_ITEM.DEPT_ID (649 NEW_DEPT_ID_TBL MAX 는 사용자 확정 유지)", "8,652/9,956/10,520 동일", "없음 (해시 동일)"),
    ("PROCEDURE", "DMF0020_GET_SCM_MST_FUTURE_PRICE_SAVE_PRICE_POPUP", "3건 수정", "292 GT_PJT SI.DEPT_ID / 370 TEAM·DEPRTMENT → SITE_ITEM.DEPT_ID / 752 SITE_PROJECT TEAM → SITE_ITEM.DEPT_ID",
     "81 2,454 동일 · 82 38,659→38,672 (+13) · 40 측정 제외(GT_DMF0020_PJT 적재 >10분)", "81 NULL 122→2 · 82 NULL 0→3,063"),
    ("PROCEDURE", "PR_DM_BPS_METAL_LINK_S", "변경대상없음", "A.DEPT_ID = DM_BPS_METAL_LINK_PRICE_C 저장값 (상류 PKG_DMF0074)", "—", "—"),
    ("PROCEDURE", "PR_DM_BPS_SALES_ETC_PLAN_S", "변경대상없음", "251 이미 D.DEPT_ID (SITE_ITEM). SU OWNER 조인은 DEPT 미사용 잔존 (사용자 확정 미변경)", "—", "—"),
    ("PROCEDURE", "PR_DM_BPS_SALES_PLAN_S", "변경대상없음", "367/441/513 이미 D.DEPT_ID. SU 잔존 동일", "—", "—"),
    ("PROCEDURE", "PR_DM_METAL_LINK_CUTOFF", "변경대상없음", "85/154 저장값 이관", "—", "—"),
    ("VIEW", "IF_DM_HI_CUST_FCST_INPUT_V", "변경대상없음", "DM_HI_CUST_FCST_INPUT.DEPT_ID passthrough", "—", "—"),
    ("VIEW", "IF_DM_HI_SALES_INPUT_V", "변경대상없음", "DM_HI_SALES_INPUT.DEPT_ID passthrough", "—", "—"),
]

thin = Side(style="thin", color="999999")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill("solid", fgColor="1F4E78")
MF = PatternFill("solid", fgColor="FFF2CC")
NF = PatternFill("solid", fgColor="E2EFDA")
RF = PatternFill("solid", fgColor="F8CBAD")


def style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HF
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BD


def style_body(ws, widths):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BD
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, w in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


wb = openpyxl.Workbook()
ws = wb.active
ws.title = "객체별 처리현황"
ws.append(["No", "유형", "객체", "비고", "수정 지점 (라인)", "건수 변화 (before→after, 40/81/82)", "값 달라진 지점"])
for i, r in enumerate(rows, 1):
    ws.append([i, *r])
style_header(ws)
style_body(ws, [5, 14, 44, 14, 60, 70, 60])
for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
    row[3].fill = NF if row[3].value == "변경대상없음" else MF
    row[3].font = Font(bold=True)
ws.auto_filter.ref = f"A1:G{len(rows)+1}"
n = len(rows) + 3
ws.cell(n, 1, "합계").font = Font(bold=True)
ws.cell(n, 3, "수정 10 객체 / 34 지점 (+MAX 유지 1) · 변경대상없음 8 객체 · 배포 10 객체 전부 VALID · USER_ERRORS 0 · 신규 INVALID 0")
ws.cell(n + 1, 3, "건수 상이 = 2 커서 (POPUP 82, standalone·PKG 동일 원인 +13: 한 SITE/PROJECT 안에서 부서가 갈리는 아이템의 부서별 분리) / 값 상이 = POPUP 81·82 DEPT_ID 컬럼 + SAVE/UPLOAD 3 지점")
ws.cell(n + 2, 3, "기준: AJS 175.211.29.206:2521 · BEFORE 2026-09-06 01:34 / 배포 08:33~08:35 / AFTER 09:34 · 15 GET SP × 3 division 동일 파라미터 (90_VERIFY)")

# 시트2: 커서별 대사 원본
ws2 = wb.create_sheet("커서별 건수 대사")
b = json.loads((PKG / "90_VERIFY" / "baseline_before.json").read_text(encoding="utf-8"))
a = json.loads((PKG / "90_VERIFY" / "baseline_after.json").read_text(encoding="utf-8"))
ws2.append(["SP", "Division", "커서", "before 행수", "after 행수", "판정", "DEPT distinct(b→a)", "DEPT NULL(b→a)", "해시", "비고"])
for k in b:
    sp, div = k.split("|")
    bb = b[k]
    aa = a.get(k, {})
    if "error" in bb or "error" in aa:
        ws2.append([sp, div, "-", "-", "-", "측정불가", "", "", "", f"before={bb.get('error','')[:70]} / after={aa.get('error','')[:70]}"])
        continue
    for cb, ca in zip(bb["cursors"], aa["cursors"]):
        rb, ra = cb.get("rows"), ca.get("rows")
        h = "=" if cb.get("hash") and cb.get("hash") == ca.get("hash") else ("≠" if cb.get("hash") and ca.get("hash") else "-")
        dd = f"{cb.get('dept_distinct')}→{ca.get('dept_distinct')}" if cb.get("dept_distinct") is not None else ""
        dn = f"{cb.get('dept_null')}→{ca.get('dept_null')}" if cb.get("dept_null") is not None else ""
        ws2.append([sp, div, cb["cursor"], rb, ra, "동일" if rb == ra else "상이", dd, dn, h, ""])
style_header(ws2)
style_body(ws2, [62, 9, 6, 11, 11, 9, 18, 16, 6, 80])
for row in ws2.iter_rows(min_row=2):
    if row[5].value == "상이":
        row[5].fill = RF
ws2.auto_filter.ref = f"A1:J{ws2.max_row}"

# 시트3: 블록 SQL 대조
ws3 = wb.create_sheet("블록SQL 대조(SAVE·UPLOAD)")
ws3.append(["객체", "지점", "대조 방식", "old", "new", "판정", "설명"])
for r in [
    ("PKG_DMF0044", "8588 SAVE_SALES_PLAN_VERSION_COPY", "SITE_ITEM 행 OWNER 체인 INNER 생존", "40 15,565 / 81 45,459 / 82 35,557", "동일", "동일", "OWNER 첫 토큰→SSO_USER_INFO→DEPT_INFO 와 SITE_ITEM.DEPT_ID→DEPT_INFO 생존 집합 일치"),
    ("PKG_DMF0044", "3084 SAVE_BIZ_PLAN_VERSION_COPY (div40)", "스칼라 값 동일성", "23,199", "23,199 동일값", "동일", "A.OWNER 첫 토큰 경유 DEPT == SITE_ITEM.DEPT_ID 전건"),
    ("PKG_DMF0044", "8718 (81 분기)", "'MDF_DEPARTMENT'||div 룩업 INNER 생존", "0", "0", "동일", "룩업 타입 MDF_DEPARTMENT81 부재(실존 MDF_DEPARTMENT_81) → 기존 사각 분기"),
    ("PKG_DMF0044", "8825 (82 분기1)", "site 부서∈DEPT_INFO INNER 생존", "38,621", "35,557", "상이", "item DEPT_ID NULL 3,064행 분기1 이탈 (분기2 UNION 에서 유지)"),
    ("PKG_DMF0082", "5066 EXCEL_UPLOAD END_CUSTOMER 후보", "DISTINCT PROJECT_CODE", "40 875 / 81 333 / 82 1,297", "875 / 340 / 1,375", "상이(상위집합)", "기존 SI.OWNER=UI.USER_ID 전체문자열 + 첫토큰 이중조건 → 단일 OWNER 행만 통과하던 결함 해소"),
    ("PKG_DMF0020", "3800 GET_OWNER_DATA V_TMPDEPT", "DISTINCT DEPT 집합", "40 13 / 81 108 / 82 62", "13 / 109 / 62", "상이(상위집합)", "A.OWNER = B.USER_ID 전체문자열 비교 → 다중 OWNER 행 부서 1개 누락 복원"),
    ("PKG_DMF0082 / PKG_DMF0040", "844·7407 / 1335 TEAM 체인", "데이터 증명", "117,948행 OWNER 첫토큰 DEPT", "== SITE_ITEM.DEPT_ID 전행", "동일", "DIFF 0 · 편측 0 · DEPT_INFO 고아 0"),
]:
    ws3.append(list(r))
style_header(ws3)
style_body(ws3, [24, 38, 34, 30, 26, 14, 70])

# 시트4: 기존 결함
ws4 = wb.create_sheet("기존결함(변경무관)")
ws4.append(["No", "객체", "증상", "근거", "조치 필요"])
for r in [
    (1, "PKG_DMF0082.P_DMF0082_GET_MAIN_DATA / _CURRENT", "ORA-00904 SORT_KEY: invalid identifier — 호출 실패", "적용 전(01:34) 부터 동일 오류", "GTT_TEMPFCSTNPRICE 적재 SELECT 의 999 AS SORT_KEY 컬럼 대조"),
    (2, "PKG_DMF0044 8718 분기", "'MDF_DEPARTMENT'||p_division_id → MDF_DEPARTMENT81 룩업 부재 → 항상 0행", "실존 LOOKUP_TYPE = MDF_DEPARTMENT_81 / DMF0044_DEPARTMENT_TEAM2_81", "언더스코어 누락 정정 여부 판단"),
    (3, "PKG_DMF0040.P_DMF0040_GET_MKT_FCST_COMPARISON|40", "호출 300초 초과 (전/후 동일)", "하네스 DPY-4024", "성능 튜닝"),
    (4, "DMF0020_GET_SCM_MST_FUTURE_PRICE_SAVE_PRICE_POPUP|40 (standalone)", "GT_DMF0020_PJT 적재 >10분 (PKG 버전은 6초)", "V$SESSION 12분 ACTIVE 실측", "standalone 구버전 폐기/PKG 배선 확인"),
]:
    ws4.append(list(r))
style_header(ws4)
style_body(ws4, [5, 52, 60, 44, 40])

wb.save(OUT)
print("saved", OUT)
