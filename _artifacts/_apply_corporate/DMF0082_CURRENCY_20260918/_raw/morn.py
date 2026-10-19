# -*- coding: utf-8 -*-
"""오전 04_TESTDATA_SET.xlsx 승계 가능성 확인 — 단계ID 분포"""
import openpyxl, collections
ws = openpyxl.load_workbook(r"C:\Projects\NSCM\Doc\통합테스트_데이터셋_20260918\04_TESTDATA_SET.xlsx")["TESTDATA_SET"]
cols = [c.value for c in ws[1]]
i = {n: k for k, n in enumerate(cols)}
cnt = collections.Counter()
for r in ws.iter_rows(min_row=2, values_only=True):
    cnt[(r[i["프로세스단계"]], r[i["화면"]], r[i["권한그룹"]])] += 1
print("오전 115행 — 단계 × 화면 × 권한 분포")
for k, v in sorted(cnt.items()):
    print("   {:8s} {:9s} {:18s} {:3d}행".format(str(k[0]), str(k[1]), str(k[2]), v))
print("\n컬럼 {}개".format(len(cols)))
