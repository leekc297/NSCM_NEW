#!/usr/bin/env python3
"""
msg_sync.py — 클라우드 MSG 원장을 사내로 옮길 때
              (1) 명칭 충돌분을 사전 적출하고
              (2) 그 코드를 쓰는 UI/소스를 역추적해
              (3) 로컬 SOURCE 의 명칭 오류분을 보완할 수 있는 판정표를 만든다.

전제 (실측 2026-08-21)
  원장    : SCM_MGR.SCM_ADM_MSG_M  PK(MSG_ID, LANG_CD)  2,634행 / 917코드 / 3언어
  코드계  : LBL / MNU / MSG / MSG_API_* / MNU_* (5자리 고정 아님 — 정규식 주의)
  파이프라인 : DB(USE_YN='Y') -> BE /admin/message -> pnpm gen:i18n(sync-i18n.mjs)
               -> nscm-fe/public/locales/{KO,EN,CN}.json -> 화면 t('LBL00112')
  DB 소비축 : SCM_ADM_MNU_M.MNU_MSG_ID(메뉴명) / SCM_ADM_API_M.API_MSG_ID
              / SCM_ADM_LOOKUP_MASTER_DETAIL.SEGMENT2
  화면ID   : features/*/src/pages/<폴더>/*.constants.ts 의 `MENU_ID = 'XXXX'`

사용
  python msg_sync.py export-corp -o snap_corp.txt
  python msg_sync.py diff  --cloud snap_cloud.txt --corp snap_corp.txt -o diff.tsv
  python msg_sync.py usage --diff diff.tsv -o usage.tsv
  python msg_sync.py report --diff diff.tsv --usage usage.tsv -o MSG_DIFF.xlsx
  python msg_sync.py run   --cloud snap_cloud.txt -o OUTDIR     # 위 4단계 일괄

  # 클라우드 스냅샷이 아직 없을 때 — 사내 DB(권위) vs 로컬 locales 2축만 점검
  python msg_sync.py run --corp-only -o OUTDIR
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import os
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

# ---------------------------------------------------------------------------
# 경로 상수 (레포 루트 기준)
# ---------------------------------------------------------------------------
HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]                      # _artifacts/_cloud_to_corp/MSG_.. -> repo root
ORA_PY = REPO / "_artifacts" / "_diagnostics" / "ora.py"
LOCALES_DIR = REPO / "nscm-fe" / "public" / "locales"
LANGS = ("KO", "EN", "CN")

SCAN_ROOTS = [
    "nscm-fe/features",
    "nscm-fe/apps",
    "nscm-fe/packages",
    "nscm-be/src/main/java",
    "nscm-be/src/main/resources",
]
SCAN_GLOBS = ["*.ts", "*.tsx", "*.java", "*.xml"]
# 제외는 포함 뒤에 와야 실제로 바인딩된다(rg 는 나중 glob 이 우선).
# *.test.ts 는 **의도적으로 포함**한다 — 옛 명칭을 단언하는 테스트도 같이 깨지므로 보완 대상이다.
SCAN_EXCLUDE = ["!node_modules", "!dist", "!build", "!.git"]

# 코드 토큰 — 5자리 고정 아님(MSG_API_DMF0002_TITLESET_POPUP 등 실재)
CODE_TOKEN_RE = re.compile(r"\b(?:LBL|MSG|MNU)[0-9A-Za-z_]{3,}\b")
MENU_ID_RE = re.compile(r"MENU_ID\s*(?::[^=]*)?=\s*['\"]([A-Za-z0-9_]+)['\"]")

SNAPSHOT_SQL = """
SELECT M.MSG_ID, M.LANG_CD, M.MSG_NM, M.MSG_DESC, M.MSG_TP_CD, M.USE_YN
  FROM SCM_MGR.SCM_ADM_MSG_M M
 ORDER BY M.MSG_ID, M.LANG_CD
"""

# ---------------------------------------------------------------------------
# 스냅샷 이스케이프 (10_EXPORT_MSG_SNAPSHOT.sql 과 **동일 규약**)
#   \ -> \\ , CR -> \r , LF -> \n , TAB -> \t , NULL -> \N
# ---------------------------------------------------------------------------
NULL_TOKEN = r"\N"


def esc(v):
    if v is None:
        return NULL_TOKEN
    s = str(v)
    s = s.replace("\\", "\\\\").replace("\r", "\\r").replace("\n", "\\n").replace("\t", "\\t")
    return s


def unesc(s):
    if s == NULL_TOKEN:
        return None
    out, i, n = [], 0, len(s)
    while i < n:
        c = s[i]
        if c == "\\" and i + 1 < n:
            nx = s[i + 1]
            out.append({"\\": "\\", "r": "\r", "n": "\n", "t": "\t"}.get(nx, "\\" + nx))
            i += 2
        else:
            out.append(c)
            i += 1
    return "".join(out)


# ---------------------------------------------------------------------------
# Oracle 접속 — ora.py 의 conn_info() 재사용 (평문 접속정보 신규 생성 금지)
# ---------------------------------------------------------------------------
def _ora_conn():
    spec = importlib.util.spec_from_file_location("ora_helper", ORA_PY)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    import oracledb

    u, p, d = mod.conn_info()
    return oracledb.connect(user=u, password=p, dsn=d)


def cmd_export_corp(args):
    out = Path(args.out)
    with _ora_conn() as con, con.cursor() as cur:
        cur.execute(SNAPSHOT_SQL)
        rows = cur.fetchall()
    lines = [
        "#FORMAT\tMSG_ID\tLANG_CD\tMSG_NM\tMSG_DESC\tMSG_TP_CD\tUSE_YN",
        f"#META\tSIDE=CORP\tROWS={len(rows)}",
    ]
    for r in rows:
        lines.append("\t".join(esc(v) for v in r))
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"[export-corp] {out}  rows={len(rows)}")
    return 0


# ---------------------------------------------------------------------------
# 스냅샷 로드
# ---------------------------------------------------------------------------
FIELDS = ("MSG_ID", "LANG_CD", "MSG_NM", "MSG_DESC", "MSG_TP_CD", "USE_YN")


def load_snapshot(path):
    recs, meta, bad = {}, [], 0
    for raw in Path(path).read_text(encoding="utf-8").splitlines():
        if not raw.strip():
            continue
        if raw.startswith("#"):
            meta.append(raw)
            continue
        parts = raw.split("\t")
        if len(parts) != 6:
            bad += 1
            continue
        vals = [unesc(p) for p in parts]
        rec = dict(zip(FIELDS, vals))
        recs[(rec["MSG_ID"], rec["LANG_CD"])] = rec
    if bad:
        print(f"  ! 필드수 불일치로 건너뛴 행 {bad}건 — 스냅샷 포맷 확인 필요", file=sys.stderr)
    return recs, meta


def norm_name(s):
    return " ".join((s or "").split()).casefold()


# ---------------------------------------------------------------------------
# diff — 분류
# ---------------------------------------------------------------------------
DIFF_COLS = [
    "분류", "MSG_ID", "LANG_CD", "CLOUD_NM", "CORP_NM",
    "CLOUD_DESC", "CORP_DESC", "CLOUD_TP", "CORP_TP", "CLOUD_USE", "CORP_USE",
    "차이필드", "LOCALE_VAL", "LOCALE_상태", "MERGE_영향",
]


def load_locales():
    loc = {}
    for lg in LANGS:
        f = LOCALES_DIR / f"{lg}.json"
        loc[lg] = json.loads(f.read_text(encoding="utf-8")) if f.exists() else {}
    return loc


def classify(cloud, corp, locales, corp_only_mode=False):
    rows, keys = [], sorted(set(cloud) | set(corp))
    for k in keys:
        c, p = cloud.get(k), corp.get(k)
        msg_id, lang = k
        lv = locales.get(lang, {}).get(msg_id)
        lv_present = msg_id in locales.get(lang, {})

        if corp_only_mode:
            kind, diff_fields = "CORP_BASE", ""
            authority = (p or {}).get("MSG_NM")
        elif c and not p:
            kind, diff_fields, authority = "CLOUD_ONLY", "", c["MSG_NM"]
        elif p and not c:
            kind, diff_fields, authority = "CORP_ONLY", "", p["MSG_NM"]
        else:
            df = [f for f in ("MSG_NM", "MSG_DESC", "MSG_TP_CD", "USE_YN") if c[f] != p[f]]
            diff_fields = ",".join(df)
            authority = c["MSG_NM"]
            if not df:
                kind = "SAME"
            elif "MSG_NM" in df:
                kind = "NAME_DIFF"
            else:
                kind = "ATTR_DIFF"

        # locales 상태 — 권위값 대비
        if not lv_present:
            lstat = "LOCALE_없음"
        elif norm_name(lv) == norm_name(authority):
            lstat = "일치"
        else:
            lstat = "LOCALE_상이"

        # MERGE 가 사내 원장에 미치는 영향
        if kind == "CLOUD_ONLY":
            impact = "INSERT(신규 유입)"
        elif kind in ("NAME_DIFF", "ATTR_DIFF"):
            impact = "UPDATE(사내 값 덮어씀)"
        elif kind == "CORP_ONLY":
            impact = "없음(MERGE 미대상)"
        else:
            impact = "없음"

        rows.append({
            "분류": kind, "MSG_ID": msg_id, "LANG_CD": lang,
            "CLOUD_NM": (c or {}).get("MSG_NM"), "CORP_NM": (p or {}).get("MSG_NM"),
            "CLOUD_DESC": (c or {}).get("MSG_DESC"), "CORP_DESC": (p or {}).get("MSG_DESC"),
            "CLOUD_TP": (c or {}).get("MSG_TP_CD"), "CORP_TP": (p or {}).get("MSG_TP_CD"),
            "CLOUD_USE": (c or {}).get("USE_YN"), "CORP_USE": (p or {}).get("USE_YN"),
            "차이필드": diff_fields, "LOCALE_VAL": lv, "LOCALE_상태": lstat,
            "MERGE_영향": impact,
        })
    return rows


SPLIT_COLS = ["LANG_CD", "명칭", "CLOUD_CODES", "CORP_CODES", "판정"]


def detect_code_split(cloud, corp):
    """같은 명칭이 양측에서 **서로 다른 코드**로 등록된 경우 (재채번 사고 계보).
    예: 사내 MSG00656 == 클라우드 MSG01015 같은 문구."""
    out = []
    for lang in LANGS:
        cmap, pmap = defaultdict(set), defaultdict(set)
        for (mid, lg), r in cloud.items():
            if lg == lang and r["MSG_NM"] and len(r["MSG_NM"].strip()) >= 3:
                cmap[norm_name(r["MSG_NM"])].add(mid)
        for (mid, lg), r in corp.items():
            if lg == lang and r["MSG_NM"] and len(r["MSG_NM"].strip()) >= 3:
                pmap[norm_name(r["MSG_NM"])].add(mid)
        for nm in sorted(set(cmap) & set(pmap)):
            cc, pc = cmap[nm], pmap[nm]
            if cc == pc:
                continue
            verdict = "코드 완전 분리(교집합 0) — 재채번 의심" if not (cc & pc) else "부분 중복 — 코드 추가/분화"
            out.append({
                "LANG_CD": lang, "명칭": nm,
                "CLOUD_CODES": ",".join(sorted(cc)), "CORP_CODES": ",".join(sorted(pc)),
                "판정": verdict,
            })
    return out


def write_tsv(path, cols, rows):
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write("\t".join(cols) + "\n")
        for r in rows:
            f.write("\t".join(esc(r.get(c)) if r.get(c) is not None else "" for c in cols) + "\n")


def read_tsv(path):
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    cols = lines[0].split("\t")
    return [dict(zip(cols, [unesc(v) for v in ln.split("\t")])) for ln in lines[1:] if ln]


def cmd_diff(args):
    corp, _ = load_snapshot(args.corp)
    if args.cloud:
        cloud, _ = load_snapshot(args.cloud)
        corp_only_mode = False
    else:
        cloud, corp_only_mode = {}, True
    locales = load_locales()
    rows = classify(cloud, corp, locales, corp_only_mode)
    write_tsv(args.out, DIFF_COLS, rows)
    if not corp_only_mode:
        write_tsv(Path(args.out).with_name("code_split.tsv"), SPLIT_COLS,
                  detect_code_split(cloud, corp))
    tally = defaultdict(int)
    for r in rows:
        tally[r["분류"]] += 1
    print(f"[diff] {args.out}  " + " / ".join(f"{k}={v}" for k, v in sorted(tally.items())))
    lt = defaultdict(int)
    for r in rows:
        lt[r["LOCALE_상태"]] += 1
    print("       locales: " + " / ".join(f"{k}={v}" for k, v in sorted(lt.items())))
    return 0


# ---------------------------------------------------------------------------
# usage — 사용처 역추적 (FE/BE 소스 · locales · 사내 DB 메타 · 하드코딩 텍스트)
# ---------------------------------------------------------------------------
def _rg(pattern_args, roots, only_matching=True):
    cmd = ["rg", "-n", "--no-heading", "-H", "--no-messages"] + (["-o"] if only_matching else [])
    for g in SCAN_GLOBS:
        cmd += ["-g", g]
    for g in SCAN_EXCLUDE:                  # 제외가 뒤 — rg 는 나중 glob 이 우선
        cmd += ["-g", g]
    cmd += pattern_args
    cmd += [r for r in roots if (REPO / r).exists()]
    p = subprocess.run(cmd, cwd=REPO, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if p.returncode not in (0, 1):
        print(f"  ! rg 실패(rc={p.returncode}): {p.stderr[:300]}", file=sys.stderr)
    return p.stdout.splitlines()


def build_menu_id_map():
    """<pages 하위 폴더> -> MENU_ID. constants 파일에서 실측."""
    m = {}
    out = _rg(["-e", r"MENU_ID\s*(:[^=]*)?=\s*['\"][A-Za-z0-9_]+['\"]"], ["nscm-fe/features", "nscm-fe/apps"])
    for ln in out:
        try:
            path, _lno, matched = ln.split(":", 2)
        except ValueError:
            continue
        g = MENU_ID_RE.search(matched)
        if g:
            m.setdefault(Path(path).parent.as_posix(), g.group(1))
    return m


def screen_of(path, menu_map):
    p = Path(path).parent
    while True:
        key = p.as_posix()
        if key in menu_map:
            return menu_map[key]
        if p.parent == p or len(p.parts) <= 1:
            return ""
        p = p.parent


def scan_db_meta(codes):
    """사내 DB 에서 MSG_ID 를 참조하는 3축 실측."""
    sql = """
      SELECT 'MNU_M.MNU_MSG_ID' AS AXIS, MNU_MSG_ID AS CODE, MNU_ID AS REF FROM SCM_MGR.SCM_ADM_MNU_M
       WHERE MNU_MSG_ID IS NOT NULL
      UNION ALL
      SELECT 'API_M.API_MSG_ID', API_MSG_ID, API_ID FROM SCM_MGR.SCM_ADM_API_M
       WHERE API_MSG_ID IS NOT NULL
      UNION ALL
      SELECT 'LOOKUP.SEGMENT2', SEGMENT2, LOOKUP_TYPE || '/' || SEGMENT1
        FROM SCM_MGR.SCM_ADM_LOOKUP_MASTER_DETAIL
       WHERE SEGMENT2 IS NOT NULL
    """
    hits = defaultdict(list)
    try:
        with _ora_conn() as con, con.cursor() as cur:
            cur.execute(sql)
            for axis, code, ref in cur.fetchall():
                if code in codes:
                    hits[code].append(f"{axis}:{ref}")
    except Exception as e:                                   # 사내 미도달 시에도 나머지는 진행
        print(f"  ! DB 메타 스캔 생략 (사내 접속 실패: {e})", file=sys.stderr)
    return hits


USAGE_COLS = ["MSG_ID", "화면ID", "사용파일", "라인", "구분"]


def cmd_usage(args):
    diff_rows = read_tsv(args.diff)
    target = {r["MSG_ID"] for r in diff_rows
              if r["분류"] not in ("SAME",) or r["LOCALE_상태"] != "일치"}
    print(f"[usage] 역추적 대상 코드 {len(target)}개")

    menu_map = build_menu_id_map()
    print(f"        MENU_ID 매핑 폴더 {len(menu_map)}개")

    rows, seen = [], set()

    # (1) 소스 코드 리터럴
    for ln in _rg(["-e", CODE_TOKEN_RE.pattern], SCAN_ROOTS):
        try:
            path, lno, matched = ln.split(":", 2)
        except ValueError:
            continue
        code = matched.strip()
        if code not in target:
            continue
        key = (code, path, lno)
        if key in seen:
            continue
        seen.add(key)
        rows.append({"MSG_ID": code, "화면ID": screen_of(path, menu_map),
                     "사용파일": path.replace("\\", "/"), "라인": lno, "구분": "소스코드"})

    # (2) locales JSON 키 보유 여부
    loc = load_locales()
    for code in sorted(target):
        have = [lg for lg in LANGS if code in loc[lg]]
        if have:
            rows.append({"MSG_ID": code, "화면ID": "", "사용파일": "nscm-fe/public/locales",
                         "라인": "", "구분": "locales:" + ",".join(have)})

    # (3) 사내 DB 메타 3축
    for code, refs in scan_db_meta(target).items():
        for ref in refs:
            rows.append({"MSG_ID": code, "화면ID": "", "사용파일": "(사내 DB 메타)",
                         "라인": "", "구분": ref})

    # (4) 하드코딩 텍스트 — NAME_DIFF 의 '사내 명칭' 이 소스에 직접 박혀 있는지
    #
    #     ⚠ 단순 부분일치는 쓰지 않는다. 실측(2026-08-21) 결과 'Column' / 'Filter' /
    #       'Request' 같은 명칭이 **식별자**에 걸려 고유명칭 338건이 잡혔고, 그중
    #       따옴표 리터럴 안에 실제로 들어있는 것은 163건뿐이었다.
    #       -> 반드시 '문자열 리터럴로 감싸인 경우' 만 인정한다. (len>=4 도 유지)
    by_name = defaultdict(list)
    for r in diff_rows:
        if r["분류"] == "NAME_DIFF" and r.get("CORP_NM") and len(r["CORP_NM"]) >= 4:
            by_name[r["CORP_NM"]].append(r["MSG_ID"])
    if by_name:
        pf = Path(args.out).with_name("_hardcode_patterns.txt")
        pf.write_text("\n".join(sorted(by_name)), encoding="utf-8")
        for ln in _rg(["-F", "-f", str(pf)], SCAN_ROOTS, only_matching=False):
            try:
                path, lno, text = ln.split(":", 2)
            except ValueError:
                continue
            for nm, codes in by_name.items():
                if not any(q + nm + q in text for q in ("'", '"', "`")):
                    continue
                for code in codes:
                    rows.append({"MSG_ID": code, "화면ID": screen_of(path, menu_map),
                                 "사용파일": path.replace("\\", "/"), "라인": lno,
                                 "구분": "하드코딩텍스트"})

    write_tsv(args.out, USAGE_COLS, rows)
    kinds = defaultdict(int)
    for r in rows:
        kinds[r["구분"].split(":")[0]] += 1
    print(f"[usage] {args.out}  총 {len(rows)}건  " + " / ".join(f"{k}={v}" for k, v in sorted(kinds.items())))
    return 0


# ---------------------------------------------------------------------------
# report — xlsx
# ---------------------------------------------------------------------------
def cmd_report(args):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    diff_rows = read_tsv(args.diff)
    usage_rows = read_tsv(args.usage) if args.usage and Path(args.usage).exists() else []
    split_path = Path(args.diff).with_name("code_split.tsv")
    split_rows = read_tsv(split_path) if split_path.exists() else []

    use_by_code = defaultdict(list)
    scr_by_code = defaultdict(set)
    for u in usage_rows:
        if u["구분"] == "소스코드" or u["구분"] == "하드코딩텍스트":
            use_by_code[u["MSG_ID"]].append(f"{u['사용파일']}:{u['라인']}")
        if u.get("화면ID"):
            scr_by_code[u["MSG_ID"]].add(u["화면ID"])
        if u["사용파일"] == "(사내 DB 메타)":
            use_by_code[u["MSG_ID"]].append(u["구분"])

    wb = openpyxl.Workbook()
    hdr_fill = PatternFill("solid", fgColor="DDEBF7")
    hdr_font = Font(bold=True)

    def sheet(name, cols, rows):
        ws = wb.create_sheet(name)
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font = hdr_fill, hdr_font
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            w = max(len(str(c)), *(len(str(r.get(c, ""))[:60]) for r in rows)) if rows else len(str(c))
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(w + 2, 10), 60)
        return ws

    def enrich(rows):
        out = []
        for r in rows:
            d = dict(r)
            d["사용화면ID"] = ",".join(sorted(scr_by_code.get(r["MSG_ID"], []))) or "(미검출)"
            hits = use_by_code.get(r["MSG_ID"], [])
            d["사용처건수"] = len(hits)
            d["사용처(상위5)"] = " ; ".join(hits[:5])
            d["조치"] = ""
            out.append(d)
        return out

    ACT = ["사용화면ID", "사용처건수", "사용처(상위5)", "조치"]
    base = ["MSG_ID", "LANG_CD", "CLOUD_NM", "CORP_NM", "LOCALE_VAL", "LOCALE_상태", "MERGE_영향"]

    by_kind = defaultdict(list)
    for r in diff_rows:
        by_kind[r["분류"]].append(r)

    ws0 = wb.active
    ws0.title = "0_요약"
    ws0.append(["분류", "건수", "의미 / 조치 원칙"])
    for c in ws0[1]:
        c.fill, c.font = hdr_fill, hdr_font
    guide = {
        "NAME_DIFF": "⭐ 같은 코드·다른 명칭 — MERGE 가 사내 명칭을 덮어씀. 권위 결정 후 제외/적용",
        "ATTR_DIFF": "명칭 동일, DESC/TP/USE_YN 만 상이 — 대개 적용 무해",
        "CLOUD_ONLY": "사내 신규 유입 — locales 재생성 필요",
        "CORP_ONLY": "사내 단독 — MERGE 미대상. 삭제 여부는 사용자 판단",
        "SAME": "동일 — 조치 불요",
        "CORP_BASE": "사내 단독 점검 모드(클라우드 스냅샷 없음) — locales 격차만 판정",
    }
    for k in ("NAME_DIFF", "ATTR_DIFF", "CLOUD_ONLY", "CORP_ONLY", "SAME", "CORP_BASE"):
        if by_kind.get(k):
            ws0.append([k, len(by_kind[k]), guide[k]])
    lt = defaultdict(int)
    for r in diff_rows:
        lt[r["LOCALE_상태"]] += 1
    ws0.append([])
    ws0.append(["locales 상태", "건수", "의미"])
    lguide = {"일치": "로컬 소스 정상", "LOCALE_상이": "⭐ 로컬 소스 명칭 오류 — gen:i18n 재생성 대상",
              "LOCALE_없음": "locales 키 부재 — 신규 코드이거나 USE_YN='N'"}
    for k, v in sorted(lt.items()):
        ws0.append([k, v, lguide.get(k, "")])
    for i, w in enumerate((22, 10, 78), 1):
        ws0.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    sheet("1_명칭충돌_NAME_DIFF", base + ["차이필드"] + ACT, enrich(by_kind.get("NAME_DIFF", [])))
    sheet("2_코드분리_CODE_SPLIT", SPLIT_COLS, split_rows)
    sheet("3_클라우드신규_CLOUD_ONLY", base + ACT, enrich(by_kind.get("CLOUD_ONLY", [])))
    sheet("4_사내단독_CORP_ONLY", base + ACT, enrich(by_kind.get("CORP_ONLY", [])))
    sheet("5_속성차이_ATTR_DIFF", base + ["차이필드"] + ACT, enrich(by_kind.get("ATTR_DIFF", [])))
    loc_bad = [r for r in diff_rows if r["LOCALE_상태"] != "일치"]
    sheet("6_locales격차", base + ACT, enrich(loc_bad))
    sheet("7_사용처인덱스", USAGE_COLS, usage_rows)

    wb.save(args.out)
    print(f"[report] {args.out}  시트 {len(wb.sheetnames)}개")

    # 01_GEN_MSG_MERGE.sql 의 P_EXCLUDE_MSG_IDS 에 붙여넣을 **후보** 목록.
    # 전량 제외가 아니라 "검토 대상"이다 — xlsx 1번 시트의 조치 판정 후 반드시 선별할 것.
    cand = sorted({r["MSG_ID"] for r in by_kind.get("NAME_DIFF", [])})
    cf = Path(args.out).with_name("exclude_candidates.txt")
    cf.write_text(
        "-- NAME_DIFF 코드 전체 (MERGE 시 사내 명칭이 덮여쓰이는 대상)\n"
        "-- 사내 명칭을 정본으로 유지할 코드만 남기고 나머지는 지운 뒤,\n"
        "-- 01_GEN_MSG_MERGE.sql 의 P_EXCLUDE_MSG_IDS 에 아래 한 줄을 붙여넣는다.\n"
        + ",".join(cand) + "\n", encoding="utf-8")
    print(f"[report] {cf}  NAME_DIFF 코드 {len(cand)}개 (전량 제외 아님 — 선별 대상)")
    return 0


def cmd_run(args):
    outdir = Path(args.out)
    outdir.mkdir(parents=True, exist_ok=True)
    corp = outdir / "snap_corp.txt"
    cmd_export_corp(argparse.Namespace(out=str(corp)))
    dif = outdir / "diff.tsv"
    cmd_diff(argparse.Namespace(cloud=args.cloud, corp=str(corp), out=str(dif)))
    usg = outdir / "usage.tsv"
    cmd_usage(argparse.Namespace(diff=str(dif), out=str(usg)))
    rep = outdir / "MSG_DIFF_REPORT.xlsx"
    cmd_report(argparse.Namespace(diff=str(dif), usage=str(usg), out=str(rep)))
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("export-corp"); p.add_argument("-o", "--out", required=True); p.set_defaults(fn=cmd_export_corp)
    p = sub.add_parser("diff")
    p.add_argument("--cloud"); p.add_argument("--corp", required=True); p.add_argument("-o", "--out", required=True)
    p.set_defaults(fn=cmd_diff)
    p = sub.add_parser("usage"); p.add_argument("--diff", required=True); p.add_argument("-o", "--out", required=True)
    p.set_defaults(fn=cmd_usage)
    p = sub.add_parser("report")
    p.add_argument("--diff", required=True); p.add_argument("--usage"); p.add_argument("-o", "--out", required=True)
    p.set_defaults(fn=cmd_report)
    p = sub.add_parser("run")
    p.add_argument("--cloud"); p.add_argument("--corp-only", action="store_true"); p.add_argument("-o", "--out", required=True)
    p.set_defaults(fn=cmd_run)

    a = ap.parse_args()
    if a.cmd == "run" and not a.cloud and not a.corp_only:
        ap.error("run 은 --cloud <스냅샷> 또는 --corp-only 중 하나가 필요합니다.")
    sys.exit(a.fn(a))


if __name__ == "__main__":
    main()
