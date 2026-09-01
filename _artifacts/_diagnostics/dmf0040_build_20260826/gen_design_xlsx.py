# -*- coding: utf-8 -*-
"""DMF0040 화면설계서 v1.0 (8시트 — hook #105 §1-A, 정본 DMF0042 v1.1 시트 구성)"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace", line_buffering=True)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = r"C:\Projects\NSCM\Doc\화면설계서\DMF0040_MktFcstComparison_화면설계서_v1.0.xlsx"
H = Font(bold=True, size=12)
HD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="4472C4")
THIN = Border(*[Side(style="thin")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = openpyxl.Workbook()
wb.remove(wb.active)


def sheet(title, widths):
    ws = wb.create_sheet(title)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    return ws


def put(ws, row, vals, header=False, title=False, merge_to=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.alignment = WRAP
        c.border = THIN
        if header:
            c.font = HD
            c.fill = FILL
        if title:
            c.font = H
            c.border = Border()
    if merge_to:
        ws.merge_cells(start_row=row, start_column=merge_to[0], end_row=row, end_column=merge_to[1])
    return row + 1


# 1.개요
ws = sheet("1.개요", [22, 70, 30])
r = put(ws, 1, ["DMF0040 (MKT FCST Comparison) — 화면설계서 v1.0"], title=True)
for k, v in [
    ("작성", "2026-08-26 — 선행: UI분석서 v1.0(10시트)·주요이벤트/제약매트릭스/Cascade 3 md·Figma 영역 매핑표"),
    ("화면 목적", "MKT FCST 확정 버전(최대 2개) 간 Waterfall 비교 조회 — 실적 구간 회색·Difference 파생행·월>주>일 드릴다운. 읽기 전용 리포트"),
    ("메뉴 경로", "Demand > Demand Waterfall > MKT FCST Comparison (MNU_ID=DMF0040, MNU00046)"),
    ("페이지 폴더", "nscm-fe/features/demand/src/pages/mkt-fcst-comparison/ (ROUTE_PATH·MNU_URL 은 이관 마지막 단계에 mkt-fcst-comparison 으로 동기 — 4중 cross-check)"),
    ("reference baseline", "단일 grid 리포트형 — DMF0035(features/report/src/pages/demand-plan-waterfall, 동계열 waterfall 읽기 전용) 1차 + 표준 DMF0006 패턴"),
    ("편집 축", "없음 — insert/delete/save=false, 셀 전부 조회 전용. 개인화(Column/Measure 설정)만 저장(titleSet 공용 경로)"),
]:
    r = put(ws, r, [k, v, ""])

# 2.전체흐름
ws = sheet("2.전체흐름", [8, 28, 62, 30])
r = put(ws, 1, ["2. 전체흐름 — DM 프로세스 내 위치 (UI분석서 2시트 승계)"], title=True)
r = put(ws, r, ["단계", "프로세스", "설명", "본 화면"], header=True)
for row in [
    ("1", "기준정보", "Division/Site/Item/Project 마스터", ""),
    ("2", "수요예측 입력", "MKT FCST 입력 (DMF0003/DMF0082)", ""),
    ("3", "판매계획", "Sales Plan (DMF0021)", ""),
    ("4", "Cutoff", "버전(MF_D/W/M) 확정 — DM_MST_VERSION", "* 비교 대상 생성"),
    ("5", "사업계획·I/F", "Waterfall 리포트 (DMF0035/DMF0040/DMF0053)", "** 본 화면"),
    ("6", "배치·정리", "이력·정리", ""),
]:
    r = put(ws, r, list(row))

# 3.LAYOUT
ws = sheet("3.LAYOUT", [20, 80])
r = put(ws, 1, ["3. LAYOUT — 표준 셸 (#132) · Figma 영역 매핑표 승계"], title=True)
r = put(ws, r, ["영역", "구성"], header=True)
for row in [
    ("타이틀", "MainTitle = 메뉴명 (MNU00046) + 브레드크럼 (Demand > Demand Plan 아님 — Demand Waterfall)"),
    ("SearchArea", "1단: Division*(단일콤보) / Team(다중) / End Customer(다중) / Sales Site(다중) — 36px (#94). 2단: Biz.Type(다중) / Plant Site(다중) / Period*(주차표기 기간) / Bucket(단일). extraConditions(접이): Project(+검색모달) / Model·Nickname(+검색모달) / Item Type(다중) / Owner(+검색모달) / Version*(readOnly+모달) / QTY(단일) / Sum(체크) / Diff.Only(체크)"),
    ("SubTitle", "총 N 건 + 우측: Excel Form 없음 — [Excel Download] [Sort(F3)] [Filter(F4)] [Condition(F6)] [Reset Condition(F10)] + Column 설정·Measure 설정 아이콘 (hook #120 순서: ColumnSet→Measure). Insert/Delete/Save 미노출"),
    ("Grid", "공용 DataGrid 단일 시트 — 좌측 정적 컬럼(frozen·세로 머지) + 동적 3단 헤더 (월>주>일+SUM, 48/24/16px #93) + 컬럼 footer 합계 + 셀 배경(실적 Gainsboro)·글자색(Difference Blue·SAT Blue/SUN Red 헤더)"),
    ("모달", "VersionByWaterfall(전용) / Column·Measure 설정(master 공용) / Project·Item·Owner 검색(공용)"),
    ("금지", "최상위 래퍼 div ⛔·DataGrid 감싸는 div ⛔ (#132) / 로컬 scss 폼 레이아웃 ⛔ (#121)"),
]:
    r = put(ws, r, list(row))

# 4.조회조건·버튼
ws = sheet("4.조회조건·버튼", [18, 16, 26, 40, 40])
r = put(ws, 1, ["4. 조회조건 · 버튼 (AS-IS Designer 실측 22종 — Cascade_Filter_전수.md 승계)"], title=True)
r = put(ws, r, ["항목", "타입", "기본값", "원천(CMBO/SP)", "동작"], header=True)
for row in [
    ("Division*", "단일 콤보", "목록 첫 실값", "MyDivision 계열 공용 CMBO (@sesUserId)", "변경 시: 하위 5콤보 재적재+텍스트 4종 클리어+Grid 초기화+기본 Version 재세팅+동적 헤더 재질의. 40=Project·Plant Site disabled"),
    ("Team", "다중 체크", "ALL", "UserTeamByDivision", "변경 시 Customer·SalesSite 재적재 (81계열·82/40 파라미터 분기)"),
    ("End Customer", "다중 체크", "ALL", "EndCustomerByDivision", "변경 시 SalesSite 재적재"),
    ("Biz Type", "다중 체크", "ALL", "BizType", ""),
    ("Sales Site", "다중 체크", "ALL", "UserSiteByEndCustomer", ""),
    ("Plant Site", "다중 체크", "ALL", "UserPlantSiteByDivision", "40 = disabled"),
    ("Project", "텍스트+검색", "빈값", "ProjectByDivisionModal(공용)", "40 = disabled"),
    ("Model/Nickname", "텍스트+검색", "빈값", "ItemByDivisionModal(공용)", ""),
    ("Item Type", "다중 체크", "ALL", "LOOKUP ITEM_TYPE_CODE", ""),
    ("Owner", "텍스트+검색", "빈값", "Owner 모달(공용)", ""),
    ("Version*", "readOnly+모달", "최근 MF_W 2개 자동", "SCM0002_GET_VERSIONBYWATERFALLPOPUP", "클릭/키 = 모달. 최대 2개(MSG00255)"),
    ("Period*", "기간(주차표기)", "당월1일~6개월 말일 (PR_GET_DAY_RANGE)", "", "역전=MSG00108 / 2년 초과=MSG00164+자동보정"),
    ("Bucket", "단일 콤보", "MONTHLY", "DMF0003_GET_REPORT_BUCKET_TYPE_CODE", "D_/W_/M_ 가시성 + ONLY_WEEKLY 헤더 전환"),
    ("QTY", "단일 콤보", "원단위", "DMF0003_GET_QTY_TYPE_CODE", "표시 스케일 ÷1/1e3/1e6/1e8"),
    ("Sum", "체크", "false", "", "상위 집계 컬럼 표시 토글"),
    ("Diff.Only", "체크", "false", "", "p_diff_only 파라미터"),
    ("(제외) Cell Qty", "체크", "", "", "AS-IS 항상 숨김(layoutControlItem10 Never) — 미구현"),
]:
    r = put(ws, r, list(row))
r = put(ws, r + 1, ["버튼", "", "", "", ""], header=True)
for row in [
    ("Search (F5)", "hotkey", "", "PKG_DMF0040 GET (18p)", "V1~V3 검증 후 조회"),
    ("Excel Download", "hotkey/버튼", "", "FE (공용 DataGrid Excel)", "콤보 Name·HideZero·소수3자리"),
    ("Sort/Filter/Condition/Reset", "기본 활성", "", "", "hotkeyConfig 기본 (§0 필수 2)"),
    ("Column 설정 / Measure 설정", "아이콘 버튼", "", "master 공용 모달 (titleSet)", "저장 시 그리드 Reset (UC-03)"),
    ("Insert/Delete/Save", "미노출", "", "", "AS-IS 빈 구현 — insert/delete/save=false"),
]:
    r = put(ws, r, list(row))

# 5.GRID
ws = sheet("5.GRID", [24, 76])
r = put(ws, 1, ["5. GRID — 단일 시트 'Mkt Fcst Waterfall'"], title=True)
r = put(ws, r, ["축", "설계"], header=True)
for row in [
    ("정적 컬럼", "SCM_MST_SHEET_COLUMNS PROGRAM_ID='DMF0040' (48행 실측 — *40 접미 쌍 SORT 충돌 해소 후) — 서버 직접 비교(#113, PR_GET_SHEET_COLUMN_LIST + divisionId) · frozen + 세로 머지(SmartColMerge 대응 = CustomMergeManager)"),
    ("동적 컬럼", "spreadTitle SP (5p) — 월(M_)>주(W_)>일(D_) 3단 헤더 + 그룹 SUM 컬럼 + B_ 접두 숨김 + SAT/SUN 헤더 글자색"),
    ("드릴다운", "M_ 헤더 더블클릭 = 그 월 W_ 토글·열린 D_ 접기 / W_ 더블클릭 = 그 주 D_ 토글 (AS-IS CellDoubleClick L1179~1216)"),
    ("Bucket 가시성", "DAILY: D_ 표시·W_/M_=Sum 체크 / WEEKLY: W_ 표시·M_=Sum / MONTHLY: M_ 만 / ONLY_WEEKLY: W_ 단순 3단 (Sum 비활성)"),
    ("행 구성", "버전(≤2) × Measure — MEASURE_GROUP/MEASURE + Difference 파생행 (SP 산출)"),
    ("셀 스타일", "실적 구간 = 행별 Version 경계까지 배경 Gainsboro (해당 월 SUM 은 White) / Difference 행 데이터 글자 Blue / 전 셀 조회 전용 (Gray 잠금색 아님 — 리포트 White+zebra 표준 #95 유지, 실적 회색만 오버레이)"),
    ("Footer", "컬럼 footer 'Sum' — D_/W_/M_ SumIgnoreHidden 집계"),
    ("QTY 스케일", "표시 값 = 원본 ÷ (1/1e3/1e6/1e8) 소수 2자리 — 원본 보존 후 전환식 (FE 표시 변환)"),
    ("No. 컬럼", "코너 헤더 No. 표준 (#108)"),
]:
    r = put(ws, r, list(row))

# 6.EVENT (선행 MD 승계 — §1-D)
ws = sheet("6.EVENT", [10, 30, 70])
r = put(ws, 1, ["6. EVENT — 목차 + 선행 MD 전문 승계 (#105 §1-D). 정본 = Doc/화면분석서/DMF0040_MktFcstComparison_주요이벤트_상세분석.md"], title=True)
r = put(ws, r, ["EID", "이벤트", "요지 (상세 = 선행 MD)"], header=True)
events = [
    ("E01", "Form_Load", "초기화 체인 — 콤보 전부 적재·기본값·동적 헤더·Division cascade 1회"),
    ("E02", "Division 변경", "헤더 재질의+cascade 5콤보+40 disabled+텍스트 클리어+기본 Version(MF_W 최근 2)"),
    ("E03", "Team 변경", "Customer·SalesSite 재적재 (분기)"),
    ("E04", "End Customer 변경", "SalesSite 재적재"),
    ("E05~E08", "SalesSite/PlantSite/Project/ItemId 변경", "SetClbVersion 사문 — TO-BE 제외"),
    ("E09", "Bucket 변경", "ONLY_WEEKLY 경계 헤더 재질의 + 컬럼 가시성"),
    ("E10", "그리드 콤보 변경", "제외 — Insert 전제 사문"),
    ("E11", "헤더 더블클릭", "M_/W_ 드릴다운 토글 (셀 더블클릭 분기는 사문)"),
    ("E12", "스크롤", "세로 머지 유지"),
    ("E13", "Ctrl+C", "클립보드 복사 (붙여넣기는 잠금이라 무효)"),
    ("E14~E15", "Project/Model 검색", "공용 모달 → 텍스트 반영"),
    ("E16", "Column 설정", "master 공용 Column 모달 (PROGRAM_ID=DMF0040) → 헤더 재구성"),
    ("E17", "Measure 설정", "master 공용 Measure 모달 — AS-IS PROGRAM_ID=DMF0035 공유 (시딩 방침 따름)"),
    ("E18", "Sum 체크", "상위 집계 컬럼 토글"),
    ("E19~E20", "Version 버튼/키", "Version 모달 (최대 2·기간 승계)"),
    ("E21", "Owner 검색", "공용 모달"),
    ("E22", "QTY 변경", "표시 스케일 전환 (원본 보존)"),
    ("FW", "Find/Excel", "조회 V1~V3 + 실적 회색·Difference 파랑 / Excel Download"),
    ("FW-X", "Insert/Delete/Edit/Save", "빈 구현·MSG00109 — 미노출"),
]
for row in events:
    r = put(ws, r, list(row))

# 7.VALIDATION
ws = sheet("7.VALIDATION", [10, 44, 14, 50])
r = put(ws, 1, ["7. VALIDATION — 제약매트릭스_상세.md 승계 (20행: A3+D10+G4+P3)"], title=True)
r = put(ws, r, ["ID", "조건·동작", "MSG", "TO-BE"], header=True)
for row in [
    ("A-01", "Division 필수", "MSG00097", "handleSubmit V1"),
    ("A-02", "기간 역전 차단", "MSG00108", "V2"),
    ("A-03", "2년 초과 — 알림+종료일 자동 보정", "MSG00164", "V3 (문구 3년 vs 로직 2년 — 로직 파리티, 문구 정정 별도)"),
    ("D-01~08", "Division cascade 가드·파라미터 분기·40 disabled", "", "SearchArea cascade + CMBO 분기"),
    ("D-09~10", "그리드 편집 전제 분기", "", "제외 (읽기 전용 사문)"),
    ("G-01", "Version 직접 타이핑 차단", "", "readOnly + 모달"),
    ("G-02~04", "실적 경계 역산·QTY 가드", "", "util"),
    ("P-01", "편집 시도 읽기전용 통보", "MSG00109", "편집 진입점 미노출 — 제외"),
    ("P-02~03", "버전 3개 이상 차단 (반환부+모달 내부)", "MSG00255", "Version 모달 이중 차단"),
]:
    r = put(ws, r, list(row))

# 8.API·SP
ws = sheet("8.API·SP", [30, 34, 60])
r = put(ws, 1, ["8. API · SP — 배선 계약 (positional — proc-meta-signature-match)"], title=True)
r = put(ws, r, ["FE 배선 (sheet:BTN)", "SP", "계약"], header=True)
for row in [
    ("Mkt Fcst Waterfall:SEARCH", "PKG_DMF0040.P_DMF0040_GET_MKT_FCST_COMPARISON", "18 IN (division/site/plant/team/customer/bizType/project/item/from/to/createdBy/version/itemType/owner/diffOnly('True'|'False' 문자열)/bucket/userId/menuId) + O_RC1(메인)/O_RESULT_CURSOR — PROC_M positional 등재·API_M 신규"),
    ("spreadTitle:SEARCH", "DMF0040_GET_SPREAD_TITLE (PKG 편입 검토)", "5 IN (start/end/createdBy/division/bucket) + 커서 — 동적 헤더"),
    ("version:SEARCH", "SCM0002_GET_VERSIONBYWATERFALLPOPUP (TO-BE 실측 재사용)", "6 IN (start/end/createdBy/division/versionType/versionNo)"),
    ("titleSet:SEARCH/SAVE", "공용 titleSet 고정 매핑 (BE MenuServiceImpl)", "Column=PROGRAM_ID DMF0040 / Measure=시딩 방침 축"),
    ("콤보 (CMBO)", "Division/Team/EndCustomer/BizType/SalesSite/PlantSite/ItemType/Bucket/QTY — 기존 공용 실측 재사용, 미등록만 신규", "@sesUserId·@divisionId 파라미터 계약 (#117)"),
    ("공용 모달", "Project/Item/Owner — master 공용 모달 자체 배선", "본 화면 배선 대상 아님"),
    ("저장 SP", "없음 (읽기 전용)", "E2E F축 no-save·G6-2/3 해당없음"),
]:
    r = put(ws, r, list(row))

wb.save(OUT)
print("저장:", OUT)
wb2 = openpyxl.load_workbook(OUT, read_only=True)
assert wb2.sheetnames == ["1.개요", "2.전체흐름", "3.LAYOUT", "4.조회조건·버튼", "5.GRID", "6.EVENT", "7.VALIDATION", "8.API·SP"], wb2.sheetnames
print("시트 8 OK:", wb2.sheetnames)
wb2.close()
