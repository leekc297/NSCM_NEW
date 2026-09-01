# -*- coding: utf-8 -*-
"""DMF0040 TestCases (TestCases4 양식 — 34열·division 시트 분리·권한별 행·DATA 시트)"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace", line_buffering=True)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = r"C:\Projects\NSCM\Doc\단위테스트\DMF0040_MktFcstComparison_TestCases.xlsx"
HD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="4472C4")
THIN = Border(*[Side(style="thin")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

HEADERS = ["No", "테스트ID", "시나리오 ID", "", "", "시나리오", "구분", "타임라인", "스케줄\n(주기)", "단계",
           "버전타입", "화면 ID", "테스트명", "점검\n구분", "테스트데이터", "예상결과", "As-IS Image ",
           "TGT 테이블(R/W)", "TGT Measure(R/W)", "SRC 테이블(READ)", "SRC Measure/원천", "설명", "삭제/복사구간",
           "비고", "실행절차", "JOB (Full Name)", "프로시저 (Full Name)", "건수\n(As-Is/MSSQL)", "검증방법\n(As-Is/MSSQL)",
           "건수\n(To-Be/Oracle)", "검증방법\n(To-Be/Oracle)", "결과\n(P/F)", "실행일", "수행자"]

MAIN_SP = "PKG_DMF0040.P_DMF0040_GET_MKT_FCST_COMPARISON"
TITLE_SP = "DMF0040_GET_SPREAD_TITLE"
VER_SP = "PKG_SCM0002.P_SCM0002_GET_VERSIONBYWATERFALLPOPUP"
SRC = "DM_PLN_MKT_FCST · DM_PLN_MKT_FCST_C · DM_PLN_MKT_FCST_RESULT(_C) · DM_MST_VERSION"
MEAS = "Version(A) · Version(B) · Difference"

# HTTP E2E 실측 rows (2026-08-27, versionNo=Current MONTHLY 20260701~20261231)
ROWS = {"40": "774", "81": "544", "82": "558"}
ACCTS = {
    "40": [("primary/ADMIN축", "minyonglee1"), ("DM ADMIN", "hongjae.kim"), ("DM Viewer", "yooah.chung"), ("SCP Planner", "dodae"), ("DM Planner(무권한)", "eunah.koo")],
    "81": [("primary/ADMIN축", "minyonglee1"), ("DM ADMIN", "jiwon.park"), ("DM Viewer", "park6151"), ("SCP Planner", "bottle"), ("DM Planner(무권한)", "bpark416")],
    "82": [("primary/ADMIN축", "minyonglee1"), ("DM ADMIN", "wonqthanq"), ("DM Viewer", "yeonjusung"), ("SCP Planner", "coeogus"), ("DM Planner(무권한)", "ys_cho")],
}

def cases(div):
    d40 = div == "40"
    rows = []
    def c(step, name, kind, data, expect, proc=MAIN_SP, tobe_cnt="", tobe_how="화면 표시 + BE 응답 rows 대조", note=""):
        rows.append([step, name, kind, data, expect, proc, tobe_cnt, tobe_how, note])
    c("조회", "화면 진입·메뉴", "(정상)", f"Div {div} / Demand > Demand Waterfall > MKT FCST Comparison",
      "새 탭 오픈·조회조건 표시. Insert/Delete/Save 버튼 미노출 (읽기 전용)", "", "", "버튼 3종 미표시 육안")
    c("조회", "Division 변경 cascade", "(정상)", f"Division {div} 선택",
      ("Team/End Customer/Biz Type/Sales Site/Plant Site 재적재 + Project·Model·Owner·Version 클리어 + 기본 Version(최근 MF_W 2) 자동 세팅"
       + (" / Project·Plant Site 비활성(40 전용)" if d40 else " / Project·Plant Site 활성")), VER_SP)
    c("조회", "기본 Version 자동 세팅", "(정상)", "Division 선택 직후 Version 칸",
      "'Current' 제외 최근 MF_W 2개가 콤마로 자동 입력", VER_SP)
    c("조회", "필수값 — Division", "(오류)", "Division 미선택 상태 Search", "MSG00097 'Division 을 선택하세요.' 차단", "")
    c("조회", "기간 역전", "(오류)", "From > To 로 설정 후 Search", "MSG00108 차단", "")
    c("조회", "기간 2년 초과", "(오류)", "From +2년 넘는 To 설정 후 Search",
      "MSG00164 알림 + To 가 From+2년으로 자동 보정 후 중단 (문구는 '3년' — 로직 2년 = AS-IS 파리티, 제약 A-03 비고)", "")
    c("조회", "버전 모달 — 조회·필터", "(정상)", "Version 칸 클릭(또는 아무 키) → 모달 / Day·Week·Manual 체크 조합·Description 필터",
      "MF_D/MF_W/MF_M 조합으로 버전 목록 조회 (표시 5컬럼: Version/Description/Type/Create By/Create Date)", VER_SP)
    c("조회", "버전 모달 — 최대 2 제한", "(오류)", "3개 이상 선택 후 OK", "MSG00255 '버전은 최대 2개까지…' 차단", VER_SP)
    c("조회", "버전 모달 — 선택 반영", "(정상)", "2개 선택 OK (또는 더블클릭 1개)",
      "Version 칸 콤마 반영 + 모달 조회기간이 Period 에 반영 (AS-IS sPeriod 파리티)", VER_SP)
    c("조회", "메인 조회", "(정상)", f"Div {div} / Version=Current / 2026.07.01~2026.12.31 / MONTHLY / Search",
      f"버전×Measure 행 표시 — MEASURE = {MEAS}. 총건수 표기", MAIN_SP, ROWS[div], "HTTP E2E 실측 rows 와 대조 (DMF0040_http_e2e_20260827_0038.md)")
    c("조회", "동적 헤더 3단", "(정상)", "조회 결과 헤더", "월>주>일 구조 + 그룹 SUM. SAT 파랑/SUN 빨강 (Daily 표시 시)", TITLE_SP)
    c("조회", "실적 구간 회색", "(정상)", "Current 행 vs Cutoff 버전 행",
      "행별 버전 기준 경계까지 배경 회색(Gainsboro) — 해당 월 SUM 은 White 예외", "")
    c("조회", "Difference 행 파랑", "(정상)", "MEASURE=Difference 행", "기간 컬럼 글자 파랑", "")
    c("조회", "드릴다운 — 월 헤더", "(정상)", "M_ 헤더 더블클릭", "그 월의 주(W_) 토글·열린 일(D_) 접기", "")
    c("조회", "드릴다운 — 주 헤더", "(정상)", "W_ 헤더 더블클릭", "그 주의 일(D_) 토글", "")
    c("조회", "Bucket 전환", "(정상)", "MONTHLY→WEEKLY→DAILY→ONLY_WEEKLY",
      "D_/W_/M_ 가시성 전환. ONLY_WEEKLY 는 Sum 체크 비활성·해제", TITLE_SP)
    c("조회", "Sum 체크", "(정상)", "DAILY/WEEKLY 에서 Sum 체크", "상위 집계 컬럼(W_/M_) 표시 토글", "")
    c("조회", "QTY 스케일", "(정상)", "QTY=1,000 선택", "표시값 ÷1,000 (소수 2자리·0 제외) + footer 합계 정합. 1 로 되돌리면 원값", "")
    c("조회", "Diff.Only", "(정상)", "Diff.Only 체크 후 Search", "p_diff_only=True 로 조회 (차이 존재 행만)", MAIN_SP)
    c("조회", "Column 설정", "(정상)", "Column Setting 아이콘 → 표시/순서 변경 저장",
      "titleSet 'C' 저장 + 기존 조회 데이터 Reset(그리드 비움 — UC-03) → Search 재조회", "SCM0002_SET_TITLESET_POPUP")
    c("조회", "Measure 설정", "(정상)", "Measure 아이콘 → 3종(Version(A)/(B)/Difference) 토글 저장",
      "titleSet 'M' 저장 + 데이터 Reset (UC-03)", "SCM0002_SET_TITLESET_POPUP")
    c("조회", "필터 유지 재조회", "(정상)", "그리드 필터 적용 상태 Search", "재조회 후 필터 유지 (UC-07 — 공용 rebindFilterColumns)", MAIN_SP)
    c("조회", "Excel Download", "(정상)", "Excel 버튼", "표시 컬럼 기준 다운로드 (UC-05)", "")
    for grp, acct in ACCTS[div]:
        deny = "무권한" in grp
        c("권한", f"권한별 접근 — {grp}", "(정상)" if not deny else "(오류)", f"{acct} 로그인 → Div {div} 조회",
          ("SEARCH 403 (메뉴권한 미등재 — AS-IS 이관 4그룹 정합. 개방 여부는 운영 결정)" if deny
           else "조회 성공 (读기 전용 — 편집 진입점 없음). SAVE 강제 호출 시 403"), MAIN_SP, "" if deny else ROWS[div],
          "HTTP E2E H축 실측 대조")
    return rows

wb = openpyxl.Workbook()
wb.remove(wb.active)
for div in ("40", "81", "82"):
    ws = wb.create_sheet(div)
    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HD; cell.fill = FILL; cell.border = THIN; cell.alignment = WRAP
    widths = {1: 5, 6: 22, 7: 6, 12: 10, 13: 26, 14: 8, 15: 34, 16: 44, 18: 26, 19: 20, 20: 26, 27: 34, 30: 12, 31: 30}
    for col, w in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    r = 2
    for i, (step, name, kind, data, expect, proc, cnt, how, note) in enumerate(cases(div), 1):
        vals = {1: i, 6: "UI 상세 (MKT FCST Comparison)", 7: "UI", 8: "주중", 10: step, 12: "DMF0040",
                13: f"MKT FCST Comparison — {name}", 14: kind, 15: data, 16: expect,
                18: "(조회 전용 — Write 없음)", 19: MEAS, 20: SRC, 21: "MKT FCST 버전 원장 + 실적",
                24: note, 27: proc, 30: cnt, 31: how}
        for col, v in vals.items():
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = THIN; cell.alignment = WRAP
        r += 1
    # DATA 시트
    wd = wb.create_sheet(f"{div}_DATA")
    for i, h in enumerate(["No", "데이터ID", "구분", "Division", "권한그룹", "테스트계정", "EDIT_YN", "대상시트·팝업"], 1):
        cell = wd.cell(row=1, column=i, value=h)
        cell.font = HD; cell.fill = FILL; cell.border = THIN
    for i, (grp, acct) in enumerate(ACCTS[div], 1):
        for j, v in enumerate([i, f"D{div}-{i:02d}", "계정", div, grp, acct, "N (조회 전용)", "mktFcstComparison / VersionModal / titleSet"], 1):
            cell = wd.cell(row=i + 1, column=j, value=v)
            cell.border = THIN
wb.save(OUT)
print("저장:", OUT)
wb2 = openpyxl.load_workbook(OUT, read_only=True)
counts = {ws.title: ws.max_row - 1 for ws in wb2.worksheets}
wb2.close()
print("시트·행:", counts)
