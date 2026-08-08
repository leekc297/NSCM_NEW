#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""asis.py — AS-IS(GSCM MSSQL, LGCE_SCMDEV) one-shot query CLI (읽기 전용).
접속정보: 프로젝트 루트 .env.asis 의 ASIS_DB_* (pyodbc + ODBC Driver 17).
사용:
  python asis.py "SELECT TOP 10 name FROM sys.tables"
  python asis.py tables MEASURE          # 이름에 MEASURE 포함 테이블
  python asis.py procs MKT_FCST          # 이름에 패턴 포함 프로시저
  python asis.py cols TABLE_NAME         # 컬럼 인벤토리
  python asis.py proc  PROC_NAME         # 프로시저 본문(sys.sql_modules)
  python asis.py xlsx out.xlsx "SQL..."  # 쿼리 결과를 xlsx 로 저장
"""
import sys, os

def _conn():
    import pyodbc
    env = {}
    p = os.path.join(os.path.dirname(__file__), "..", "..", ".env.asis")
    p = os.path.abspath(p)
    for ln in open(p, encoding="utf-8-sig"):
        ln = ln.strip()
        if ln and not ln.startswith("#") and "=" in ln:
            k, v = ln.split("=", 1); env[k.strip()] = v.strip()
    cs = (f"DRIVER={{{env.get('ASIS_DB_DRIVER','ODBC Driver 17 for SQL Server')}}};"
          f"SERVER={env['ASIS_DB_HOST']},{env['ASIS_DB_PORT']};DATABASE={env['ASIS_DB_NAME']};"
          f"UID={env['ASIS_DB_USER']};PWD={env['ASIS_DB_PASS']};TrustServerCertificate=yes;Encrypt=no")
    return pyodbc.connect(cs, timeout=15)

def _rows(cur):
    cols = [d[0] for d in cur.description]
    return cols, cur.fetchall()

def _print(cols, rows, limit=None):
    """전량 출력이 기본. ⚠️ 2026-08-08 이전에는 limit=200 으로 **말없이 잘라** 출력하면서
    `(N rows)` 는 전체 건수를 찍어, 200행 넘는 조회에서 조용한 누락이 발생할 수 있었다.
    절단이 필요하면 limit 을 명시하고, 잘린 사실을 반드시 고지한다."""
    print(" | ".join(cols))
    print("-+-".join("-" * len(c) for c in cols))
    shown = rows if limit is None else rows[:limit]
    for r in shown:
        print(" | ".join("" if v is None else str(v) for v in r))
    cut = len(rows) - len(shown)
    print(f"({len(rows)} rows" + (f" — {len(shown)} 표시 · ⚠ {cut} 행 생략" if cut else "") + ")")

def main():
    a = sys.argv[1:]
    if not a:
        print(__doc__); return
    cn = _conn(); cur = cn.cursor()
    cmd = a[0]
    if cmd == "tables":
        pat = a[1] if len(a) > 1 else ""
        cur.execute("SELECT s.name+'.'+t.name AS tbl FROM sys.tables t JOIN sys.schemas s ON t.schema_id=s.schema_id WHERE t.name LIKE ? ORDER BY tbl", f"%{pat}%")
        _print(*_rows(cur))
    elif cmd == "procs":
        pat = a[1] if len(a) > 1 else ""
        cur.execute("SELECT s.name+'.'+o.name AS objname, o.type_desc FROM sys.objects o JOIN sys.schemas s ON o.schema_id=s.schema_id WHERE o.type IN ('P','FN','TF','IF') AND o.name LIKE ? ORDER BY objname", f"%{pat}%")
        _print(*_rows(cur))
    elif cmd == "cols":
        cur.execute("SELECT c.name, ty.name AS datatype, c.max_length, c.is_nullable FROM sys.columns c JOIN sys.types ty ON c.user_type_id=ty.user_type_id WHERE c.object_id=OBJECT_ID(?) ORDER BY c.column_id", a[1])
        _print(*_rows(cur))
    elif cmd == "proc":
        cur.execute("SELECT m.definition FROM sys.sql_modules m WHERE m.object_id=OBJECT_ID(?)", a[1])
        row = cur.fetchone()
        print(row[0] if row else "(not found)")
    elif cmd == "xlsx":
        out, sql = a[1], a[2]
        cur.execute(sql); cols, rows = _rows(cur)
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.append(cols)
        for r in rows: ws.append(["" if v is None else v for v in r])
        wb.save(out); print(f"[xlsx] {out} — {len(rows)} rows, {len(cols)} cols")
    else:
        cur.execute(cmd);
        if cur.description: _print(*_rows(cur))
        else: print("OK (no result set)")
    cn.close()

if __name__ == "__main__":
    main()
