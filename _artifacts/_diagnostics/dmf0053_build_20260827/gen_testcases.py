# -*- coding: utf-8 -*-
"""DMF0053 TestCases (TestCases4 양식 — 34열·division 시트 분리·권한별 행·DATA 시트, DMF0040 생성기 계보)"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace", line_buffering=True)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = r"C:\Projects\NSCM\Doc\단위테스트\DMF0053_CustFcstWaterfall_TestCases.xlsx"
HD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="4472C4")
THIN = Border(*[Side(style="thin")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

HEADERS = ["No", "테스트ID", "시나리오 ID", "", "", "시나리오", "구분", "타임라인", "스케줄\n(주기)", "단계",
           "버전타입", "화면 ID", "테스트명", "점검\n구분", "테스트데이터", "예상결과", "As-IS Image ",
           "TGT 테이블(R/W)", "TGT Measure(R/W)", "SRC 테이블(READ)", "SRC Measure/원천", "설명", "삭제/복사구간",
           "비고", "실행절차", "JOB (Full Name)", "프로시저 (Full Name)", "건수\n(As-Is/MSSQL)", "검증방법\n(As-Is/MSSQL)",
           "건수\n(To-Be/Oracle)", "검증방법\n(To-Be/Oracle)", "결과\n(P/F)", "실행일", "수행자"]

MAIN_SP = "PKG_DMF0053.P_DMF0053_GET_CUST_FCST_WATERFALL"
TITLE_SP = "PKG_DMF0053.P_DMF0053_GET_SPREAD_TITLE"
SRC = "DM_PLN_CUST_FCST_H · DM_DNM_EDI_CALLOFF_ITEM · DM_MST_CUSTOMER_FCST · SCM_MST_TITLESET_USER"
MEAS = "Cust. FCST(CUST_FCST) · EDI(CUST_FCST_EDI 하위 7종)"

# 데이터 실측 (2026-08-27): EDI 원장 = 81 계열 전용 — 40/82 는 0행이 데이터 정상.
ROWS = {"40": "0 (EDI 원장 없음 — 데이터 실측 정상)", "81": "(E2E 리포트 rows 대조)", "82": "0 (EDI 원장 없음)"}
ACCTS = {
    "40": [("primary/ADMIN축", "minyonglee1"), ("DM ADMIN", "hongjae.kim"), ("DM Planner", "eunah.koo"), ("DM Viewer", "yooah.chung"), ("SCP Planner(무권한)", "dodae")],
    "81": [("primary/ADMIN축", "minyonglee1"), ("DM ADMIN", "jiwon.park"), ("DM Planner", "bpark416"), ("DM Viewer", "park6151"), ("SCP Planner(무권한)", "bottle")],
    "82": [("primary/ADMIN축", "minyonglee1"), ("DM ADMIN", "wonqthanq"), ("DM Planner", "ys_cho"), ("DM Viewer", "yeonjusung"), ("SCP Planner(무권한)", "coeogus")],
}

def cases(div):
    is81 = div == "81"
    rows = []
    def c(step, name, kind, data, expect, proc=MAIN_SP, tobe_cnt="", tobe_how="화면 표시 + BE 응답 rows 대조", note=""):
        rows.append([step, name, kind, data, expect, proc, tobe_cnt, tobe_how, note])
    c("조회", "화면 진입·메뉴", "(정상)", f"Div {div} / Demand > Waterfall > Customer FCST Waterfall",
      "새 탭 오픈·조회조건 표시. Insert/Delete/Save 버튼 미노출 (읽기 전용)", "", "", "버튼 3종 미표시 육안")
    c("조회", "Division 변경 cascade", "(정상)", f"Division {div} 선택",
      ("Team/End Customer/Biz Type/Sales Site/Plant Site 재적재 + PJT·Model·Owner 클리어 + Grid 초기화"
       + (" / PJT·Plant Site 활성(81 계열·82)" if div != "40" else " / PJT·Plant Site 비활성(40 전용)")), "")
    c("조회", "필수값 — Division", "(오류)", "Division 미선택 상태 Search", "MSG00097 'Division 을 선택하세요.' 차단", "")
    c("조회", "Period 역전", "(오류)", "Period From > To 설정 후 Search", "MSG00108 차단", "")
    c("조회", "Period 3년 초과", "(오류)", "From +3년 넘는 To 설정 후 Search",
      "MSG00164 알림 + To 가 From+3년으로 자동 보정 후 중단 (문구·로직 3년 정합)", "")
    c("조회", "Input Date 역전", "(오류)", "Input Date From > To 설정 후 Search", "MSG00108 차단", "")
    c("조회", "Input Date 3년 초과", "(오류)", "Input Date +3년 초과 설정 후 Search", "MSG00164 알림 + 보정 후 중단", "")
    c("조회", "메인 조회", "(정상)", f"Div {div} / Period 2026.01.01~06.30 / Input Date 2026.01.01~03.31 / MONTHLY / Search",
      f"Measure 행 표시 — {MEAS}. 총건수 표기" + ("" if is81 else ". 0행 = EDI 원장 부재(데이터 정상 — 결함 아님)"),
      MAIN_SP, ROWS[div], "HTTP E2E 실측 rows 와 대조")
    c("조회", "동적 헤더 3단", "(정상)", "조회 결과 헤더", "월>주>일 구조 + 그룹 SUM. SAT 파랑/SUN 빨강 (Daily 표시 시)", TITLE_SP)
    c("조회", "드릴다운 — 월 헤더", "(정상)", "M_ 헤더 더블클릭", "그 월의 주(W_) 토글·열린 일(D_) 접기", "")
    c("조회", "드릴다운 — 주 헤더", "(정상)", "W_ 헤더 더블클릭", "그 주의 일(D_) 토글", "")
    c("조회", "Bucket 전환", "(정상)", "MONTHLY→WEEKLY→DAILY→ONLY_WEEKLY",
      "D_/W_/M_ 가시성 전환 + 타이틀 재조회. ONLY_WEEKLY 는 Sum 체크 비활성·해제", TITLE_SP)
    c("조회", "Sum 체크", "(정상)", "DAILY/WEEKLY 에서 Sum 체크", "상위 집계 컬럼(W_/M_) 표시 토글", "")
    c("조회", "QTY 스케일", "(정상)", "QTY=1,000 선택", "표시값 ÷1,000 (소수 2자리·0 제외) + footer 합계 정합. 1 로 되돌리면 원값", "")
    c("조회", "PJT 검색 모달", "(정상)" if div != "40" else "(제약)", "PJT 검색 아이콘",
      ("ProjectByDivisionModal 오픈 → 선택 반영" if div != "40" else "40 은 PJT 칸 비활성 (AS-IS B-04)"), "")
    c("조회", "Model/Nickname 검색 모달", "(정상)", "Model 검색 아이콘", "ItemByDivisionModal 오픈 → 선택 반영", "")
    c("조회", "Owner 검색 모달", "(정상)", "Owner 검색 아이콘", "OwnerByDivisionModal 오픈 → 선택 반영", "")
    c("조회", "Column 설정", "(정상)", "Column Setting 아이콘 → 표시/순서 변경 저장",
      "titleSet 'C' 저장 + 기존 조회 데이터 Reset(그리드 비움 — UC-03) → Search 재조회", "SCM0002_SET_TITLESET_POPUP")
    c("조회", "Measure 설정", "(정상)", "Measure 아이콘 → 2종(Cust. FCST/EDI) 토글 저장",
      "titleSet 'M' 저장 + 데이터 Reset (UC-03)", "SCM0002_SET_TITLESET_POPUP")
    c("조회", "필터 유지 재조회", "(정상)", "그리드 필터 적용 상태 Search", "재조회 후 필터 유지 (UC-07)", MAIN_SP)
    c("조회", "콤마 다중 조회", "(정상)", "Model 칸에 'A,B' 콤마 입력 후 Search", "전건 조회 (UC-06)", MAIN_SP)
    c("조회", "Excel Download", "(정상)", "Excel 버튼", "표시 컬럼 기준 다운로드 (UC-05)", "")
    for grp, acct in ACCTS[div]:
        deny = "무권한" in grp
        c("권한", f"권한별 접근 — {grp}", "(정상)" if not deny else "(오류)", f"{acct} 로그인 → Div {div} 조회",
          ("SEARCH 403 (메뉴권한 미등재 — AS-IS 4그룹 정합: ADMIN/DM ADMIN/DM Planner/DM Viewer)" if deny
           else "조회 성공 (읽기 전용 — 편집 진입점 없음). SAVE 강제 호출 시 403"), MAIN_SP, "" if deny else ROWS[div],
          "HTTP E2E H축 실측 대조")
    return rows

wb = openpyxl.Workbook()
wb.remove(wb.active)
for div in ("40", "81", "82"):
    ws = wb.create_sheet(div)
    for i, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HD; cell.fill = FILL; cell.border = THIN; cell.alignment = WRAP
    widths = {1: 5, 6: 22, 7: 6, 12: 10, 13: 28, 14: 8, 15: 34, 16: 46, 18: 26, 19: 22, 20: 30, 27: 36, 30: 16, 31: 30}
    for col, w in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    r = 2
    for i, (step, name, kind, data, expect, proc, cnt, how, note) in enumerate(cases(div), 1):
        vals = {1: i, 6: "UI 상세 (Customer FCST Waterfall)", 7: "UI", 8: "주중", 10: step, 12: "DMF0053",
                13: f"Customer FCST Waterfall — {name}", 14: kind, 15: data, 16: expect,
                18: "(조회 전용 — Write 없음. titleSet 개인화만 SCM_MST_TITLESET_USER W)", 19: MEAS, 20: SRC,
                21: "Customer FCST 원장 + EDI CALLOFF", 24: note, 27: proc, 30: cnt, 31: how}
        for col, v in vals.items():
            cell = ws.cell(row=r, column=col, value=v)
            cell.border = THIN; cell.alignment = WRAP
        r += 1
    wd = wb.create_sheet(f"{div}_DATA")
    for i, h in enumerate(["No", "데이터ID", "구분", "Division", "권한그룹", "테스트계정", "EDIT_YN", "대상시트·팝업"], 1):
        cell = wd.cell(row=1, column=i, value=h)
        cell.font = HD; cell.fill = FILL; cell.border = THIN
    for i, (grp, acct) in enumerate(ACCTS[div], 1):
        for j, v in enumerate([i, f"D{div}-{i:02d}", "계정", div, grp, acct, "N (조회 전용)",
                                "custFcstWaterfall / Project·Item·Owner 모달 / titleSet"], 1):
            cell = wd.cell(row=i + 1, column=j, value=v)
            cell.border = THIN
wb.save(OUT)
print("저장:", OUT)
wb2 = openpyxl.load_workbook(OUT, read_only=True)
counts = {ws.title: ws.max_row - 1 for ws in wb2.worksheets}
wb2.close()
print("시트·행:", counts)
