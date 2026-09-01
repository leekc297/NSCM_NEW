# -*- coding: utf-8 -*-
"""DMF0053 화면설계서 v1.0 (8시트 — hook #105 §1-A·§1-D·§1-E, DMF0040 생성기 계보)"""
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace", line_buffering=True)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OUT = r"C:\Projects\NSCM\Doc\화면설계서\DMF0053_CustFcstWaterfall_화면설계서_v1.0.xlsx"
H = Font(bold=True, size=12)
HD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="4472C4")
THIN = Border(*[Side(style="thin")] * 4)
WRAP = Alignment(vertical="top", wrap_text=True)

wb = openpyxl.Workbook()
wb.remove(wb.active)


def sheet(title, widths):
    ws = wb.create_sheet(title)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    return ws


def put(ws, row, vals, header=False, title=False):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.alignment = WRAP
        c.border = THIN
        if header:
            c.font = HD
            c.fill = FILL
        if title:
            c.font = H
            c.border = Border()
    return row + 1


# 1.개요
ws = sheet("1.개요", [22, 70, 30])
r = put(ws, 1, ["DMF0053 (Customer FCST Waterfall) — 화면설계서 v1.0"], title=True)
for k, v in [
    ("작성", "2026-08-27 — UI분석서 v1.0 (10시트) 승계"),
    ("화면 ID / 명", "DMF0053 / Customer FCST Waterfall (MNU00048)"),
    ("메뉴 경로", "Demand > Waterfall > Customer FCST Waterfall (ROUTE_PATH: demand/demand-waterfall/cust-fcst-waterfall — 이관 마지막 단계에 의미형 갱신)"),
    ("페이지 폴더", "nscm-fe/features/demand/src/pages/cust-fcst-waterfall/ (SCM_ADM_MNU_M walk-up: Demand_Waterfall → Demand — 라이브 DB 실측)"),
    ("성격", "읽기 전용 Waterfall 조회 리포트 — Insert/Delete/Save/Upload 미노출"),
    ("reference", "1차 mkt-fcst-comparison(DMF0040 — 직전 정본·동일 부모 메뉴) + demand-plan-waterfall(DMF0035)"),
    ("권한", "ADMIN / DM ADMIN / DM Planner / DM Viewer (4그룹 — 전 그룹 조회 동일)"),
]:
    r = put(ws, r, [k, v, ""])

# 2.전체흐름
ws = sheet("2.전체흐름", [8, 28, 62, 30])
r = put(ws, 1, ["2. 전체흐름 (UI분석서 2시트 승계)"], title=True)
r = put(ws, r, ["단계", "프로세스", "설명", "본 화면"], header=True)
for row in [
    ("1", "EDI 수신·Customer FCST 적재", "DM_DNM_EDI_CALLOFF_ITEM + DM_PLN_CUST_FCST_H (DMF0014 Upload)", "* 원장"),
    ("2", "Waterfall 리포트", "Customer FCST vs EDI 를 M>W>D 3단 기간 축으로 표시", "** 본 화면"),
    ("3", "수요계획 소비", "MKT FCST/Sales Plan 수립 참고", ""),
]:
    r = put(ws, r, list(row))

# 3.LAYOUT
ws = sheet("3.LAYOUT", [20, 80])
r = put(ws, 1, ["3. LAYOUT — 표준 셸 (#132) · Figma 영역 매핑표 승계 (Doc/_figma/screens/DMF0053.md)"], title=True)
r = put(ws, r, ["영역", "구성"], header=True)
for row in [
    ("MainTitle", "타이틀 + Search Area 토글 (공용 셸 — 래퍼 div 금지)"),
    ("SearchArea 1행", "Div.(단일*) / Team(다중 ALL) / End Customer(다중 ALL) / Sales Site(다중 ALL) + Search(F5)"),
    ("SearchArea 2행", "Biz. Type(다중 ALL) / Input Date(범위*) / Period(범위*) / Bucket(단일 — 기본 Monthly)"),
    ("SearchArea 3행", "PJT(텍스트+검색 — 40 비활성) / Model/Nickname(텍스트+검색) / Plant Site(다중 ALL — 40 비활성) / Item Type(다중 ALL)"),
    ("SearchArea 4행", "Owner(텍스트+검색) / QTY(단일 — 기본 1) / Sum(체크)"),
    ("SubTitle", "총건수 + Excel + 기본 버튼(Sort/Filter/Condition/Reset 자동) + customButtons: Column Setting → Measure (hook #120 순서·아이콘)"),
    ("Grid", "공용 DataGrid 단일 시트 — 정적 66컬럼(hook #113 서버 직접 비교) + 동적 waterfall(M_/W_/D_ 3단 헤더+SUM) + footer Sum + Freeze"),
    ("모달", "ProjectByDivisionModal / ItemByDivisionModal / OwnerByDivisionModal / ColumnLayoutModal('C') / MeasureSelectModal — 전부 master 공용 정본 (사본 금지)"),
]:
    r = put(ws, r, list(row))

# 4.조회조건·버튼
ws = sheet("4.조회조건·버튼", [18, 16, 26, 40, 40])
r = put(ws, 1, ["4. 조회조건 · 버튼 (#112 카탈로그 17행 전 행 대조 — Cascade_Filter_전수.md §1 승계)"], title=True)
r = put(ws, r, ["항목", "유형", "데이터 소스", "동작·cascade", "비고"], header=True)
for row in [
    ("Div.*", "단일 콤보", "master DivisionField", "변경 시 Team/EndCustomer/BizType/SalesSite/PlantSite 재적재 + PJT·Model·Owner 클리어 + Grid 초기화·CRUD reset (§1.4)", "필수 (A-01)"),
    ("Team", "다중+검색+ALL", "TeamField(division)", "→ End Customer, Sales Site", ""),
    ("End Customer", "다중+검색+ALL", "EndCustomerField(division, team)", "→ Sales Site", ""),
    ("Sales Site", "다중+검색+ALL", "SalesSiteField(division, endCustomer[, team=40만])", "", "81*/82 는 team 미전달 (B-06/08)"),
    ("Biz. Type", "다중+검색+ALL", "BizTypeField(division)", "", ""),
    ("Input Date*", "날짜 범위", "PR_GET_DAY_RANGE p_month=1", "역전 MSG00108 / 3년 초과 MSG00164+보정", "A-04/A-05"),
    ("Period*", "날짜 범위", "PR_GET_DAY_RANGE p_month=6", "역전 MSG00108 / 3년 초과 MSG00164+보정", "A-02/A-03"),
    ("Bucket", "단일 콤보", "DMF0003_GET_REPORT_BUCKET_TYPE_CODE", "변경 시 타이틀 재조회 + D_/W_/M_ 가시성 전환. ONLY_WEEKLY = Sum 비활성+해제", "기본 MONTHLY"),
    ("PJT", "텍스트+검색 모달", "ProjectByDivisionModal", "40 비활성 (81*/82 활성)", "명칭 Project→PJT (#112)"),
    ("Model/Nickname", "텍스트+검색 모달", "ItemByDivisionModal", "", ""),
    ("Plant Site", "다중+검색+ALL", "PlantSiteField(division)", "40 비활성", ""),
    ("Item Type", "다중+검색+ALL", "LOOKUP ITEM_TYPE_CODE", "", ""),
    ("Owner", "텍스트+검색 모달", "OwnerByDivisionModal", "", ""),
    ("QTY", "단일 콤보", "DMF0003_GET_QTY_TYPE_CODE", "표시값 ÷{1, 1,000, 1,000,000, 100,000,000} 소수2·0 제외 (표시 사본 — 원본 보존)", "C-02"),
    ("Sum", "체크", "로컬 상태", "D_/W_ 상위 집계 컬럼 표시 토글", ""),
    ("(버튼) Column Setting", "아이콘", "ColumnLayoutModal columnType='C'", "저장 = titleSet 'C' + 그리드 비움 (UC-03) + Toast", "#120 1번째"),
    ("(버튼) Measure", "아이콘", "MeasureSelectModal", "저장 = titleSet 'M' + UC-03 + Toast", "#120 2번째"),
    ("(제외) Version", "-", "-", "AS-IS btnVersion 숨김 + Figma·카탈로그 부재 → 미구현. versionNo='' 고정", "E21/E22/D-01"),
    ("(공통) hotkeyConfig", "-", "-", "search=handleSubmit / insert·delete·save·upload·manual=false (읽기 전용)", ""),
]:
    r = put(ws, r, list(row))

# 5.GRID
ws = sheet("5.GRID", [24, 76])
r = put(ws, 1, ["5. GRID — 단일 시트 'CUST Fcst Waterfall'"], title=True)
r = put(ws, r, ["축", "설계"], header=True)
for row in [
    ("정적 컬럼", "SHEET_COLUMNS 66행 (PROGRAM_ID='DMF0053') — hook #113 PR_GET_SHEET_COLUMN_LIST 서버 직접 비교 (divisionId 전송·division 변경 재fetch). DIVISION_ID 토큰: 40 전용 20종·81 계열 전용 10종·82 전용 2종"),
    ("Measure 행 축", "MEASURE_GROUP + MEASURE (SP 커서 산출 — Cust. FCST / EDI) — 좌측 고정·세로 머지"),
    ("동적 컬럼", "spreadTitle SP 산출 — M_/W_/D_ (+B_ 는 항상 숨김). 3단 헤더: 월>주>일 + 그룹 SUM 컬럼. SAT Blue/SUN Red (row2)"),
    ("드릴다운", "M_ 헤더 더블클릭 = 해당 월 W_ 토글(열린 D_ 접기) / W_ 더블클릭 = 해당 주 D_ 토글 (E11 — DMF0040 HeaderMergeManager 계보)"),
    ("Bucket 가시성", "MONTHLY: M_ / WEEKLY: W_(+Sum 시 M_) / DAILY: D_(+Sum 시 W_·M_) / ONLY_WEEKLY: W_ 단순 3단"),
    ("footer", "GroupRow Sum — D_/W_/M_ 컬럼 Aggregate.Sum (표시 컬럼만)"),
    ("셀 색", "기본 White+zebra (#95) — 편집 없음(전열 잠금·읽기 전용). 실적 회색 없음 (AS-IS 주석 — DMF0040 과 차이)"),
    ("편집", "없음 — isReadOnly. 헤더 높이 #93 (3단 48px 규약)"),
    ("QTY 표시", "QTY≠1 이면 표시 사본에 divisor 적용 (원본 데이터 불변 — Excel 은 표시 기준)"),
    ("Excel", "UC-05 — excelFileName='CustFcstWaterfall', 표시 컬럼 기준"),
    ("필터 유지", "UC-07 rebindFilterColumns (재조회 후 필터 유지) + UC-06 콤마 다중 조회"),
]:
    r = put(ws, r, list(row))

# 6.EVENT
ws = sheet("6.EVENT", [10, 30, 70])
r = put(ws, 1, ["6. EVENT — 목차 + 선행 MD 전문 승계 (#105 §1-D). 정본 = Doc/화면분석서/DMF0053_CustFcstWaterfall_주요이벤트_상세분석.md"], title=True)
r = put(ws, r, ["E#", "핸들러", "TO-BE 판정 (요약 — 상세는 정본 md)"], header=True)
for row in [
    ("E01", "Form_Load", "구현 — store 초기화 + Division 자동선택 + Bucket=MONTHLY"),
    ("E02", "cboDivision_SelectedIndexChanged", "구현 — cascade + 40=PJT·PlantSite 비활성 + 텍스트 클리어 + Grid 초기화"),
    ("E03", "clbTeam_SelectedIndexChanged", "구현 — EndCustomer/SalesSite cascade"),
    ("E04", "clbCustomer_SelectedIndexChanged", "구현 — SalesSite cascade"),
    ("E05~E08", "SalesSite/PlantSite/Project/ItemId 핸들러 4종", "제외 — 미배선 사문 (본문 SetClbVersion=no-op)"),
    ("E09", "cboBucketType_SelectedIndexChanged", "구현 — 타이틀 재조회+가시성 (AS-IS 이중 배선은 1회로)"),
    ("E10", "spdMain_ComboSelChange", "제외 — dead path (전열 잠금)"),
    ("E11", "spdMain_CellDoubleClick(헤더)", "구현 — 드릴다운"),
    ("E12", "spdMain_CellDoubleClick(셀)", "제외 — dead path (ACTION='I' 전용)"),
    ("E13", "spdMain_TopChange", "구현 — wijmo 기본 머지"),
    ("E14", "spdMain_KeyUp", "부분 — Copy=wijmo 기본 / Paste 제외 (읽기 전용)"),
    ("E15", "btnProject_Click", "구현 — ProjectByDivisionModal"),
    ("E16", "btnItemId_Click", "구현 — ItemByDivisionModal"),
    ("E17", "btnColumn_Click", "구현 — ColumnLayoutModal 'C' + UC-03"),
    ("E18", "btnMeasure_Click", "구현 — MeasureSelectModal + UC-03"),
    ("E19", "chkSum_CheckedChanged", "구현 — 상위 집계 가시성"),
    ("E20", "chkColor_CheckedChanged", "제외 — 미배선 사문 (실적 회색 주석)"),
    ("E21/E22", "btnVersion_Click / txtVersion_KeyDown", "제외 — Version 조회조건 미구현"),
    ("E23", "btnOwner_Click", "구현 — OwnerByDivisionModal"),
    ("E23+", "cboQTY_SelectedIndexChanged", "구현 — QTY 표시 사본 divisor"),
]:
    r = put(ws, r, list(row))

# 7.VALIDATION
ws = sheet("7.VALIDATION", [10, 44, 14, 50])
r = put(ws, 1, ["7. VALIDATION — 제약매트릭스_상세.md 승계 (A6+B10+C2+D1)"], title=True)
r = put(ws, r, ["ID", "조건", "MSG", "TO-BE"], header=True)
for row in [
    ("A-01", "Division 미선택", "MSG00097", "차단 — 'Division 을 선택하세요.'"),
    ("A-02", "Period 역전", "MSG00108", "차단"),
    ("A-03", "Period 3년 초과", "MSG00164", "알림 + to=from+3년 보정 후 중단 (문구·로직 3년 정합)"),
    ("A-04", "Input Date 역전", "MSG00108", "차단"),
    ("A-05", "Input Date 3년 초과", "MSG00164", "보정 후 중단"),
    ("A-06", "편집 시도 (F2)", "MSG00109", "제외 — 편집 진입점 미노출"),
    ("B-01~08", "Division 분기 (cascade·활성) 8건", "", "구현 — 접두 '81'/등호 '82' 판정 + 40 비활성"),
    ("B-09/10", "그리드 셀 콤보·Insert dedup", "", "제외 — dead path/dead code"),
    ("C-01", "전주 일요일 역산", "", "제외 — 소비처 주석"),
    ("C-02", "QTY=1 원본 복원", "", "구현"),
    ("D-01", "Version 키 입력", "", "제외 — Version 미구현"),
]:
    r = put(ws, r, list(row))

# 8.API·SP
ws = sheet("8.API·SP", [30, 34, 60])
r = put(ws, 1, ["8. API · SP — 배선 계약 (positional — proc-meta-signature-match)"], title=True)
r = put(ws, r, ["축", "값", "상세"], header=True)
for row in [
    ("BTN → API (신규 시딩 4+2)", "custFcstWaterfall:SEARCH / spreadTitle:SEARCH / titleSet:SEARCH·SAVE", "SCM_ADM_MNU_BTN_MPG_M 4행 (전부 API_ID 실배선 — DMF0040 E20260827-TITLESETAPI 계보) + API_M 4본"),
    ("main API", "DMF0053_custFcstWaterfall_SEARCH → PKG_DMF0053.P_DMF0053_GET_CUST_FCST_WATERFALL", "PROC_SELECT. IN 21 positional — camelCase payload 키 (divisionId/deptId/endCustomer/scmSiteId/projectCode/itemId/itemType/fromYyyymmdd/toYyyymmdd/createdBy(=BE sesUserId)/plantSiteId/bizType/owner/versionNo(''')/bucket/userDeptId/groupId/fromCreatedate/toCreatedate/userId/menuId) — 정확 명칭·순서는 Phase 5 ALL_ARGUMENTS 실측으로 확정"),
    ("spreadTitle API", "DMF0053_spreadTitle_SEARCH → PKG_DMF0053.P_DMF0053_GET_SPREAD_TITLE", "IN 7 (startYyyymmdd/endYyyymmdd/createdBy/divisionId/bucket/userId/menuId)"),
    ("titleSet API", "DMF0053_titleSet_SAVE → SCM0002_SET_TITLESET_POPUP / _SEARCH → DMF0002_GET_TITLESET_POPUP", "공유 PROC_M 기등재 재사용 — programId='DMF0053', sheetName='CUST Fcst Waterfall'"),
    ("그룹권한", "4그룹 × 4버튼 = 16행", "ADMIN / DM ADMIN / DM Planner / DM Viewer (SCM_ADM_USER_GR_MNU_BTN_MPG_M)"),
    ("procedures.d.ts", "DMF0053 계약 블록", "custFcstWaterfall 21키 / spreadTitle 7키 — SSOT (로컬 재정의 금지)"),
    ("GRANT", "PKG_DMF0053 EXECUTE → SCMP_APP", "+ 공유 콤보 CMBO 는 기이관 재사용"),
    ("응답 커서", "O_RC1(데이터)+O_RC2(ISO_WEEK)+O_RC3(DATA_ORDER)+O_RESULT_CURSOR", "O_RC3 정렬식은 SP 산출 — FE 는 수신 순서 유지(서버 정렬) 원칙, R 스모크로 실측 확정"),
]:
    r = put(ws, r, list(row))

wb.save(OUT)
print("저장:", OUT)
wb2 = openpyxl.load_workbook(OUT, read_only=True)
names = wb2.sheetnames
assert len(names) == 8, names
wb2.close()
print("시트 8 OK:", names)
