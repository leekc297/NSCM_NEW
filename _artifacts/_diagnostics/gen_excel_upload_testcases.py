"""
gen_excel_upload_testcases.py — Excel Upload Validation 케이스별 테스트 Excel 생성기 (2026-08-27)

  DMF0074 Metal Link Price (division 40 · DMF0074_40 53열 양식)
    원천: Doc/화면분석서/DMF0074_MetalLinkPrice_ExcelUpload_Validation_v1.1.xlsx [1_Validation목록] 53행 (No 1~53)
  DMF0014 Customer FCST Upload (division 81 · 1단 헤더 13식별열 + 일자 4열)
    원천: Doc/화면분석서/DMF0014_Excel_Upload_Validation_Rules_20260825.xlsx [01_검증규칙] 18행 + [03_수량컬럼규칙]

  산출:
    Doc/단위테스트/DMF0074_ExcelUpload_ValidationCases_20260827.xlsx  (00_INDEX + 케이스별 Sheet)
    Doc/단위테스트/DMF0014_ExcelUpload_ValidationCases_20260827.xlsx  (00_INDEX + 케이스별 Sheet)
    Doc/단위테스트/_excel_upload_cases/DMF0074/<No>_<ID>.xlsx  (케이스별 단일 시트 — BE ExcelUtils 는 첫 시트만 읽으므로 실업로드용)
    Doc/단위테스트/_excel_upload_cases/DMF0014/<No>_<ID>.xlsx

  ⚠ 값은 전부 문자열 셀(텍스트 서식)로 넣는다 — Period 를 숫자로 넣으면 서버 파서가 '6.0' 으로 읽어
    'Invalid INTEGER 0 ~ 12.' 가 된다(2026-08-18 실측). 일자 헤더도 문자열 '20260831'.
"""
import os, sys, json, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = r"C:\Projects\NSCM"
OUT_DIR = os.path.join(ROOT, "Doc", "단위테스트")
CASE_DIR = os.path.join(OUT_DIR, "_excel_upload_cases")
STAMP = datetime.date.today().strftime("%Y%m%d")  # [E20260831-REGEN] 생성일 동적 (구 하드코딩 '20260827')

# ─────────────────────────────────────────────────────────────────────────────
# DMF0074 (division 40 · 53열)
# ─────────────────────────────────────────────────────────────────────────────
H74 = ["Div", "Plan Date", "Team", "End Customer", "Model Code", "Metal Link", "Sales Site", "Currency", "FCST type",
       "Exception", "Energy",
       "Key_Ni", "Key_Co", "Key_Li2CO3", "Key_LIOH", "Key_Al", "Key_Mn", "Key_Cu",
       "PeriodFr_Ni", "PeriodFr_Co", "PeriodFr_Li2CO3", "PeriodFr_LIOH", "PeriodFr_Al", "PeriodFr_Mn", "PeriodFr_Cu",
       "PeriodTo_Ni", "PeriodTo_Co", "PeriodTo_Li2CO3", "PeriodTo_LIOH", "PeriodTo_Al", "PeriodTo_Mn", "PeriodTo_Cu",
       "ContractPrice_Ni", "ContractPrice_Co", "ContractPrice_Li2CO3", "ContractPrice_LIOH", "ContractPrice_Al", "ContractPrice_Mn", "ContractPrice_Cu",
       "ChangePrice_Ni", "ChangePrice_Co", "ChangePrice_Li2CO3", "ChangePrice_LIOH", "ChangePrice_Al", "ChangePrice_Mn", "ChangePrice_Cu",
       "MetalIndex_Ni", "MetalIndex_Co", "MetalIndex_Li2CO3", "MetalIndex_LIOH", "MetalIndex_Al", "MetalIndex_Mn", "MetalIndex_Cu"]
assert len(H74) == 53
METALS = ["Ni", "Co", "Li2CO3", "LIOH", "Al", "Mn", "Cu"]
KEY_OK = {"Ni": "M035", "Co": "M037", "Li2CO3": "M043", "LIOH": "M045", "Al": "M050", "Mn": "M053", "Cu": "M056"}   # 연동 코드 (LME/Fastmarket)
KEY_UNUSED = {"Ni": "M036", "Co": "M040", "Li2CO3": "M044", "LIOH": "M058", "Al": "M051", "Mn": "M055", "Cu": "M057"}  # '미연동'
PLAN_OK = "202610"  # 당월(2026-08) 이후 — 'Past Date.' 회피. 실행 시점이 2026-10 을 넘기면 갱신할 것

# 기준행 (2026-08-27 사내 실측 — DM_PLN_SALES_PLAN_C SP_W_20260820_V002 · SITE_ITEM VALID · END_CUSTOMER 룩업 · 환율 USD · ITEM_TYPE P)
BASE74 = {"Div": "40", "Plan Date": PLAN_OK, "Team": "", "End Customer": "Isuzu", "Model Code": "MPCISZISLD1V", "Metal Link": "N",
          "Sales Site": "Isuzu_JP_TYO", "Currency": "USD", "FCST type": "Mass", "Exception": "N", "Energy": ""}

def row74(marker, **over):
    r = {h: "" for h in H74}
    r.update(BASE74); r["Team"] = marker; r.update(over)
    return r

def metal_full(m, **over):
    """연동 광물 m 1종 완전입력 (Key + 5값)"""
    d = {f"Key_{m}": KEY_OK[m], f"PeriodFr_{m}": "6", f"PeriodTo_{m}": "0", f"ContractPrice_{m}": "18.5",
         f"ChangePrice_{m}": "0.75", f"MetalIndex_{m}": "0.42"}
    d.update(over); return d

# (No, 검증ID, 제목, 기대 메시지, 행 목록 or None(N/A), 비고)
CASES74 = []
def add74(no, cid, title, expect, rows, note=""):
    CASES74.append(dict(no=no, cid=cid, title=title, expect=expect, rows=rows, note=note))

add74(0, "OK-N", "정상 (Metal Link=N)", "(정상 — Error 0)", [row74("TC00A")],
      "기준행. C19 는 금주 MRF(SP_W ATTRIBUTE7=Y) 버전의 SPW 원장 존재에 의존 — 사내 2026-08-27 실측: MRF 버전(SP_W_20260728_V002)에 PLAN_C/RESULT_C 0행 → 이 상태에선 'Invalid Sales Plan Waterfall MODEL.' 로 반려된다(데이터 상태 이슈·규칙 결함 아님)")
add74(0, "OK-Y", "정상 (Metal Link=Y, Ni 1종 완전입력)", "(정상 — Error 0)", [row74("TC00B", **{"Metal Link": "Y"}, **metal_full("Ni"))], "OK-N 과 동일 전제")
add74(1, "C2", "Div 공란", "(SKIP 1건 — 오류 아님)", [row74("TC01", Div="")],
      "[E20260831-REGEN] 기대값을 실측 기준으로 정정 — SP 의 'NULL DIVISION.' 체크는 실존하나(BODY L1272), BE 가 Division 공란 행을 스테이징에서 UPLOAD_STATUS_CODE='SKIP' 으로 선처리해 SP 검증(UPLOAD 행 대상)에 도달하지 않음(2026-08-27 실측 SKIP:1). 화면 표기는 SKIP 카운트")
add74(2, "C2", "Plan Date 공란", "NULL PLAN DATE.", [row74("TC02", **{"Plan Date": ""})])
add74(3, "C2", "End Customer 공란", "NULL END CUSTOMER.", [row74("TC03", **{"End Customer": ""})])
add74(4, "C2", "Project Code 공란", "NULL PROJECT CODE.", None, "division 40 양식(DMF0074_40)에는 Project Code 열이 없어 대상 외 — 81/82(55열 DMF0074 양식)에서만 발생")
add74(5, "C2", "Sales Site 공란", "NULL SALES SITE.", [row74("TC05", **{"Sales Site": ""})])
add74(6, "C2", "Currency 공란", "NULL CURRENCY.", [row74("TC06", Currency="")])
add74(7, "C2", "FCST type 공란", "NULL FCST TYPE.", [row74("TC07", **{"FCST type": ""})])
add74(8, "C2", "Metal Link 공란", "NULL Metal Link.", [row74("TC08", **{"Metal Link": ""})])
add74(9, "C2", "Team 공란 (체크 제외)", "(정상 — Team 은 필수 체크 제외)", [row74("")], "[PJ2024-50461] Team 미관리 — 오류 없이 통과해야 함(마커 없음)")
add74(10, "C3", "금주 MRF 버전 없음", "There is no version of Monthly Rolling Forest:Y in Sales Plan Waterfall this week.", None,
      "DM_MST_VERSION(SP_W, ATTRIBUTE7='Y') 에 해당 division 버전이 없어야 발생 — 데이터 상태로만 재현 가능(40/81/82 모두 MRF 버전 보유). 파일로 유발 불가")
add74(11, "C4", "Metal Link=Y + Key 7종 전부 공란", "NULL APPLY PRICE TYPE CODE.", [row74("TC11", **{"Metal Link": "Y"})])
for i, m in enumerate(METALS):
    add74(12 + i, f"C{5 + i}", f"Key_{m} 미등재 코드", f"Invalid KEY CODE : {m}",
          [row74(f"TC{12 + i}", **{"Metal Link": "Y", f"Key_{m}": "BOGUS"})])
add74(19, "C12", "8키 동일 중복 2행", "Duplication Model.", [row74("TC19"), row74("TC19")], "결함 D3 정정(2026-08-16) 후 40 에서도 검출")
add74(20, "C13", "Plan Date 형식 오류", "Invalid Date Format.", [row74("TC20", **{"Plan Date": "2026-10"})])
add74(21, "C14", "Plan Date 과거월", "Past Date.", [row74("TC21", **{"Plan Date": "202601"})])
add74(22, "C15", "End Customer 미등재", "Invalid END CUSTOMER.", [row74("TC22", **{"End Customer": "ZZ_NO_CUST"})])
add74(23, "C15", "Project Code 미등재", "Invalid PROJECT CODE.", None, "division 40 대상 외 (Project Code 열 없음)")
add74(24, "C16", "Model 미등재", "Invalid MODEL.", [row74("TC24", **{"Model Code": "ZZ_NO_MODEL"})], "결함 D1 정정 후 발동")
add74(25, "C17", "82 ENERGY 없는 모델", "Invalid MODEL : NULL ENERGY", None, "division 82 전용 — 40 양식 대상 외")
add74(26, "C18", "제품(P) 아닌 모델", "Invalid MODEL : TYPE P.", [row74("TC26", **{"End Customer": "IM", "Model Code": "APCCA3173A", "Sales Site": "IM_US_TXT"})],
      "APCCA3173A = SITE_ITEM(IM/IM_US_TXT) 등재(C16 통과) + DM_MST_ITEM_SINGLE_V ITEM_TYPE='H'(제품 P 아님) — 2026-08-27 실측")
add74(27, "C19", "SPW 조합 불일치 (Currency 만 다름)", "Invalid Sales Plan Waterfall MODEL.", [row74("TC27", Currency="EUR")],
      "EUR 은 환율 등재 통화라 C22 통과 → SPW 조합(USD)만 불일치. D5 정정 후에도 C19 로 표기되는 정상 케이스")
add74(28, "C20", "Metal Link 코드 오류", "Invalid Metal Link.", [row74("TC28", **{"Metal Link": "X"})])
add74(29, "C21", "Sales Site 미등재", "Invalid Sales Site.", [row74("TC29", **{"Sales Site": "ZZ_NO_SITE"})], "D5 정정 전엔 C19 메시지로 마스킹(TC14) — 정정 후 구체 메시지")
add74(30, "C22", "Currency 미등재", "Invalid CURRENCY.", [row74("TC30", Currency="ZZZ")], "D5 정정 전엔 C19 마스킹(TC15)")
add74(31, "C23", "FCST type 미등재", "Invalid FCST TYPE.", [row74("TC31", **{"FCST type": "ZZ_TYPE"})], "D5 정정 전엔 C19 마스킹(TC16)")
add74(32, "C24", "Exception 코드 오류", "Invalid EXCEPTION.", [row74("TC32", Exception="Q")], "결함 D2 정정 후 발동")
add74(33, "C25", "(GTT 적재 단계)", "-", None, "검증이 아니라 C26~C39 판정 기준(미사용/미연동 코드) 수집 단계 — 케이스 없음")
for i, m in enumerate(METALS):
    add74(34 + i, f"C{26 + i}", f"Key_{m} 미연동인데 값 입력", f"NEED EMPTY VALUE - {m} : Period(Fr)",
          [row74(f"TC{34 + i}", **{"Metal Link": "Y", f"Key_{m}": KEY_UNUSED[m], f"PeriodFr_{m}": "6"})],
          "메시지는 '먼저 걸린 항목명' 접미 — Period(Fr) 만 입력해 Period(Fr) 로 고정")
for i, m in enumerate(METALS):
    add74(41 + i, f"C{33 + i}", f"Key_{m} 유효인데 5값 공란", f"NULL {m} : Period(Fr)",
          [row74(f"TC{41 + i}", **{"Metal Link": "Y", f"Key_{m}": KEY_OK[m]})],
          "division 앞2자리 81 제외(40 은 대상). 메시지 접미 = 먼저 걸린 항목 Period(Fr)")
add74(48, "C40", "Period 비숫자", "Invalid NUMERIC.", [row74("TC48", **{"Metal Link": "Y"}, **metal_full("Ni", PeriodFr_Ni="abc"))])
add74(49, "C41", "Period 음수", "Negative Number.", [row74("TC49", **{"Metal Link": "Y"}, **metal_full("Ni", PeriodFr_Ni="-1"))])
add74(50, "C42", "Period 12 초과", "Invalid INTEGER 0 ~ 12.", [row74("TC50", **{"Metal Link": "Y"}, **metal_full("Ni", PeriodFr_Ni="13"))])
add74(51, "C43", "Period(Fr) <= Period(To)", "From Value must be larger than To Value.",
      [row74("TC51", **{"Metal Link": "Y"}, **metal_full("Ni", PeriodFr_Ni="1", PeriodTo_Ni="3"))])
add74(52, "C44", "(결과 반영 MERGE)", "-", None, "GTT 판정을 SCM_ADM_EXCEL_UPLOAD 에 반영하는 단계 — 케이스 없음")
add74(53, "GATE", "오류 잔존 시 [OK] 적재 차단", "NULL PLAN DATE.",
      [row74("TC53", **{"Plan Date": ""}), row74("TC53B")],
      "[E20260831-REGEN] 러너/1단계([Upload]) 실측 기대 = 'NULL PLAN DATE.' 1건(Error 유발이 목적). 본 케이스의 진짜 판정은 **수동 축**: Error 잔존 상태에서 [OK](P_MODE=1) 호출 시 CODE='-1' 'Validation errors remain' + 원장 Δ0 ([E20260827-D4-NOLOADGUARD]). FE 는 Error>0 이면 [OK] 를 막으므로 API 직접 호출로 확인")

# ─────────────────────────────────────────────────────────────────────────────
# DMF0014 (division 81 · 1단 헤더 13식별열 + 주차 시작일 4열)
# ─────────────────────────────────────────────────────────────────────────────
H14 = ["Division*", "Sales Site*", "Plant Site*", "Customer Project*", "Customer Model*", "Incoterms*", "FCST Type*",
       "Customs Handbook*", "Measure Group*", "Measure*", "Upload UOM*", "Project", "Model"]
# [E20260831-REGEN] 일자 동적화 — 구 하드코딩("20260831"~)은 생성 주가 지나면 회색(과거) 주차가 되어
#   4번째 케이스(03_BKT 의 DAILY 헤더 = 생성주 월~목)부터 의도와 다른 오류를 유발했다(사용자 신고 2026-08-31).
#   다음 주 월요일부터 시작해 파일 시효를 1주 확보한다.
_next_mon = datetime.date.today() + datetime.timedelta(days=(7 - datetime.date.today().weekday()) % 7 or 7)
WEEKS = [(_next_mon + datetime.timedelta(weeks=i)).strftime("%Y%m%d") for i in range(4)]   # 월요일 7일 간격 → WEEKLY
DAYS = [(_next_mon + datetime.timedelta(days=i)).strftime("%Y%m%d") for i in range(4)]     # 1일 간격 → DAILY (BKT 위반 유발)
# 기준행 (2026-08-27 사내 실측 — DM_MST_CUSTOMER_MODEL/CUSTOMER_FCST(WEEKLY·CELL 17행·타 버킷 0)/SCM_MST_SITE PLANT/룩업)
BASE14 = {"Division*": "81", "Sales Site*": "GM_US_BQN", "Plant Site*": "EG01_81-W", "Customer Project*": "SGM Modified",
          "Customer Model*": "N2.2", "Incoterms*": "FOB", "FCST Type*": "Mass", "Customs Handbook*": "D",
          "Measure Group*": "CUST_FCST", "Measure*": "Customer FCST", "Upload UOM*": "Cell",
          "Project": "A17LG8001", "Model": "ACEN4065I-A1"}

def row14(qty=("100", "200", "300", "400"), **over):
    r = dict(BASE14); r.update(over); r["_qty"] = list(qty); return r

CASES14 = []
def add14(no, cid, title, expect, rows, note="", periods=None):
    CASES14.append(dict(no=no, cid=cid, title=title, expect=expect, rows=rows, note=note, periods=periods or WEEKS))

# [E20260923-DMF0014VALMSG] 메시지 규격 '<규칙> [<헤더>: <입력값>]' + 연쇄 오류 가드(근본 원인 1건) + 기간 칸 셀 단위(Invalid Numeric/Decimal/Please Input QTY)
add14(0, "OK", "정상 (WEEKLY · CELL · 4주 수량)", "(정상 — Error 0)", [row14()], "기준행")
add14(1, "DUP", "키 10개 조합 동일 2행", "Duplication Key [Sales Site*: GM_US_BQN, Customer Project*: SGM Modified, Customer Model*: N2.2]",
      [row14(), row14(qty=("1", "2", "3", "4"))], "ATTRIBUTE1~10 동일 (UOM/Project/Model 은 키 밖) — 2행 모두 표기")
add14(2, "DIV", "Division 미등재", "Invalid Division [Division*: 99]", [row14(**{"Division*": "99"})],
      "[E20260923-DMF0014VALMSG] 연쇄(Sales Site/Plant Site/Bucket…) 미표기 — 상위 검증 통과 행만 하위 검증")
add14(3, "BKT", "헤더 일자 간격 DAILY (등록 WEEKLY)", "Invalid Measure Upload Bucket [Measure*: Customer FCST]", [row14()],
      "일자 헤더 1일 간격 → DAILY ≠ 등록 WEEKLY. DAILY 헤더는 마스터 밖 주차라 빈값 검사(Please Input QTY) 대상 아님", periods=DAYS)
add14(4, "MUOM", "Upload UOM 불일치 (등록 CELL)", "Invalid Measure Upload UOM [Upload UOM*: Pack]", [row14(**{"Upload UOM*": "Pack"})])
add14(5, "SITE", "Sales Site 미등재", "Invalid Sales Site [Sales Site*: ZZ_SITE]", [row14(**{"Sales Site*": "ZZ_SITE"})],
      "[E20260923-DMF0014VALMSG] 종전 'Invalid Measure Upload Bucket' 마스킹 해소 — Bucket/UOM/Measure 검증은 Sales Site 통과 행만")
add14(6, "PLNT", "Plant Site 미등재", "Invalid Plant Site [Plant Site*: ZZ_PLANT]", [row14(**{"Plant Site*": "ZZ_PLANT"})])
add14(7, "CPRJ", "Customer Project 미등재 ([Project]/[Model] SITE_ITEM 유효 → 통과)", "(정상 — Error 0)", [row14(**{"Customer Project*": "ZZ_PRJ"})],
      "[E20260930-DMF0014ASISPM] AS-IS 파리티 — Customer Project 검증은 AS-IS 주석(L399~427). Customer Model 은 고객모델 마스터 OR SITE_ITEM([Project]/[Model]) 중 하나만 맞아도 통과(AS-IS L429~470)")
add14(8, "CMDL", "Customer Model 미등재 ([Project]/[Model] SITE_ITEM 유효 → 통과)", "(정상 — Error 0)", [row14(**{"Customer Model*": "ZZ_MODEL"})],
      "[E20260930-DMF0014ASISPM] AS-IS L447~455 SITE_ITEM 경로 복원 — 기준행 [Project]/[Model](A17LG8001/ACEN4065I-A1) 유효라 통과")
add14(9, "INCO", "Incoterms 미등재", "Invalid Incoterms [Incoterms*: ZZZ]", [row14(**{"Incoterms*": "ZZZ"})])
add14(10, "FTYP", "FCST Type 미등재", "Invalid FCST Type [FCST Type*: ZZ_TYPE]", [row14(**{"FCST Type*": "ZZ_TYPE"})])
add14(11, "CSTM", "Customs Handbook 미등재", "Invalid Customs Handbook [Customs Handbook*: X]", [row14(**{"Customs Handbook*": "X"})],
      "[E20260915-DMF0014MSTCHK] 로 부활한 규칙(D/E 룩업)")
add14(12, "MGRP", "Measure Group ≠ CUST_FCST", "Invalid Measure Group [Measure Group*: PSI]", [row14(**{"Measure Group*": "PSI"})])
add14(13, "QSTR", "기간 칸 문자 혼입", f"Please Input QTY [{WEEKS[1]}: abc]", [row14(qty=("100", "abc", "300", "400"))],
      "[E20260923-DMF0014BLANK] 09-24 정책 반전 — 비숫자 = 'Please Input QTY [기간: 값]' (구 'Invalid Numeric' 폐기)")
add14(14, "QDEC", "기간 칸 소수점", f"Invalid Decimal [{WEEKS[1]}: 130.2]", [row14(qty=("100", "130.2", "300", "400"))],
      "정수만 허용(E20230119-000186) — 셀 단위 표기")
add14(15, "UOM", "Upload UOM 룩업 미등재", "Invalid Upload UOM [Upload UOM*: BOX]", [row14(**{"Upload UOM*": "BOX"})],
      "[E20260923-DMF0014VALMSG] 종전 MUOM 마스킹 해소 — 룩업 미등재는 이 메시지 1건만")
add14(16, "UOMC", "(Division, Sales Site, Measure) 고객 FCST 미등재", "Invalid Measure [Sales Site*: GM_US_BQN, Measure*: ZZ_MEASURE]",
      [row14(**{"Measure*": "ZZ_MEASURE"})], "[E20260923-DMF0014VALMSG] 종전 BKT 마스킹 해소 + 'Invalid Upload UOM (Customer FCST)' → 'Invalid Measure' 개명")
add14(17, "PRJ", "[Project] DMF0005 마스터(SITE_ITEM) 미등재 · [Model] 정상", "Invalid Project [Project: ZZ_PRJCODE]", [row14(Project="ZZ_PRJCODE")],
      "[E20260930-DMF0014MDLIND] Project 오류 행의 Model 은 (Site, Model) 자체 등재 여부로 판정 — 기준행 Model(ACEN4065I-A1)은 GM_US_BQN SITE_ITEM 등재라 Model 오류 없음(오탐 방지)")
add14(18, "MDL", "Model 품목 마스터(ITEM_SN) 미등재", "Invalid Model [Model: DMAT_0000066]",
      [row14(**{"Customer Project*": "ZZ_PRJ", "Customer Model*": "ZZ_MODEL", "Project": "D18GM0001", "Model": "DMAT_0000066"})],
      "SITE_ITEM(GM_US_BQN/D18GM0001/DMAT_0000066)은 유효 → Customer Model 은 SITE_ITEM 경로로 통과, 품목이 SCM_MST_ITEM_SN 에 없어 품목분류 검증(E20260915)만 오류. [E20260930-DMF0014ASISPM] Customer Project 검증 비활성으로 1건만")
add14(19, "QEMPTY", "기간 칸 빈값 (입력가능 주차)", "(정상 — Error 0)", [row14(qty=("100", "", "300", ""))],
      "[E20260923-DMF0014BLANK] 09-24 정책 반전 — 빈칸 = 오류 아님")
add14(20, "QGRAY", "회색(입력불가) 주차 빈값 = 통과", "(정상 — Error 0)", [row14(qty=("100", "200", "300", ""))],
      "4번째 헤더 = 마스터 밖 주차(20270104) → GRAYZERO 정책상 빈값 검사 제외", periods=WEEKS[:3] + ["20270104"])
add14(21, "MULTI", "복수 오류 누적 (Incoterms + 문자 + 빈값)",
      f"Invalid Incoterms [Incoterms*: ZZZ] / Please Input QTY [{WEEKS[1]}: abc]",
      [row14(qty=("100", "abc", "300", ""), **{"Incoterms*": "ZZZ"})], "한 행의 사유 전부 ' / ' 누적 (E20260911-DMF0014MULTI). 빈칸은 09-24 이후 오류 아님")
add14(22, "PMEMPTY", "[Project]/[Model] 빈칸", "Invalid Project [Project: (empty)] / Invalid Model [Model: (empty)]", [row14(Project="", Model="")],
      "[E20260930-DMF0014MDLIND] 빈칸도 오류 — Project 오류 행도 Model 판정(빈 Model 은 SITE_ITEM 미등재) → 2건 누적")
add14(23, "PMDL", "[Model] SITE_ITEM 조합 미등재 ([Project] 유효)", "Invalid Model [Model: ZZ_ITEM]", [row14(Model="ZZ_ITEM")],
      "[E20260930-DMF0014PMSI] (Division, Sales Site, [Project], [Model]) SCM_MST_SITE_ITEM 조합 존재 검증")
add14(24, "CPMX", "Customer Project 미등재 + [Project] SITE_ITEM 미등재 (두 경로 모두 불일치)", "Invalid Customer Model [Customer Model*: N2.2] / Invalid Project [Project: ZZ_PRJCODE]",
      [row14(**{"Customer Project*": "ZZ_PRJ"}, Project="ZZ_PRJCODE")],
      "[E20260930-DMF0014ASISPM] AS-IS L429~470 — 고객모델 마스터·SITE_ITEM 모두 불일치 → 'Invalid Customer Model' + [E20260930-DMF0014PMSI] [Project] 단독 검증 누적")
add14(25, "CMDLX", "Customer Model 미등재 + [Model] SITE_ITEM 미등재 (두 경로 모두 불일치)", "Invalid Customer Model [Customer Model*: ZZ_MODEL] / Invalid Model [Model: ZZ_ITEM]",
      [row14(**{"Customer Model*": "ZZ_MODEL"}, Model="ZZ_ITEM")],
      "[E20260930-DMF0014ASISPM] Customer Project 통과 가드 제거로 Customer Model 오류 직접 표기 + [E20260930-DMF0014PMSI] [Model] 단독 검증 누적")
add14(26, "AAABBB", "[Project]=AAA · [Model]=BBB (사용자 실측 재현 — 고객모델 마스터 유효)", "Invalid Project [Project: AAA] / Invalid Model [Model: BBB]",
      [row14(Project="AAA", Model="BBB")],
      "[E20260930-DMF0014PMSI] 재활성 + [E20260930-DMF0014MDLIND] Project 오류 행도 Model 판정 — 사용자 지적 2026-09-30 두 건 모두 표기")
# ─────────────────────────────────────────────────────────────────────────────
HDR_FILL = PatternFill("solid", fgColor="DDEBF7"); BOLD = Font(bold=True)
def text_cell(ws, r, c, v):
    cell = ws.cell(row=r, column=c, value=("" if v is None else str(v))); cell.number_format = "@"; return cell

def write_sheet74(ws, case):
    for c, h in enumerate(H74, 1): text_cell(ws, 1, c, h).fill = HDR_FILL
    for i, row in enumerate(case["rows"], 2):
        for c, h in enumerate(H74, 1): text_cell(ws, i, c, row.get(h, ""))
    ws.freeze_panes = "A2"

def write_sheet14(ws, case):
    hdr = H14 + case["periods"]
    for c, h in enumerate(hdr, 1): text_cell(ws, 1, c, h).fill = HDR_FILL
    for i, row in enumerate(case["rows"], 2):
        vals = [row.get(h, "") for h in H14] + row["_qty"]
        for c, v in enumerate(vals, 1): text_cell(ws, i, c, v)
    ws.freeze_panes = "A2"

def sheet_name(case):
    nm = f"{case['no']:02d}_{case['cid']}"
    return nm[:31]

def build(program, cases, writer, header_note):
    wb = openpyxl.Workbook(); idx = wb.active; idx.title = "00_INDEX"
    idx.append([f"{program} Excel Upload Validation 케이스별 테스트 데이터 — 생성 {STAMP} (gen_excel_upload_testcases.py)"])
    idx.append([header_note]); idx.append([])
    idx.append(["Sheet", "No", "검증ID", "케이스", "기대 메시지(MESSAGE_TEXT)", "행수", "비고 / N/A 사유"])
    for c in range(1, 8): idx.cell(row=4, column=c).font = BOLD; idx.cell(row=4, column=c).fill = HDR_FILL
    os.makedirs(os.path.join(CASE_DIR, program), exist_ok=True)
    for case in cases:
        sn = sheet_name(case)
        idx.append([sn, case["no"], case["cid"], case["title"], case["expect"], (len(case["rows"]) if case["rows"] else 0), case["note"]])
        ws = wb.create_sheet(sn)
        if case["rows"] is None:
            ws.append([f"[N/A] {case['title']} — {case['expect']}"]); ws.append([case["note"]]); continue
        writer(ws, case)
        # 케이스별 단일 시트 파일 (실업로드용 — 첫 시트만 읽는 BE 계약)
        wb1 = openpyxl.Workbook(); ws1 = wb1.active; ws1.title = program if program == "DMF0014" else "Metal Link Price"
        writer(ws1, case); wb1.save(os.path.join(CASE_DIR, program, f"{sn}.xlsx"))
    for col, w in zip("ABCDEFG", (12, 5, 8, 40, 60, 6, 90)): idx.column_dimensions[col].width = w
    idx.freeze_panes = "A5"
    out = os.path.join(OUT_DIR, f"{program}_ExcelUpload_ValidationCases_{STAMP}.xlsx"); wb.save(out); return out

if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    p74 = build("DMF0074", CASES74, write_sheet74,
                "양식: DMF0074_40(53열, division 40). 업로드 = 화면 [Upload](P_MODE=0 검증) → 그리드 MESSAGE_TEXT 확인. BE 는 첫 번째 시트만 읽으므로 케이스 시트를 1번 위치로 옮기거나 _excel_upload_cases/DMF0074/*.xlsx(단일 시트) 를 사용. 값은 전부 텍스트 서식(Period 숫자 서식 금지).")
    p14 = build("DMF0014", CASES14, write_sheet14,
                "양식: DMF0014 'FCST Upload' 1단 헤더(13 식별열 + 주차 시작일 헤더). division 81 기준행. 검증 전부 수행 후 행별 사유를 ' / ' 로 누적 표시(E20260911-DMF0014MULTI). 메시지 = '<규칙> [<헤더>: <입력값>]' (E20260923-DMF0014VALMSG). BE 는 첫 시트만 읽음 → _excel_upload_cases/DMF0014/*.xlsx 사용.")
    print(p74); print(p14)
    print("DMF0074 sheets:", len(CASES74), "| testable:", sum(1 for c in CASES74 if c["rows"]), "| N/A:", sum(1 for c in CASES74 if not c["rows"]))
    print("DMF0014 sheets:", len(CASES14))
    json.dump({"DMF0074": [{k: v for k, v in c.items() if k != "rows"} | {"rows": len(c["rows"]) if c["rows"] else 0} for c in CASES74],
               "DMF0014": [{k: v for k, v in c.items() if k != "rows"} | {"rows": len(c["rows"])} for c in CASES14]},
              open(os.path.join(CASE_DIR, f"cases_{STAMP}.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
