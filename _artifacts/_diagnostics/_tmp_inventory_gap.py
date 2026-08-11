# -*- coding: utf-8 -*-
"""v2.2 취합에 '배선/설정 메타 테이블'이 들어갔는지 정밀 확인 + 누락 family 도출"""
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from ora import conn_info  # noqa: E402
import oracledb  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

P = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI_DB객체_전수취합_v2.2.xlsx")
wb = load_workbook(P, read_only=True, data_only=True)

# 시트1 객체명 / 시트3 TABLE명 집합
s1 = wb["1_UI별_DB객체"]
objs = set()
for r in list(s1.iter_rows(values_only=True))[1:]:
    if r and len(r) > 6 and r[6]:
        objs.add(str(r[6]).strip().upper())
s3 = wb["3_TABLE_이관SOURCE"]
tabs = set()
for r in list(s3.iter_rows(values_only=True))[1:]:
    if r and len(r) > 1 and r[1]:
        tabs.add(str(r[1]).strip().upper())
wb.close()
print("시트1 고유 객체 %d개 / 시트3 고유 TABLE %d개" % (len(objs), len(tabs)))

# ── 후보: BE 범용 코드가 읽는 '배선/설정 메타' 테이블 (SP 의존성에 안 잡히는 계열) ──
META = [
    ("SCM_ADM_CMBO_SQL_M", "콤보 동적 SQL 본문 (/admin/combo)"),
    ("SCM_ADM_API_M", "API_ID → SP/URL 배선"),
    ("SCM_ADM_MNU_BTN_MPG_M", "메뉴×시트×버튼 → API 배선"),
    ("SCM_ADM_PROC_M", "SP 파라미터 순서/타입 메타"),
    ("SCM_MST_SHEET_COLUMNS", "그리드 컬럼 메타"),
    ("SCM_ADM_MSG_M", "메시지 마스터"),
    ("SCM_ADM_MNU_M", "메뉴 마스터"),
    ("SCM_ADM_USER_GR_MNU_BTN_MPG_M", "그룹×버튼 권한"),
    ("SCM_ADM_USER_GR_MPG_M", "사용자×그룹"),
    ("SCM_ADM_USER_GR_UNIT_AUTH_MPG_M", "그룹×조직단위 권한"),
    ("SCM_ADM_USER_MNU_EDIT_AUTH_V", "편집권한 뷰"),
    ("SCM_ADM_LOOKUP_MASTER_DETAIL", "룩업 상세"),
    ("SCM_COM_SYS_LOG", "시스템 로그"),
    ("SCM_COM_SQL_LOG", "SQL 로그"),
    ("SCM_MST_SSO_USER_INFO", "SSO 사용자"),
    ("SCM_MST_SSO_DEPT_INFO", "SSO 부서"),
]
print("\n== 배선/설정 메타 테이블의 취합 포함 여부 ==")
print("  %-34s %-8s %-8s %s" % ("TABLE", "시트1", "시트3", "역할"))
missing = []
for t, desc in META:
    in1 = "O" if t in objs else "—"
    in3 = "O" if t in tabs else "—"
    if in1 == "—" and in3 == "—":
        missing.append((t, desc))
    print("  %-34s %-8s %-8s %s" % (t, in1, in3, desc))

print("\n== 취합에서 완전히 빠진 것 (%d개) ==" % len(missing))
for t, desc in missing:
    print("  ⛔ %-34s %s" % (t, desc))

# ── 콤보 SQL 이 실제로 참조하는 테이블이 취합에 있는지 (2차 누락 확인) ──
u, p, d = conn_info()
con = oracledb.connect(user=u, password=p, dsn=d)
cur = con.cursor()
cur.execute("SELECT SQL_ID, SQL_DTL_CTN FROM SCM_MGR.SCM_ADM_CMBO_SQL_M WHERE USE_YN='Y'")
import re
ref = {}
for sid, body in cur.fetchall():
    txt = (body.read() if hasattr(body, "read") else (body or "")).upper()
    for m in re.finditer(r"\b(?:FROM|JOIN)\s+([A-Z0-9_]+)", txt):
        ref.setdefault(m.group(1), set()).add(sid)
cur.execute("SELECT COUNT(*) FROM SCM_MGR.SCM_ADM_CMBO_SQL_M WHERE USE_YN='Y'")
print("\n== 콤보 SQL %d본이 참조하는 테이블 중 취합 미포함 ==" % cur.fetchone()[0])
gap = sorted((t, len(s)) for t, s in ref.items() if t not in objs and t not in tabs and not t.startswith("DUAL"))
for t, n in gap:
    print("  ⛔ %-38s 콤보 %d본에서 사용" % (t, n))
print("  (참조 테이블 총 %d종 중 미포함 %d종)" % (len(ref), len(gap)))
con.close()
