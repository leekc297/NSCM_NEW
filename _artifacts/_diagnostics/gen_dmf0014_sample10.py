"""DMF0014 Customer FCST Upload — 정상 샘플 10건 생성 + 로컬 BE 업로드 검증(P_MODE=0, 원장 저장 없음).

행 선정(사내 ZIOUI1 실측, 생성 시점):
  DM_MST_CUSTOMER_FCST (81, Sales Site, Measure) VALID · UPLOAD_BUCKET_TYPE=WEEKLY · 헤더 4주 SUMMARY_BUCKET_VALUE 전부 개방
  × DM_MST_CUSTOMER_MODEL (같은 Div/Site) VALID × SCM_MST_SITE(Sales Site) VALID × UPLOAD_UOM 룩업 VALID
  → Sales Site 10곳에서 1건씩 (CELL/PACK/MODULE 혼합). 공통값 = 생성기 기준행(Plant EG01_81-W / FOB / Mass / D / CUST_FCST).
usage: python gen_dmf0014_sample10.py [--upload]
"""
import datetime, os, sys, argparse
import openpyxl
from openpyxl.styles import PatternFill

sys.path.insert(0, r"C:\Projects\NSCM\_artifacts\_diagnostics")
import ora, oracledb

ROOT = r"C:\Projects\NSCM"
OUT_DIR = os.path.join(ROOT, "Doc", "단위테스트", "_excel_upload_cases", "DMF0014_SAMPLE")
H14 = ["Division*", "Sales Site*", "Plant Site*", "Customer Project*", "Customer Model*", "Incoterms*", "FCST Type*",
       "Customs Handbook*", "Measure Group*", "Measure*", "Upload UOM*", "Project", "Model"]
PICK_SITES = ["GM_US_BQN", "HKMC_CZ_NVS", "SGM_CN_PDG", "Siemens_NO_TRD", "MOBIS_KR_CHU",
              "FCA_CA_WND", "FCA_US_SEH", "Renault_FR_HRB", "Ford_US_RDF", "Renault_ES_VLL"]
UOM_LABEL = {"CELL": "Cell", "PACK": "Pack", "MODULE": "Module"}

today = datetime.date.today()
next_mon = today + datetime.timedelta(days=(7 - today.weekday()) % 7 or 7)
WEEKS = [(next_mon + datetime.timedelta(weeks=i)).strftime("%Y%m%d") for i in range(4)]
STAMP = today.strftime("%Y%m%d")


def pick_rows(cur):
    wk = ",".join(f"'{w}'" for w in WEEKS)
    cur.execute(f"""
    WITH CF AS (
      SELECT C.DIVISION_ID, C.SCM_SITE_ID, C.MEASURE, MAX(C.UPLOAD_UOM) UOM
        FROM SCM_MGR.DM_MST_CUSTOMER_FCST C
       WHERE C.DIVISION_ID='81' AND C.VALID_YN='Y' AND C.UPLOAD_BUCKET_TYPE='WEEKLY' AND C.SUMMARY_BUCKET_VALUE IN ({wk})
       GROUP BY C.DIVISION_ID, C.SCM_SITE_ID, C.MEASURE
      HAVING COUNT(DISTINCT C.SUMMARY_BUCKET_VALUE) = 4 AND COUNT(DISTINCT C.UPLOAD_UOM) = 1)
    SELECT CF.SCM_SITE_ID, CF.MEASURE, CF.UOM, CM.CUSTOMER_PROJECT_CODE, CM.CUSTOMER_MODEL, CM.PROJECT_CODE, CM.ITEM_ID
      FROM CF
      JOIN SCM_MGR.SCM_MST_SITE S ON S.DIVISION_ID = CF.DIVISION_ID AND S.SCM_SITE_ID = CF.SCM_SITE_ID AND S.VALID_YN='Y'
      JOIN SCM_MGR.DM_MST_CUSTOMER_MODEL CM ON CM.DIVISION_ID = CF.DIVISION_ID AND CM.SCM_SITE_ID = CF.SCM_SITE_ID AND CM.VALID_YN='Y'
     WHERE EXISTS (SELECT 1 FROM SCM_MGR.SCM_ADM_LOOKUP_MASTER_DETAIL L WHERE L.LOOKUP_TYPE='UPLOAD_UOM_TYPE_CODE' AND L.SEGMENT1 = UPPER(CF.UOM) AND L.VALID_YN='Y')
       AND CM.CUSTOMER_PROJECT_CODE <> 'ABCDE'
     ORDER BY CF.SCM_SITE_ID, CM.CUSTOMER_PROJECT_CODE, CM.CUSTOMER_MODEL""")
    first = {}
    for site, measure, uom, cprj, cmdl, prj, item in cur.fetchall():
        first.setdefault(site, (site, measure, uom, cprj, cmdl, prj, item))
    rows = []
    for i, site in enumerate(PICK_SITES):
        assert site in first, f"{site} — 이번 주차 기준 유효 조합 없음(재선정 필요)"
        site, measure, uom, cprj, cmdl, prj, item = first[site]
        rows.append({"Division*": "81", "Sales Site*": site, "Plant Site*": "EG01_81-W", "Customer Project*": cprj,
                     "Customer Model*": cmdl, "Incoterms*": "FOB", "FCST Type*": "Mass", "Customs Handbook*": "D",
                     "Measure Group*": "CUST_FCST", "Measure*": measure, "Upload UOM*": UOM_LABEL.get(uom.upper(), uom),
                     "Project": prj or "", "Model": item or "",
                     "_qty": [str(100 * (i + 1) + k * 10) for k in range(4)]})
    return rows


def write(rows):
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, f"DMF0014_Upload_Sample10_{STAMP}.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "FCST Upload"
    fill = PatternFill("solid", fgColor="DDEBF7")
    for c, h in enumerate(H14 + WEEKS, 1):
        cell = ws.cell(row=1, column=c, value=h); cell.number_format = "@"; cell.fill = fill
    for r, row in enumerate(rows, 2):
        for c, v in enumerate([row[h] for h in H14] + row["_qty"], 1):
            ws.cell(row=r, column=c, value=v).number_format = "@"
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def upload(path, base="http://localhost:8080", user="minyonglee1"):
    import requests
    s = requests.Session()
    r = s.post(f"{base}/user/login", json={"userId": user}, timeout=30); r.raise_for_status()
    with open(path, "rb") as fh:
        resp = s.post(f"{base}/DMF0014/FCST Upload/UPLOAD/excel-upload",
                      files={"file": (os.path.basename(path), fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, timeout=300)
    print("HTTP", resp.status_code, "| code", (resp.json() if resp.headers.get("content-type", "").startswith("application/json") else {}).get("code"))
    return resp


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--upload", action="store_true"); a = ap.parse_args()
    u, p, d = ora.conn_info()
    with oracledb.connect(user=u, password=p, dsn=d) as con:
        cur = con.cursor()
        rows = pick_rows(cur)
        path = write(rows)
        print("weeks", WEEKS, "->", path)
        for r in rows:
            print("  ", r["Sales Site*"], "|", r["Customer Project*"], "|", r["Customer Model*"], "|", r["Measure*"], "|", r["Upload UOM*"], "|", r["_qty"])
        if a.upload:
            upload(path)
            cur.execute("""SELECT NVL(UPLOAD_STATUS_CODE,'-'), NVL(MESSAGE_TEXT,'(none)'), COUNT(*) FROM SCM_MGR.SCM_ADM_EXCEL_UPLOAD
                            WHERE PROGRAM_ID='DMF0014' AND SHEET_NAME='FCST Upload' AND CREATED_BY='minyonglee1'
                              AND CREATION_DATE = (SELECT MAX(CREATION_DATE) FROM SCM_MGR.SCM_ADM_EXCEL_UPLOAD WHERE PROGRAM_ID='DMF0014' AND SHEET_NAME='FCST Upload' AND CREATED_BY='minyonglee1')
                              AND NVL(UPLOAD_STATUS_CODE,'-') NOT IN ('V_TITLE','@TITLE')
                            GROUP BY UPLOAD_STATUS_CODE, MESSAGE_TEXT ORDER BY 1""")
            print("staging result:", cur.fetchall())


if __name__ == "__main__":
    main()
