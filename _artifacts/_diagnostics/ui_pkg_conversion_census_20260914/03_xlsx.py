"""03_xlsx.py — 02_build 산출 CSV → Doc/화면분석서/TOBE_UI별_PKG전환대상_20260914.xlsx"""
import csv, pathlib, collections, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = pathlib.Path(__file__).resolve().parent
OUT = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI별_PKG전환대상_20260914.xlsx")
R = lambda n: list(csv.DictReader((D / f"{n}.csv").open(encoding="utf-8-sig")))
sp, ui, ledger, sim = R("out_sp"), R("out_ui"), R("out_ledger"), R("out_similar")

HEAD = PatternFill("solid", fgColor="1F4E78"); HF = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {"PKG전환대상": "E2EFDA", "PKG전환대상(내부호출)": "FFF2CC", "공통": "DDEBF7", "미사용": "F2F2F2", "배치": "FCE4D6"}

wb = Workbook(); wb.remove(wb.active)
def sheet(name, rows, cols=None, widths=None, fill_key=None, note=None):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        for line in note:
            ws.cell(r0, 1, line).font = Font(italic=True, color="595959"); r0 += 1
        r0 += 1
    cols = cols or (list(rows[0].keys()) if rows else ["(해당 없음)"])
    for j, c in enumerate(cols, 1):
        x = ws.cell(r0, j, c); x.fill = HEAD; x.font = HF; x.alignment = Alignment(vertical="center", wrap_text=True); x.border = BORDER
    for i, row in enumerate(rows, r0 + 1):
        color = None
        if fill_key:
            v = row.get(fill_key, "")
            color = next((FILL[k] for k in FILL if v.startswith(k)), None)
        for j, c in enumerate(cols, 1):
            val = row.get(c, "")
            if isinstance(val, str) and val.isdigit() and len(val) < 10: val = int(val)
            x = ws.cell(i, j, val); x.border = BORDER; x.alignment = Alignment(vertical="top", wrap_text=False)
            if color: x.fill = PatternFill("solid", fgColor=color)
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{max(r0, r0 + len(rows))}"
    for j, c in enumerate(cols, 1):
        w = (widths or {}).get(c)
        if w is None:
            m = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:300]])
            w = min(max(10, m * 1.1), 60)
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws

# ── 0_요약 ──
cnt = collections.Counter(r["판정"] for r in sp)
conv = [r for r in sp if r["판정"].startswith("PKG전환대상")]
summary = []
order = ["PKG전환대상", "PKG전환대상(내부호출)", "등록됐으나 FE 미사용(배선 정리·삭제 후보)", "등록됐으나 FE 미사용(타 UI·배치 사용 — 보류)", "FE 화면 미구현 메뉴 사용(전환 보류)",
         "공통SP", "공통SP(텍스트참조 UI 포함)", "공통SP(UI+배치)", "프레임워크공통(BE 매퍼 직접호출)",
         "범위외메뉴전용", "배치전용(UI 미도달)", "API_M/EXCEL 등록만(버튼·메뉴 미배선)", "내부호출·콤보·로그 참조만(UI·배치 미도달)", "미사용"]
desc = {
 "PKG전환대상": "FE 가 실제 호출하는(시트명 확인) 등록 API 가 standalone 을 **직접** 호출. 사용 UI 1개·배치 JOB 미도달 → PKG_{UI} 합류 후 standalone 삭제",
 "PKG전환대상(내부호출)": "FE 실호출 등록 API 의 PKG/SP 본문이 내부에서 호출(멤버 단위 특정). 사용 UI 1개·배치 JOB 미도달 → 합류 후보(호출부 PKG 한정명 변경 필요)",
 "등록됐으나 FE 미사용(배선 정리·삭제 후보)": "메뉴 버튼에는 등록돼 있으나 FE 코드에서 해당 시트 호출이 없음 → 전환이 아니라 BTN_MPG/API_M 정리 후 삭제 검토",
 "등록됐으나 FE 미사용(타 UI·배치 사용 — 보류)": "FE 미사용 등록이지만 다른 UI 등록 또는 배치 JOB 도달 → 삭제 보류",
 "FE 화면 미구현 메뉴 사용(전환 보류)": "등록 메뉴가 FE 화면 미구현(SCP0001·SCP0006·SCP0007 등) → 화면 구현 시 재판정",
 "공통SP": "사용 UI 2개 이상(범위 외 SDM·비활성 메뉴 포함) → PKG 전환 대상 제외",
 "공통SP(텍스트참조 UI 포함)": "의존 딕셔너리 밖 텍스트 참조(동적 SQL 등)로 다른 UI 객체도 사용 → 공통",
 "공통SP(UI+배치)": "UI 1개 + 스케줄러 JOB 도달 → UI PKG 로 옮기면 배치가 깨짐 → 공통",
 "프레임워크공통(BE 매퍼 직접호출)": "nscm-be sql/*.xml 에서 직접 {call} → 판정 제외",
 "범위외메뉴전용": "SDM·비활성 화면 등 범위 밖 메뉴만 사용",
 "배치전용(UI 미도달)": "UI 에서 도달하지 않고 스케줄러 JOB 에서만 도달",
 "API_M/EXCEL 등록만(버튼·메뉴 미배선)": "API_M/EXCEL_COLUMNS 에는 있으나 어떤 메뉴 버튼에도 배선 안 됨",
 "내부호출·콤보·로그 참조만(UI·배치 미도달)": "다른 객체 본문·콤보 SQL·실행 로그·텍스트에만 등장, UI·배치 경로 없음",
 "미사용": "API_M·PROC_M 제외 모든 축(메뉴·Excel·콤보·의존·텍스트·JOB·로그) 참조 0",
}
for k in order + [k for k in cnt if k not in order]:
    summary.append({"판정": k, "건수": cnt.get(k, 0), "의미": desc.get(k, "")})
summary.append({"판정": "합계(SCM_MGR standalone PROCEDURE)", "건수": len(sp), "의미": ""})
by_pkg = collections.Counter(r["통합대상PKG"] for r in conv)
ws = sheet("0_요약", summary, ["판정", "건수", "의미"], {"판정": 42, "건수": 8, "의미": 120}, note=[
    f"UI별 PKG 전환대상 목록 — 사내 ZIOUI1 SCM_MGR 기준 / 산출 {datetime.date.today()} / DB 조회만(변경 없음)",
    "범위: SCM_ADM_MNU_M SCREEN·USE_YN='Y'·SDM 제외 50화면 + 그 화면 FE 코드가 참조하는 팝업 메뉴(COM0001~0012, DMF0074_40) = 63 메뉴",
    "판정 단위 = SP 가 등록된 메뉴(MNU_ID) — 이름 접두어로 소유 UI 를 추정하지 않음. DMF0043_01 → UI DMF0043",
    "근거 축: A1 등록 API(BTN_MPG→API_M) / A2 FE 실호출(주석 제거 코드의 메뉴ID·시트명 리터럴) / A3 EXCEL_COLUMNS.STORAGE_SQL_ID / A4 ALL_DEPENDENCIES+ALL_SOURCE 멤버 단위 호출 /",
    "        A5 SCM_COM_SQL_LOG(09-03~) / A6 스케줄러 JOB 도달(113 JOB, 객체 의존 폐포) / A7 텍스트 참조(SCM_MGR 전체 소스·nscm-be 코드, 동적 SQL 포착)",
])
r = ws.max_row + 2
ws.cell(r, 1, "통합 대상 PKG 별 전환 건수").font = Font(bold=True); r += 1
for k, v in sorted(by_pkg.items()):
    ws.cell(r, 1, k); ws.cell(r, 2, v)
    ws.cell(r, 3, "PKG 존재" if next(x for x in conv if x["통합대상PKG"] == k)["PKG존재"] == "Y" else "PKG 신규 생성 필요"); r += 1

# ── 0_UI별요약 ──
sheet("0_UI별요약", ui, widths={"사용PKG": 40, "직접standalone": 50, "내부호출standalone": 50, "전환대상SP": 60, "공통SP": 50, "설명": 28})

# ── 1_UI별_사용객체 (범위 메뉴 배선, 프레임워크 제외) ──
l1 = [x for x in ledger if x["SCOPE"] == "Y" and x["FRAMEWORK"] == "N"]
l1.sort(key=lambda x: (x["UI"], x["MNU_ID"], x["SHT_NM"], x["BTN"]))
sheet("1_UI별_사용객체", l1, ["UI", "MNU_ID", "MNU_TP", "SRC", "SHT_NM", "BTN", "BTN_USE", "API_ID", "API_TP", "API_USE", "API_URL", "KIND", "PKG", "MEMBER", "STANDALONE",
                            "PROC_M", "FE_MENU", "FE_SHEET", "LOG_MENU", "LOG_ALL"], widths={"API_URL": 55, "API_ID": 40})

# ── 2_PKG전환대상 ──
c2 = ["통합대상PKG", "PKG존재", "현재멤버수", "SP", "판정", "STATUS", "FE확인메뉴", "등록시트", "근거", "유사명_PKG멤버", "열어보는화면", "LOG건수", "LOG최종"]
conv.sort(key=lambda x: (x["통합대상PKG"], x["판정"], x["SP"]))
sheet("2_PKG전환대상", conv, c2, {"근거": 90, "SP": 50, "등록시트": 60}, fill_key="판정", note=[
    "PKG전환대상 = FE 실호출 등록 API 가 standalone 을 직접 호출 / PKG전환대상(내부호출) = 그 PKG·SP 본문이 호출. 둘 다 사용 UI 1개 + 배치 JOB 미도달 + 타 UI 텍스트 참조 없음",
    "등록시트 (FE N) = 같은 SP 가 FE 에서 쓰지 않는 시트에도 등록돼 있음 → 전환 시 함께 정리",
    "유사명_PKG멤버 = 같은 이름의 PKG 멤버가 이미 존재 → 신규 합류 대신 기존 멤버 재배선 가능(본문 동일성 확인, 4_유사명_중복판정 참조)",
    "전환 방식(07-23 DMF0021 선례): 본문 무수정으로 PKG 멤버 추가 → API_M/EXCEL/호출부를 PKG 한정명으로 재배선 → PROC_M 정리 → 참조 0 확인 후 standalone DROP",
])

# ── 3_통합후_삭제대상SP ── (전환 후 삭제 + FE 미사용 등록 정리 + 유사명 standalone 삭제 가능)
d3 = []
sim_del = {r["STANDALONE"]: r for r in sim if r["판정"].startswith("PKG 멤버만 FE 사용")}
fe_unused = [x for x in sp if x["판정"] == "등록됐으나 FE 미사용(배선 정리·삭제 후보)"]
extra = [x for x in sp if x["SP"] in sim_del and not x["판정"].startswith("PKG전환대상") and x not in fe_unused]
for x in conv + fe_unused + extra:
    kind = ("전환 후 삭제" if x["판정"].startswith("PKG전환대상") else
            "FE 미사용 등록 정리 후 삭제" if x in fe_unused else "유사명 PKG 멤버가 실사용 — standalone 삭제")
    need = []
    if x["REF_API_M"]: need.append(f"API_M 재배선({x['REF_API_M']})")
    if int(x["REF_PROC_M"] or 0): need.append(f"PROC_M {x['REF_PROC_M']}행 → PKG 한정명 재등록")
    if x["REF_EXCEL"]: need.append(f"EXCEL STORAGE_SQL_ID 재배선({x['REF_EXCEL']})")
    if x["REF_CMBO"]: need.append(f"콤보 SQL 수정({x['REF_CMBO']})")
    if x["REF_호출객체_LIVE"]: need.append(f"호출부 PKG 한정명 변경({x['REF_호출객체_LIVE']})")
    if x["REF_텍스트_의존외"]: need.append(f"텍스트 참조 확인({x['REF_텍스트_의존외']})")
    if x["SCMP_APP_시노님"] == "Y": need.append("SCMP_APP 시노님 DROP(PKG 시노님·EXECUTE 권한 확인)")
    if x["배치JOB_도달수"] not in ("", "0"): need.append(f"배치 JOB {x['배치JOB_도달수']}개 도달 — 삭제 불가")
    d3.append(dict(삭제구분=kind, 삭제대상SP=x["SP"], 통합대상PKG=x["통합대상PKG"], 판정=x["판정"], 유사명_PKG멤버=x["유사명_PKG멤버"],
        REF_API_M=x["REF_API_M"], REF_PROC_M=x["REF_PROC_M"], REF_EXCEL=x["REF_EXCEL"], REF_CMBO=x["REF_CMBO"],
        REF_호출객체_LIVE=x["REF_호출객체_LIVE"], REF_의존만=x["REF_의존만"], REF_텍스트_의존외=x["REF_텍스트_의존외"],
        배치JOB_도달수=x["배치JOB_도달수"], SCMP_APP_시노님=x["SCMP_APP_시노님"], LOG건수=x["LOG건수"],
        DROP전_선행작업=" / ".join(need) or "없음(참조 0)", 삭제조건="선행작업 완료 후 참조 0 재확인(API_M·PROC_M·EXCEL·CMBO·DEPENDENCIES·ALL_SOURCE·JOB)"))
sheet("3_통합후_삭제대상SP", d3, widths={"DROP전_선행작업": 90, "삭제조건": 60, "삭제대상SP": 50}, fill_key="판정")

# ── 4_유사명_중복판정 ──
sheet("4_유사명_중복판정", sim, widths={"standalone_근거": 70, "멤버_근거": 70, "판정": 55}, note=[
    "대상 = standalone X 와 PKG 멤버 X / P_X 가 동시에 존재하는 쌍. 근거 = 메뉴 등록(A1/A3)·내부 호출(A4 멤버 단위)·로그(A5)",
    "사용 쪽 유지·미사용 쪽 삭제 가능. 둘 다 사용 = 배선 이원화(한쪽으로 재배선 후 정리)",
])

# ── 5_공통SP ──
c5 = [x for x in sp if x["판정"].startswith("공통") or x["판정"].startswith("프레임워크") or x["판정"] == "범위외메뉴전용"]
sheet("5_공통SP", c5, ["SP", "판정", "사용UI", "범위내메뉴", "범위외메뉴", "텍스트참조_추가UI", "배치JOB_도달수", "배치JOB_직접", "근거", "REF_API_M", "REF_호출객체_LIVE",
                      "REF_텍스트_전체", "LOG건수"], {"근거": 80, "사용UI": 50, "REF_텍스트_전체": 60}, fill_key="판정")

# ── 6_미사용SP(참고) ──
c6 = [x for x in sp if x["판정"] in ("미사용", "배치전용(UI 미도달)", "API_M/EXCEL 등록만(버튼·메뉴 미배선)", "내부호출·콤보·로그 참조만(UI·배치 미도달)",
                                   "FE 화면 미구현 메뉴 사용(전환 보류)", "등록됐으나 FE 미사용(타 UI·배치 사용 — 보류)")]
c6.sort(key=lambda x: (order.index(x["판정"]) if x["판정"] in order else 99, x["SP"]))
sheet("6_UI미사용SP_참고", c6, ["SP", "판정", "STATUS", "사용UI", "등록시트", "유사명_PKG멤버", "REF_API_M", "REF_PROC_M", "REF_EXCEL", "REF_CMBO", "REF_호출객체_LIVE", "REF_의존만", "REF_텍스트_전체",
                            "배치JOB_도달수", "배치JOB_직접", "배치로그건수", "LOG건수", "SCMP_APP_시노님"], {"REF_텍스트_전체": 60, "REF_호출객체_LIVE": 50}, fill_key="판정",
      note=["삭제 판단은 사용자 몫 — '미사용'도 AS-IS 이관 배치·IF 가능성은 JOB·텍스트 참조로만 확인했음. 배치전용은 삭제 대상 아님"])

# ── 7_근거_배선원장 ──
sheet("7_근거_배선원장", sorted(ledger, key=lambda x: (x["SCOPE"] != "Y", x["UI"], x["MNU_ID"], x["SHT_NM"], x["BTN"])), widths={"API_URL": 55, "API_ID": 40})

# ── 8_등록객체부재 ──
m8 = [x for x in ledger if x["KIND"] in ("MISSING", "PKG_MEMBER_MISSING", "UNKNOWN_QUALIFIED", "API_M_NONE", "EMPTY")]
sheet("8_등록객체부재", m8, ["SCOPE", "UI", "MNU_ID", "SRC", "SHT_NM", "BTN", "BTN_USE", "API_ID", "API_USE", "API_URL", "KIND", "FE_SHEET", "LOG_ALL"],
      {"API_URL": 55, "API_ID": 40}, note=["메뉴 버튼·Excel 에 등록됐으나 DB 객체가 없음(또는 API_M 행 없음·URL 공란) — 호출 시 실패하는 배선. 참고 목록"])

# ── 9_방법·한계 ──
lim = [
 {"항목": "UI↔메뉴", "내용": "FE 화면 = pages/<MNU_URL> 폴더(viewRegistry 키). 팝업 메뉴 = 주석 제거 FE 코드에서 메뉴ID 리터럴 사용 시 범위 포함"},
 {"항목": "FE 시트 확인", "내용": "메뉴ID 리터럴 파일이 속한 pages 폴더/컴포넌트 폴더 전체에서 시트명 리터럴 탐색. 상수 조합·동적 문자열로 만든 시트명은 N 으로 나올 수 있음(참고 플래그)"},
 {"항목": "내부 호출", "내용": "ALL_DEPENDENCIES 로 호출 후보 → ALL_SOURCE 본문(주석 제거)에서 멤버 단위로 호출 위치 특정 → 멤버가 등록 API 가 아니면 같은 패키지/타 객체 호출부로 역추적"},
 {"항목": "배치", "내용": "ALL_SCHEDULER_JOBS 113개(사내 전부 ENABLED=FALSE — AS-IS 이관 배치) JOB_ACTION 호출 객체 → 객체 의존 폐포. 패키지 단위라 넓게 잡힘(이동·삭제를 막는 보수 방향)"},
 {"항목": "배치 로그", "내용": "SCM_SYS_LOG_BATCH.PROGRAM_NAME 은 UI SP 도 기록 → 배치 근거로 쓰지 않고 참고 컬럼만"},
 {"항목": "텍스트 참조", "내용": "SCM_MGR 전체 소스·nscm-be(java/xml/yml) 에서 이름 단어 일치(주석 제거). 의존 딕셔너리 밖 참조가 타 UI 패키지면 공통으로 격상"},
 {"항목": "미확인 축", "내용": "타 DB/DB Link·외부 IF 시스템·클라우드 전용 객체·FE 외 클라이언트 호출은 확인 불가. 클라우드는 사내와 객체 차이가 있을 수 있음 → 전환 전 클라우드 동일 조회 필요"},
 {"항목": "로그 기간", "내용": "SCM_COM_SQL_LOG 2026-09-03 이후(사내 로컬 BE 실행분) — 0 이 미사용을 뜻하지 않음"},
]
sheet("9_방법_한계", lim, ["항목", "내용"], {"항목": 16, "내용": 150})

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("saved", OUT, {ws.title: ws.max_row for ws in wb.worksheets})
