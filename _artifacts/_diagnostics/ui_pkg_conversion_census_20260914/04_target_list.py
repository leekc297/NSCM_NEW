"""04_target_list.py — 사용자 확인용 'PKG 취합 대상 목록' EXCEL + MD (02_build 산출 CSV 기반, DB 는 권한 조회만)
출력: Doc/화면분석서/TOBE_UI별_PKG전환대상_20260914.xlsx / .md
"""
import csv, pathlib, collections, datetime, re
import oracledb
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = pathlib.Path(__file__).resolve().parent
OUT_X = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI별_PKG전환대상_20260914.xlsx")
OUT_M = OUT_X.with_suffix(".md")
R = lambda n: list(csv.DictReader((D / f"{n}.csv").open(encoding="utf-8-sig")))
sp, ui, ledger, sim, mnu = R("out_sp"), R("out_ui"), R("out_ledger"), R("out_similar"), R("mnu_m")
mdesc = {r["MNU_ID"]: r["MNU_DESC"] for r in mnu}
TODAY = datetime.date.today().isoformat()

# ── DB: 권한·상태 (조회만) ──
con = oracledb.connect(user="SCM_MGR", password="zionex1!", dsn="175.211.29.206:2521/ZIOUI1"); cur = con.cursor()
cur.execute("SELECT TABLE_NAME, GRANTEE, PRIVILEGE FROM ALL_TAB_PRIVS WHERE TABLE_SCHEMA='SCM_MGR' AND GRANTEE='SCMP_APP' AND PRIVILEGE='EXECUTE'")
app_exec = {t for t, _, _ in cur}
cur.execute("SELECT SYNONYM_NAME FROM ALL_SYNONYMS WHERE OWNER='SCMP_APP' AND TABLE_OWNER='SCM_MGR'")
app_syn = {r[0] for r in cur}
cur.execute("SELECT OBJECT_NAME FROM ALL_OBJECTS WHERE OWNER='SCM_MGR' AND OBJECT_TYPE='PACKAGE'")
pkgs = {r[0] for r in cur}
cur.execute("SELECT NAME, COUNT(*) FROM ALL_ERRORS WHERE OWNER='SCM_MGR' GROUP BY NAME")
errs = dict(cur.fetchall())
cur.execute("SELECT SYS_CONTEXT('USERENV','CON_NAME'), SYS_CONTEXT('USERENV','SERVER_HOST') FROM DUAL")
dbid = cur.fetchone()

sim_by = collections.defaultdict(list)
for r in sim: sim_by[r["STANDALONE"]].append(r)
ui_of = lambda m: re.sub(r"_\d+$", "", m)

# ───────────────────────── 1. PKG 취합 대상 ─────────────────────────
conv = sorted([r for r in sp if r["판정"].startswith("PKG전환대상")], key=lambda r: (r["통합대상PKG"], r["판정"], r["SP"]))
work = []
for i, r in enumerate(conv, 1):
    tgt = r["통합대상PKG"]
    member = r["SP"] if r["SP"].startswith("P_") else "P_" + r["SP"]
    same = [s for s in sim_by.get(r["SP"], []) if s["PKG"] == tgt]
    other = [s for s in sim_by.get(r["SP"], []) if s["PKG"] != tgt]
    if same and same[0]["래퍼"] == "Y":
        wtype, member_note = "래퍼 로직을 기존 멤버로 흡수", f"{tgt}.{same[0]['MEMBER']} (기존)"
    elif same:
        wtype, member_note = "기존 동일명 멤버로 재배선", f"{tgt}.{same[0]['MEMBER']} (기존 — 본문 동일성 확인)"
    elif other:
        wtype, member_note = "멤버 신규 추가 (타 PKG 동일명 멤버 존재 — 재배선 대안)", f"{tgt}.{member} (안) / 대안 {other[0]['PKG']}.{other[0]['MEMBER']}"
    else:
        wtype, member_note = "멤버 신규 추가", f"{tgt}.{member} (안)"
    rewire = []
    if r["REF_API_M"]: rewire.append(f"API_M: {r['REF_API_M']}")
    if int(r["REF_PROC_M"] or 0): rewire.append(f"PROC_M {r['REF_PROC_M']}행")
    if r["REF_EXCEL"]: rewire.append(f"EXCEL: {r['REF_EXCEL']}")
    if r["REF_CMBO"]: rewire.append(f"콤보: {r['REF_CMBO']}")
    if r["REF_호출객체_LIVE"]: rewire.append(f"호출부 변경: {r['REF_호출객체_LIVE']}")
    if r["REF_텍스트_의존외"]: rewire.append(f"텍스트참조 확인: {r['REF_텍스트_의존외']}")
    fe_unused = [x for x in r["등록시트"].split(",") if "(FE N)" in x]
    prep = []
    if tgt not in pkgs: prep.append(f"{tgt} 신규 생성 + SCMP_APP EXECUTE 권한·시노님")
    elif tgt not in app_exec: prep.append(f"{tgt} SCMP_APP EXECUTE 권한")
    if r["STATUS"] != "VALID": prep.append(f"standalone INVALID(ALL_ERRORS {errs.get(r['SP'], 0)}건) — 재컴파일로 정상 여부 먼저 확인")
    work.append({
        "No": i, "통합대상 PKG": tgt, "PKG 상태": "기존" if tgt in pkgs else "신규 생성",
        "UI": ui_of(r["FE확인메뉴"].split(",")[0]) if r["FE확인메뉴"] else tgt[4:],
        "화면명": mdesc.get(tgt[4:], ""),
        "standalone SP": r["SP"], "SP 상태": r["STATUS"],
        "호출 구분": "화면 API 직접 호출" if r["판정"] == "PKG전환대상" else "UI PKG 내부 호출",
        "작업 유형": wtype, "PKG 멤버(안)": member_note,
        "재배선 대상": " / ".join(rewire) or "-",
        "FE 미사용 등록(함께 정리)": ", ".join(x.replace("(FE N)", "") for x in fe_unused) or "-",
        "선행 작업": " / ".join(prep) or "-",
        "호출 경로(근거)": r["근거"][:400],
        "실행로그(09-03~)": int(r["LOG건수"] or 0),
    })

# ───────────────────────── 2. DROP 대상 ─────────────────────────
drop = []
def ref_counts(r):
    return dict(API_M=len([x for x in r["REF_API_M"].split(",") if x]), PROC_M=int(r["REF_PROC_M"] or 0),
                EXCEL=len([x for x in r["REF_EXCEL"].split(",") if x]), 콤보=len([x for x in r["REF_CMBO"].split(",") if x]),
                호출객체=len([x for x in r["REF_호출객체_LIVE"].split(",") if x]), 배치JOB=int(r["배치JOB_도달수"] or 0),
                텍스트=len([x for x in r["REF_텍스트_의존외"].split(",") if x]))
for r in conv:
    drop.append((1, "① PKG 취합 후 DROP", r, f"{r['통합대상PKG']} 재배선 완료 후"))
for r in sp:
    if r["판정"] == "등록됐으나 FE 미사용(배선 정리·삭제 후보)":
        drop.append((2, "② FE 미사용 등록 정리 후 DROP", r, "BTN_MPG·API_M 등록 정리 후 (FE 호출 없음)"))
for s in sim:
    r = next(x for x in sp if x["SP"] == s["STANDALONE"])
    if s["판정"].startswith("PKG 멤버만 FE 사용") and not r["판정"].startswith(("PKG전환대상", "등록됐으나 FE 미사용(배선")):
        drop.append((3, "③ 동일명 PKG 멤버가 실사용 — standalone DROP", r, f"실사용 = {s['PKG']}.{s['MEMBER']}"))
drop_rows = []
for i, (_, kind, r, cond) in enumerate(sorted(drop, key=lambda x: (x[0], x[2]["통합대상PKG"], x[2]["SP"])), 1):
    rc = ref_counts(r)
    row = {"No": i, "DROP 구분": kind, "standalone SP": r["SP"], "SP 상태": r["STATUS"], "관련 PKG/UI": r["통합대상PKG"] or r["사용UI"], "DROP 전제": cond}
    row.update({f"현재참조_{k}": v for k, v in rc.items()})
    row["SCMP_APP 시노님"] = "Y(함께 DROP)" if r["SP"] in app_syn else "N"
    row["DROP 가능 조건"] = "선행 재배선 후 현재참조 전 축 0 재확인 (API_M·PROC_M·EXCEL·콤보·호출객체·배치JOB·텍스트) → 백업 DDL 보관 후 DROP"
    drop_rows.append(row)

# ───────────────────────── 3. 제외(공통) / 4. 보류 / 5. 참고 ─────────────────────────
reason = {"공통SP": "여러 UI 에서 사용", "공통SP(UI+배치)": "UI 1개 + 스케줄러 JOB 에서도 호출 — UI PKG 로 옮기면 배치 영향",
          "공통SP(텍스트참조 UI 포함)": "다른 UI 객체에서 동적 참조", "프레임워크공통(BE 매퍼 직접호출)": "BE 매퍼(sql/*.xml)가 직접 호출하는 공용"}
excl = [{"No": i, "standalone SP": r["SP"], "분류": r["판정"], "제외 사유": reason.get(r["판정"], ""),
         "사용 UI": r["사용UI"], "배치 JOB 수": int(r["배치JOB_도달수"] or 0), "배치 JOB(직접)": r["배치JOB_직접"], "SP 상태": r["STATUS"]}
        for i, r in enumerate(sorted([x for x in sp if x["판정"] in reason], key=lambda x: (x["판정"], x["SP"])), 1)]
hold_kinds = {"FE 화면 미구현 메뉴 사용(전환 보류)": "등록 메뉴의 FE 화면이 아직 없음 — 화면 구현 후 재판정",
              "등록됐으나 FE 미사용(타 UI·배치 사용 — 보류)": "FE 호출은 없으나 다른 UI 등록·배치 JOB 도달",
              "범위외메뉴전용": "SDM·비활성 화면 등 조사 범위 밖 메뉴만 사용"}
hold = [{"No": i, "standalone SP": r["SP"], "분류": r["판정"], "보류 사유": hold_kinds[r["판정"]], "등록 메뉴": r["범위내메뉴"] or r["범위외메뉴"],
         "호출 경로(근거)": r["근거"][:300], "배치 JOB 수": int(r["배치JOB_도달수"] or 0), "SP 상태": r["STATUS"]}
        for i, r in enumerate(sorted([x for x in sp if x["판정"] in hold_kinds], key=lambda x: (x["판정"], x["SP"])), 1)]
ref_kinds = ["배치전용(UI 미도달)", "API_M/EXCEL 등록만(버튼·메뉴 미배선)", "내부호출·콤보·로그 참조만(UI·배치 미도달)", "미사용"]
ref = [{"No": i, "standalone SP": r["SP"], "분류": r["판정"], "SP 상태": r["STATUS"], "배치 JOB 수": int(r["배치JOB_도달수"] or 0),
        "API_M": r["REF_API_M"], "호출 객체": r["REF_호출객체_LIVE"], "텍스트 참조": r["REF_텍스트_전체"][:200], "실행로그": int(r["LOG건수"] or 0)}
       for i, r in enumerate(sorted([x for x in sp if x["판정"] in ref_kinds], key=lambda x: (ref_kinds.index(x["판정"]), x["SP"])), 1)]
simrows = [{"No": i, "standalone SP": s["STANDALONE"], "동일명 PKG 멤버": f"{s['PKG']}.{s['MEMBER']}", "래퍼": s["래퍼"],
            "판정": s["판정"], "standalone 사용 근거": s["standalone_근거"], "PKG 멤버 사용 근거": s["멤버_근거"]}
           for i, s in enumerate(sim, 1)]
uirows = [{"메뉴": u["MNU_ID"], "UI": u["UI"], "유형": u["유형"], "화면명": u["설명"], "FE 화면": u["FE페이지"],
           "사용 PKG": u["사용PKG"], "취합 대상 수": int(u["PKG전환대상"] or 0), "취합 대상 SP": u["전환대상SP"],
           "공통 SP 사용": u["공통SP"], "목표 PKG": u["목표PKG"], "목표 PKG 존재": u["목표PKG존재"], "등록객체 부재": int(u["등록객체부재"] or 0)}
          for u in ui]

# ───────────────────────── EXCEL ─────────────────────────
HEAD = PatternFill("solid", fgColor="1F4E78"); HF = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF"); BD = Border(left=thin, right=thin, top=thin, bottom=thin)
BANDS = ["FFFFFF", "EEF3F8"]
wb = Workbook(); wb.remove(wb.active)

def sheet(name, rows, note=None, widths=None, band_key=None, color_key=None, colors=None, tab=None):
    ws = wb.create_sheet(name)
    if tab: ws.sheet_properties.tabColor = tab
    r0 = 1
    for line in (note or []):
        ws.cell(r0, 1, line).font = Font(color="404040", italic=True); r0 += 1
    if note: r0 += 1
    cols = list(rows[0].keys()) if rows else ["(해당 없음)"]
    for j, c in enumerate(cols, 1):
        x = ws.cell(r0, j, c); x.fill = HEAD; x.font = HF; x.border = BD; x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 30
    band, prev = 0, None
    for i, row in enumerate(rows, r0 + 1):
        if band_key:
            if row[band_key] != prev: band ^= 1; prev = row[band_key]
            fill = PatternFill("solid", fgColor=BANDS[band])
        else:
            fill = None
        if color_key and colors:
            v = str(row.get(color_key, ""))
            col = next((c for k, c in colors.items() if v.startswith(k)), None)
            if col: fill = PatternFill("solid", fgColor=col)
        for j, c in enumerate(cols, 1):
            x = ws.cell(i, j, row.get(c, "")); x.border = BD; x.alignment = Alignment(vertical="top", wrap_text=True)
            if fill: x.fill = fill
    ws.freeze_panes = ws.cell(r0 + 1, 3 if len(cols) > 3 else 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(cols))}{r0 + max(len(rows), 1)}"
    for j, c in enumerate(cols, 1):
        w = (widths or {}).get(c)
        if w is None:
            m = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
            w = min(max(8, m * 1.15), 45)
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws

# 0. 요약
cnt = collections.Counter(r["판정"] for r in sp)
ws = wb.create_sheet("0_요약"); ws.sheet_properties.tabColor = "1F4E78"
ws["A1"] = "UI별 PKG 취합 대상 목록"; ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"기준: 사내 {dbid[0]}(host {dbid[1]}) SCM_MGR / 산출 {TODAY} / 조회만 — DB 변경 없음"
ws["A3"] = "범위: 활성 SCREEN 50화면(SDM 제외) + 화면이 여는 팝업 메뉴 13 / 판정 = FE 가 실제 호출하는 등록 API 기준(이름 추정 없음)"
r = 5
def block(title, items, head=("구분", "건수", "설명")):
    global r
    ws.cell(r, 1, title).font = Font(bold=True, size=12); r += 1
    for j, h in enumerate(head, 1):
        x = ws.cell(r, j, h); x.fill = HEAD; x.font = HF; x.border = BD
    r += 1
    for it in items:
        for j, v in enumerate(it, 1):
            x = ws.cell(r, j, v); x.border = BD; x.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    r += 1
g = lambda *ks: sum(cnt.get(k, 0) for k in ks)
block("① 한눈에 보기 (standalone SP 전체 208)", [
    ("PKG 취합 대상", len(conv), f"시트 1 — UI PKG 로 옮길 SP (화면 API 직접 호출 {g('PKG전환대상')} + UI PKG 내부 호출 {g('PKG전환대상(내부호출)')})"),
    ("FE 미사용 등록 — 정리 후 DROP", g("등록됐으나 FE 미사용(배선 정리·삭제 후보)"), "시트 2 ② — 메뉴 버튼엔 등록됐으나 FE 호출 없음 → 취합 아님, 등록 정리 후 DROP"),
    ("제외 — 공통", len(excl), "시트 3 — 여러 UI·배치 JOB·BE 공용 → PKG 취합 안 함"),
    ("보류", len(hold), "시트 4 — FE 화면 미구현 / 타 UI·배치 사용 / 범위 밖 메뉴"),
    ("참고 — UI 미사용", len(ref), f"시트 6 — 배치 전용 {g('배치전용(UI 미도달)')} · 미사용 {g('미사용')} · 기타 {g('API_M/EXCEL 등록만(버튼·메뉴 미배선)', '내부호출·콤보·로그 참조만(UI·배치 미도달)')}"),
    ("합계", len(conv) + g("등록됐으나 FE 미사용(배선 정리·삭제 후보)") + len(excl) + len(hold) + len(ref), f"= standalone PROCEDURE 전체 {len(sp)}"),
    ("(별도) DROP 대상", len(drop_rows), "시트 2 — ①취합 후 DROP 43 ②FE 미사용 등록 정리 후 DROP ③동일명 PKG 멤버 실사용(standalone 미사용) — 위 분류와 겹침"),
])
bypkg = collections.OrderedDict()
for w in work: bypkg.setdefault(w["통합대상 PKG"], []).append(w)
block("② 통합 대상 PKG별 취합 건수", [(p, len(v), f"{v[0]['PKG 상태']} · 직접 {sum(1 for x in v if x['호출 구분'].startswith('화면'))} / 내부 {sum(1 for x in v if x['호출 구분'].startswith('UI'))}"
                                       + (" · ⚠ SP INVALID 포함" if any(x['SP 상태'] != 'VALID' for x in v) else "")) for p, v in bypkg.items()])
wt = collections.Counter(w["작업 유형"].split(" (")[0] for w in work)
block("③ 작업 유형", [(k, v, "") for k, v in wt.most_common()])
block("④ 시트 안내", [
    ("1_PKG취합대상", "", "PKG 별로 묶음(줄무늬). 작업 유형·PKG 멤버(안)·재배선 대상·선행 작업"),
    ("2_DROP대상", "", "현재참조_* = 지금 남아 있는 참조 수. 재배선 후 전 축 0 이어야 DROP"),
    ("3_제외_공통SP / 4_보류", "", "취합하지 않는 SP 와 사유"),
    ("5_동일명중복", "", "standalone 과 같은 이름의 PKG 멤버가 둘 다 있는 쌍 — 실사용 쪽 판정"),
    ("6_참고_UI미사용SP", "", "UI 가 쓰지 않는 SP (배치 전용은 삭제 대상 아님)"),
    ("7_UI별현황", "", "메뉴별 사용 PKG·취합 대상·공통 SP"),
    ("8~ 근거", "", "배선 원장·등록객체 부재·판정 방법과 한계"),
])
for c, w in zip("ABC", (26, 8, 110)): ws.column_dimensions[c].width = w

PKG_COL = {"화면 API 직접 호출": None}
sheet("1_PKG취합대상", work, tab="70AD47", band_key="통합대상 PKG", note=[
    "PKG 별로 묶어 표시. 작업 방식(07-23 DMF0021 선례): standalone 본문 무수정 → PKG 멤버 → 재배선 대상을 PKG 한정명으로 변경 → 참조 0 확인 후 standalone DROP",
    "작업 유형: 멤버 신규 추가 / 기존 동일명 멤버로 재배선(본문 동일성 확인) / 래퍼 로직 흡수",
], widths={"standalone SP": 44, "작업 유형": 30, "PKG 멤버(안)": 48, "재배선 대상": 55, "선행 작업": 40, "호출 경로(근거)": 55, "FE 미사용 등록(함께 정리)": 30, "화면명": 22})
sheet("2_DROP대상", drop_rows, tab="C00000", color_key="DROP 구분",
      colors={"①": "E2EFDA", "②": "FFF2CC", "③": "FCE4D6"},
      widths={"standalone SP": 44, "DROP 구분": 30, "DROP 전제": 40, "DROP 가능 조건": 60, "관련 PKG/UI": 20})
sheet("3_제외_공통SP", excl, tab="5B9BD5", widths={"standalone SP": 48, "제외 사유": 45, "사용 UI": 45, "배치 JOB(직접)": 35})
sheet("4_보류", hold, tab="FFC000", widths={"standalone SP": 44, "보류 사유": 45, "호출 경로(근거)": 60})
sheet("5_동일명중복", simrows, tab="7030A0", widths={"standalone SP": 44, "동일명 PKG 멤버": 52, "판정": 55, "standalone 사용 근거": 60, "PKG 멤버 사용 근거": 60})
sheet("6_참고_UI미사용SP", ref, tab="A5A5A5", color_key="분류", colors={"배치": "FCE4D6", "미사용": "F2F2F2"},
      widths={"standalone SP": 48, "분류": 34, "텍스트 참조": 45, "호출 객체": 40})
sheet("7_UI별현황", uirows, tab="1F4E78", widths={"취합 대상 SP": 60, "공통 SP 사용": 45, "사용 PKG": 30})
led = sorted([x for x in ledger if x["SCOPE"] == "Y"], key=lambda x: (x["UI"], x["MNU_ID"], x["SHT_NM"], x["BTN"]))
sheet("8_근거_배선원장", [{k: x[k] for k in ("UI", "MNU_ID", "SRC", "SHT_NM", "BTN", "API_ID", "API_URL", "KIND", "PKG", "MEMBER", "STANDALONE", "PROC_M", "FE_SHEET", "LOG_ALL")} for x in led],
      widths={"API_URL": 55, "API_ID": 38})
sheet("9_근거_등록객체부재", [{k: x[k] for k in ("SCOPE", "UI", "MNU_ID", "SHT_NM", "BTN", "API_ID", "API_URL", "KIND", "FE_SHEET")} for x in ledger
                         if x["KIND"] in ("MISSING", "PKG_MEMBER_MISSING", "UNKNOWN_QUALIFIED", "API_M_NONE", "EMPTY")], widths={"API_URL": 55, "API_ID": 38})
sheet("10_판정방법_한계", [
    {"항목": "사용 판정", "내용": "메뉴 버튼 등록(BTN_MPG→API_M) + FE 코드(주석 제외)에서 해당 메뉴·시트 호출 확인. Excel 업로드(EXCEL_COLUMNS) 포함"},
    {"항목": "내부 호출", "내용": "ALL_DEPENDENCIES → ALL_SOURCE 본문에서 호출 멤버까지 특정, 등록 API 까지 역추적"},
    {"항목": "배치", "내용": "스케줄러 JOB 113개(사내 비활성 상태·AS-IS 이관 배치) 호출 폐포. UI 1개만 써도 JOB 이 호출하면 공통으로 제외"},
    {"항목": "동적 참조", "내용": "SCM_MGR 전체 소스·nscm-be 코드에서 이름 텍스트 검색(동적 SQL 포착)"},
    {"항목": "공통 판정", "내용": "SDM·비활성 메뉴 등록까지 포함해 사용 UI 2개 이상이면 공통(이동·삭제를 막는 보수 방향)"},
    {"항목": "확인 불가", "내용": "DB Link·외부 IF·클라우드 전용 객체. 클라우드는 사내와 객체가 다를 수 있어 작업 전 동일 조회 필요"},
    {"항목": "기준 DB", "내용": f"사내 ZIOUI1 — 2026-09-14 14:10 재생성 PDB(host ol9-db), 당일 적용분 재반영(16:42) 후 조회"},
], widths={"항목": 14, "내용": 140})
OUT_X.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT_X)

# ───────────────────────── MD ─────────────────────────
def mdtable(rows, cols):
    esc = lambda v: str(v).replace("|", "\\|").replace("\n", " ")
    out = ["| " + " | ".join(cols) + " |", "|" + "|".join("---" for _ in cols) + "|"]
    out += ["| " + " | ".join(esc(r.get(c, "")) for c in cols) + " |" for r in rows]
    return "\n".join(out)
md = [f"# UI별 PKG 취합 대상 목록 ({TODAY})", "",
      f"> 기준: 사내 {dbid[0]} (host {dbid[1]}) · SCM_MGR standalone PROCEDURE 208개 · 조회만(DB 변경 없음)",
      "> 범위: 활성 SCREEN 50화면(SDM 제외) + 화면이 여는 팝업 메뉴 13 · 판정 = FE 가 실제 호출하는 등록 API 기준",
      f"> 상세·필터: [TOBE_UI별_PKG전환대상_20260914.xlsx]({OUT_X.name})", "",
      "## 1. 요약", "",
      mdtable([{"구분": "PKG 취합 대상", "건수": len(conv)},
               {"구분": "FE 미사용 등록 — 정리 후 DROP", "건수": g("등록됐으나 FE 미사용(배선 정리·삭제 후보)")},
               {"구분": "제외(공통)", "건수": len(excl)}, {"구분": "보류", "건수": len(hold)}, {"구분": "참고(UI 미사용)", "건수": len(ref)},
               {"구분": "**합계 (standalone 전체)**", "건수": len(sp)},
               {"구분": "(별도) DROP 대상 — 위와 겹침", "건수": len(drop_rows)}], ["구분", "건수"]), "",
      "## 2. PKG 취합 대상 — 통합 대상 PKG별", ""]
for p, v in bypkg.items():
    md += [f"### {p} ({v[0]['PKG 상태']}) — {len(v)}건", "",
           mdtable(v, ["standalone SP", "SP 상태", "호출 구분", "작업 유형", "PKG 멤버(안)", "재배선 대상", "선행 작업"]), ""]
md += ["## 3. DROP 대상", "", "재배선 후 현재참조 전 축(API_M·PROC_M·EXCEL·콤보·호출객체·배치JOB·텍스트) 0 재확인 → 백업 DDL 보관 → DROP", "",
       mdtable(drop_rows, ["No", "DROP 구분", "standalone SP", "관련 PKG/UI", "DROP 전제", "현재참조_API_M", "현재참조_PROC_M", "현재참조_호출객체", "SCMP_APP 시노님"]), "",
       "## 4. 제외 — 공통 SP (취합하지 않음)", "", mdtable(excl, ["standalone SP", "분류", "제외 사유", "사용 UI", "배치 JOB 수"]), "",
       "## 5. 보류", "", mdtable(hold, ["standalone SP", "분류", "보류 사유", "등록 메뉴"]), "",
       "## 6. 동일명 중복 (standalone ↔ PKG 멤버)", "", mdtable(simrows, ["standalone SP", "동일명 PKG 멤버", "래퍼", "판정"]), "",
       "## 7. 취합 작업 시 절차 (참고 — 이번에는 실행하지 않음)", "",
       "1. 작업 직전 이 목록을 사내·클라우드에서 다시 조회(객체 차이·신규 등록 확인)",
       "2. 신규 PKG(PKG_DMF0008·PKG_SYS0012): SPEC/BODY 생성 + SCMP_APP EXECUTE 권한·시노님",
       "3. standalone 본문 무수정으로 PKG 멤버 추가 (동일명 멤버가 있으면 본문 동일성 확인 후 재배선만, 래퍼는 로직 흡수)",
       "4. 재배선: API_M.API_URL · PROC_M(PROC_NM) · EXCEL STORAGE_SQL_ID · 호출부 본문을 PKG 한정명으로",
       "5. FE 미사용 등록(시트 표기)은 BTN_MPG·API_M 정리",
       "6. 화면 실측(조회·저장) → 참조 0 재확인 → 백업 DDL 보관 → standalone·SCMP_APP 시노님 DROP",
       "7. INVALID standalone(DMF0008 4건·DMF0043_GET_CONTRIBUTION_SALES_LIST)은 먼저 재컴파일로 정상 여부 확인",
       "8. 클라우드: 멱등 SQL 이관 패키지(사내 전값 WHERE 금지·MERGE)", "",
       "## 8. 판정 한계", "", "- DB Link·외부 IF·클라우드 전용 객체는 확인 불가", "- 스케줄러 JOB 은 패키지 단위 의존 폐포라 넓게 잡힘(제외 쪽으로 보수적)",
       "- FE 시트 확인은 문자열 리터럴 기준 — 동적으로 만든 시트명은 FE 미사용으로 보일 수 있음(해당 행은 시트 1·2 에 표기)", ""]
OUT_M.write_text("\n".join(md), encoding="utf-8")
print("saved", OUT_X, "/", OUT_M)
print({"취합": len(conv), "DROP": len(drop_rows), "제외": len(excl), "보류": len(hold), "참고": len(ref), "동일명": len(simrows)})
print("drop kinds", collections.Counter(x[1] for x in drop))
print("work types", wt)
