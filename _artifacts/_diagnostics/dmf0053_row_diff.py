"""dmf0053_row_diff.py — ① 로컬 SP ↔ AS-IS(클라우드) 엑셀 행·수량 대사  ② AS-IS ↔ TO-BE(클라우드) 엑셀 차이 Row 추출(생성일자 포함). 2026-09-14"""
import json, collections, openpyxl, datetime
from openpyxl.styles import PatternFill, Font
n = lambda v: "" if v is None else str(v).strip()
# 표기 정규화 — EDI SEQ 센티널 '-' = 공백 (TO-BE PK 센티널 표준) / FCST Type 코드→표시명 (LOOKUP SALES_TYPE_CODE)
import oracledb
_cur = oracledb.connect(user="SCM_MGR", password="zionex1!", dsn="175.211.29.206:2521/ZIOUI1").cursor()
FCST_NAME = {c: nm for c, nm in _cur.execute("SELECT SEGMENT1, SEGMENT2 FROM SCM_ADM_LOOKUP_MASTER_DETAIL WHERE LOOKUP_TYPE='SALES_TYPE_CODE' AND VALID_YN='Y'").fetchall()}
def nk(v, col=None):
    s = n(v)
    if col == "SEQ" and s == "-": return ""
    if col == "FCST": return FCST_NAME.get(s, s)
    return s
def num(v):
    if v in (None, "", " "): return 0.0
    try: return round(float(v), 4)
    except: return 0.0
MONTHS = ["2026-07","2026-08","2026-09","2026-10","2026-11","2026-12","2027-01","2027-02"]
mmap = {"Cust. FCST": "CUST_FCST", "└ EDI": "CUST_FCST_EDI", "└ Commercial_Agreement": "Commercial_Agreement", "└ EXCEL": "Cust.FCST_EXCEL"}
def load_x(tag, skip):
    ws = openpyxl.load_workbook(rf"C:\TRANFILE\RECEIVE\CustFcstWaterfall_20260914_102610_{tag}.xlsx", read_only=True, data_only=True).worksheets[0]
    all_ = list(ws.iter_rows(values_only=True)); hdr = list(all_[0]); out = []
    for i, r in enumerate(all_[skip:], start=skip + 1):
        r = list(r) + [None] * (30 - len(r))
        if any(v not in (None, "", " ") for v in r[:22]): out.append((i, r))
    return hdr, out
KEYIDX = [0,1,2,3,4,5,6,7,8,9,11,12,14,15]   # EDI Sender~Model, Incoterms, FCST Type, Input Date, EDI SEQ
def xkey(r): return tuple(nk(r[i], "SEQ" if i == 15 else None) for i in KEYIDX) + (mmap.get(n(r[21]), n(r[21])),)
def xqty(r): return tuple(num(v) for v in r[22:30])
ha, A = load_x("ASIS", 3); ht, T = load_x("TOBE", 2)
# 로컬 SP
R = json.load(open(r"C:\Users\ejene\AppData\Local\Temp\claude\c--Projects-NSCM\0ea4cbc9-b2ef-4b42-8284-e360f3928fc5\scratchpad\dmf0053_sp_rows.json", encoding="utf-8"))
spk = collections.Counter(); spq = {}
for r in R:
    k = tuple(nk(r[c], "FCST" if c == "FCST_TYPE" else "SEQ" if c == "EDI_IF_SEQ" else None) for c in ["SENDER_NAME","DIVISION_ID","CUSTOMER_PLANT","UNLOADING_POINT","CUST_ITEM_ID","END_CUSTOMER","SCM_SITE_ID","PROJECT_CODE","PLANT_SITE_ID","ITEM_ID","INCOTERMS","FCST_TYPE","CREATE_DATE","EDI_IF_SEQ"]) + (n(r["MEASURE"]),)
    spk[k] += 1; spq.setdefault(k, []).append(tuple(num(r.get("M_" + m.replace("-", ""))) for m in MONTHS))
ak = collections.Counter(xkey(r) for _, r in A); aq = {}
for _, r in A: aq.setdefault(xkey(r), []).append(xqty(r))
# ① 로컬 SP ↔ AS-IS 엑셀 (로직 판정) — 배치 일자 차이 제외 구간(Input Date <= 20260830)
cut = "20260830"
a1 = {k: v for k, v in ak.items() if k[12] <= cut}; s1 = {k: v for k, v in spk.items() if k[12] <= cut}
only_a = sum((collections.Counter(a1) - collections.Counter(s1)).values()); only_s = sum((collections.Counter(s1) - collections.Counter(a1)).values())
qty_ng = sum(1 for k in set(a1) & set(s1) if sorted(aq[k]) != sorted(spq[k]))
print(f"① Input Date<={cut}  AS-IS엑셀 {sum(a1.values())}행 / 로컬SP {sum(s1.values())}행 | AS-IS만 {only_a} · 로컬SP만 {only_s} · 키 같고 수량 다름 {qty_ng}")
ex = [k for k in set(a1) & set(s1) if sorted(aq[k]) != sorted(spq[k])][:5]
for k in ex: print("   qty diff", k[5], k[9], k[12], k[14], aq[k], spq[k])
# ② AS-IS ↔ TO-BE 엑셀 차이 Row
tk = collections.Counter(xkey(r) for _, r in T)
onlyA = collections.Counter(ak) - tk; onlyT = tk - collections.Counter(ak)
print(f"② 엑셀 차이 — AS-IS만 {sum(onlyA.values())}행 · TO-BE만 {sum(onlyT.values())}행 (AS-IS {len(A)} / TO-BE {len(T)})")
out = []
used = collections.Counter()
for side, rows, only in (("AS-IS에만 있음", A, onlyA), ("TO-BE에만 있음", T, onlyT)):
    for i, r in rows:
        k = xkey(r)
        if used[(side, k)] < only.get(k, 0):
            used[(side, k)] += 1; out.append((side, i, r))
summ = collections.Counter((s, n(r[14]), n(r[18]), mmap.get(n(r[21]), n(r[21])), n(r[17])) for s, _, r in out)
print("   구분 / Input Date / Last Update By / Measure / Last Update → 행수")
for k, v in sorted(summ.items()): print("  ", k, v)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "차이Row"
ws.append(["구분", "원본 엑셀 행번호"] + ha[:30])
for s, i, r in out:
    ws.append([s, i] + [r[j] for j in range(30)])
for c in ws[1]: c.font = Font(bold=True)
s2 = wb.create_sheet("요약")
s2.append(["구분", "Input Date(CREATE_DATE)", "Last Update By", "Measure", "Last Update", "행수"])
for k, v in sorted(summ.items()): s2.append(list(k) + [v])
p = r"C:\Projects\NSCM\_artifacts\_diagnostics\DMF0053_CustFcstWaterfall_차이Row_20260914.xlsx"; wb.save(p); print("saved", p)
