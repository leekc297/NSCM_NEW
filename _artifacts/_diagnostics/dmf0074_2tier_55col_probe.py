"""dmf0074_2tier_55col_probe.py — 40 외 division(55열, formId DMF0074) 2단 양식 위치 매핑 확인. 2026-09-14"""
import pathlib, requests, openpyxl, oracledb
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L
OUT = pathlib.Path(r"C:\Users\ejene\AppData\Local\Temp\claude\c--Projects-NSCM\0ea4cbc9-b2ef-4b42-8284-e360f3928fc5\scratchpad\dmf0074_2tier_55.xlsx")
base = ["Div","Plan Date","Team","End Customer","Project Code","Project","Model Code","Metal Link","Sales Site","Currency","FCST type","Exception","Energy"]
groups = ["Key","Period(Fr)","Period(To)","Contract Price(/kg)","Change Price(/kg)","Metal Index (kg/Biz type)"]
el = ["Ni","Co","Li2CO3","LIOH","Al","Mn","Cu"]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Metal Link Price"
for i, h in enumerate(base, 1): ws.cell(1, i, h); ws.merge_cells(start_row=1, start_column=i, end_row=2, end_column=i)
c = len(base) + 1
for g in groups:
    ws.cell(1, c, g); ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c+6)
    for k, e in enumerate(el): ws.cell(2, c+k, e)
    c += 7
vals = ["82", "202611", "T55", "CUST_X", "PRJ_X", "PRJNAME", "MODEL_X", "Y", "SITE_X", "USD", "Mass", "N", "100"] + [f"v{i}" for i in range(42)]
for i, v in enumerate(vals, 1): ws.cell(3, i, v)
wb.save(OUT)
s = requests.Session(); s.post("http://localhost:8080/user/login", json={"userId": "minyonglee1"}, timeout=30)
with open(OUT, "rb") as fh:
    body = s.post("http://localhost:8080/DMF0074/Metal Link Price/UPLOAD/excel-upload", files={"file": (OUT.name, fh)}, timeout=600).json()
print("code", body.get("code"), body.get("message"), body["result"]["O_CUR_CNT"])
print("LIST", [(x.get("EXCEL_ROW"), x.get("MESSAGE_TEXT")) for x in body["result"]["O_CUR_LIST"]])
con = oracledb.connect(user="SCM_MGR", password="zionex1!", dsn="175.211.29.206:2521/ZIOUI1"); cur = con.cursor()
cols = ", ".join(f"ATTRIBUTE{i}" for i in range(1, 56))
cur.execute(f"""SELECT UPLOAD_STATUS_CODE, {cols} FROM SCM_ADM_EXCEL_UPLOAD WHERE PROGRAM_ID='DMF0074' AND SHEET_NAME='Metal Link Price' AND CREATED_BY='minyonglee1'
               AND CREATION_DATE=(SELECT MAX(CREATION_DATE) FROM SCM_ADM_EXCEL_UPLOAD WHERE PROGRAM_ID='DMF0074' AND SHEET_NAME='Metal Link Price' AND CREATED_BY='minyonglee1') ORDER BY UPLOAD_ID""")
rows = cur.fetchall()
for r in rows: print(r[0], "A1..5", r[1:6], "A13..15", r[13:16], "A55", r[55])
data = [r for r in rows if r[1] == "82"][0]
print("positional match 55:", list(data[1:56]) == vals)
