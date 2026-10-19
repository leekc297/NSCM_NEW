"""TO-BE 시트 검증 — 산출 xlsx 를 직접 읽어 건수 대사 + 표본 출력"""
import os, glob, random, collections
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(glob.glob(os.path.join(HERE, "NULL*대상취합*.xlsx"))[0], read_only=True)
for ws in wb.worksheets:
    print(ws.title, ws.max_row - 1 if ws.title != "요약" else "")
random.seed(3)
for name in ("4_TOBE_NVL적용대상", "5_TOBE_함께검토", "7_TOBE_제외·검토컬럼", "6_TOBE_컬럼미특정"):
    rows = list(wb[name].iter_rows(min_row=2, values_only=True))
    print("=====", name, len(rows))
    print("  쓰임", collections.Counter(r[5] for r in rows).most_common(8))
    print("  변경안 비어있음", sum(1 for r in rows if not r[11]), "/ 동적", sum(1 for r in rows if r[7] == "Y"))
    if name.startswith("6"):
        print("  미특정 사유", collections.Counter((r[14] or "").split("]")[0] for r in rows).most_common(12))
    for r in random.sample(rows, min(10, len(rows))):
        print("   -", r[2], r[3], "L%s" % r[4], "|", r[5], "|", (r[8] or r[12] or "")[:70], "|", r[9][:40] if r[9] else "")
        print("       원문:", (r[10] or "")[:170])
        if r[11]:
            print("       변경:", r[11][:170])
ws = wb["요약"]
for r in ws.iter_rows(values_only=True):
    if r[0]:
        print(r[0][:200], r[1] if r[1] is not None else "", r[2] if r[2] is not None else "")
