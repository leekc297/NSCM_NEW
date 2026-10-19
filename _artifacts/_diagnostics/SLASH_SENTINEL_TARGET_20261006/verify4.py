import glob,collections,random,sys
from openpyxl import load_workbook
wb=load_workbook(glob.glob('NULL*작업대상*.xlsx')[0],read_only=True)
ws=wb['3_TOBE_변경목록']; h=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]; rows=list(ws.iter_rows(min_row=2,values_only=True))
A,C,W,T,CH=h.index('액션'),h.index('확신도'),h.index('사유'),h.index('원문'),h.index('변경안')
print(collections.Counter(r[A] for r in rows)); print(collections.Counter(r[C] for r in rows))
print('미확정 사유',collections.Counter((r[W] or '').split(' · ')[0][:45] for r in rows if r[A]=='미확정(검토)').most_common(6))
print("리터럴 사유",collections.Counter((r[W] or '').split(' · ')[0][:45] for r in rows if r[A]=="리터럴→'/'").most_common(8))
random.seed(9)
for act in ("리터럴→'/'","유지(AS-IS 값)"):
    xs=[r for r in rows if r[A]==act and "''" in (r[W] or '')]; print('==',act,"('' 근거)",len(xs))
    for r in random.sample(xs,min(4,len(xs))): print('  ',r[2],r[4],'|',r[W][:70],'|',(r[T] or '')[:100])
eq=list(wb['5_동치정규화(ASIS·C#)'].iter_rows(min_row=2,values_only=True)); print('동치',len(eq),'TO-BE 미발견',sum(1 for r in eq if r[7] and r[7].startswith('(TO-BE')),"대응에 '-' 없음 표시",sum(1 for r in eq if r[7] and "'-' 없음" in r[7]))
fe=list(wb['6_FE_대시기본값'].iter_rows(min_row=2,values_only=True)); print('FE',collections.Counter(r[2] for r in fe))
