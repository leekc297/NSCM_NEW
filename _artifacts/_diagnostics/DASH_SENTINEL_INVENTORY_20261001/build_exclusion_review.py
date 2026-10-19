"""NULL 센티널 '-' → '/' — 제외대상 후보 검토용 xlsx (읽기 전용 · 캐시만 사용, DB 접속 없음)

extract_dash_inventory.py 가 만든 _cache/*.json 을 읽어, 자동 판정이 "제외(AS-IS '-' 사용)" 인 후보만 모은다.
최종 제외 여부는 사용자가 판단 → 각 시트 '사용자 판정' 열에 입력.

사용: python build_exclusion_review.py
"""
import os, re, json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(HERE, "_cache")
OUT = os.path.join(HERE, "NULL센티널_제외대상_검토용_20261006.xlsx")
EXCL = "제외(AS-IS '-' 사용)"
PATTERNS = ["ISNULL(col,'-')", "ISNULL(...,'-') AS col", "col = ISNULL(..,'-')", "col op '-'", "'-' op col", "col IN (..'-'..)", "'-' AS col"]
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def load(name):
    return json.load(open(os.path.join(CACHE, name + ".json"), encoding="utf-8"))


def clean(v):
    return ILLEGAL.sub(" ", v) if isinstance(v, str) else v


def main():
    judged = load("judged")
    cols = [r for r in judged["cols"] if r["verdict"] == EXCL]
    srcs = [r for r in judged["src"] if r["verdict"] == EXCL]
    code_ev = load("asis_code")
    adata = load("asis_data_partial")
    meta = load("asis_meta")
    tcols = load("tobe_cols")
    names = sorted({r["column"] for r in cols})
    nameset = set(names)

    # AS-IS 객체별 참조 토큰 — ③ 근거를 "그 테이블을 다루는 객체" 로 좁혀 보기 위함
    need_objs = {(e["db"], e["object"]) for e in code_ev if e["column"] in nameset}
    tokens = {}
    for m in load("asis_modules"):
        k = (m["db"], m["name"])
        if k in need_objs:
            tokens[k] = set(re.findall(r"\w+", m["definition"].upper()))

    ev_by_col = collections.defaultdict(list)
    for e in code_ev:
        if e["column"] in nameset:
            ev_by_col[e["column"]].append(e)

    def pat_counts(evs):
        c = collections.Counter(e["pattern"] for e in evs)
        return [c.get(p, 0) for p in PATTERNS]

    def samples(evs, n):
        # 패턴이 다양하게 보이도록 패턴별 1개씩 먼저
        out, seen = [], set()
        for p in PATTERNS:
            for e in evs:
                if e["pattern"] == p and e["snippet"] not in seen:
                    out.append(f"[{e['object']}] {e['snippet']}"); seen.add(e["snippet"]); break
        return out[:n] + [""] * max(0, n - len(out))

    # ── 시트 1: 컬럼명별
    by_name = collections.defaultdict(list)
    for r in cols:
        by_name[r["column"]].append(r)
    src_by_col = collections.Counter(r["column"] for r in srcs)
    asis_def = collections.defaultdict(set)
    for d in meta["defaults"]:
        if d["is_dash"]:
            asis_def[d["column"].upper()].add(f"{d['db']}.{d['table']}")
    asis_dash = collections.defaultdict(list)      # 컬럼명 → [(db|table, dash, total)]
    for k, v in adata.items():
        if v.get("error"):
            continue
        for cn, x in v["cols"].items():
            if x["dash"]:
                asis_dash[cn.upper()].append((k, x["dash"], v["total"]))

    h1 = ["컬럼명", "TO-BE 제외후보 컬럼 수", "TO-BE 테이블 예시", "TO-BE 소스 제외후보 줄 수",
          "①AS-IS 기본값'-' 테이블 수", "②AS-IS 데이터에 '-' 있는 테이블 수", "②'-' 행수 합계", "②'-' 최다 테이블(행수/전체)",
          "③AS-IS 코드 객체 수"] + ["③" + p for p in PATTERNS] + ["③발췌 1", "③발췌 2", "③발췌 3", "사용자 판정", "비고"]
    rows1 = []
    for n in names:
        rs = by_name[n]; evs = ev_by_col.get(n, [])
        dl = sorted(asis_dash.get(n, []), key=lambda x: -x[1])
        top = f"{dl[0][0].replace('|', '.')} ({dl[0][1]:,}/{dl[0][2]:,})" if dl else ""
        rows1.append([n, len(rs), ", ".join(sorted({r["table"] for r in rs})[:6]) + (" …" if len({r["table"] for r in rs}) > 6 else ""),
                      src_by_col.get(n, 0), len(asis_def.get(n, ())), len(dl), sum(x[1] for x in dl), top,
                      len({(e["db"], e["object"]) for e in evs})] + pat_counts(evs) + samples(evs, 3) + ["", ""])
    rows1.sort(key=lambda r: -r[1])

    # ── 시트 2: 제외후보 컬럼 (테이블 단위)
    h2 = ["스키마", "테이블", "컬럼", "데이터형", "길이", "NULLABLE", "DEFAULT ON NULL", "AS-IS DB", "AS-IS 테이블",
          "①AS-IS 기본값'-'", "AS-IS 행수", "②AS-IS '-'행수", "③코드 객체 수(컬럼명 기준·전체)", "③코드 객체 수(이 테이블을 참조하는 객체만)"] + \
         ["③(테이블한정)" + p for p in PATTERNS] + ["③(테이블한정) 발췌 1", "③(테이블한정) 발췌 2", "자동 판정 근거", "사용자 판정", "비고"]
    rows2 = []
    for r in cols:
        evs = ev_by_col.get(r["column"], [])
        tb = (r["asis_table"] or r["table"]).upper()
        scoped = [e for e in evs if tb in tokens.get((e["db"], e["object"]), ())]
        rows2.append([r["owner"], r["table"], r["column"], r["dtype"], r["dlen"], r["nullable"], r["on_null"], r["asis_db"], r["asis_table"],
                      r["e1"], r["asis_rows"], r["e2_dash"], len({(e["db"], e["object"]) for e in evs}),
                      len({(e["db"], e["object"]) for e in scoped})] + pat_counts(scoped) + samples(scoped, 2) + [r["why"], "", ""])
    rows2.sort(key=lambda x: (x[2], x[0], x[1]))

    # ── 시트 3: AS-IS 코드 근거 원장 (제외후보 컬럼명 한정)
    h3 = ["컬럼명", "AS-IS DB", "객체", "객체유형", "패턴", "발췌"]
    rows3 = [[e["column"], e["db"], e["object"], e["type"], e["pattern"], e["snippet"]]
             for n in names for e in sorted(ev_by_col.get(n, []), key=lambda e: (PATTERNS.index(e["pattern"]), e["object"]))]

    # ── 시트 4: 제외후보 소스 줄
    h4 = ["스키마", "객체유형", "객체명", "서브프로그램", "줄", "대상컬럼(추정)", "용도", "용도 상세", "원문", "이관작업 마커",
          "AS-IS 대응 객체", "자동 판정 근거", "사용자 판정", "비고"]
    rows4 = [[r["owner"], r["type"], r["name"], r["sub"], r["line"], r["column"], r["usage"], r["detail"], r["text"], r["markers"],
              r["asis_cp"], r["why"], "", ""] for r in sorted(srcs, key=lambda r: (r["column"] or "", r["owner"], r["name"], r["line"]))]

    wb = Workbook()
    ws = wb.active; ws.title = "안내"
    info = [
        "NULL 센티널 '-' → '/' — 제외대상 후보 검토용",
        "",
        f"TO-BE 기준: 사내 ZIOUI1 {tcols['fetched']} 조회 / AS-IS 근거: LGCE_SCMDEV·LGCE_FPDEV {str(meta['fetched'])[:19]} 수집",
        "이 파일에는 자동 판정이 \"제외(AS-IS '-' 사용)\" 인 후보만 들어 있다. 최종 제외 여부는 각 시트 '사용자 판정' 열에 입력한다 (제외 / 작업대상).",
        "",
        f"1_컬럼명별   : 제외후보 컬럼명 {len(rows1):,}개 요약 — 컬럼명 단위로 먼저 판단할 때",
        f"2_제외후보_컬럼 : 테이블.컬럼 {len(rows2):,}개 — 테이블마다 판단이 갈릴 때",
        f"3_ASIS코드근거 : ③ 근거 원문 발췌 {len(rows3):,}행 — 컬럼명 필터로 실제 쓰임 확인",
        f"4_제외후보_소스줄: TO-BE PKG/SP/VIEW 의 '-' 사용 줄 {len(rows4):,}개",
        "",
        "근거 ① AS-IS 컬럼 기본값 제약이 '-'   ② AS-IS 실데이터에 '-' 값 존재   ③ AS-IS SP/VIEW/FUNCTION 코드가 그 컬럼명에 '-' 사용",
        "주의: ③ 은 컬럼명만으로 맞춘다(테이블 구분 없음). 2_제외후보_컬럼 의 '이 테이블을 참조하는 객체만' 열은",
        "      ③ 근거 객체 중 정의문에 해당 AS-IS 테이블명이 나오는 것만 센 값이다 — 0 이면 그 테이블과 무관한 코드가 근거일 가능성이 높다.",
        "주의: 패턴 \"'-' AS col\" 은 '-' 바로 뒤 단어를 컬럼으로 본 느슨한 매칭이라 오탐이 섞여 있다 — 발췌로 확인.",
        "주의: AS-IS 의 ISNULL(col,'-') 는 AS-IS 도 '-' 를 빈값 대체로 쓴 경우일 수 있다 — '의미 있는 사용' 인지는 발췌로 판단.",
        "AS-IS 데이터 근거 ② 는 DEV DB 기준(운영 스냅샷 아님). TO-BE 기존 데이터의 '-' 는 이번 범위 밖(별도 DB 작업).",
    ]
    for t in info:
        ws.append([t])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 150

    hdr_fill = PatternFill("solid", fgColor="DDEBF7"); in_fill = PatternFill("solid", fgColor="FFF2CC")

    def sheet(title, headers, rows, widths, wrap_cols=()):
        s = wb.create_sheet(title)
        s.append(headers)
        for r in rows:
            s.append([clean(v) for v in r])
        for c in s[1]:
            c.font = Font(bold=True); c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center")
        s.freeze_panes = "D2" if len(headers) > 6 else "A2"
        s.auto_filter.ref = s.dimensions
        for i, w in enumerate(widths, 1):
            s.column_dimensions[get_column_letter(i)].width = w
        jc = headers.index("사용자 판정") + 1 if "사용자 판정" in headers else None
        if jc and rows:
            L = get_column_letter(jc)
            dv = DataValidation(type="list", formula1='"제외,작업대상"', allow_blank=True)
            s.add_data_validation(dv); dv.add(f"{L}2:{L}{len(rows) + 1}")
            s[f"{L}1"].fill = in_fill; s[f"{get_column_letter(jc + 1)}1"].fill = in_fill
        return s

    sheet("1_컬럼명별", h1, rows1, [24, 9, 46, 9, 9, 9, 12, 44, 9] + [9] * len(PATTERNS) + [70, 70, 70, 12, 30])
    sheet("2_제외후보_컬럼", h2, rows2, [9, 34, 24, 10, 6, 6, 8, 13, 30, 8, 11, 11, 10, 10] + [9] * len(PATTERNS) + [70, 70, 44, 12, 30])
    sheet("3_ASIS코드근거", h3, rows3, [24, 13, 44, 22, 22, 120])
    sheet("4_제외후보_소스줄", h4, rows4, [9, 13, 34, 34, 6, 24, 14, 22, 90, 18, 30, 44, 12, 30])
    wb.save(OUT)

    zero = sum(1 for r in rows2 if r[13] == 0 and not r[9] and not r[11])
    print("columns", len(rows2), "names", len(rows1), "code_ev", len(rows3), "src", len(rows4))
    print("③ 테이블한정 0 이고 ①② 도 없는 컬럼", zero)
    print(OUT)


if __name__ == "__main__":
    main()
