"""DMF0074 Metal Link Price Upload (Division 40, Form DMF0074_40) — 정상 샘플 10건 생성 + 로컬 BE 업로드 검증(P_MODE=0, 원장 저장 없음).

목적: 업로드 시 누락되던 2단 헤더 영역 Key / Period(Fr) / Period(To) 7개 금속(Ni·Co·Li2CO3·LIOH·Al·Mn·Cu) 전부에 값을 채워 누락 여부 테스트.
양식: AS-IS 2단 병합 헤더 원본(C:\\TRANFILE\\RECEIVE\\DMF0074_MetalLinkPricePopup_ASIS.xlsx) 그대로 사용, 3행부터 데이터.
행 선정(사내 ZIOUI1 실측, 생성 시점): MRF 버전 판매계획/실적 조합 × END_CUSTOMER 룩업(40) × SITE_ITEM(P, 활성) × SCM_MST_SITE × 환율통화 × FCST type(≠FREE)
  → End Customer 중복 없이 10건 (FCST type = Mass).
검증 규칙(PKG_DMF0074.P_DMF0074_EXCEL_UPLOAD): Key = METAL_KEY_VALUE 룩업(40), Period 정수 0~12 · Fr > To, Plan Date >= 당월.
usage: python gen_dmf0074_sample10.py [--upload]
"""
import datetime, json, os, sys, argparse
import openpyxl

sys.path.insert(0, r"C:\Projects\NSCM\_artifacts\_diagnostics")
import ora, oracledb

ROOT = r"C:\Projects\NSCM"
SRC = r"C:\TRANFILE\RECEIVE\DMF0074_MetalLinkPricePopup_ASIS.xlsx"
OUT_DIR = os.path.join(ROOT, "Doc", "단위테스트", "_excel_upload_cases", "DMF0074_SAMPLE")
METALS = ["Ni", "Co", "Li2CO3", "LIOH", "Al", "Mn", "Cu"]
GROUP_START = {"Key": "L", "Period(Fr)": "S", "Period(To)": "Z", "Contract Price(/kg)": "AG",
               "Change Price(/kg)": "AN", "Metal Index (kg/Biz type)": "AU"}
MRF_VERSION = "SP_W_20260728_V002"
# (Fr, To) 쌍 — 행·금속별로 순환해 다양한 조합 (모두 0~12, Fr > To)
PERIOD_PAIRS = [(6, 0), (3, 1), (12, 6), (9, 3), (4, 2), (2, 1), (1, 0), (11, 5), (8, 4), (7, 2)]

today = datetime.date.today()
PLAN_DATE = (today.replace(day=1) + datetime.timedelta(days=32)).strftime("%Y%m")  # 익월
STAMP = today.strftime("%Y%m%d")


def col_offset(letter, n):
    return openpyxl.utils.get_column_letter(openpyxl.utils.column_index_from_string(letter) + n)


def load_keys(cur):
    cur.execute("""SELECT UPPER(SEGMENT2), SEGMENT3 FROM SCM_MGR.SCM_ADM_LOOKUP_MASTER_DETAIL
                    WHERE LOOKUP_TYPE='METAL_KEY_VALUE' AND SEGMENT1='40' AND VALID_YN='Y'
                      AND SEGMENT4 <> '미연동' ORDER BY SEGMENT3  /* 미연동 Key 는 Period 공란 필수(NEED EMPTY VALUE) — Key·Period 동시 입력 시험 대상에서 제외 */""")
    return cur.fetchall()


def pick_rows():
    cands = json.load(open(os.path.join(os.environ["SCRATCH"], "dmf0074_cands.json"), encoding="utf-8"))
    seen, rows = set(), []
    for cust, item, site, ccy, fcst in cands:
        if fcst != "Mass" or cust in seen or site in {r[2] for r in rows}:
            continue
        if not cust.isascii():
            continue
        seen.add(cust); rows.append((cust, item, site, ccy, fcst))
        if len(rows) == 10:
            break
    assert len(rows) == 10, len(rows)
    return rows


def write(rows, keys):
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, f"DMF0074_40_Upload_Sample10_{STAMP}.xlsx")
    wb = openpyxl.load_workbook(SRC); ws = wb.active
    for c in ws[3]:
        c.value = None
    for i, (cust, item, site, ccy, fcst) in enumerate(rows):
        r = 3 + i
        base = {"A": "40", "B": PLAN_DATE, "C": "", "D": cust, "E": item, "F": "Y", "G": site, "H": ccy, "I": fcst, "J": "N", "K": ""}
        for k, v in base.items():
            ws[f"{k}{r}"] = v
        for m, metal in enumerate(METALS):
            fr, to = PERIOD_PAIRS[(i + m) % len(PERIOD_PAIRS)]
            vals = {"Key": keys[metal][i % len(keys[metal])], "Period(Fr)": str(fr), "Period(To)": str(to),
                    "Contract Price(/kg)": f"{10 + i + m * 1.5:.2f}", "Change Price(/kg)": f"{0.5 + (i + m) * 0.05:.2f}",
                    "Metal Index (kg/Biz type)": f"{0.1 + (i * 7 + m) * 0.01:.3f}"}
            for grp, start in GROUP_START.items():
                ws[f"{col_offset(start, m)}{r}"] = vals[grp]
        for c in ws[r]:
            c.number_format = "@"
    wb.save(path)
    return path


def upload(path, base="http://localhost:8080", user="minyonglee1"):
    import requests
    s = requests.Session()
    s.post(f"{base}/user/login", json={"userId": user}, timeout=30).raise_for_status()
    with open(path, "rb") as fh:
        body = s.post(f"{base}/DMF0074_40/Metal Link Price/UPLOAD/excel-upload",
                      files={"file": (os.path.basename(path), fh)}, timeout=600).json()
    res = body.get("result") or {}
    print("code", body.get("code"), "| O_CUR_CNT", res.get("O_CUR_CNT"))
    for x in res.get("O_CUR_LIST") or []:
        print("   row", x.get("EXCEL_ROW"), "|", x.get("UPLOAD_STATUS_CODE"), "|", x.get("MESSAGE_TEXT"))
    return body


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--upload", action="store_true"); a = ap.parse_args()
    u, p, d = ora.conn_info()
    with oracledb.connect(user=u, password=p, dsn=d) as con:
        cur = con.cursor()
        lk = load_keys(cur)
        print("METAL_KEY_VALUE(40):", lk)
        # 금속별 유효 Key 전부(룩업 원천) — 행마다 순환해 여러 Key 코드를 시험
        keys = {m: [code for seg, code in lk if seg == m.upper()] for m in METALS}
        assert all(keys.values()), keys
        rows = pick_rows()
        path = write(rows, keys)
        print("plan date", PLAN_DATE, "->", path)
        for r in rows:
            print("  ", " | ".join(r))
        if a.upload:
            upload(path)
            cur.execute("""SELECT HEADER_CONTENT FROM SCM_MGR.SCM_ADM_EXCEL_COLUMNS WHERE PROGRAM_ID='DMF0074_40' AND SHEET_NAME='Metal Link Price'""")
            hdr = cur.fetchone()[0]
            hdr = (hdr.read() if hasattr(hdr, "read") else hdr).split("|")
            cols = ", ".join(f"ATTRIBUTE{n}" for n in range(1, len(hdr) + 1))
            cur.execute(f"""SELECT UPLOAD_ID, UPLOAD_STATUS_CODE, {cols} FROM SCM_MGR.SCM_ADM_EXCEL_UPLOAD
                            WHERE PROGRAM_ID='DMF0074_40' AND SHEET_NAME='Metal Link Price' AND CREATED_BY='minyonglee1'
                              AND CREATION_DATE = (SELECT MAX(CREATION_DATE) FROM SCM_MGR.SCM_ADM_EXCEL_UPLOAD
                                                    WHERE PROGRAM_ID='DMF0074_40' AND SHEET_NAME='Metal Link Price' AND CREATED_BY='minyonglee1')
                              AND NVL(UPLOAD_STATUS_CODE,'-') NOT IN ('V_TITLE','@TITLE') ORDER BY UPLOAD_ID""")
            data = cur.fetchall()
            print("staging rows", len(data))
            watch = [i for i, h in enumerate(hdr) if h.split("_")[0] in ("Key", "Period(Fr)", "Period(To)")]
            for row in data:
                missing = [hdr[i] for i in watch if row[2 + i] in (None, "")]
                print("   row", row[0], row[1], "| Key/Period null:", missing or 0,
                      "| sample", {hdr[i]: row[2 + i] for i in watch[::7]})


if __name__ == "__main__":
    main()
