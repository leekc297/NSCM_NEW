#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""07_build_outputs.py — 지식사전 산출물 생성 (MD 정본 + xlsx 뷰).

MD   Doc/_rules_v2/_dictionary/{00_INDEX,01_search_conditions,02_sheet_columns,03_sheet_measures}.md
     → AI 가 작업 시 직접 Read 하는 정본. 전 행 수록(절단 금지).
xlsx Doc/화면분석서/UI메타_지식사전_v1.0.xlsx
     → 사람이 필터링·정렬해 검토하는 뷰. 9 시트.

서식은 기존 정본 `_gen_DMF0074_validation_xlsx.py` 계보를 그대로 쓴다.
"""
import os
import sys

import cfg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
ERR_FILL = PatternFill("solid", fgColor="FCE4E4")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
OK_FILL = PatternFill("solid", fgColor="E7F4E4")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE_FONT = Font(size=9)
TITLE_FONT = Font(bold=True, size=13, color="1F4E79")


def style_sheet(ws, widths, header_row=2, freeze="A3"):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[header_row]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for row in ws.iter_rows(min_row=header_row + 1):
        for c in row:
            c.font = BASE_FONT
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
    ws.freeze_panes = freeze
    if ws.max_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"


def title(ws, text, span):
    ws.cell(row=1, column=1, value=text).font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)


def md_esc(v):
    """표 셀용 이스케이프 — 파이프와 개행이 표를 깬다."""
    s = "" if v is None else str(v)
    return s.replace("|", "\\|").replace("\n", " ").replace("\r", " ")


def head(source_label, subtitle):
    return (f"> **스냅샷** {cfg.SNAPSHOT_DATE} · **AS-IS 원천** {source_label}\n"
            f"> **재생성** `_artifacts/_diagnostics/ui_meta_dictionary_20260821/` "
            f"(01~05 실행 후 07_build_outputs.py)\n"
            f"> ⚠ 이 파일은 생성물이다. 직접 수정하지 말고 파이프라인을 고쳐 재생성한다.\n"
            f">\n> {subtitle}\n\n")


# ─────────────────────────────────────────────────────────────────────────────
def build_md_search(sc):
    m = sc["meta"]
    L = ["# 조회조건 지식사전 (항목 축)\n",
         head(m["asis_source"],
              "신규/보완 화면의 조회조건 작성 전 해당 항목 행을 대조한다 (hook #112 와 병행)."),
         "## 0. 읽는 법\n",
         "- **binding**: `ARG`=콤보 SQL 이 `@토큰` 요구(FE 전달 의무·미전달 시 조용한 0행) / "
         "`CASCADE`=상위 변경 시 하위 재조회·reset / `FILTER`=전량 fetch 후 클라 필터.\n",
         f"- **우선순위**: {m['binding_priority']}\n",
         f"- **division**: {m['division_axis_limit']}\n",
         f"- **근거**: 1차 {m['primary_evidence']} / 2차 {m['secondary_evidence']} / "
         f"콤보 {m['combo_authority']}\n",
         f"- **TO-BE 공용 필드** {len(m['searchfields'])}종: {', '.join(m['searchfields'])}\n\n",
         "## 1. 항목 표준 (전 34종)\n\n",
         "| # | 조회조건 | 사용 UI | TO-BE 필드 | binding | 콤보(CMBO) | @토큰 | "
         "div40 | div81 | div82 | 카탈로그 화면 |\n",
         "|---:|---|---:|---|---|---|---|---|---|---|---:|\n"]
    for i, c in enumerate(sc["conditions"], 1):
        dv = c["division_usage_standard"]
        L.append("| {} | **{}** | {} | {} | `{}` | {} | {} | {} | {} | {} | {} |\n".format(
            i, md_esc(c["cond"]), c["ui_count"],
            md_esc(", ".join(c["tobe_fields"])), md_esc(c["binding_primary"]),
            md_esc(", ".join(c["combo_sql_ids"])),
            md_esc(", ".join(c["combo_tokens"]) or "-"),
            md_esc(dv["40"]), md_esc(dv["81"]), md_esc(dv["82"]), c["catalog_screens"]))
    L.append("\n## 2. 항목별 화면 상세 (154 조건 전수)\n")
    for c in sc["conditions"]:
        L.append(f"\n### {c['cond']}  ({c['ui_count']} UI · {c['coverage']})\n\n")
        L.append("| UI | TO-BE 라벨 | 컨트롤 | 콤보 | binding | 필수 | 기본값 | "
                 "div40 | div81 | div82 |\n")
        L.append("|---|---|---|---|---|---|---|---|---|---|\n")
        for s in c["screens"]:
            d = s["division_usage"]
            L.append("| {} | {} | {} | {} | {} | {} | {} | {} | {} | {} |\n".format(
                s["ui"], md_esc(s["label_tobe"]), md_esc(s["control_tobe"]),
                md_esc(s["combo_sql_id"]), md_esc(s["binding"]["primary"]),
                md_esc(s["required"]), md_esc(s["default"]),
                md_esc(d["40"]["use"]), md_esc(d["81"]["use"]), md_esc(d["82"]["use"])))
    return "".join(L)


def build_md_columns(dc):
    m = dc["meta"]
    cols = dc["columns"]
    L = ["# SHEET_COLUMN 지식사전 (COLUMN_NAME 축)\n",
         head(m["asis_source"],
              "신규/보완 화면의 컬럼 메타 작성 전 해당 컬럼 행을 대조한다. "
              "`confidence` 가 low/none 이면 표준으로 쓰지 말고 AS-IS 를 확인한다."),
         "## 0. 이 사전이 왜 2단인가 (구조 사실)\n\n"]
    for f in m["structural_facts"]:
        L.append(f"- {f}\n")
    L.append(f"\n- **confidence**: {m['confidence_rule']}\n")
    L.append(f"- **허용 컨트롤** {len(m['allowed_controls'])}종: "
             f"{', '.join(m['allowed_controls'])}\n")
    L.append(f"  - FE enum 전용(= AS-IS lookup 미등록): {', '.join(m['controls_fe_only'])}\n")
    L.append(f"  - AS-IS lookup 전용(= FE 미해석): {', '.join(m['controls_lookup_only'])}\n")
    L.append(f"- **콤보 격차**: {m['combo_gap_note']}\n")
    L.append(f"- **교대쌍**: {m['swap_pair_note']}\n\n")

    L.append("## 1. 컬럼 표준 (전 {}종)\n\n".format(len(cols)))
    L.append("| COLUMN | AS-IS 행/화면 | TO-BE 행/화면 | 라벨 표준 | 정렬 | 폭 | 표시 컨트롤 | "
             "편집 컨트롤 | 편집 콤보 | 등록 | 모달키 | conf | 교대쌍 |\n")
    L.append("|---|---:|---:|---|---|---:|---|---|---|---|---|---|---|\n")
    for c in cols:
        u, e = c["usage"], c["edit_standard"]
        pair = c["swap_pair"]
        L.append("| `{}` | {}/{} | {}/{} | {} | {} | {} | {} | {} | {} | {} | {} | {} | {} |\n".format(
            md_esc(c["key"]), u["asis_rows"], u["asis_programs"], u["tobe_rows"], u["tobe_programs"],
            md_esc(c["label_standard"]["value"]), md_esc(c["align_standard"]["value"]),
            md_esc(c["width_standard"]["value"]), md_esc(c["display_standard"]["value"]),
            md_esc(e["control"]["value"]), md_esc(e["combo_sql_id"]["value"]),
            "O" if e["combo_registered"] else ("-" if cfg.is_unknown(e["combo_sql_id"]["value"]) else "X"),
            md_esc(e["modal_key"]["value"]), e["confidence"],
            md_esc(f'{pair["role"]}↔{pair["partner"]}' if pair else "-")))

    L.append("\n## 2. 교대쌍 (콤보 ↔ 텍스트)\n\n")
    L.append("| 콤보측 | 텍스트측 | 화면 | 시트 | 라벨 | 콤보측 컨트롤 | 콤보측 VIS | 텍스트측 VIS | 원천 |\n")
    L.append("|---|---|---|---|---|---|---|---|---|\n")
    for sp in dc["swap_pairs"]:
        for h in sp["hits"]:
            L.append("| `{}` | `{}` | {} | {} | {} | {} | {} | {} | {} |\n".format(
                md_esc(sp["combo_side"]), md_esc(sp["text_side"]), h["program"], md_esc(h["sheet"]),
                md_esc(h["label"]), md_esc(h["combo_side_ctrl"]),
                h["combo_side_visible"], h["text_side_visible"], h["side"].upper()))

    L.append("\n## 3. 적출 결함\n\n")
    L.append("| 코드 | 심각도 | 설명 | 건수 |\n|---|---|---|---:|\n")
    for code, desc in sorted(m["defect_codes"].items()):
        n = m["defects_by_code"].get(code, 0)
        sev = next((d["severity"] for d in dc["defects"] if d["code"] == code), "-")
        L.append(f"| {code} | {sev} | {md_esc(desc)} | {n} |\n")
    L.append(f"\n합계 — ERROR {m['defects_by_severity'].get('ERROR', 0)} · "
             f"WARN {m['defects_by_severity'].get('WARN', 0)} · "
             f"INFO {m['defects_by_severity'].get('INFO', 0)}\n")
    L.append("\n### 3.1 ERROR 전수\n\n")
    L.append("| 코드 | 화면 | 시트 | 컬럼 | 내용 |\n|---|---|---|---|---|\n")
    for d in sorted([x for x in dc["defects"] if x["severity"] == "ERROR"],
                    key=lambda x: (x["code"], x["program"], x["sheet"], x["column"])):
        L.append("| {} | {} | {} | `{}` | {} |\n".format(
            d["code"], d["program"], md_esc(d["sheet"]), md_esc(d["column"]), md_esc(d["detail"])))

    L.append("\n## 4. 표준 콤보 TO-BE 미등록 ({}컬럼)\n\n".format(
        len(m["combo_unregistered_standard"])))
    L.append("| COLUMN | 표준 콤보(AS-IS) | TO-BE 실사용 대안 |\n|---|---|---|\n")
    byk = {c["key"]: c for c in cols}
    for k in m["combo_unregistered_standard"]:
        e = byk[k]["edit_standard"]
        L.append("| `{}` | {} | {} |\n".format(
            md_esc(k), md_esc(e["combo_sql_id"]["value"]),
            md_esc(", ".join(e["combo_tobe_alternatives"]) or "-")))
    return "".join(L)


def build_md_measures(dm, md):
    m, mm = dm["meta"], md["meta"]
    L = ["# SHEET_MEASURE 지식사전 + AS-IS↔TO-BE 대사\n",
         head(m["asis_source"],
              "AS-IS 를 정본으로 두고 TO-BE 를 전수 대사한다. 차이는 누락(MISSING)/확장(EXTRA)/"
              "속성불일치(DIFF)로 나뉜다."),
         "## 0. 대사 규약\n\n",
         f"- **키**: {' + '.join(mm['key'])} — {mm['key_evidence']}\n",
         f"- **division 비교**: {mm['division_compare']}\n",
         f"- **대사 속성**: {', '.join(mm['compare_columns'])}\n",
         f"- **범위**: DM {len(mm['scope'])} PROGRAM_ID\n",
         f"- {m['note']}\n\n",
         "## 1. 대사 결과\n\n",
         "| 구분 | DM 범위 | 전 도메인 |\n|---|---:|---:|\n"]
    for cls in ("SAME", "DIFF", "MISSING", "EXTRA"):
        L.append(f"| {cls} | {mm['scoped_summary'].get(cls, 0)} | "
                 f"{mm['all_summary'].get(cls, 0)} |\n")
    L.append("\n### 1.1 DIFF 계통 (DM 범위)\n\n| 계통 | 건수 | 뜻 |\n|---|---:|---|\n")
    for kind, n in mm["diff_kinds"].items():
        L.append(f"| {kind} | {n} | {md_esc(mm['diff_kind_rule'].get(kind, '-'))} |\n")
    L.append(f"\n⇒ **설명되지 않는 실차이(REAL) {mm['real_diff']}건** 이 확인 대상이다.\n")
    L.append("\n### 1.2 속성별 DIFF 분포\n\n| 속성 | 건수 |\n|---|---:|\n")
    for c, n in mm["diff_columns"].items():
        L.append(f"| {c} | {n} |\n")
    L.append("\n### 1.3 화면별 요약 (SAME 외)\n\n| 화면 | 내역 |\n|---|---|\n")
    for p, c in mm["program_summary"].items():
        L.append(f"| {p} | {md_esc(c)} |\n")

    L.append("\n## 2. DM 범위 차이 전수 (SAME 제외)\n\n")
    L.append("| 화면 | 시트 | MEASURE | 구분 | 계통 | 내용/사유 |\n|---|---|---|---|---|---|\n")
    for r in md["rows"]:
        if not r["in_scope"] or r["class"] == "SAME":
            continue
        detail = r["reason"]
        if r["class"] == "DIFF":
            detail = "; ".join(
                f"{x['col']}: {x['asis']} → {x['tobe']}" for x in r["diffs"])
        L.append("| {} | {} | `{}` | {} | {} | {} |\n".format(
            r["program"], md_esc(r["sheet"]), md_esc(r["column"]), r["class"],
            md_esc("+".join(r["diff_kinds"]) or "-"), md_esc(detail)))

    L.append("\n## 3. MEASURE 표준 (전 {}종)\n\n".format(len(dm["measures"])))
    L.append("| MEASURE | AS-IS 행/화면 | TO-BE 행/화면 | 라벨 표준 | DEFAULT | 선택가능 | "
             "색(ATTR3) | division |\n")
    L.append("|---|---:|---:|---|---|---|---|---|\n")
    for x in dm["measures"]:
        u = x["usage"]
        L.append("| `{}` | {}/{} | {}/{} | {} | {} | {} | {} | {} |\n".format(
            md_esc(x["key"]), u["asis_rows"], u["asis_programs"], u["tobe_rows"], u["tobe_programs"],
            md_esc(x["label_standard"]["value"]), md_esc(x["default_flag_dist"]),
            md_esc(x["user_selectable_dist"]), md_esc(x["measure_color_attr3"] or "-"),
            md_esc(",".join(sorted(x["division_usage"])) or "-")))
    return "".join(L)


def build_index(sc, dc, dm, md, vr):
    m = dc["meta"]
    return f"""# UI 메타 지식사전 — 진입점

{head(m['asis_source'], '조회조건 / SHEET_COLUMN / SHEET_MEASURE 3종 지식사전의 목차와 사용법.')}
## 무엇을 언제 읽나

| 상황 | 읽을 파일 |
|---|---|
| 신규/보완 화면의 **조회조건** 작성 전 | [01_search_conditions.md](01_search_conditions.md) |
| 신규/보완 화면의 **그리드 컬럼 메타** 작성 전 | [02_sheet_columns.md](02_sheet_columns.md) |
| 화면의 **MEASURE** 등록·대조 시 | [03_sheet_measures.md](03_sheet_measures.md) |

## 규모

| 사전 | 항목 | 근거 |
|---|---:|---|
| 조회조건 | {len(sc['conditions'])}종 (154 조건 레코드 / 13UI 실측) | {md_esc(sc['meta']['primary_evidence'])} |
| SHEET_COLUMN | {len(dc['columns'])}종 | AS-IS {m['asis_rows']:,}행 + TO-BE {m['tobe_rows']:,}행 |
| SHEET_MEASURE | {len(dm['measures'])}종 | AS-IS {dm['meta']['asis_rows']:,}행 + TO-BE {dm['meta']['tobe_rows']:,}행 |

## 현재 적출 현황

- 컬럼 결함 — ERROR **{m['defects_by_severity'].get('ERROR', 0)}** · WARN {m['defects_by_severity'].get('WARN', 0)} · INFO {m['defects_by_severity'].get('INFO', 0)}
- 표준 콤보 TO-BE 미등록 **{len(m['combo_unregistered_standard'])}컬럼** (매핑 판단 대상)
- MEASURE 대사 (DM) — DIFF {md['meta']['scoped_summary'].get('DIFF', 0)} / MISSING {md['meta']['scoped_summary'].get('MISSING', 0)} / EXTRA {md['meta']['scoped_summary'].get('EXTRA', 0)} · 설명되지 않는 실차이 **{md['meta']['real_diff']}건**
- 게이트: **{vr['result']}** (FAIL {len(vr['fail'])} / WARN {len(vr['warn'])})

## 사용 규약

1. 사전은 **대조·검출용**이다. 사전 값으로 `SCM_MST_SHEET_COLUMNS` 를 자동 INSERT 하지 않는다
   (`ui-category-setup-no-arbitrary-data`). 등록은 운영자·사용자 결정이다.
2. `confidence` 가 `low`/`none` 인 컬럼은 표준이 아니다 — AS-IS 화면을 확인하고 결정한다.
3. 화면 작업 종료 전 `python 06_verify.py --program {{MENU_ID}}` 로 해당 화면 ERROR 0 을 확인한다.
4. 원천이 바뀌면 파이프라인을 재실행한다. MD 를 직접 고치지 않는다.

## 한계

1. AS-IS 근거는 **DEV**(`LGCE_SCMDEV`) 스냅샷이다. QA(10.34.221.91:7433)는 본 PC 에서 도달 불가로
   확인됐다. QA 개방 시 환경변수만 바꿔 재실행하고 차이를 대사한다.
   ⚠ DEV 는 작업 중 변동한다 — 2026-08-21 당일 테이블 재적재가 실측됐다
   (`asis_volatility_note.json`).
2. **편집·모달 축은 AS-IS 메타로 검증 불가**하다 (AS-IS `INPUT_TP_CD` 전행 'S' · `FPModal` 0건).
3. 조회조건은 13UI 만 소스 실측이다. 그 밖은 카탈로그(2026-01-22 스냅샷) 커버리지로 표기된다.
4. 컨트롤·정렬 허용값 정본 = **TO-BE `SCM_ADM_LOOKUP_MASTER_DETAIL`**(64,500행) — AS-IS 와 동일
   (컨트롤 13종·align C/L/R)임을 교차검증했다. 허용 컨트롤 = FE enum ∪ lookup.
   ⚠ 2026-08-21 15:57 적재 창에서 이 테이블이 **일시적으로 0행**이었다. 그 상태로 사전을 만들면
   허용값이 비어 전 컬럼이 위반으로 뒤집히므로, `02_dump_tobe.py` 는 0행이면 **중단**한다.
"""


# ─────────────────────────────────────────────────────────────────────────────
def build_xlsx(sc, dc, dm, md, path):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("0.요약")
    title(ws, "UI 메타 지식사전 v1.0 — 요약", 4)
    ws.append([])
    ws.append(["구분", "항목", "값", "비고"])
    rows = [
        ("스냅샷", "기준일", cfg.SNAPSHOT_DATE, ""),
        ("스냅샷", "AS-IS 원천", dc["meta"]["asis_source"], "DEV — 작업 중 변동 실측"),
        ("규모", "조회조건 항목", len(sc["conditions"]), "154 조건 레코드 / 13UI 실측"),
        ("규모", "COLUMN 종", len(dc["columns"]),
         f'AS-IS {dc["meta"]["asis_rows"]}행 + TO-BE {dc["meta"]["tobe_rows"]}행'),
        ("규모", "MEASURE 종", len(dm["measures"]),
         f'AS-IS {dm["meta"]["asis_rows"]}행 + TO-BE {dm["meta"]["tobe_rows"]}행'),
        ("결함", "COLUMN ERROR", dc["meta"]["defects_by_severity"].get("ERROR", 0), ""),
        ("결함", "COLUMN WARN", dc["meta"]["defects_by_severity"].get("WARN", 0), ""),
        ("결함", "COLUMN INFO", dc["meta"]["defects_by_severity"].get("INFO", 0),
         "FE 전달 의무 고지"),
        ("격차", "표준 콤보 TO-BE 미등록", len(dc["meta"]["combo_unregistered_standard"]),
         "매핑 판단 대상"),
        ("대사", "MEASURE DIFF(DM)", md["meta"]["scoped_summary"].get("DIFF", 0), ""),
        ("대사", "MEASURE MISSING(DM)", md["meta"]["scoped_summary"].get("MISSING", 0), "이관 누락"),
        ("대사", "MEASURE EXTRA(DM)", md["meta"]["scoped_summary"].get("EXTRA", 0), ""),
        ("대사", "설명되지 않는 실차이(REAL)", md["meta"]["real_diff"], "확인 대상"),
    ]
    for r in rows:
        ws.append(list(r))
    style_sheet(ws, [14, 30, 16, 60])

    ws = wb.create_sheet("1.조회조건_항목")
    title(ws, "조회조건 지식사전 — 항목 축", 11)
    ws.append([])
    ws.append(["조회조건", "사용 UI", "TO-BE 필드", "binding", "콤보(CMBO)", "@토큰",
               "div40", "div81", "div82", "카탈로그 화면", "커버리지"])
    for c in sc["conditions"]:
        d = c["division_usage_standard"]
        ws.append([c["cond"], c["ui_count"], ", ".join(c["tobe_fields"]), c["binding_primary"],
                   ", ".join(c["combo_sql_ids"]), ", ".join(c["combo_tokens"]),
                   d["40"], d["81"], d["82"], c["catalog_screens"], c["coverage"]])
    style_sheet(ws, [22, 8, 26, 22, 28, 22, 14, 14, 14, 12, 20])

    ws = wb.create_sheet("2.조회조건_화면별")
    title(ws, "조회조건 — 화면별 상세 (154 조건)", 12)
    ws.append([])
    ws.append(["조회조건", "UI", "AS-IS 라벨", "TO-BE 라벨", "TO-BE 컨트롤", "콤보", "binding",
               "필수", "기본값", "div40", "div81", "div82"])
    for c in sc["conditions"]:
        for s in c["screens"]:
            d = s["division_usage"]
            ws.append([c["cond"], s["ui"], s["label_asis"], s["label_tobe"], s["control_tobe"],
                       s["combo_sql_id"], s["binding"]["primary"], s["required"], s["default"],
                       d["40"]["use"], d["81"]["use"], d["82"]["use"]])
    style_sheet(ws, [20, 10, 24, 24, 30, 24, 14, 20, 22, 10, 10, 10])

    ws = wb.create_sheet("3.COLUMN_표준")
    title(ws, "SHEET_COLUMN 지식사전 — 표시/편집 2단 표준", 14)
    ws.append([])
    ws.append(["COLUMN", "AS-IS행", "AS-IS화면", "TO-BE행", "TO-BE화면", "라벨 표준", "정렬", "폭",
               "표시 컨트롤", "편집 컨트롤", "편집 콤보", "콤보등록", "모달키", "confidence"])
    for c in dc["columns"]:
        u, e = c["usage"], c["edit_standard"]
        reg = "-" if cfg.is_unknown(e["combo_sql_id"]["value"]) else ("O" if e["combo_registered"] else "X")
        ws.append([c["key"], u["asis_rows"], u["asis_programs"], u["tobe_rows"], u["tobe_programs"],
                   c["label_standard"]["value"], c["align_standard"]["value"],
                   c["width_standard"]["value"], c["display_standard"]["value"],
                   e["control"]["value"], e["combo_sql_id"]["value"], reg,
                   e["modal_key"]["value"], e["confidence"]])
        if reg == "X":
            ws.cell(row=ws.max_row, column=12).fill = WARN_FILL
    style_sheet(ws, [26, 8, 9, 8, 9, 24, 6, 6, 18, 16, 26, 8, 20, 11])

    ws = wb.create_sheet("4.COLUMN_결함")
    title(ws, "SHEET_COLUMN 적출 결함", 6)
    ws.append([])
    ws.append(["심각도", "코드", "화면", "시트", "컬럼", "내용"])
    for d in sorted(dc["defects"], key=lambda x: (x["severity"] != "ERROR", x["code"],
                                                  x["program"], x["sheet"], x["column"])):
        ws.append([d["severity"], d["code"], d["program"], d["sheet"], d["column"], d["detail"]])
        f = ERR_FILL if d["severity"] == "ERROR" else (WARN_FILL if d["severity"] == "WARN" else None)
        if f:
            ws.cell(row=ws.max_row, column=1).fill = f
    style_sheet(ws, [10, 8, 12, 26, 28, 90])

    ws = wb.create_sheet("5.MEASURE_표준")
    title(ws, "SHEET_MEASURE 지식사전", 9)
    ws.append([])
    ws.append(["MEASURE", "AS-IS행", "AS-IS화면", "TO-BE행", "TO-BE화면", "라벨 표준",
               "DEFAULT 분포", "선택가능 분포", "division"])
    for x in dm["measures"]:
        u = x["usage"]
        ws.append([x["key"], u["asis_rows"], u["asis_programs"], u["tobe_rows"], u["tobe_programs"],
                   x["label_standard"]["value"], str(x["default_flag_dist"]),
                   str(x["user_selectable_dist"]), ",".join(sorted(x["division_usage"]))])
    style_sheet(ws, [26, 8, 9, 8, 9, 30, 18, 18, 30])

    ws = wb.create_sheet("6.MEASURE_대사")
    title(ws, "SHEET_MEASURE AS-IS↔TO-BE 대사 (전 도메인)", 7)
    ws.append([])
    ws.append(["화면", "시트", "MEASURE", "구분", "DM범위", "계통", "내용/사유"])
    for r in md["rows"]:
        if r["class"] == "SAME":
            continue
        detail = r["reason"]
        if r["class"] == "DIFF":
            detail = "; ".join(f"{x['col']}: {x['asis']} → {x['tobe']}" for x in r["diffs"])
        ws.append([r["program"], r["sheet"], r["column"], r["class"],
                   "Y" if r["in_scope"] else "", "+".join(r["diff_kinds"]), detail])
        f = ERR_FILL if r["class"] == "MISSING" else (
            WARN_FILL if r["class"] == "EXTRA" else None)
        if f:
            ws.cell(row=ws.max_row, column=4).fill = f
    style_sheet(ws, [12, 26, 26, 10, 8, 20, 80])

    ws = wb.create_sheet("7.교대쌍")
    title(ws, "교대쌍 (콤보 ↔ 텍스트) — AS-IS .cs 가 모드별로 번갈아 노출", 9)
    ws.append([])
    ws.append(["콤보측", "텍스트측", "원천", "화면", "시트", "라벨", "콤보측 컨트롤",
               "콤보측 VISIBLE", "텍스트측 VISIBLE"])
    for sp in dc["swap_pairs"]:
        for h in sp["hits"]:
            ws.append([sp["combo_side"], sp["text_side"], h["side"].upper(), h["program"],
                       h["sheet"], h["label"], h["combo_side_ctrl"],
                       h["combo_side_visible"], h["text_side_visible"]])
    style_sheet(ws, [22, 22, 8, 12, 24, 20, 18, 14, 14])

    ws = wb.create_sheet("8.콤보_미등록표준")
    title(ws, "표준 콤보 TO-BE 미등록 — 매핑 판단 대상", 3)
    ws.append([])
    ws.append(["COLUMN", "표준 콤보(AS-IS)", "TO-BE 실사용 대안"])
    byk = {c["key"]: c for c in dc["columns"]}
    for k in dc["meta"]["combo_unregistered_standard"]:
        e = byk[k]["edit_standard"]
        ws.append([k, e["combo_sql_id"]["value"], ", ".join(e["combo_tobe_alternatives"])])
    style_sheet(ws, [30, 34, 50])

    wb.save(path)
    return path


def main():
    sc = cfg.jload("dict_search_conditions.json")
    dc = cfg.jload("dict_sheet_columns.json")
    dm = cfg.jload("dict_sheet_measures.json")
    md = cfg.jload("measure_diff.json")
    vr = cfg.jload("verify_report.json")

    os.makedirs(cfg.DICT_DIR, exist_ok=True)
    files = {
        "00_INDEX.md": build_index(sc, dc, dm, md, vr),
        "01_search_conditions.md": build_md_search(sc),
        "02_sheet_columns.md": build_md_columns(dc),
        "03_sheet_measures.md": build_md_measures(dm, md),
    }
    for name, body in files.items():
        p = os.path.join(cfg.DICT_DIR, name)
        with open(p, "w", encoding="utf-8", newline="\n") as f:
            f.write(body)
        print(f"  MD  {p}  ({len(body):,} bytes / {body.count(chr(10)):,} lines)")

    os.makedirs(cfg.XLSX_DIR, exist_ok=True)
    xp = build_xlsx(sc, dc, dm, md, os.path.join(cfg.XLSX_DIR, cfg.XLSX_NAME))
    print(f"  XLSX {xp}")


if __name__ == "__main__":
    sys.exit(main())
