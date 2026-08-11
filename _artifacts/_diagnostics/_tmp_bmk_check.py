# -*- coding: utf-8 -*-
"""SCM_ADM_USER_BMK_M 누락 여부 + 성격 확인, 그리고 같은 계열(BE 범용 코드가 읽는 테이블) 재탐색"""
import sys
import pathlib
import re

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from ora import conn_info  # noqa: E402
import oracledb  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

P = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI_DB객체_전수취합_v2.2.xlsx")
wb = load_workbook(P, read_only=True, data_only=True)
objs, tabs = set(), set()
for r in list(wb["1_UI별_DB객체"].iter_rows(values_only=True))[1:]:
    if r and len(r) > 6 and r[6]:
        objs.add(str(r[6]).strip().upper())
for r in list(wb["3_TABLE_이관SOURCE"].iter_rows(values_only=True))[1:]:
    if r and len(r) > 1 and r[1]:
        tabs.add(str(r[1]).strip().upper())
wb.close()

T = "SCM_ADM_USER_BMK_M"
print("== %s ==" % T)
print("  시트1 포함:", "O" if T in objs else "— 누락")
print("  시트3 포함:", "O" if T in tabs else "— 누락")

u, p, d = conn_info()
con = oracledb.connect(user=u, password=p, dsn=d)
cur = con.cursor()
cur.execute("""SELECT COUNT(*) FROM ALL_TABLES WHERE OWNER='SCM_MGR' AND TABLE_NAME=:t""", {"t": T})
print("  DB 실재:", "O" if cur.fetchone()[0] else "X")
cur.execute("""SELECT COLUMN_NAME, DATA_TYPE, NULLABLE FROM ALL_TAB_COLUMNS
                WHERE OWNER='SCM_MGR' AND TABLE_NAME=:t ORDER BY COLUMN_ID""", {"t": T})
cols = cur.fetchall()
print("  컬럼:", [c[0] for c in cols])
try:
    cur.execute("SELECT COUNT(*) FROM SCM_MGR.%s" % T)
    print("  행수:", cur.fetchone()[0])
except Exception as e:  # noqa: BLE001
    print("  행수 조회 실패:", str(e)[:80])

print("\n== 이 테이블을 참조하는 DB 오브젝트(SP/PKG/VIEW) 존재하는가 ==")
cur.execute("""SELECT NAME, TYPE, COUNT(*) FROM ALL_SOURCE
                WHERE OWNER='SCM_MGR' AND INSTR(UPPER(TEXT), :t) > 0
                GROUP BY NAME, TYPE ORDER BY 1""", {"t": T})
rows = cur.fetchall()
for r in rows:
    print("  %-34s %-14s %d줄" % r)
if not rows:
    print("  (없음) → SP 의존폐포로는 절대 안 잡힘 = API_M 기점 취합의 사각지대")

print("\n== 콤보 SQL 본문에서의 참조 ==")
cur.execute("SELECT SQL_ID FROM SCM_MGR.SCM_ADM_CMBO_SQL_M WHERE UPPER(SQL_DTL_CTN) LIKE '%'||:t||'%'", {"t": T})
r2 = [x[0] for x in cur.fetchall()]
print("  ", r2 or "(없음)")

# ── SCM_ADM_* / SCM_COM_* 전 테이블 중 취합 미포함 전수 ──
print("\n== SCM_ADM_* / SCM_COM_* 테이블 중 취합 미포함 전수 ==")
cur.execute("""SELECT TABLE_NAME FROM ALL_TABLES
                WHERE OWNER='SCM_MGR' AND (TABLE_NAME LIKE 'SCM_ADM_%' OR TABLE_NAME LIKE 'SCM_COM_%')
                ORDER BY TABLE_NAME""")
allmeta = [r[0] for r in cur.fetchall()]
miss = [t for t in allmeta if t not in objs and t not in tabs]
print("  전체 %d개 중 미포함 %d개" % (len(allmeta), len(miss)))
for t in miss:
    cur.execute("""SELECT COUNT(*) FROM ALL_SOURCE WHERE OWNER='SCM_MGR' AND INSTR(UPPER(TEXT), :t) > 0""", {"t": t})
    insrc = cur.fetchone()[0]
    try:
        cur.execute("SELECT COUNT(*) FROM SCM_MGR.%s" % t)
        n = cur.fetchone()[0]
    except Exception:  # noqa: BLE001
        n = "?"
    print("  ⛔ %-38s 행수=%-9s SP참조=%s" % (t, n, "있음(%d줄)" % insrc if insrc else "없음"))
con.close()
