# -*- coding: utf-8 -*-
"""TOBE_UI_DB객체_전수취합 v2.2 의 시트/컬럼/수집 대상 확인 — 무엇이 들어갔고 무엇이 빠졌나"""
import sys
import pathlib

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from openpyxl import load_workbook  # noqa: E402

P = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI_DB객체_전수취합_v2.2.xlsx")
wb = load_workbook(P, read_only=True, data_only=True)
print("파일:", P.name)
print("시트:", wb.sheetnames)

TARGETS = ("SCM_ADM_CMBO_SQL_M", "CMBO", "SCM_ADM_MSG_M", "SCM_ADM_LBL", "SCM_ADM_MNU",
           "LOOKUP_MASTER", "USER_GR", "SHEET_COLUMNS", "PROC_M", "API_M", "BTN_MPG")

for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    print("\n=== [%s] %d행 x %d열 ===" % (name, len(rows), ws.max_column or 0))
    if not rows:
        continue
    hdr = [str(c) if c is not None else "" for c in rows[0]]
    print("  헤더:", hdr[:14])
    # 대상 토큰이 이 시트 어디에 나오는지
    hits = {}
    for r in rows[1:]:
        flat = " ".join(str(c) for c in r if c is not None).upper()
        for t in TARGETS:
            if t.upper() in flat:
                hits[t] = hits.get(t, 0) + 1
    print("  토큰 등장:", hits if hits else "없음")
    for r in rows[1:4]:
        print("   샘플:", [str(c)[:26] if c is not None else "" for c in r][:9])
wb.close()
