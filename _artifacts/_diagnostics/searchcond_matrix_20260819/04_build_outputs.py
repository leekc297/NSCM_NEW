#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""04_build_outputs.py — DMF 13UI 조회조건 MATRIX 산출물 3종 생성 (플랜 §3)

입력: extract/{MENU_ID}.json 13본 + cmbo_meta.json + verify_report.json
출력: Doc\화면분석서\DMF_13UI_조회조건_MATRIX_v1.0.{xlsx,md,html}
서식: _gen_DMF0074_validation_xlsx.py 의 style_sheet/FILL/BORDER 이식 (정본 서식)
"""
import html as html_mod
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = r"C:\Projects\NSCM\Doc\화면분석서"
BASE = "DMF_13UI_조회조건_MATRIX_v1.0"

UIS = ["DMF0003", "DMF0004", "DMF0005", "DMF0006", "DMF0007", "DMF0014", "DMF0020",
       "DMF0021", "DMF0042", "DMF0048", "DMF0072", "DMF0074", "DMF0082"]

# 가로축 표준 순서 (hook #112 rename) — 미등재 조건은 뒤에 최초 출현 순
COND_ORDER = ["Div.", "Team", "End Customer", "Sold Customer", "Sales Site", "Plant Site",
              "PJT", "Model/Nickname", "Customer PJT", "Customer Model", "Item Type",
              "Biz. Type", "Data Type", "Bucket", "Period", "Ver.", "Owner", "Sender Code",
              "DOM/EXP", "Active(YesOrNo)", "Search Over Range", "Detail",
              "End Customer(Report) Null", "Project Null",
              "Sum(체크)", "Color(체크)", "Qty≥0(체크)", "Qty≠0(체크)", "Cell Qty(체크)",
              "Include Internal Txn(체크)", "AS(체크)", "MKT FCST 미등록 제외(체크)",
              "Price 미등록 제외(체크)", "Closed(숨김)"]

# 조건 표준명 정규화 — 추출 에이전트 간 이형 표기 통일 (원본 extract JSON 은 불변)
COND_ALIAS = {
    "Sum": "Sum(체크)",
    "Sum (표준명 외 — 체크박스)": "Sum(체크)",
    "Sum(체크)": "Sum(체크)",
    "Color": "Color(체크)",
    "Color (표준명 외 — 체크박스)": "Color(체크)",
    "Qty≥0(체크)": "Qty≥0(체크)",
    "Qty≥0 (표준명 외 — 체크박스)": "Qty≥0(체크)",
    "Qty≠0": "Qty≠0(체크)",
    "Cell Qty (표준명 외)": "Cell Qty(체크)",
    "Cell Qty (표준명 외 — 체크박스)": "Cell Qty(체크)",
    "Cell Qty (표준명 외 — 숨김 dead)": "Cell Qty(체크)",
    "Include Internal Txn": "Include Internal Txn(체크)",
    "Include Internal Txn (표준명 외 — 체크박스)": "Include Internal Txn(체크)",
    "Active(YesOrNo) — AS 체크": "AS(체크)",
    "MKT FCST 미등록 제외 (표준명 외 — 체크박스, Invalid)": "MKT FCST 미등록 제외(체크)",
    "Price 미등록 제외": "Price 미등록 제외(체크)",
    "Closed": "Closed(숨김)",
}

# ── 서식 (정본 이식: _gen_DMF0074_validation_xlsx.py) ────────────────────────
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
SEC_FILL = PatternFill("solid", fgColor="DDEBF7")
ASIS_FILL = PatternFill("solid", fgColor="FFF2CC")   # AS-IS 상단 행
TOBE_FILL = PatternFill("solid", fgColor="E7F4E4")   # TO-BE 하단 행
NA_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE_FONT = Font(size=9)
TITLE_FONT = Font(bold=True, size=13, color="1F4E79")


def style_sheet(ws, widths, freeze="A3", header_row=2):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[header_row]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for row in ws.iter_rows(min_row=header_row + 1):
        for c in row:
            if c.font is None or not c.font.bold:
                c.font = BASE_FONT
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = BORDER
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}"


def title(ws, text, span):
    ws.cell(row=1, column=1, value=text).font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)


# ── 데이터 로드 ───────────────────────────────────────────────────────────────
def load():
    data = {}
    for mid in UIS:
        with open(os.path.join(HERE, "extract", f"{mid}.json"), encoding="utf-8") as f:
            data[mid] = json.load(f)
        # 조건 표준명 정규화 + UI 내 중복 방지 assert
        seen = set()
        for c in data[mid]["conditions"]:
            c["cond"] = COND_ALIAS.get(c["cond"], c["cond"])
            assert c["cond"] not in seen, f"{mid}: 정규화 후 조건명 중복 '{c['cond']}'"
            seen.add(c["cond"])
    cmbo = json.load(open(os.path.join(HERE, "cmbo_meta.json"), encoding="utf-8"))
    verify = json.load(open(os.path.join(HERE, "verify_report.json"), encoding="utf-8"))
    return data, cmbo, verify


def cond_axis(data):
    axis = list(COND_ORDER)
    for mid in UIS:
        for c in data[mid]["conditions"]:
            if c["cond"] not in axis:
                axis.append(c["cond"])
    used = {c["cond"] for mid in UIS for c in data[mid]["conditions"]}
    return [c for c in axis if c in used]


def div_lines(s):
    d40, d81, d82 = (s.get("div40", ""), s.get("div81", ""), s.get("div82", ""))
    if d40 == d81 == d82 == "공통":
        return []
    return [f"[40] {d40}", f"[81] {d81}", f"[82] {d82}"]


def cell_text(s, with_evidence=False):
    """side dict → 매트릭스 셀 멀티라인 텍스트"""
    if not s:
        return "N/A(없음)"
    lines = [f"〔{s.get('control', '?')}〕 {s.get('label', '')}".rstrip()]
    combo = s.get("combo", "")
    if combo and combo not in ("없음", "N/A"):
        lines.append(f"COMBO: {combo}")
    params = s.get("params", "")
    if params and params not in ("없음", "N/A"):
        lines.append(f"변수: {params}")
    filt = s.get("filter", "")
    if filt and filt not in ("없음", "N/A"):
        lines.append(f"FILTER: {filt}")
    casc = s.get("cascade", "")
    if casc and casc not in ("없음", "N/A"):
        lines.append(f"CASCADE: {casc}")
    lines += div_lines(s)
    req, dft = s.get("required", ""), s.get("default", "")
    tail = []
    if req and req != "N":
        tail.append(f"필수:{req}")
    if dft and dft not in ("없음", "N/A"):
        tail.append(f"기본:{dft}")
    if tail:
        lines.append(" · ".join(tail))
    if with_evidence and s.get("evidence"):
        lines.append("근거: " + ", ".join(s["evidence"]))
    return "\n".join(lines)


def find_cond(data, mid, cond):
    for c in data[mid]["conditions"]:
        if c["cond"] == cond:
            return c
    return None


# ── SIMPLE 버전: 콤보ID만 추출 ────────────────────────────────────────────────
import re  # noqa: E402

ASIS_ID_PATTERNS = [
    re.compile(r"GetDMInnerTable\('(\w+)'\)"),
    re.compile(r"Data\.Common\(Data\.Type\.(\w+)"),
    re.compile(r"@Code\s*=\s*'(\w+)'"),
    re.compile(r"UP_FX\s+'(\w+)'"),
    re.compile(r"Data\.Type\.(\w+)"),
    re.compile(r"Data\.(\w+)(?:\s|/|\(|$)"),
    re.compile(r"\b(UP_FX_\w+)"),
    re.compile(r"\bSP\s+([\w.]+)"),
    re.compile(r"\b(DM\.DMF\w+)"),
]


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def _extract_ids(text, side, cmbo_ids):
    """콤보 원천 문자열에서 콤보ID 후보 추출 (TO-BE=CMBO 실덤프 case-insensitive 매칭 /
    AS-IS=호출 패턴). 대소문자 이형·중복은 최초 표기로 dedupe."""
    ids = []

    def add(x):
        if _norm(x) not in {_norm(i) for i in ids}:
            ids.append(x)

    if side == "tobe":
        # 대소문자 이형 CMBO 쌍(BizTypeByDivision/BizTypebyDivision) 존재 — 토큰 원표기 우선,
        #   lower 매칭은 유일할 때만 canonical 치환. FE 상수 토큰은 제외.
        stop = {"DIVISION", "MOBILITY", "CMBO", "Form", "ALL"}
        by_lower = {}
        for c in cmbo_ids:
            by_lower.setdefault(c.lower(), []).append(c)
        for w in re.findall(r"[A-Za-z_][A-Za-z0-9_]{2,}", text or ""):
            if w in stop:
                continue
            if w in cmbo_ids:
                add(w)
            elif w.lower() in by_lower and len(by_lower[w.lower()]) == 1:
                add(by_lower[w.lower()][0])
    else:
        for pat in ASIS_ID_PATTERNS:
            for m in pat.findall(text or ""):
                if m not in ("Common", "Type"):
                    add(m)
    return ids


def _match_div_ids(div_text, base_ids):
    """division 서술 텍스트의 '토큰'이 base ID 와 동치면 그 ID 반환. 토큰 단위 비교 —
    정규화 완전 동치(exact) 우선, exact 0건일 때만 접미 동치(SP 스키마 접두 생략 케이스) 허용."""
    exact, suffix = [], []
    tokens = re.findall(r"[A-Za-z_][A-Za-z0-9_.]{3,}", div_text or "")
    for bid in base_ids:
        nb = _norm(bid)
        for tk in tokens:
            nt = _norm(tk)
            if nb == nt:
                if bid not in exact:
                    exact.append(bid)
                break
            if nb.endswith(nt) or nt.endswith(nb):
                if bid not in suffix:
                    suffix.append(bid)
    return exact if exact else suffix


def _ctl_tag(control):
    c = str(control or "")
    for key, tag in (("팝업", "팝업"), ("달력", "달력"), ("체크", "체크"),
                     ("텍스트", "텍스트"), ("라디오", "라디오")):
        if key in c:
            return tag
    return "—"


def simple_cell(s, side, cmbo_ids):
    """SIMPLE 셀: 콤보ID만. division 별 상이 시 '40: X / 81: Y / 82: Z' (동일 division 병합)."""
    if not s:
        return ""
    control = str(s.get("control", ""))
    if not control.startswith("콤보"):
        return _ctl_tag(control)
    base_ids = _extract_ids(str(s.get("combo", "")), side, cmbo_ids)
    # division 분기: combo 에서 뽑은 base ID 가 div 서술 토큰과 동치일 때만 그 division 에 귀속
    #   (자유 서술 속 테이블명 등 오매칭 방지). '81 동일' 류 승계 반영.
    per_div = {}
    for div, key in (("40", "div40"), ("81", "div81"), ("82", "div82")):
        txt = str(s.get(key, ""))
        if txt == "공통":
            continue
        hits = _match_div_ids(txt, base_ids)
        if hits:
            per_div[div] = " / ".join(hits)
        elif div == "82" and "81" in txt and re.search(r"동일|같", txt) and "81" in per_div:
            per_div["82"] = per_div["81"]
        elif div == "81" and "40" in txt and re.search(r"동일|같", txt) and "40" in per_div:
            per_div["81"] = per_div["40"]
    # 전 division 동일 값이면 분기 불필요
    if per_div and len(set(per_div.values())) == 1 and len(per_div) == 3:
        return per_div["40"]
    if per_div and len(base_ids) <= 1 and len(set(per_div.values())) == 1:
        per_div = {}
    if per_div:
        full = {}
        for div in ("40", "81", "82"):
            full[div] = per_div.get(div) or (" / ".join(base_ids) if base_ids else "?")
        # 동일 값 division 병합 (예: 81·82 동일)
        groups = []
        for div in ("40", "81", "82"):
            for g in groups:
                if g["val"] == full[div]:
                    g["divs"].append(div)
                    break
            else:
                groups.append({"divs": [div], "val": full[div]})
        if len(groups) == 1:
            return groups[0]["val"]
        return " / ".join(f"{'·'.join(g['divs'])}: {g['val']}" for g in groups)
    if base_ids:
        return " / ".join(base_ids)
    # ID 미추출 콤보 (정적 옵션·프로시저 시트 등) — 원문 축약
    raw = str(s.get("combo", "")).strip()
    return raw if len(raw) <= 40 else raw[:37] + "…"


# ── xlsx ─────────────────────────────────────────────────────────────────────
def build_xlsx(data, cmbo, verify, axis):
    wb = Workbook()

    # 1.MATRIX
    ws = wb.active
    ws.title = "1.MATRIX"
    title(ws, "DMF 13UI 조회조건 MATRIX — 가로:조회조건 / 세로:UI×2 (AS-IS 상단·TO-BE 하단) / 실측 2026-08-19", 4 + len(axis))
    hdr = ["UI", "화면명", "구분"] + axis
    ws.append(hdr)
    for mid in UIS:
        d = data[mid]
        for side, side_nm in (("asis", "AS-IS"), ("tobe", "TO-BE")):
            row = [mid, d.get("ui_name_asis", ""), side_nm]
            for cond in axis:
                c = find_cond(data, mid, cond)
                row.append(cell_text((c or {}).get(side)) if c else "-")
            ws.append(row)
            fill = ASIS_FILL if side == "asis" else TOBE_FILL
            r = ws.max_row
            for col in (1, 2, 3):
                ws.cell(row=r, column=col).fill = fill
            for col in range(4, 4 + len(axis)):
                if ws.cell(row=r, column=col).value in ("-", "N/A(없음)"):
                    ws.cell(row=r, column=col).fill = NA_FILL
    style_sheet(ws, [9, 16, 7] + [34] * len(axis), freeze="D3", header_row=2)
    for r in range(3, ws.max_row + 1):
        ws.row_dimensions[r].height = 150

    # 1S.MATRIX_SIMPLE — 콤보ID만 (division 별 상이 시 분기 표기)
    cmbo_ids = {r["SQL_ID"] for r in cmbo}
    ws1s = wb.create_sheet("1S.MATRIX_SIMPLE", 1)
    title(ws1s, "SIMPLE — 각 조회조건의 콤보ID만 표기 (division 별 상이 시 40/81/82 분기 · 콤보 아님=유형 태그)", 4 + len(axis))
    ws1s.append(hdr)
    for mid in UIS:
        d = data[mid]
        for side, side_nm in (("asis", "AS-IS"), ("tobe", "TO-BE")):
            row = [mid, d.get("ui_name_asis", ""), side_nm]
            for cond in axis:
                c = find_cond(data, mid, cond)
                row.append(simple_cell((c or {}).get(side), side, cmbo_ids) if c else "-")
            ws1s.append(row)
            fill = ASIS_FILL if side == "asis" else TOBE_FILL
            r = ws1s.max_row
            for col in (1, 2, 3):
                ws1s.cell(row=r, column=col).fill = fill
            for col in range(4, 4 + len(axis)):
                if ws1s.cell(row=r, column=col).value in ("-", "", "—"):
                    ws1s.cell(row=r, column=col).fill = NA_FILL
    style_sheet(ws1s, [9, 16, 7] + [24] * len(axis), freeze="D3", header_row=2)
    for r in range(3, ws1s.max_row + 1):
        ws1s.row_dimensions[r].height = 42

    # 2.콤보정의
    ws2 = wb.create_sheet("2.콤보정의")
    title(ws2, "사용 콤보 정의 — TO-BE SCM_ADM_CMBO_SQL_M 실덤프 + AS-IS 원천", 7)
    ws2.append(["콤보 ID", "구분", "MDU", "@토큰", "SQL_NM (등록 설명 원문)", "사용 UI", "비고"])
    cmbo_by_id = {}
    for r in cmbo:
        cmbo_by_id.setdefault(r["SQL_ID"], []).append(r)
    used_tobe, used_asis = {}, {}
    for mid in UIS:
        for c in data[mid]["conditions"]:
            t = c.get("tobe")
            if t:
                for cid in cmbo_by_id:
                    if cid in str(t.get("combo", "")):
                        used_tobe.setdefault(cid, set()).add(mid)
            a = c.get("asis")
            if a and str(a.get("control", "")).startswith("콤보"):
                key = str(a.get("combo", ""))
                used_asis.setdefault(key, set()).add(mid)
    for cid in sorted(used_tobe):
        for rec in cmbo_by_id[cid]:
            ws2.append([cid, "TO-BE", rec["MDU_DIVS_CD"], ", ".join(rec["TOKENS"]) or "-",
                        rec["SQL_NM"] or "-", ", ".join(sorted(used_tobe[cid])), ""])
    for key in sorted(used_asis):
        ws2.append([key, "AS-IS", "-", "-", "-", ", ".join(sorted(used_asis[key])),
                    "AS-IS 원천 호출 (BigFilter.xml / UP_FX)"])
    style_sheet(ws2, [34, 7, 6, 30, 70, 26, 30])

    # 3.CASCADE
    ws3 = wb.create_sheet("3.CASCADE")
    title(ws3, "UI 별 CASCADE 체인 (상위 변경 → 하위 clear/재조회)", 3)
    ws3.append(["UI", "화면명", "CASCADE 체인"])
    for mid in UIS:
        d = data[mid]
        chains = d.get("cascade_chains") or ["미기재"]
        ws3.append([mid, d.get("ui_name_asis", ""), "\n".join(chains)])
    style_sheet(ws3, [9, 18, 120])

    # 4.검증대사
    ws4 = wb.create_sheet("4.검증대사")
    title(ws4, "검증 게이트 결과 (03_verify.py — hook #112 카탈로그 대사 + 실존 검증)", 8)
    ws4.append(["UI", "조건 수", "AS-IS", "TO-BE", "카탈로그 화면명", "카탈로그 As-Is/To-Be 행",
                "대조 사유", "비고"])
    for mid in UIS:
        u = verify["ui"].get(mid, {})
        ws4.append([mid, u.get("conds"), u.get("asis"), u.get("tobe"),
                    u.get("catalog_screen") or "카탈로그 미수록",
                    f"{u.get('catalog_asis_rows')}/{u.get('catalog_tobe_rows')}",
                    u.get("reason") or "-", ""])
    ws4.append([])
    ws4.append([f"FAIL {len(verify['fails'])}건 / WARN {len(verify['warns'])}건", "", "", "", "", "", "", ""])
    ws4.append(["※ WARN 해설", "① 'evidence 파일 미해석/축약' = 근거 표기가 범용 파일명(Designer.cs 등)이라 "
                "자동 대조기가 못 찾은 것 — 각 UI 폴더의 {화면명}.Designer.cs 실측 라인임. "
                "② 'AS-IS 콤보 원천 미확인' = GetDMInnerTable 콤보명은 BigFilter.xml 원문이 아닌 "
                "DataService.cs L3682 switch(Tables[n] 매핑)에 존재 — 추출 에이전트가 switch 실측으로 확인함. "
                "둘 다 데이터 결함 아님(도구 한계).", "", "", "", "", "", ""])
    for w in verify["warns"]:
        ws4.append(["WARN", w, "", "", "", "", "", ""])
    style_sheet(ws4, [9, 8, 8, 8, 24, 18, 50, 40])

    # 5.제외범위
    ws5 = wb.create_sheet("5.제외범위")
    title(ws5, "범위 밖 선언 — 그리드 셀 콤보·모달 내부 콤보 등 (조회조건 아님)", 2)
    ws5.append(["UI", "제외 항목"])
    for mid in UIS:
        exc = data[mid].get("excluded") or ["기재 없음"]
        ws5.append([mid, "\n".join(exc)])
    style_sheet(ws5, [9, 140])

    out = os.path.join(OUT_DIR, BASE + ".xlsx")
    wb.save(out)
    return out


# ── MD ───────────────────────────────────────────────────────────────────────
def build_md(data, verify, axis):
    L = []
    L.append(f"# DMF 13UI 조회조건 MATRIX (AS-IS ↔ TO-BE) v1.0")
    L.append("")
    L.append("> 생성 2026-08-19 · 원천 = AS-IS `.cs`(주석 제거 전처리본, hook #138) + `BigFilter.xml`/`UP_FX_GetCommonDataList` + TO-BE `*Form.tsx`/`SearchFields.tsx` + `SCM_ADM_CMBO_SQL_M` 실덤프(165건) · 대사 = hook #112 카탈로그")
    L.append("> 범위 = 조회조건(SearchArea)만. 그리드 셀·모달 내부 콤보 제외(각 UI '제외' 참조). 각 UI: **AS-IS 상단 / TO-BE 하단**.")
    L.append("")
    L.append("## 0. 요약 매트릭스 (유형만 — 상세는 UI 별 섹션)")
    L.append("")
    L.append("| UI | 구분 | " + " | ".join(axis) + " |")
    L.append("|---|---|" + "---|" * len(axis))
    for mid in UIS:
        for side, nm in (("asis", "AS-IS"), ("tobe", "TO-BE")):
            cells = []
            for cond in axis:
                c = find_cond(data, mid, cond)
                s = (c or {}).get(side)
                cells.append(s.get("control", "?") if s else ("－" if not c else "N/A"))
            L.append(f"| {mid} | {nm} | " + " | ".join(cells) + " |")
    L.append("")
    for mid in UIS:
        d = data[mid]
        L.append(f"## {mid} — {d.get('ui_name_asis', '')}")
        L.append("")
        for side, nm in (("asis", "AS-IS (상단)"), ("tobe", "TO-BE (하단)")):
            L.append(f"### {nm}")
            L.append("")
            L.append("| 조회조건 | 유형 | COMBO(원천) | 변수 | FILTER | CASCADE | 40 | 81 | 82 | 필수 | 기본값 | 근거 |")
            L.append("|---|---|---|---|---|---|---|---|---|---|---|---|")
            for c in d["conditions"]:
                s = c.get(side)
                if not s:
                    L.append(f"| {c['cond']} | N/A(없음) | | | | | | | | | | |")
                    continue

                def e(v):
                    return str(v or "").replace("|", "¦").replace("\n", "<br>")
                L.append("| " + " | ".join([
                    c["cond"], e(s.get("control")), e(s.get("combo")), e(s.get("params")),
                    e(s.get("filter")), e(s.get("cascade")), e(s.get("div40")), e(s.get("div81")),
                    e(s.get("div82")), e(s.get("required")), e(s.get("default")),
                    e(", ".join(s.get("evidence") or [])),
                ]) + " |")
            L.append("")
        chains = d.get("cascade_chains") or []
        if chains:
            L.append("**CASCADE 체인**: " + " / ".join(chains))
            L.append("")
        exc = d.get("excluded") or []
        if exc:
            L.append("**제외(범위 밖)**: " + " · ".join(exc))
            L.append("")
        notes = d.get("notes") or []
        if notes:
            L.append("**특기**: " + " · ".join(notes))
            L.append("")
    L.append("## 검증 대사 (hook #112 카탈로그)")
    L.append("")
    L.append("| UI | 조건 수 | AS-IS | TO-BE | 카탈로그 화면 | 카탈로그 행(As-Is/To-Be) | 사유 |")
    L.append("|---|---|---|---|---|---|---|")
    for mid in UIS:
        u = verify["ui"].get(mid, {})
        L.append(f"| {mid} | {u.get('conds')} | {u.get('asis')} | {u.get('tobe')} | "
                 f"{u.get('catalog_screen') or '미수록'} | {u.get('catalog_asis_rows')}/{u.get('catalog_tobe_rows')} | "
                 f"{u.get('reason') or '-'} |")
    L.append("")
    L.append(f"검증 게이트: FAIL {len(verify['fails'])}건 / WARN {len(verify['warns'])}건 (상세 verify_report.json)")
    L.append("")
    L.append("### 한계 명시")
    L.append("- To-Be 필수값은 원본 카탈로그(2026-01-22 v0.4) 전체 공란(미확정) — 본 매트릭스의 TO-BE 필수는 **현행 코드 실측** 기준.")
    L.append("- 카탈로그와 카운트가 다른 화면은 소스 실측 우선 원칙 적용(사유 병기).")
    out = os.path.join(OUT_DIR, BASE + ".md")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
    return out


# ── HTML ─────────────────────────────────────────────────────────────────────
def build_html(data, verify, axis, cmbo):
    cmbo_ids = {r["SQL_ID"] for r in cmbo}
    simple = {}
    for mid in UIS:
        simple[mid] = {"asis": {}, "tobe": {}}
        for c in data[mid]["conditions"]:
            for side in ("asis", "tobe"):
                simple[mid][side][c["cond"]] = simple_cell(c.get(side), side, cmbo_ids)
    payload = {"axis": axis, "uis": UIS,
               "data": {mid: data[mid] for mid in UIS},
               "simple": simple,
               "verify": verify["ui"],
               "gate": {"fails": len(verify["fails"]), "warns": len(verify["warns"])}}
    pj = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    doc = """<!DOCTYPE html>
<html lang="ko"><head><meta charset="utf-8">
<title>DMF 13UI 조회조건 MATRIX v1.0</title>
<style>
 body{margin:0;font-family:'Malgun Gothic',sans-serif;background:#0f1720;color:#dbe4ee;font-size:13px}
 header{position:sticky;top:0;z-index:30;background:#1F4E79;color:#fff;padding:10px 16px}
 header h1{margin:0;font-size:16px} header .sub{font-size:11px;opacity:.85;margin-top:3px}
 .bar{display:flex;gap:10px;align-items:center;padding:8px 16px;background:#16202b;position:sticky;top:52px;z-index:29;flex-wrap:wrap}
 .bar input{background:#0f1720;color:#dbe4ee;border:1px solid #33465c;border-radius:4px;padding:5px 8px;width:260px}
 .bar label{font-size:12px;cursor:pointer}
 .wrap{overflow:auto;max-height:calc(100vh - 110px)}
 table{border-collapse:collapse;min-width:100%}
 th,td{border:1px solid #26364a;padding:5px 7px;vertical-align:top;min-width:170px;max-width:300px}
 thead th{position:sticky;top:0;background:#1F4E79;color:#fff;z-index:20;font-size:12px}
 td.ui,th.ui{position:sticky;left:0;background:#16202b;z-index:10;min-width:70px;font-weight:bold}
 td.side,th.side{position:sticky;left:86px;background:#16202b;z-index:10;min-width:52px}
 tr.asis td.side{color:#f6c744} tr.tobe td.side{color:#7ad48a}
 tr.asis{background:#1a1f16} tr.tobe{background:#131f18}
 td .ctl{font-weight:bold;color:#9ec5ea} td .ln{display:block;margin-top:2px}
 td .div{color:#f0a35e} td.na{color:#4a5a6d;background:#141b23}
 td:hover{outline:2px solid #4a90d9;cursor:pointer}
 .hit{outline:2px solid #f6c744}
 #detail{position:fixed;right:0;top:0;height:100vh;width:420px;background:#16202b;border-left:2px solid #1F4E79;
   padding:14px;overflow:auto;display:none;z-index:40;white-space:pre-wrap;font-size:12px}
 #detail h2{font-size:14px;color:#9ec5ea;margin:0 0 8px}
 #detail .close{float:right;cursor:pointer;color:#f66}
 .gate{margin-left:auto;font-size:12px;color:#9ec5ea}
 .tabs{display:inline-flex;gap:4px}
 .tab{background:#0f1720;color:#9ec5ea;border:1px solid #33465c;border-radius:6px 6px 0 0;
   padding:6px 14px;cursor:pointer;font-size:13px}
 .tab.on{background:#1F4E79;color:#fff;border-color:#1F4E79;font-weight:bold}
 body.simple td{min-width:120px;max-width:220px}
 body.simple td .sv{color:#9ec5ea;font-weight:bold}
 body.simple td .sv.tag{color:#4a5a6d;font-weight:normal}
</style></head><body>
<header><h1>DMF 13UI 조회조건 MATRIX (AS-IS 상단 · TO-BE 하단)</h1>
<div class="sub">실측 2026-08-19 · 원천: AS-IS .cs(주석제거)+BigFilter/UP_FX · TO-BE Form.tsx+SCM_ADM_CMBO_SQL_M(165) · 대사: hook #112 · 범위: 조회조건(SearchArea)만</div></header>
<div class="bar">
 <span class="tabs">
  <button id="tabDetail" class="tab on" onclick="setTab('detail')">상세 MATRIX</button>
  <button id="tabSimple" class="tab" onclick="setTab('simple')">SIMPLE (콤보ID)</button>
 </span>
 <input id="q" placeholder="검색 (콤보ID/변수/필터/조건명…)" oninput="render()">
 <label><input type="checkbox" id="fAsis" checked onchange="render()"> AS-IS</label>
 <label><input type="checkbox" id="fTobe" checked onchange="render()"> TO-BE</label>
 <label><input type="checkbox" id="fDiff" onchange="render()"> division 분기 있는 셀만 강조</label>
 <span class="gate" id="gate"></span>
</div>
<div class="wrap"><table id="tbl"></table></div>
<div id="detail"></div>
<script>
const P = __PAYLOAD__;
let TAB='detail';
function setTab(t){TAB=t;document.body.classList.toggle('simple',t==='simple');
 document.getElementById('tabDetail').classList.toggle('on',t==='detail');
 document.getElementById('tabSimple').classList.toggle('on',t==='simple');render();}
function esc(s){return String(s==null?'':s).replace(/[&<>]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;'}[c]));}
function cellHtml(s){
 if(!s) return '<td class="na">N/A</td>';
 let h='<span class="ctl">'+esc(s.control)+'</span>';
 if(s.combo&&s.combo!=='없음') h+='<span class="ln">C: '+esc(s.combo)+'</span>';
 if(s.params&&s.params!=='없음') h+='<span class="ln">@ '+esc(s.params)+'</span>';
 const dv=!(s.div40==='공통'&&s.div81==='공통'&&s.div82==='공통');
 if(dv) h+='<span class="ln div">[40] '+esc(s.div40)+' / [81] '+esc(s.div81)+' / [82] '+esc(s.div82)+'</span>';
 return {html:h,div:dv};
}
function detail(mid,cond,side,s){
 const d=document.getElementById('detail');
 if(!s){d.style.display='none';return;}
 d.innerHTML='<span class="close" onclick="this.parentNode.style.display=\\'none\\'">✕ 닫기</span>'
  +'<h2>'+mid+' · '+esc(cond)+' · '+side.toUpperCase()+'</h2>'
  +'라벨: '+esc(s.label)+'\\n유형: '+esc(s.control)+'\\nCOMBO: '+esc(s.combo)+'\\n변수: '+esc(s.params)
  +'\\nFILTER: '+esc(s.filter)+'\\nCASCADE: '+esc(s.cascade)
  +'\\n[40] '+esc(s.div40)+'\\n[81] '+esc(s.div81)+'\\n[82] '+esc(s.div82)
  +'\\n필수: '+esc(s.required)+'\\n기본값: '+esc(s.default)
  +'\\n근거: '+esc((s.evidence||[]).join(', '));
 d.style.display='block';
}
function render(){
 const q=document.getElementById('q').value.toLowerCase();
 const showA=document.getElementById('fAsis').checked, showT=document.getElementById('fTobe').checked;
 const diffOnly=document.getElementById('fDiff').checked;
 let h='<thead><tr><th class="ui">UI</th><th class="side">구분</th>';
 for(const c of P.axis) h+='<th>'+esc(c)+'</th>';
 h+='</tr></thead><tbody>';
 for(const mid of P.uis){
  const conds=P.data[mid].conditions;
  for(const [side,cls] of [['asis','asis'],['tobe','tobe']]){
   if(side==='asis'&&!showA) continue; if(side==='tobe'&&!showT) continue;
   h+='<tr class="'+cls+'"><td class="ui">'+mid+'</td><td class="side">'+side.toUpperCase()+'</td>';
   for(const cond of P.axis){
    const c=conds.find(x=>x.cond===cond); const s=c?c[side]:null;
    if(!s){h+='<td class="na">'+(c?'N/A':'－')+'</td>';continue;}
    const txt=JSON.stringify(s).toLowerCase();
    const hit=q&&(txt.indexOf(q)>=0||cond.toLowerCase().indexOf(q)>=0);
    if(TAB==='simple'){
     const sv=(P.simple[mid][side]||{})[cond]||'';
     const isTag=['팝업','달력','체크','텍스트','라디오','—',''].indexOf(sv)>=0;
     const dv=!(s.div40==='공통'&&s.div81==='공통'&&s.div82==='공통');
     const cl=(hit?'hit ':'')+((diffOnly&&dv)?'hit ':'');
     h+='<td class="'+cl+(isTag?' na':'')+'"><span class="sv'+(isTag?' tag':'')+'">'+esc(sv||'—')+'</span></td>';
     continue;
    }
    const r=cellHtml(s);
    const cl=(hit?'hit ':'')+((diffOnly&&r.div)?'hit ':'');
    h+='<td class="'+cl+'">'+r.html+'</td>';
   }
   h+='</tr>';
  }
 }
 h+='</tbody>';
 document.getElementById('tbl').innerHTML=h;
 document.getElementById('gate').textContent='검증 게이트: FAIL '+P.gate.fails+' / WARN '+P.gate.warns;
 // 셀 클릭 → 상세 패널 (렌더 후 일괄 바인딩)
 const tds=document.querySelectorAll('#tbl td:not(.ui):not(.side)');
 const flat=[];
 for(const mid of P.uis){ for(const side of ['asis','tobe']){
   if(side==='asis'&&!showA) continue; if(side==='tobe'&&!showT) continue;
   for(const cond of P.axis){ const c=P.data[mid].conditions.find(x=>x.cond===cond);
     flat.push({mid,cond,side,s:c?c[side]:null}); } } }
 tds.forEach((td,ix)=>{ const f=flat[ix]; td.onclick=()=>detail(f.mid,f.cond,f.side,f.s); });
}
render();
</script></body></html>
"""
    doc = doc.replace("__PAYLOAD__", pj)
    out = os.path.join(OUT_DIR, BASE + ".html")
    with open(out, "w", encoding="utf-8") as f:
        f.write(doc)
    return out


def main():
    data, cmbo, verify = load()
    axis = cond_axis(data)
    print(f"가로축 조건 {len(axis)}: {axis}")
    x = build_xlsx(data, cmbo, verify, axis)
    m = build_md(data, verify, axis)
    h = build_html(data, verify, axis, cmbo)
    for p in (x, m, h):
        print("생성:", p, f"({os.path.getsize(p):,} bytes)")


if __name__ == "__main__":
    main()
