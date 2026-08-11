#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""tobe_scope_table_inventory.py — 공통 + 개발된 UI 한정 TABLE 전수 취합 (2026-08-10)

사용자 지시: *"현재 전체 생성본 중에 다른 모듈에서 사용할 TABLE 이 많습니다.
             공통과 현재 개발된 UI 에만 한정해서 TABLE 목록을 뽑고 싶습니다."*

## 왜 새로 뽑는가 — v2.2 누락의 근본원인 (2026-08-10 실측 규명)

v2.2 취합식 = `closure( API_M ROOT ∩ 취합대상 UI 43개 )` 단일 기점.
이 식에 안 걸리면 무조건 누락되며, 3층에서 구멍이 났다:

| 층 | 사유 | 실측 |
|---|---|---|
| L1 | UI 가 취합 대상 목록에서 빠짐 | 배선 52 UI 중 **9개 제외**(COM0007/0010/0011, DMF0035, DMF0043_01, DMF0068, DMF0074_40, SCP0013, SYS0017) |
| L2 | SP 는 있는데 **API_M 미등재**(프레임워크 공통) | DB 프로그램 1,125 중 **1,042개가 API_M 밖** — 즐겨찾기(`SCM_ADM_USER_BMK_M`), 검색조건, 업로드, 공지, VOC |
| L3 | **참조 SP 자체가 없음**(BE 범용 코드 직접 접근) | `SCM_ADM_CMBO_SQL_M`(콤보 SQL 156본), 로그 2종 |

⇒ 본 스크립트는 기점을 **3축 합집합**으로 바꾼다:
  · 축A **개발된 UI 전체** = `SCM_ADM_MNU_BTN_MPG_M` 배선 보유 메뉴 (43 → 52, L1 해소)
  · 축B **공통 프로그램** = BE MyBatis XML/Java 가 호출하는 SP + 프레임워크 명명 계열 (L2 해소)
  · 축C **BE 직접 접근 테이블** = BE 소스에서 SQL 로 직접 읽는 테이블 (L3 해소)

각 축에서 `ALL_DEPENDENCIES` 재귀 폐포를 돌려 TABLE 을 모으고,
**다른 모듈 전용 테이블은 자동으로 빠진다**(기점에서 도달 불가하므로).

산출: `Doc/화면분석서/TOBE_공통UI_TABLE_전수_{YYYYMMDD}.xlsx`
"""
import pathlib
import re
import sys
import datetime

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from ora import conn_info  # noqa: E402
import oracledb  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = pathlib.Path(__file__).resolve().parents[2]
BE = ROOT / "nscm-be" / "src" / "main"
OUT = ROOT / "Doc" / "화면분석서"

u, p, d = conn_info()
con = oracledb.connect(user=u, password=p, dsn=d)
cur = con.cursor()

# ── 사전: DB 실재 객체 ──────────────────────────────────────────────
cur.execute("""SELECT OBJECT_NAME, OBJECT_TYPE FROM ALL_OBJECTS
                WHERE OWNER='SCM_MGR' AND OBJECT_TYPE IN
                      ('PACKAGE','PROCEDURE','FUNCTION','TABLE','VIEW')""")
db_obj = {}
for n, t in cur.fetchall():
    db_obj.setdefault(n.strip().upper(), t)
PROGS = {n for n, t in db_obj.items() if t in ("PACKAGE", "PROCEDURE", "FUNCTION")}
TABLES = {n for n, t in db_obj.items() if t == "TABLE"}
VIEWS = {n for n, t in db_obj.items() if t == "VIEW"}
print("DB: 프로그램 %d / TABLE %d / VIEW %d" % (len(PROGS), len(TABLES), len(VIEWS)))

# ── 축A: 개발된 UI (배선 보유) → API_M ROOT ─────────────────────────
cur.execute("""SELECT DISTINCT B.MNU_ID, A.API_URL
                 FROM SCM_MGR.SCM_ADM_MNU_BTN_MPG_M B
                 JOIN SCM_MGR.SCM_ADM_API_M A ON A.API_ID = B.API_ID""")
ui_root = {}
for mnu, url in cur.fetchall():
    s = (url or "").strip().upper()
    if not s:
        continue
    r = s.split("(")[0].split(".")[0]
    if r in PROGS:
        ui_root.setdefault(r, set()).add((mnu or "").strip().upper())
UI_MENUS = sorted({m for v in ui_root.values() for m in v})
print("축A 개발된 UI %d개 / ROOT 프로그램 %d개" % (len(UI_MENUS), len(ui_root)))

# ── 축B/C: BE 소스에서 SP·TABLE 직접 참조 추출 ──────────────────────
be_files = [q for q in BE.rglob("*") if q.suffix.lower() in (".xml", ".java") and q.is_file()]
be_text = []
for q in be_files:
    try:
        be_text.append((q, q.read_text(encoding="utf-8", errors="replace")))
    except Exception:  # noqa: BLE001
        pass
print("BE 소스 %d파일 스캔" % len(be_text))

ident = re.compile(r"\b([A-Z][A-Z0-9_]{3,})\b")
be_progs, be_tables = {}, {}
for q, txt in be_text:
    rel = str(q.relative_to(ROOT))
    for m in ident.finditer(txt.upper()):
        nm = m.group(1)
        if nm in PROGS:
            be_progs.setdefault(nm, set()).add(rel)
        elif nm in TABLES or nm in VIEWS:
            be_tables.setdefault(nm, set()).add(rel)
print("축B BE 호출 프로그램 %d개 / 축C BE 직접참조 TABLE·VIEW %d개" % (len(be_progs), len(be_tables)))

# ── 프레임워크 공통 명명 계열 보강 (BE 가 동적 호출해 문자열로 안 잡히는 것) ──
FW_PAT = re.compile(r"^(PKG_NSCM_UI_|UP_FX_|PR_GET_MENU|PR_SCM_ADM_|PKG_SCM0000$|FN_GET_VOC_)")
fw_progs = {n for n in PROGS if FW_PAT.search(n)}
print("축B' 프레임워크 명명 계열 %d개" % len(fw_progs))

# ── 폐포: 프로그램 → 참조 TABLE/VIEW (재귀) ─────────────────────────
cur.execute("""SELECT NAME, TYPE, REFERENCED_NAME, REFERENCED_TYPE
                 FROM ALL_DEPENDENCIES
                WHERE OWNER='SCM_MGR' AND REFERENCED_OWNER IN ('SCM_MGR','FP_MGR')""")
dep = {}
for n, t, rn, rt in cur.fetchall():
    dep.setdefault(n.strip().upper(), set()).add((rn.strip().upper(), rt))


def closure(seeds):
    seen, tabs, stack = set(), {}, list(seeds)
    while stack:
        cur_o = stack.pop()
        if cur_o in seen:
            continue
        seen.add(cur_o)
        for rn, rt in dep.get(cur_o, ()):
            if rt in ("TABLE", "VIEW"):
                tabs.setdefault(rn, set()).add(cur_o)
            if rt in ("PACKAGE", "PROCEDURE", "FUNCTION", "PACKAGE BODY") and rn not in seen:
                stack.append(rn)
    return tabs, seen


axA_tabs, axA_seen = closure(ui_root.keys())
axB_tabs, axB_seen = closure(set(be_progs) | fw_progs)
print("\n폐포 결과: 축A TABLE %d / 축B TABLE %d" % (len(axA_tabs), len(axB_tabs)))

# ── 합집합 + 출처 표기 ──────────────────────────────────────────────
allt = {}
for t, via in axA_tabs.items():
    allt.setdefault(t, {"axis": set(), "via": set()})["axis"].add("A:개발UI")
    allt[t]["via"] |= via
for t, via in axB_tabs.items():
    allt.setdefault(t, {"axis": set(), "via": set()})["axis"].add("B:공통SP")
    allt[t]["via"] |= via
for t, files in be_tables.items():
    allt.setdefault(t, {"axis": set(), "via": set()})["axis"].add("C:BE직접")
    allt[t]["via"] |= {pathlib.Path(f).name for f in files}

BAK = re.compile(r"(_BAK|_BK\d|_BACKUP|_\d{8}$|BAK_)", re.I)
rows = []
for t in sorted(allt):
    if t not in TABLES and t not in VIEWS:
        continue
    kind = "VIEW" if t in VIEWS else "TABLE"
    is_gtt = ""
    cls = "백업/스냅샷" if BAK.search(t) else "대상"
    rows.append((t, kind, cls, ";".join(sorted(allt[t]["axis"])),
                 ";".join(sorted(list(allt[t]["via"]))[:6])))

cur.execute("""SELECT TABLE_NAME FROM ALL_TABLES WHERE OWNER='SCM_MGR' AND TEMPORARY='Y'""")
gtts = {r[0].strip().upper() for r in cur.fetchall()}
print("\n== 결과 요약 ==")
tgt = [r for r in rows if r[2] == "대상"]
print("  합집합 TABLE/VIEW %d개 (대상 %d / 백업 %d)" % (len(rows), len(tgt), len(rows) - len(tgt)))
from collections import Counter
print("  축 조합 분포:", dict(Counter(r[3] for r in tgt)))
print("  그중 GTT:", len([r for r in tgt if r[0] in gtts]))

# v2.2 대비 추가분
try:
    from openpyxl import load_workbook
    wb = load_workbook(OUT / "TOBE_UI_DB객체_전수취합_v2.2.xlsx", read_only=True, data_only=True)
    prev = set()
    for r in list(wb["1_UI별_DB객체"].iter_rows(values_only=True))[1:]:
        if r and len(r) > 6 and r[6]:
            prev.add(str(r[6]).strip().upper())
    for r in list(wb["3_TABLE_이관SOURCE"].iter_rows(values_only=True))[1:]:
        if r and len(r) > 1 and r[1]:
            prev.add(str(r[1]).strip().upper())
    wb.close()
    added = [r for r in tgt if r[0] not in prev]
    print("\n== v2.2 대비 신규 포착 %d개 ==" % len(added))
    for r in added:
        print("  + %-40s %-6s 축=%-22s via=%s" % (r[0], r[1], r[3], r[4][:60]))
except Exception as e:  # noqa: BLE001
    print("v2.2 대조 실패:", str(e)[:100])

# ── xlsx 산출 ──────────────────────────────────────────────────────
try:
    from openpyxl import Workbook
    wb2 = Workbook()
    ws = wb2.active
    ws.title = "공통UI_TABLE"
    ws.append(["TABLE/VIEW", "구분", "분류", "포착축", "근거(경유 객체·파일)", "GTT"])
    for r in rows:
        ws.append([r[0], r[1], r[2], r[3], r[4], "Y" if r[0] in gtts else ""])
    ws2 = wb2.create_sheet("개발UI목록")
    ws2.append(["MENU_ID"])
    for m in UI_MENUS:
        ws2.append([m])
    stamp = datetime.datetime.now().strftime("%Y%m%d")
    dst = OUT / ("TOBE_공통UI_TABLE_전수_%s.xlsx" % stamp)
    wb2.save(dst)
    print("\n[산출] %s" % dst)
except Exception as e:  # noqa: BLE001
    print("xlsx 저장 실패:", str(e)[:120])
con.close()
