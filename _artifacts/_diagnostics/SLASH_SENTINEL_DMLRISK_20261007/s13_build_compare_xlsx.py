"""[xlsx 생성] '/' 적용 시 의심 대상 — 변환 전/후 비교 보고 (사용자 지시 10-07: 전체가 아니라 의심 대상만 · 실효 없는 변경은 하지 않음)
입력: _risk.json · _partner_scan.json · _dupcheck.json(s11) · _scan_counts.json(s1/s12) · _scan_stats.json · _pop_extra.json
출력: 슬래시전환_변환전후비교_20261007.xlsx
"""
import sys, json, re, pathlib, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE)); import s1_data_scan as s
R = json.loads((HERE / "_risk.json").read_text(encoding="utf-8"))
PS = json.loads((HERE / "_partner_scan.json").read_text(encoding="utf-8"))
DUP = json.loads((HERE / "_dupcheck.json").read_text(encoding="utf-8")) if (HERE / "_dupcheck.json").exists() else {}
CN = json.loads(s.COUNTS.read_text(encoding="utf-8"))
ST = {(r["table"], r["column"]): r for r in json.loads((HERE / "_scan_stats.json").read_text(encoding="utf-8"))["rows"]}
EX = json.loads((HERE / "_pop_extra.json").read_text(encoding="utf-8"))["rows"]
measured = {(k.split(".", 1)[1], c): v for k, e in CN.items() for c, v in e.get("cols", {}).items()}
pop = s.population()
CLS = {(c["table"], c["column"]): c for c in pop}
for e in EX:
    if e.get("kind") != "new_table": CLS.setdefault((e["table"], e["column"]), dict(table=e["table"], column=e["column"], cls="추가분(정제전 탈락) — 일괄 변환", held=False, use="(미분류)"))
# AS-IS 실측(작업대상 xlsx 1_컬럼판정) — s4 와 같은 방식으로 읽는다
import shutil
tmp = HERE / "_src_copy.xlsx"; shutil.copy2(s.SRC, tmp)
r1 = list(openpyxl.load_workbook(tmp, read_only=True, data_only=True)["1_컬럼판정"].iter_rows(values_only=True)); h1 = [str(c) for c in r1[0]]
def num(v):
    try: return int(str(v).replace(",", ""))
    except Exception: return None
ASIS = {}
for r in r1[1:]:
    if r and r[1]: ASIS[(r[1], r[2])] = dict(dash=num(r[h1.index("AS-IS '-' 행수(10-01)")]), blank=num(r[h1.index("AS-IS 빈값 행수(10-01, ''+' ' 합산)")]), space=num(r[h1.index("AS-IS 공백만(' ') 행수(10-07)")]))
for t, e in PS.items():
    for c, a in e.get("asis", {}).items():
        if a and "missing" not in a and (t, c) not in ASIS: ASIS[(t, c)] = dict(dash=a["dash"], blank=a["null"] + a["empty"] + a["space"], space=a["space"])

wb = openpyxl.Workbook()
def sheet(name, head, rows, widths, freeze="B2"):
    w = wb.create_sheet(name); w.append(head)
    for c in w[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496"); c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows: w.append([("" if v is None else (json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v)) for v in r])
    for i, wd in enumerate(widths, 1): w.column_dimensions[get_column_letter(i)].width = wd
    w.freeze_panes = freeze
    if rows: w.auto_filter.ref = w.dimensions

# ── 1. 코드 '/' 변환 ↔ 데이터 '-' 유지 (87) — 전/후 ──
S1 = []
for h in R["H_code_to_slash_data_stays"]:
    m = re.search(r"([A-Z_][\w$#]*)\.([A-Z_][\w$#]*) \[", h["cols"]); t, c = (m.group(1), m.group(2)) if m else ("", "")
    tb = (PS.get(t) or {}).get("tobe", {}).get(c) or {}; a = (PS.get(t) or {}).get("asis", {}).get(c) or {}
    before = tb.get("dash"); after = tb.get("slash", 0) if tb else None
    neg = bool(re.search(r"(<>|!=)\s*'-'", h["text"]))
    S1.append([h["obj"], h["sub"], h["line"], h["use"], f"{t}.{c}", h["cols"].split("[")[-1].rstrip("]"), a.get("dash") if "missing" not in a else "AS-IS 없음", (a.get("null", 0) + a.get("empty", 0) + a.get("space", 0)) if a and "missing" not in a else None,
               h["text"], before, after,
               ("부등(<>) 비교 — '/' 로 바꾸면 현재 제외되던 '-' 행이 전부 포함됨" if neg else "등치 비교 — '/' 로 바꾸면 현재 잡히던 행이 0 이 됨") if before else "현재 '-' 행 0 — 바꿔도 결과 같음(그래도 변경 불필요)",
               "코드 '-' 유지 (변경 안 함)" if before else "변경 불필요"])
sheet("1_코드변환_데이터유지(87)", ["객체", "서브프로그램", "줄", "쓰임", "비교 컬럼(변환 범위 밖)", "컬럼 실태", "AS-IS '-' 행", "AS-IS 빈값 행", "원문", "현재: 컬럼 '-' 행(코드가 잡는 상한)", "코드를 '/' 로 바꾼 뒤: '/' 행", "영향", "결론"],
      S1, [34, 30, 8, 16, 40, 30, 12, 12, 90, 16, 14, 50, 22])

# ── 2. 유니크 키 중복 시뮬레이션 ──
S2 = []
for b in R["B_unique"]:
    if not b["need_check"]: continue
    d = DUP.get(f'{b["table"]}|{b["name"]}') or {}
    if b["table"] in s.SKIP: verdict = "1억 초과 — 실측 제외(사용자 지시)"
    elif "error" in d: verdict = "실측 실패: " + d["error"][:80]
    elif not d: verdict = "미측정"
    elif d["dup_groups"] == 0: verdict = "이상 없음 — '-'·공백 전부를 '/' 로 봐도 키 중복 0"
    else: verdict = f"중복 가능 — {d['dup_groups']:,}그룹 / {d['dup_rows']:,}행 (상한: 실제 변환은 AS-IS 빈값 유래 행만) → 해당 키 행 단위 확인 필요"
    S2.append([b["table"], b["name"], b["type"], ", ".join(b["cols"]), ", ".join(b["conv_cols"]), ", ".join(b["blank_cols"]), ", ".join(b["both_cols"]), d.get("rows") or b.get("rows"), d.get("affected_rows_upper"), d.get("dup_groups"), d.get("dup_rows"), d.get("sec"), verdict])
sheet("2_유니크키_중복시뮬", ["테이블", "키", "종류", "키 컬럼", "변환 대상 키 컬럼", "공백만 값 보유 컬럼", "'-'·'/' 공존 컬럼", "행수", "영향 행 상한('-' 또는 공백)", "변환 후 중복 그룹", "중복 행", "측정(s)", "판정"],
      S2, [30, 34, 10, 50, 36, 24, 16, 12, 16, 14, 10, 8, 70])

# ── 3. DML 실효성 분류 (변환 대상 1,854) ──
S3 = []; cat = collections.Counter()
for (t, c), p in sorted(CLS.items()):
    if p.get("held") or t == "SCM_ADM_LOOKUP_MASTER_DETAIL": continue
    a = ASIS.get((t, c)) or {}; m = measured.get((t, c)); st = ST.get((t, c)) or {}
    if m is not None: tobe_dash, tobe_blank, src = m["dash"], m["blank"], "실측"
    elif t in s.SKIP: tobe_dash, tobe_blank, src = None, None, "1억 초과 — 실측 제외"
    else: tobe_dash, tobe_blank, src = {"none": 0}.get(st.get("flag")), None, "통계(" + str(st.get("flag")) + ", " + str(st.get("analyzed")) + ")"
    cls = str(p["cls"]); coexist = cls.startswith("공존") or (a.get("dash") or 0) > 0
    if tobe_dash == 0 and not (tobe_blank or 0): verdict = "DML 생략 — TO-BE 에 '-'·공백 행 없음(실효 없음)"
    elif tobe_dash is None: verdict = "미확인 — DML 은 WHERE 조건부라 실행해도 무해(0행), 통계 갱신 후 재판정 권장"
    elif coexist: verdict = "행 단위 — AS-IS 실제 '-' 와 빈값 유래 '-' 혼재 → 일괄 금지, AS-IS 대조 키 집합으로만"
    else: verdict = "일괄 DML 유효 — '-' 전부가 빈값 유래(AS-IS 실제 '-' 0)"
    gen = None
    if tobe_dash is not None and a.get("blank") is not None and not coexist:
        gen = tobe_dash - (a.get("blank") or 0)
    cat[verdict.split(" — ")[0]] += 1
    S3.append([t, c, cls, a.get("dash"), a.get("blank"), a.get("space"), tobe_dash, tobe_blank, src, gen if (gen or 0) > 0 else None, verdict])
sheet("3_DML실효성(변환대상)", ["테이블", "컬럼", "분류", "AS-IS '-' 행", "AS-IS 빈값 행", "AS-IS 공백만 행", "TO-BE '-' 행", "TO-BE 공백만 행", "TO-BE 근거", "TO-BE 가 AS-IS 빈값보다 많은 행(TO-BE 생성분 추정)", "판정"],
      S3, [30, 20, 34, 12, 12, 12, 12, 12, 30, 18, 70])

# ── 4. 코드 '-' 유지 ↔ 데이터 변환 확정 11 ──
S4 = []
for g in R["G_code_keeps_dash"]:
    if g["level"] != "확정": continue
    m = re.search(r"([A-Z_][\w$#]*)\.([A-Z_][\w$#]*)\(", g["cols"]); t, c = (m.group(1), m.group(2)) if m else ("", "")
    a = ASIS.get((t, c)) or {}
    real = (a.get("dash") or 0) > 0
    S4.append([g["obj"], g["line"], g["action"], g["use"], f"{t}.{c}", a.get("dash"), a.get("blank"), g["text"],
               "AS-IS 에 실제 '-' 가 있는 컬럼 — 코드의 '-' 는 실제 값 적재/분기이고 데이터 변환은 빈값 유래 행만 → 충돌 없음" if real else "AS-IS 실제 '-' 0 — 코드가 새 '-' 를 적재하면 변환 후 혼재 → 코드 '/' 로 변경 검토",
               "코드 유지(변경 안 함)" if real else "코드 '/' 변경 검토"])
sheet("4_코드유지_데이터변환(확정11)", ["객체", "줄", "작업대상 xlsx 액션", "쓰임", "관련 변환 컬럼", "AS-IS '-' 행", "AS-IS 빈값 행", "원문", "판단", "결론"], S4, [36, 8, 16, 16, 40, 12, 12, 90, 70, 22])

# ── 5. 코드 마스터 ──
LK = [[c, (ASIS.get(("SCM_ADM_LOOKUP_MASTER_DETAIL", c)) or {}).get("dash"), (ASIS.get(("SCM_ADM_LOOKUP_MASTER_DETAIL", c)) or {}).get("blank"), len([l for l in R["L_lookup_join"] if l["lookup_col"] == c])] for c in ("SEGMENT1", "SEGMENT2", "SEGMENT3", "SEGMENT4", "SEGMENT5")]
sheet("5_코드마스터", ["컬럼", "AS-IS 실제 '-' 행", "AS-IS 빈값 유래 행(TO-BE '-' 로 유입)", "NOT NULL 변환 컬럼과 조인하는 코드 줄"],
      LK + [[], ["결론", "코드 NVL(x,'/') 198줄 = 변경 불필요 (컬럼이 NOT NULL + DEFAULT 라 NULL 이 없어 NVL 은 무동작)"], ["", "회귀 가능 = 빈값 유래 행만 (합계 " + str(sum((x[2] or 0) for x in LK)) + "행) — 이 행만 '/' 로 정정할지 결정 대기(공존 컬럼 행 단위 규칙과 동일). 전체 일괄 변환은 하지 않음"], ["", "앞으로 들어오는 빈값은 DEFAULT '/' 가 처리"]],
      [14, 16, 30, 30])

# ── 6. 보류 51 (다른 세션 작업대상 xlsx 1b 시트 그대로) + 0. 결정 목록 통합 ──
wb_src = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
r1b = list(wb_src["1b_보류51컬럼_검토"].iter_rows(values_only=True))
h1b = [str(c) for c in r1b[0]]; rows1b = [list(r) for r in r1b[1:] if r and r[0]]
sheet("6_보류51(타세션 1b)", h1b, rows1b, [30, 20, 40, 12, 12, 10, 10, 10, 30, 10, 10, 40, 60, 60, 60, 60, 12, 12, 12, 18, 30])
held_names = {r[1] for r in rows1b}
D = []
for r in rows1b:
    D.append(["보류 51", f"{r[0]}.{r[1]}", r[2], f"AS-IS '-' {r[4]} / 전체 {r[3]} ({r[5]}%) · '' {r[6]} · ' ' {r[7]}", "유지 '-' / 슬래시", r[19] or "", "AS-IS 코드가 이 컬럼의 '-' 를 비교·분기 — 빈값 없이 '-' 만 있어 '/' 로 바꾸면 그 분기가 깨짐. 데이터 변환 범위에서 제외할지 결정"])
for e in EX:
    if e.get("kind") != "new_table" and e["column"] in held_names:
        D.append(["추가분(정제전 탈락) — 동명 보류 컬럼", f'{e["table"]}.{e["column"]}', "① 1차 목록에 없던 컬럼 · 같은 이름이 보류 51 에 있음", "AS-IS 실측 미실시", "보류 51 과 동일 결정 권고", "", "DEFAULT '/' 는 적용, 데이터 변환은 보류 51 결정을 따름"])
for x in LK:
    if x[2]: D.append(["코드 마스터 빈값 유래 행", f"SCM_ADM_LOOKUP_MASTER_DETAIL.{x[0]}", "코드 마스터 — 일괄 변환 제외", f"AS-IS 실제 '-' {x[1]} · 빈값 유래 {x[2]}행", "빈값 유래 행만 '/' 정정 / 그대로 둠", "", "AS-IS 에서 ''='' 로 맞던 조인이 '/'≠'-' 로 끊기는 행 = 빈값 유래 행뿐. 정정하면 해소, 두면 그 행만 회귀"])
for r in S2:
    if r[9]: D.append(["유니크 키 중복 가능", f"{r[0]} ({r[1]})", "변환 시뮬레이션에서 키 중복", f"중복 {r[9]:,}그룹 / {r[10]:,}행 (상한)", "행 단위 확인 후 변환", "", r[12]])
D.append(["신규 테이블 의심 복사본", "(DEFAULT범위_확대후보 xlsx 시트 3 참조)", "백업·테스트·CDC 복사본 의심 테이블", "3,454 컬럼 / 555 테이블", "DEFAULT 제외(권고) / 포함", "", "이름 패턴 기준 보류 — 운영 테이블이면 확대"])
sheet("0_결정목록(통합)", ["구분", "대상", "근거 유형", "수치", "선택지", "결정(기입)", "설명"], D, [30, 50, 50, 44, 30, 14, 90], freeze="C2")

# ── 요약 ──
ws = wb.active; ws.title = "요약"
S = [["'/' 적용 시 의심 대상 — 변환 전/후 비교 (최소 변경 원칙)"], [], ["작성", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")],
     ["원칙(사용자 10-07)", "전체가 아니라 의심 대상만 실측 · 실효 없는 변경(코드·DML)은 하지 않음 · 변환 행 = AS-IS 빈값(NULL·''·' ') 유래만 · TO-BE NULL 허용 컬럼 제외 · 코드 마스터 제외 · 신규 테이블은 DEFAULT 만"],
     ["DB 변경", "없음 — 전부 조회"], [],
     ["의심 묶음", "건수", "결과"],
     ["1 코드 '/' 변환 ↔ 데이터 '-' 유지", len(S1), f"현재 '-' 행 있음 {sum(1 for r in S1 if r[9])} / 0행 {sum(1 for r in S1 if r[9] == 0)} / 미실측 {sum(1 for r in S1 if r[9] is None)} → 전부 코드 '-' 유지(변경 안 함). 작업대상 xlsx 의 '비교→/' 액션을 이 줄들에서 되돌려야 함"],
     ["2 유니크 키 중복 시뮬레이션", len(S2), f"중복 0 = {sum(1 for r in S2 if r[9] == 0)} · 중복 있음 = {sum(1 for r in S2 if r[9])} · 미측정/제외 = {sum(1 for r in S2 if r[9] is None)}"],
     ["3 DML 실효성(변환 대상 컬럼)", len(S3), json.dumps(dict(cat), ensure_ascii=False)],
     ["4 코드 '-' 유지 ↔ 데이터 변환(확정)", len(S4), f"코드 유지 {sum(1 for r in S4 if r[9].startswith('코드 유지'))} · 변경 검토 {sum(1 for r in S4 if not r[9].startswith('코드 유지'))}"],
     ["5 코드 마스터", 5, "코드 NVL 198줄 변경 불필요 · 빈값 유래 행 정정 여부 결정 대기"],
     ["6 보류 51 (타 세션 1b 시트)", len(rows1b), "AS-IS 코드가 '-' 를 비교·분기하는 컬럼(빈값 없이 '-' 만) — 결정 대기"],
     ["0 결정 목록(통합)", len(D), f"보류 51 {sum(1 for d in D if d[0] == '보류 51')} · 추가분 동명 보류 {sum(1 for d in D if d[0].startswith('추가분'))} · 코드 마스터 {sum(1 for d in D if d[0].startswith('코드 마스터'))} · 유니크 키 {sum(1 for d in D if d[0].startswith('유니크'))} · 신규 테이블 의심 1"], [],
     ["읽는 법"],
     ["  시트 3 판정 'DML 생략' = TO-BE 에 바꿀 행이 없는 컬럼 → UPDATE 문 자체를 만들지 않는다(최소 변경). '일괄 DML 유효' = WHERE col='-' 일괄 가능. '행 단위' = AS-IS 대조 키 집합으로만."],
     ["  시트 3 'TO-BE 생성분 추정' = TO-BE '-' 행이 AS-IS 빈값 행보다 많은 만큼 — 이관 후 TO-BE 코드(DEFAULT '-'·NVL(x,'-'))가 만든 빈값 자리표시. 일괄 DML 이면 같이 '/' 가 되고(바람직), AS-IS 키 대조 방식이면 남는다 → '일괄 DML 유효' 컬럼은 일괄로 가는 것이 맞음"],
     ["  시트 1·4 는 코드 쪽 결론 — 전부 '변경 안 함' 이면 작업대상 xlsx 3_TOBE_변경목록의 해당 줄 액션을 '유지' 로 바꾼다"], [],
     ["한계"],
     ["  시트 1 의 '현재 행' 은 컬럼 전체의 '-' 행수(술어의 다른 조건 미반영) — 결론(코드 유지)은 수치와 무관하게 성립"],
     ["  시트 2 는 '-'·공백 전부를 '/' 로 본 상한. 중복 0 이면 안전, 0 보다 크면 실제 변환 행(AS-IS 빈값 유래)으로 다시 봐야 함"],
     ["  시트 3 의 TO-BE 근거가 '통계' 인 컬럼은 LAST_ANALYZED 시점 값 — 실측 아님. 1억 초과 2 테이블·3천만 초과 테이블은 실측 생략"],
     ["  AS-IS 빈값 행수는 DEV(10-01) 기준이며 NULL 미포함(조인 상대 실측분은 NULL 포함)"]]
for r in S: ws.append(r)
ws["A1"].font = Font(bold=True, size=13); ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 10; ws.column_dimensions["C"].width = 130
wb.move_sheet("0_결정목록(통합)", offset=-(len(wb.sheetnames) - 2))
out = HERE / "슬래시전환_변환전후비교_20261007.xlsx"; wb.save(out); print(out); print(dict(cat)); print("결정목록", len(D), "보류51", len(rows1b))
