"""[조회 전용 + 파일 생성] AS-IS 에 없는 TO-BE 신규 테이블의 NOT NULL 문자 컬럼 → DEFAULT '/' 대상 추가 (사용자 지시 2026-10-07)
규칙: SCM_MGR 영구 테이블(GTT·뷰·BIN$ 제외) 중 AS-IS(접속 가능한 모든 DB 의 dbo) 에 같은 이름이 없는 테이블의
      NOT NULL VARCHAR2/CHAR/NVARCHAR2/NCHAR 컬럼. 표준 감사 컬럼 4종 제외(1차 목록 기준과 동일).
      기존 DEFAULT 가 있고 '-' 가 아니면 유지, 없거나 '-' 이면 '/'.
출력: _pop_extra.json 에 'new_table' 행 추가(기존 115 유지) · DEFAULT범위_확대후보_20261007.xlsx 시트 3_신규테이블_NOTNULL
"""
import sys, json, re, pathlib, importlib.util, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
HERE = pathlib.Path(__file__).resolve().parent
RAW = pathlib.Path(r"C:\Projects\NSCM\Doc\_analysis\asis_delta_20261002\_raw")
sys.path.insert(0, str(RAW)); sys.path.insert(0, str(HERE))
spec = importlib.util.spec_from_file_location("snap", RAW / "01_snapshot.py"); snap = importlib.util.module_from_spec(spec); spec.loader.exec_module(snap)
from common3 import env
import s1_data_scan as s

AUDIT = {"CREATED_BY", "CREATION_DATE", "LAST_UPDATED_BY", "LAST_UPDATE_DATE"}
# 백업·테스트·임시 복사본 — DEFAULT 대상에서 제외하고 시트에만 남긴다 (기준 파일 '제외테이블' 92건과 같은 성격)
#   __ct = CDC(복제) 변경 추적 테이블 · attrep_* = 복제 도구 내부 · _KHJ 등 개인 복사본
SUSPECT = re.compile(r"(BAK|_BK\d*|BACKUP|_KEEP$|^ZZ|^ZT|TEMP|TMP|^DROP|TEST|^TB_|^BUF_|_OLD|COPY|\d{4,}|_T\d{2}(_|$)|_MF$|__ct$|^attrep|_KHJ$)", re.I)

# AS-IS 테이블 전수 (접속 가능한 DB 전부)
a = snap.asis_conn(env()).cursor()
a.execute("SELECT name FROM sys.databases WHERE database_id > 4")
dbs = [r[0] for r in a.fetchall()]
asis_tabs = {}
for db in dbs:
    try:
        a.execute(f"SELECT TABLE_NAME FROM [{db}].INFORMATION_SCHEMA.TABLES")
        for (t,) in a.fetchall(): asis_tabs.setdefault(t.upper(), set()).add(db)
    except Exception as e:
        print("skip", db, str(e)[:80])
a.execute("SELECT CONVERT(varchar(19), GETDATE(), 120)"); asis_at = a.fetchone()[0]
print("AS-IS DB", dbs, "테이블", len(asis_tabs), asis_at)

# 기존 목록
pop = {(c["table"], c["column"]) for c in s.population()}
EXF = HERE / "_pop_extra.json"; ex = json.loads(EXF.read_text(encoding="utf-8"))
ex["rows"] = [r for r in ex["rows"] if r.get("kind") != "new_table"]
have = pop | {(r["table"], r["column"]) for r in ex["rows"]}

cur = s.connect().cursor()
cur.execute("SELECT TO_CHAR(SYSDATE,'YYYY-MM-DD HH24:MI:SS') FROM dual"); tobe_at = cur.fetchone()[0]
cur.execute("""SELECT C.TABLE_NAME, C.COLUMN_NAME, C.COLUMN_ID, C.DATA_TYPE, C.DATA_LENGTH, C.NULLABLE, C.DEFAULT_ON_NULL, C.DATA_DEFAULT, T.NUM_ROWS,
                      TO_CHAR(O.CREATED,'YYYY-MM-DD'), TO_CHAR(O.LAST_DDL_TIME,'YYYY-MM-DD')
                 FROM ALL_TAB_COLUMNS C JOIN ALL_TABLES T ON T.OWNER=C.OWNER AND T.TABLE_NAME=C.TABLE_NAME
                 JOIN ALL_OBJECTS O ON O.OWNER=C.OWNER AND O.OBJECT_NAME=C.TABLE_NAME AND O.OBJECT_TYPE='TABLE'
                WHERE C.OWNER='SCM_MGR' AND T.TEMPORARY='N' AND C.TABLE_NAME NOT LIKE 'BIN$%'
                  AND C.NULLABLE='N' AND C.DATA_TYPE IN ('VARCHAR2','CHAR','NVARCHAR2','NCHAR')
                ORDER BY C.TABLE_NAME, C.COLUMN_ID""")
rows = []; stat = collections.Counter()
for t, c, cid, dt, dl, nu, on, dd, nr, cr, ld in cur:
    if t in asis_tabs or c in AUDIT: continue
    dflt = str(dd).strip() if dd is not None else None
    keep = dflt is not None and dflt != "'-'"
    apply = dflt if keep else "'/'"
    inlist = (t, c) in have
    rows.append(dict(table=t, column=c, col_id=cid, dtype=f"{dt}({dl})", default=dflt, on_null=on, apply=apply, rows=nr, created=cr, last_ddl=ld,
                     suspect=bool(SUSPECT.search(t)), inlist=inlist,
                     ddl=f"ALTER TABLE SCM_MGR.{t} MODIFY ({c} DEFAULT ON NULL {apply});" if not keep else "(기존 DEFAULT 유지)"))
    stat["이미 목록에 있음" if inlist else ("기존 DEFAULT 유지" if keep else ("DEFAULT '-' → '/'" if dflt == "'-'" else "DEFAULT 없음 → '/' 신설"))] += 1
added = 0
for r in rows:
    if r["inlist"] or r["suspect"]: continue
    ex["rows"].append({"owner": "SCM_MGR", "table": r["table"], "column": r["column"], "col_id": r["col_id"], "dtype": r["dtype"].split("(")[0], "nullable": "N",
                       "file_default": r["default"] or "[NULL]", "file_ddl": r["ddl"], "kind": "new_table", "apply_default": r["apply"],
                       "added": "2026-10-07 사용자 지시 — AS-IS 에 없는 신규 테이블의 NOT NULL 문자 컬럼", "where": "정제전에도 없음(신규 테이블)"})
    added += 1
ex["rule"] += " | 신규 테이블(AS-IS 전 DB 에 동명 없음) NOT NULL 문자 컬럼(감사 4종 제외) — 기존 DEFAULT('-' 아님) 유지, 그 외 '/'"
EXF.write_text(json.dumps(ex, ensure_ascii=False, indent=1), encoding="utf-8")

X = HERE / "DEFAULT범위_확대후보_20261007.xlsx"; wb = openpyxl.load_workbook(X)
if "3_신규테이블_NOTNULL" in wb.sheetnames: del wb["3_신규테이블_NOTNULL"]
w = wb.create_sheet("3_신규테이블_NOTNULL")
H = ["테이블", "컬럼", "순번", "데이터형", "현재 DEFAULT", "DEFAULT ON NULL", "적용 DEFAULT", "행수(통계)", "테이블 생성일", "최종 DDL", "이름상 백업·임시 의심", "이미 목록에 있음", "권고 DDL (실행하지 않음)", "결정", "비고"]
w.append(H)
for c in w[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496"); c.alignment = Alignment(wrap_text=True, vertical="center")
for r in rows:
    w.append([r["table"], r["column"], r["col_id"], r["dtype"], r["default"], r["on_null"], r["apply"], r["rows"], r["created"], r["last_ddl"], "의심" if r["suspect"] else "", "예" if r["inlist"] else "", r["ddl"],
              "" if r["inlist"] else ("보류 — 백업·테스트 복사본 의심(이름 기준) · 필요 시 확대" if r["suspect"] else "확대(사용자 지시 10-07)"), f"AS-IS {len(dbs)}개 DB 테이블 전수({asis_at}) 에 동명 없음 · TO-BE 사전 {tobe_at}"])
for i, wd in enumerate([34, 26, 6, 16, 12, 10, 12, 12, 12, 12, 12, 10, 70, 22, 50], 1): w.column_dimensions[get_column_letter(i)].width = wd
w.freeze_panes = "C2"; w.auto_filter.ref = w.dimensions
ws = wb["요약"]; ws.append([]); ws.append([f"③ 신규 테이블(AS-IS 없음) NOT NULL 문자 컬럼 (시트 3_신규테이블_NOTNULL) — 사용자 지시 10-07 확대", len(rows), f"{len({r['table'] for r in rows})} 테이블 · " + json.dumps(dict(stat), ensure_ascii=False)])
ws.append([f"   그중 운영 테이블 → 확대(추가분 파일에 포함)", sum(1 for r in rows if not r["suspect"] and not r["inlist"]), f"{len({r['table'] for r in rows if not r['suspect'] and not r['inlist']})} 테이블"])
ws.append([f"   이름상 백업·테스트 복사본 의심 → 보류(추가분 미포함)", sum(1 for r in rows if r["suspect"]), f"{len({r['table'] for r in rows if r['suspect']})} 테이블 — 이름 패턴(BAK/BK/TEMP/TMP/TEST/날짜 등)"])
ws.append(["   NUMBER·DATE 등 비문자 NOT NULL 컬럼은 '/' 를 넣을 수 없어 제외", "", ""])
wb.save(X)
print("신규 테이블 NOT NULL 문자 컬럼", len(rows), "/ 테이블", len({r["table"] for r in rows}), dict(stat), "· 추가분 파일에 추가", added, "· 추가분 합계", len(ex["rows"]))
print("의심 이름:", sorted({r["table"] for r in rows if r["suspect"]})[:40])
print("상위 테이블:", collections.Counter(r["table"] for r in rows).most_common(15))
