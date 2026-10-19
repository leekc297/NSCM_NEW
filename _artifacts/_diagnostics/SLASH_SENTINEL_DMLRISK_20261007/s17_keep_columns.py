"""[xlsx 시트 추가] 4_유지 대상 테이블·컬럼 목록 + 건수 (코드 줄 수 · AS-IS '-' 행수 9월 이전/이후 · TO-BE '-' 행수)"""
import json, pathlib, collections, re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
H = pathlib.Path(r"C:\Projects\NSCM\_artifacts\_diagnostics\SLASH_SENTINEL_DMLRISK_20261007")
f = H / "슬래시전환_DEFAULT_DB객체_정리_20261007_v2.xlsx"
if not f.exists(): f = H / "슬래시전환_DEFAULT_DB객체_정리_20261007.xlsx"
BD = {}
for p in H.glob("_asis_dash_bydate*.json"):
    for t, cols in json.loads(p.read_text(encoding="utf-8")).items():
        for c, e in cols.items():
            if "before_sep" in e: BD[(t, c)] = e
PS = json.loads((H / "_partner_scan.json").read_text(encoding="utf-8"))
CN = json.loads((H / "_scan_counts.json").read_text(encoding="utf-8"))
tobe = {(k.split(".", 1)[1], c): v for k, e in CN.items() for c, v in e.get("cols", {}).items()}
for t, e in PS.items():
    for c, v in e.get("tobe", {}).items(): tobe.setdefault((t, c), v)
wb = openpyxl.load_workbook(f)
rows = list(wb["4_DB객체_유지"].iter_rows(values_only=True))[1:]
per_col = collections.defaultdict(lambda: dict(lines=0, direct=0, via=0, objs=set(), uses=collections.Counter()))
for r in rows:
    cols = [x.strip() for x in str(r[8]).split(",") if "." in x]
    # 유지 판정 근거가 된 컬럼 = 9월 이전 AS-IS '-' 실재 컬럼만 (후보 여러 개면 그중 실재하는 것)
    hit = [c for c in cols if BD.get(tuple(c.split(".", 1)), {}).get("before_sep", 0) > 0] or cols
    for c in hit:
        d = per_col[c]; d["lines"] += 1; d["objs"].add(r[2]); d["uses"][r[5]] += 1
        d["via" if r[7] else "direct"] += 1
out = []
for c, d in sorted(per_col.items(), key=lambda x: -x[1]["lines"]):
    t, col = c.split(".", 1); b = BD.get((t, col), {}); tb = tobe.get((t, col), {})
    out.append([t, col, d["lines"], d["direct"], d["via"], len(d["objs"]), ", ".join(sorted(d["objs"]))[:300],
                "; ".join(f"{k} {v}" for k, v in d["uses"].most_common()), b.get("dash"), b.get("before_sep"), b.get("after_sep"), b.get("min_date"), b.get("max_date"),
                ", ".join(f"{x[0]}({x[1]:,})" for x in (b.get("created_by") or [])[:3]), tb.get("dash")])
if "5_유지_컬럼목록" in wb.sheetnames: del wb["5_유지_컬럼목록"]
w = wb.create_sheet("5_유지_컬럼목록")
head = ["테이블", "컬럼", "유지 코드 줄 수", "  직접 참조", "  경유", "객체 수", "객체", "쓰임별 줄 수", "AS-IS '-' 행 전체", "AS-IS '-' 9월 이전", "AS-IS '-' 9월 이후", "'-' 최소 생성일", "'-' 최대 생성일", "'-' 생성 주체 상위", "TO-BE '-' 행(실측·일부)"]
w.append(head)
for x in w[1]: x.font = Font(bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="305496"); x.alignment = Alignment(wrap_text=True, vertical="center")
for r in out: w.append(["" if v is None else v for v in r])
w.append([]); w.append(["합계", f"{len(out)} 컬럼 / {len({r[0] for r in out})} 테이블", sum(r[2] for r in out), sum(r[3] for r in out), sum(r[4] for r in out), len({o for d in per_col.values() for o in d['objs']})])
for i, wd in enumerate([34, 24, 12, 10, 8, 8, 60, 40, 14, 14, 14, 18, 18, 40, 14], 1): w.column_dimensions[get_column_letter(i)].width = wd
w.freeze_panes = "C2"; w.auto_filter.ref = f"A1:{get_column_letter(len(head))}{len(out) + 1}"
ws = wb["요약"]; ws.append([]); ws.append(["4_유지 컬럼 목록 (시트 5)", f"{len(out)} 컬럼 / {len({r[0] for r in out})} 테이블", "9월 이전 AS-IS '-' 실재 컬럼별 유지 줄 수·AS-IS '-' 행수·생성 주체"])
try: wb.save(f)
except PermissionError:
    f = H / "슬래시전환_DEFAULT_DB객체_정리_20261007_v3.xlsx"; wb.save(f); print("열려 있는 파일을 피해 _v3 로 저장")
print(f); print("컬럼", len(out), "테이블", len({r[0] for r in out}), "줄 합계", sum(r[2] for r in out), "(직접", sum(r[3] for r in out), "경유", sum(r[4] for r in out), ")")
for r in out[:25]: print(f"  {r[0]}.{r[1]}: {r[2]}줄 (직접 {r[3]}/경유 {r[4]}) · AS-IS '-' 9월 이전 {r[9]} / 이후 {r[10]} · 주체 {r[13][:40]}")
