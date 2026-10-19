# -*- coding: utf-8 -*-
"""v4 보정 — 시트 2 에 '적용 액션(규칙)'·'같은 줄 행 수'·'원문 잘림' 열 추가, 원 분류 열은 참고용으로 명시, 빈 3_DB객체_확인필요 시트 삭제.
입력 v3(다른 세션 산출) → 출력 v4. 데이터·DB 무변경 — 이 스크립트는 DB 에 접속하지 않는다(xlsx 라벨 보정만).
FRESH-GATE: 실제 적용 도구가 적용 직전 ALL_SOURCE 원문을 재조회해 시트 '현재' 텍스트와 대조한다(SLASH_SENTINEL_TARGET_20261006/zt/zt_case.py fresh_gate) — 시트의 줄 번호는 참고값."""
import re, collections
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

SRC = "슬래시전환_DEFAULT_DB객체_정리_20261007_v3.xlsx"; OUT = "슬래시전환_DEFAULT_DB객체_정리_20261007_v4.xlsx"
A_NVL, A_CMP, A_VAR, A_LIT = "NVL→'/'", "비교→'/'", "변형정규화→'/'", "리터럴→'/'"
wb = load_workbook(SRC); ws = wb["2_DB객체_변경"]
h = [c.value for c in ws[1]]; I = {n: i for i, n in enumerate(h)}
rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
VARRX = re.compile(r"IN\s*\(\s*'-'\s*,\s*'--'", re.I)


def rule_action(r):
    use = r[I["쓰임"]] or ""; cur = r[I["현재"]] or ""; orig = r[I["작업대상 xlsx 액션"]] or ""
    if orig == A_VAR or VARRX.search(cur):
        return A_VAR
    if use.startswith("NVL") or orig == A_NVL:
        return A_NVL
    if use in ("조건절 비교", "분기(IF/CASE/DECODE/NULLIF)"):
        return A_CMP
    return A_LIT


line_cnt = collections.Counter((r[I["객체명"]], r[I["줄"]]) for r in rows)
col = I["작업대상 xlsx 액션"] + 1
ws.insert_cols(col + 1, 3)
ws.cell(1, col).value = "원 분류(작업대상 xlsx · 참고용 — 적용 여부 아님)"
ws.cell(1, col + 1).value = "적용 액션(규칙)"; ws.cell(1, col + 2).value = "같은 줄 행 수(합쳐서 적용)"; ws.cell(1, col + 3).value = "원문 500자 잘림('-' 는 뒤쪽)"
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
fill = {A_NVL: "C6E0B4", A_CMP: "BDD7EE", A_VAR: "9BC2E6", A_LIT: "E2EFDA"}
cnt = collections.Counter(); trunc = 0
for i, r in enumerate(rows, start=2):
    a = rule_action(r); cnt[a] += 1
    n = line_cnt[(r[I["객체명"]], r[I["줄"]])]
    cur = r[I["현재"]] or ""; t = "'-'" not in cur and "''-''" not in cur; trunc += t
    ws.cell(i, col + 1).value = a; ws.cell(i, col + 1).fill = PatternFill("solid", fgColor=fill[a])
    ws.cell(i, col + 2).value = n; ws.cell(i, col + 3).value = "Y" if t else ""
for L, w in (("H", 16), ("I", 10), ("J", 12)):
    ws.column_dimensions[L].width = w
# 빈 3 시트 삭제 (사용자 지시 10-07)
if "3_DB객체_확인필요" in wb.sheetnames and wb["3_DB객체_확인필요"].max_row <= 1:
    del wb["3_DB객체_확인필요"]
multi = sum(1 for v in line_cnt.values() if v > 1)
s = wb["요약"]; s.append([])
s.append(["v4 보정 (10-07, 시트 2)", "'작업대상 xlsx 액션'(원 분류)은 이전 분석의 분류를 참고로 옮긴 것일 뿐 적용 여부가 아니다. 적용 여부 = 시트 배치(2=변경 · 4=유지, 규칙: 9월 이전 AS-IS 실데이터에 '-' 사용 여부). '적용 액션(규칙)' 열이 실제 변경 방식이다. 3_DB객체_확인필요 는 0건이라 삭제."])
s.append(["", "적용 액션: %s %d · %s %d · %s %d · %s %d / 같은 줄에 행 2개 이상 %d줄(적용 시 한 줄로 합침) / 원문 500자 잘림 %d행('현재'·'변경 후' 셀이 잘린 것 — 적용은 DB 원문 전체 줄에 규칙 적용)"
          % (A_NVL, cnt[A_NVL], A_CMP, cnt[A_CMP], A_VAR, cnt[A_VAR], A_LIT, cnt[A_LIT], multi, trunc)])
s.append(["", "적용 절차: 객체별 적용 당일 DB 원문 재조회 → '현재' 텍스트(앞부분)로 줄 위치 확인(줄 번호는 참고) → 같은 줄 행 합쳐 규칙 적용 → CREATE OR REPLACE → ALL_ERRORS 0 · VALID · 줄 수 동일 → 재스캔으로 남은 '-' = 시트 4 와 일치 확인 → VIEW 변경분은 의존 패키지 재컴파일"])
wb.save(OUT)
print(OUT, "| 적용 액션", dict(cnt), "| 다중행 줄", multi, "| 잘림", trunc, "| sheets", wb.sheetnames)
