"""[조회 전용 + xlsx 생성] DEFAULT '/' 설정 범위 확대 후보 — 1차 목록(처리대상이력 1,797)에 없는 컬럼
① 사내 ZIOUI1 SCM_MGR 에서 현재 DEFAULT '-' 인데 1차 목록에 없는 컬럼 (확대 후보)
② 기준 파일 '정제전' 에 있었으나 '처리대상이력' 에 없는 컬럼 전체 (탈락 경위 추적용)
출력 = DEFAULT범위_확대후보_20261007.xlsx
"""
import sys, json, shutil, pathlib, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE)); import s1_data_scan as s
SRC = pathlib.Path(r"C:\TRANFILE\RECEIVE\20260820_NOTNULL컬럼_DEFAULT처리(품질DB전체).xlsx")
tmp = HERE / "_pop_src_copy.xlsx"; shutil.copy2(SRC, tmp)
wb0 = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
def rows(n): return [r for r in list(wb0[n].iter_rows(values_only=True))[1:] if r and r[0]]
pre = {(r[1], r[2]): r for r in rows("정제전")}; hist = {(r[1], r[2]) for r in rows("처리대상이력")}
tgt = {r[0] for r in rows("대상테이블")}; exc = {r[0] for r in rows("제외테이블")}
pop = s.population(); POP = {(c["table"], c["column"]): c for c in pop}
held_names = collections.defaultdict(list)
for c in pop:
    if c["held"]: held_names[c["column"]].append(c["table"])
stats = {(r["table"], r["column"]): r for r in json.loads((HERE / "_scan_stats.json").read_text(encoding="utf-8"))["rows"]}

con = s.connect(); cur = con.cursor()
cur.execute("SELECT sys_context('userenv','db_name'), TO_CHAR(SYSDATE,'YYYY-MM-DD HH24:MI:SS') FROM dual"); dbn, at = cur.fetchone()
cur.execute("""SELECT C.TABLE_NAME, C.COLUMN_NAME, C.COLUMN_ID, C.DATA_TYPE, C.DATA_LENGTH, C.NULLABLE, C.DEFAULT_ON_NULL, C.DATA_DEFAULT, T.TEMPORARY, T.NUM_ROWS,
                      TO_CHAR(O.CREATED,'YYYY-MM-DD'), TO_CHAR(O.LAST_DDL_TIME,'YYYY-MM-DD')
                 FROM ALL_TAB_COLUMNS C JOIN ALL_TABLES T ON T.OWNER=C.OWNER AND T.TABLE_NAME=C.TABLE_NAME
                 JOIN ALL_OBJECTS O ON O.OWNER=C.OWNER AND O.OBJECT_NAME=C.TABLE_NAME AND O.OBJECT_TYPE='TABLE'
                WHERE C.OWNER='SCM_MGR'""")
cols = {}
for t, c, cid, dt, dl, nu, on, dd, tp, nr, cr, ld in cur:
    cols[(t, c)] = dict(cid=cid, dtype=dt, dlen=dl, nullable=nu, on_null=on, default=(str(dd).strip() if dd is not None else None), temp=tp, rows=nr, created=cr, last_ddl=ld)
cur.execute("""SELECT CC.TABLE_NAME, CC.COLUMN_NAME, C.CONSTRAINT_TYPE FROM ALL_CONSTRAINTS C JOIN ALL_CONS_COLUMNS CC ON CC.OWNER=C.OWNER AND CC.CONSTRAINT_NAME=C.CONSTRAINT_NAME
                WHERE C.OWNER='SCM_MGR' AND C.CONSTRAINT_TYPE IN ('P','U')""")
keycol = collections.defaultdict(set)
for t, c, ty in cur: keycol[(t, c)].add("PK" if ty == "P" else "UK")

def stage(t, c):
    if (t, c) in pre: return "정제전에 있었으나 탈락" + (" (테이블이 '제외테이블' 시트에 있음)" if t in exc else " (제외 사유 기록 없음)")
    if any(tt == t for tt, _ in pre): return "테이블은 정제전에 있으나 이 컬럼은 없음"
    return "정제전에도 없음 (테이블 생성 " + str(cols[(t, c)]["created"]) + ")"

H1 = ["테이블", "컬럼", "순번", "데이터형", "NULL 허용", "현재 DEFAULT", "DEFAULT ON NULL", "테이블 구분", "행수(통계)", "테이블 생성일", "최종 DDL",
      "기준 파일에서의 위치", "PK/UK 컬럼", "같은 테이블의 1차 목록 컬럼 수", "'-' 값 보유(컬럼 통계)", "같은 이름의 보류 51 컬럼(테이블)", "권고 DDL (실행하지 않음 — 검토용)", "결정(확대 / 제외)", "비고"]
cand = []
poptab = collections.Counter(t for t, _ in POP)
for (t, c), v in sorted(cols.items()):
    if v["default"] != "'-'" or (t, c) in hist: continue
    st = stats.get((t, c))
    flag = {"dash": "있음", "none": "없음", "nostat": "통계 없음", "unknown(low<'-')": "판정 불가(최솟값이 '-' 보다 앞)"}.get((st or {}).get("flag"), "미조회(1차 목록 컬럼과 이름이 다름)")
    cand.append([t, c, v["cid"], f'{v["dtype"]}({v["dlen"]})', v["nullable"], v["default"], v["on_null"], "임시테이블" if v["temp"] == "Y" else "영구", v["rows"], v["created"], v["last_ddl"],
                 stage(t, c), "/".join(sorted(keycol.get((t, c), []))), poptab.get(t, 0), flag, ", ".join(held_names.get(c, []))[:200],
                 f"ALTER TABLE SCM_MGR.{t} MODIFY ({c} DEFAULT ON NULL '/');" if v["on_null"] == "YES" else f"ALTER TABLE SCM_MGR.{t} MODIFY ({c} DEFAULT '/');", "", ""])

H2 = ["테이블", "컬럼", "정제전 NULLABLE", "정제전 DATA_DEFAULT", "테이블이 '제외테이블' 시트에 있음", "테이블이 '대상테이블' 시트에 있음", "현재 사내 DB 상태", "현재 DEFAULT", "DEFAULT ON NULL", "행수(통계)", "비고"]
drop = []
for (t, c), r in sorted(pre.items()):
    if (t, c) in hist: continue
    v = cols.get((t, c))
    state = "테이블·컬럼 없음" if v is None else ("DEFAULT '-' — 확대 후보" if v["default"] == "'-'" else ("DEFAULT 없음" if v["default"] is None else "다른 DEFAULT"))
    drop.append([t, c, r[5], r[6], "예" if t in exc else "", "예" if t in tgt else "", state, (v or {}).get("default"), (v or {}).get("on_null"), (v or {}).get("rows"), ""])

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "요약"
perm = [x for x in cand if x[7] == "영구"]
S = [["DEFAULT '/' 설정 범위 확대 후보 — 1차 목록(처리대상이력 1,797 컬럼)에 없는 컬럼"], [],
     ["기준"], [f"  사내 {dbn} SCM_MGR 컬럼 사전 직접 조회 {at}"], [f"  기준 파일 {SRC.name} (정제전 {len(pre):,} 컬럼 / 제외테이블 {len(exc)} / 대상테이블 {len(tgt)} / 처리대상이력 {len(hist):,})"],
     ["  1차 목록 = 작업대상 xlsx 1_컬럼판정 1,797 컬럼 (보류 51 포함)"], [],
     ["① 현재 DEFAULT '-' 인데 1차 목록에 없는 컬럼 (시트 1_확대후보)", len(cand), f"{len({x[0] for x in cand})} 테이블"],
     ["    영구 테이블", len(perm), f"{len({x[0] for x in perm})} 테이블"],
     ["      정제전에 있었으나 탈락 (제외 사유 기록 없음)", sum(1 for x in perm if "기록 없음" in x[11])],
     ["      정제전에 있었으나 탈락 (제외테이블 시트에 있음)", sum(1 for x in perm if "'제외테이블' 시트에 있음" in x[11])],
     ["      테이블은 정제전에 있으나 컬럼 없음", sum(1 for x in perm if x[11].startswith("테이블은"))],
     ["      정제전에도 없음 (파일 작성 뒤 생성된 테이블 등)", sum(1 for x in perm if x[11].startswith("정제전에도"))],
     ["    임시테이블", len(cand) - len(perm)],
     ["    그중 같은 이름의 보류 51 컬럼이 있는 것 (보류 결정과 함께 봐야 함)", sum(1 for x in cand if x[15])],
     ["    그중 PK/UK 컬럼", sum(1 for x in cand if x[12])], [],
     ["② 정제전에 있었으나 처리대상이력에 없는 컬럼 (시트 2_정제전_탈락전체)", len(drop), f"{len({x[0] for x in drop})} 테이블"],
     ["    테이블이 제외테이블 시트에 있음", sum(1 for x in drop if x[4])],
     ["    제외 사유 기록 없음", sum(1 for x in drop if not x[4]), f"{len({x[0] for x in drop if not x[4]})} 테이블"],
     ["    현재 상태별", json.dumps(collections.Counter(x[6] for x in drop), ensure_ascii=False)], [],
     ["읽는 법"], ["  1_확대후보 의 '결정' 열에 확대/제외 를 적으면 그 기준으로 변환 목록·위험 목록을 다시 만든다."],
     ["  '같은 이름의 보류 51 컬럼' 이 있는 행은 AS-IS 코드가 그 컬럼명의 '-' 를 비교·분기한다는 근거가 다른 테이블에 있다는 뜻 — 보류 51 과 같은 결정을 따르는 것이 일관적이다."],
     ["  \"'-' 값 보유\" 는 컬럼 통계(최솟값) 기준 1차 선별이며 실측이 아니다."], [],
     ["한계"], ["  탈락 사유는 기준 파일에 기록이 없어 확인 불가 — 파일 작성자 확인 필요."],
     ["  DEFAULT 가 '-' 가 아닌(없는) NOT NULL 컬럼은 이 목록에 없다 — 코드의 NVL(x,'-') 로만 '-' 가 들어가는 컬럼은 별도."],
     ["  DB 객체·데이터 변경 없음. DDL 은 검토용 문자열."]]
for r in S: ws.append(r)
ws["A1"].font = Font(bold=True, size=13); ws.column_dimensions["A"].width = 86; ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 60
def sheet(name, head, data, widths):
    w = wb.create_sheet(name); w.append(head)
    for c in w[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496"); c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in data: w.append(r)
    for i, wd in enumerate(widths, 1): w.column_dimensions[get_column_letter(i)].width = wd
    w.freeze_panes = "C2"; w.auto_filter.ref = w.dimensions
sheet("1_확대후보", H1, cand, [34, 24, 6, 16, 8, 10, 10, 10, 12, 12, 12, 44, 8, 10, 16, 40, 70, 14, 30])
sheet("2_정제전_탈락전체", H2, drop, [36, 26, 10, 12, 14, 14, 26, 12, 10, 12, 30])
out = HERE / "DEFAULT범위_확대후보_20261007.xlsx"; wb.save(out)
print(out); print("확대후보", len(cand), "영구", len(perm), "· 탈락전체", len(drop), collections.Counter(x[6] for x in drop))
print("보류51 동명:", sum(1 for x in cand if x[15]), "· PK/UK:", sum(1 for x in cand if x[12]), "· '-' 보유:", collections.Counter(x[14] for x in cand))
