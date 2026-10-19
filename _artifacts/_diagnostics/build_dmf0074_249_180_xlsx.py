"""DMF0074 Excel Upload 249행 → 조회 180행 과정 정리 xlsx 생성
원천: 업로드 파일 C:\\TRANFILE\\RECEIVE\\DMF0074_Template_20261001.xlsx (정적 분석)
      + 클라우드 실행 결과 (2026-10-01 18:54:40 배치, minyonglee1) — 대사 SQL [2][3][4][7][8] 사용자 공유분
"""
import openpyxl, collections
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"C:\TRANFILE\RECEIVE\DMF0074_Template_20261001.xlsx"
OUT = r"C:\Projects\NSCM\_artifacts\_diagnostics\DMF0074_업로드249행_조회180행_정리_20261001.xlsx"
METALS = ["Ni", "Co", "Li2CO3", "LIOH", "Al", "Mn", "Cu"]
HIDDEN = [("A15FD5002", "Ford_MX_CTN", "AMFMCC727-A1-021"), ("A15FD5002", "Ford_MX_CTN", "AMFMCC727-B1-021"),
          ("A15FD5002", "Ford_MX_CTN", "AMFMCC727-F1-021"), ("A15FD5002", "Ford_MX_TPZ", "AMFMCC727-B1-021")]

HFILL = PatternFill("solid", fgColor="1F4E78"); HFONT = Font(bold=True, color="FFFFFF")
SUB = PatternFill("solid", fgColor="DDEBF7"); RED = PatternFill("solid", fgColor="FCE4D6"); GRN = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def table(ws, r0, headers, rows, widths=None, fills=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h); c.fill = HFILL; c.font = HFONT; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + i, j, v); c.border = BOX; c.alignment = WRAP
            if fills and fills[i - 1]:
                c.fill = fills[i - 1]
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + len(rows) + 1


def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = Font(bold=True, size=14)
    if sub:
        ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="595959")


# ---------- 파일 정적 분석 ----------
src = openpyxl.load_workbook(SRC, data_only=True).worksheets[0]
rows = [r for r in src.iter_rows(min_row=3, values_only=True) if any(v not in (None, "") for v in r)]
e = lambda v: v in (None, "")
detail = []
for n, r in enumerate(rows, 1):
    full, part, partial_detail = 0, 0, []
    for m in range(7):
        key = r[13 + m]; vals = [r[20 + m], r[27 + m], r[34 + m], r[41 + m], r[48 + m]]
        if e(key):
            continue
        if all(not e(v) for v in vals):
            full += 1
        else:
            part += 1
            names = ["Period(Fr)", "Period(To)", "Contract", "Change", "Index"]
            partial_detail.append(f"{METALS[m]}:" + "/".join(nm for nm, v in zip(names, vals) if e(v)))
    detail.append([n + 2, r[0], r[1], r[3], r[4], r[5], r[6] or "(공란)", r[7], r[8], r[9], r[10], r[11],
                   full, part, "; ".join(partial_detail)])

wb = openpyxl.Workbook()

# ---------- 1. 요약 ----------
ws = wb.active; ws.title = "1_요약"
title(ws, "DMF0074 Excel Upload — 249행 업로드 → 조회 180행 정리",
      "클라우드 2026-10-01 18:54:40 업로드(minyonglee1) 기준 · PKG_DMF0074 EMPTYASIS 적용본(18:22:23) · 조회 기간 202610~202612")
r = table(ws, 4, ["단계", "건수", "월별(×3개월)", "줄어든 수", "사유", "AS-IS 동일 여부"], [
    ["① 엑셀 업로드", "249 행", "83 행 × 3", "", "Div 81 · Plan 202610/202611/202612 · 83개 조합(Project·Sales Site·통화·FCST) × 3개월 · Model Code 전부 공란 · 검증 오류 0 (SUCCESS 249)", "-"],
    ["② 모델 확장 대상 없음", "159 행 → 0", "53 행 × 3", "-159 행", "Model 공란 행은 최신 MRF=Y Sales Plan(계획+실적)에 있는 모델로만 펼쳐짐. 53개 조합은 해당 모델이 없어 오류 없이 0건", "동일 (AS-IS 2-1 ITEMID 없는것 처리)"],
    ["③ 모델 확장", "90 행 → 192 모델행", "30 행 → 64 모델 × 3", "", "30개 조합이 64개 모델로 펼쳐짐 (조합당 1~N 모델)", "동일"],
    ["④ 원장 저장", "192 모델행 (금속 1,344행)", "64 모델 × 7금속 = 448 × 3", "", "DM_BPS_METAL_LINK_PRICE 에 모델당 금속 7행 MERGE · 빈 금속칸 = 키/기간 '-' · 단가/수량 0 (EMPTYASIS)", "동일 (AS-IS '' / 0 저장과 같은 형태)"],
    ["⑤ 화면 조회", "180 행", "60 행 × 3", "-12 행", "조회 SP 최종 INNER JOIN SCM_MST_SSO_DEPT_INFO — SSPM(SCM_MST_SITE_ITEM.DEPT_ID) 이 부서 마스터에 없는 4개 모델 × 3개월 미조회", "구조 동일 (AS-IS 도 부서 INNER JOIN) · 원인은 클라우드 SSPM DEPT_ID 데이터"],
], widths=[20, 22, 24, 12, 80, 34], fills=[None, RED, GRN, GRN, RED])
ws.cell(r + 1, 1, "결론").font = Font(bold=True)
ws.cell(r + 2, 1, "· 249 → 192: 업로드 SP 정상 동작. Model 공란 행의 확장 규칙(AS-IS 동일)으로 Sales Plan 에 모델이 없는 53개 조합(월)이 0건, 나머지 30개 조합이 64개 모델로 펼쳐짐.")
ws.cell(r + 3, 1, "· 192 → 180: 클라우드 SSPM 데이터. 4개 모델의 SCM_MST_SITE_ITEM.DEPT_ID 가 부서 마스터(SCM_MST_SSO_DEPT_INFO)에 없어 조회에서 제외(원장에는 저장됨). 사내 동일 모델은 70021298(미주마케팅2팀) 정상.")
ws.cell(r + 4, 1, "· 조치: 클라우드 SSPM DEPT_ID 보정 / (선택) 업로드 시 SSPM DEPT 검증 신설 — 결정 대기")

# ---------- 2. 처리단계(SP) ----------
ws = wb.create_sheet("2_업로드SP_처리단계")
title(ws, "PKG_DMF0074.P_DMF0074_EXCEL_UPLOAD — Model 공란 행 처리 단계", "저장(P_MODE=1) 기준")
table(ws, 4, ["순서", "SP 단계", "처리 내용", "조건/규칙", "이번 업로드 결과"], [
    ["1", "검증 (C1~C43)", "필수값·날짜·키코드·숫자·기간 등 검증", "Model 공란은 오류 아님 · 동적 금속칸 공란 검증은 AS-IS 와 동일", "249행 전부 SUCCESS"],
    ["2", "기준 버전 (GT_DM74_TEMP_VERSION_NO)", "사업부별 최신 SP_W 버전 중 MRF(ATTRIBUTE7)='Y'", "계획구간 RESULT_TO~PLAN_TO(52주) / 실적구간 RESULT_FROM~RESULT_TO", "81 = SP_W_20261001_V001"],
    ["3", "2) 후보 모델 (GT_DM74_TEMP_SP_ITEM)", "RTF_SYS_RP 수량<>0 인 모델 수집", "DM_PLN_SALES_PLAN_C(계획구간) UNION DM_PLN_SALES_PLAN_RESULT_C(실적구간)", "-"],
    ["4", "2-1) Model 공란 확장", "엑셀 1행 → 조건 만족 모델 수만큼 복제 (금속값은 엑셀 행 그대로)",
     "SSPM(SCM_MST_SITE_ITEM) 같은 Project·Sales Site · VALID=Y / FCST type 표시명→코드(FREE 제외) / 제품(P)·활성 / 3)의 후보에 같은 Project·Site·통화·FCST 존재 / (82만) Energy 보유",
     "월 83행 중 30행 → 64모델, 53행 → 0"],
    ["5", "2-2)/2-3) Model 지정 행 병합", "Model 지정 행이 있으면 같은 키 확장행에 값 있는 칸만 덮어씀", "키 = Div·Plan·EndCustomer·Project·Plant·Model·Site·통화·FCST", "해당 없음 (지정 행 0)"],
    ["6", "3) 금속 UNPIVOT", "모델 1건 → 금속 7행 (Ni·Co·Li2CO3·LIOH·Al·Mn·Cu)", "UNPIVOT INCLUDE NULLS [E20261001-EMPTYASIS] · 빈칸 키/기간/Exception '-' · 단가/수량/Energy 0", "64 × 7 = 448 (월)"],
    ["7", "4)/5) 원장·이력 저장", "DM_BPS_METAL_LINK_PRICE MERGE + _H 이력", "MERGE 키 = 위 키 + METAL_KIND_CODE", "1,344 금속행 / 192 모델행"],
], widths=[6, 30, 45, 70, 26])

# ---------- 3. 조회 단계 ----------
ws = wb.create_sheet("3_조회SP_필터")
title(ws, "PKG_DMF0074.P_DMF0074_GET_METAL_LINK_PRICE — 원장 → 화면", "그리드 1행 = Plan·Team·EndCustomer·Project·Plant·Model·Site·통화·FCST·Exception·Energy")
table(ws, 4, ["순서", "처리", "조건", "이번 조회 영향"], [
    ["1", "원장 조회", "DIVISION_ID · PLAN_YYYYMM 기간 · Team/EndCustomer/Site/Plant/Project/Model 필터 (기본 ALL)", "192 모델행 (202610~202612)"],
    ["2", "팀 부여", "LEFT JOIN SCM_MST_SITE_ITEM(SSPM) → SITE_ITEM.DEPT_ID", "4개 모델 DEPT_ID 가 부서 마스터에 없는 값"],
    ["3", "최종 SELECT", "INNER JOIN SCM_MST_SSO_DEPT_INFO DI ON DI.DEPT_ID = 팀 (AS-IS L782 동일)", "4 모델 × 3개월 = 12행 제외 → 180"],
    ["4", "FCST 유효성", "INNER JOIN SALES_TYPE_CODE (FREE 제외)", "영향 없음"],
], widths=[6, 18, 80, 40])

# ---------- 4. 미조회 12건 ----------
ws = wb.create_sheet("4_미조회12건")
title(ws, "원장 저장 · 화면 미조회 12건 (클라우드 [8] 결과)", "SSPM DEPT_ID 가 부서 마스터(SCM_MST_SSO_DEPT_INFO)에 없음 — 사내 동일 모델 DEPT_ID = 70021298 (자동차.마케팅.마케팅1.미주마케팅2팀)")
hid = [[ym, p, s, i, "USD", "MPS", "(부서 마스터에 없는 값)", "70021298", "N"] for (p, s, i) in HIDDEN for ym in ("202610", "202611", "202612")]
table(ws, 4, ["Plan", "Project", "Sales Site", "Model", "통화", "FCST", "클라우드 SSPM DEPT_ID", "사내 SSPM DEPT_ID", "화면 조회"], hid,
      widths=[9, 12, 14, 20, 7, 7, 26, 18, 9], fills=[RED] * len(hid))

# ---------- 5. 월별 대사 ----------
ws = wb.create_sheet("5_월별대사")
title(ws, "월별 건수 대사 (클라우드 [2][4][7][8] 실측)")
mon = []
for ym in ("202610", "202611", "202612"):
    mon.append([ym, 83, 53, 30, 64, 448, 64, 4, 60])
mon.append(["합계", 249, 159, 90, 192, 1344, 192, 12, 180])
table(ws, 3, ["Plan", "엑셀 행", "확장 0 행", "확장된 행(=조합)", "저장 모델행", "저장 금속행", "원장 그리드행", "부서 미등록 제외", "화면 조회"], mon,
      widths=[9, 10, 11, 16, 12, 12, 14, 16, 11], fills=[None, None, None, SUB])
ws.cell(9, 1, "※ 확장된 행 30 = [4] COMBOS(저장된 Project·Site·통화·FCST 조합 수). 파일은 월별로 조합당 1행이므로 확장 0 행 = 83 - 30 = 53.")
ws.cell(10, 1, "※ 대사 SQL [3] 의 예상치(확장 0 = 55행, 54모델)는 실적(RESULT_C) 후보 모델을 빠뜨린 값 — SP 는 계획+실적 UNION. 실측 [4] 가 정답.")

# ---------- 6. 파일 행별 ----------
ws = wb.create_sheet("6_파일행별_금속입력")
title(ws, "업로드 파일 249행 — 행별 금속 입력 상태 (파일 정적 분석)",
      "부분입력 금속 = 키는 있고 기간/단가/Index 중 빈칸 있음 → EMPTYASIS 이후 '-'/0 으로 저장 (이전 판본은 해당 금속 탈락)")
fills = [RED if d[12] == 0 and d[13] > 0 else None for d in detail]
table(ws, 4, ["엑셀 행", "Div", "Plan", "End Customer", "Project Code", "Project", "Model", "Metal Link", "Sales Site", "Currency", "FCST",
              "Exception", "완전입력 금속 수", "부분입력 금속 수", "부분입력 상세(빈 칸)"], detail,
      widths=[7, 5, 8, 10, 12, 18, 9, 9, 16, 8, 8, 9, 10, 10, 60], fills=fills)
ws.freeze_panes = "A5"; ws.auto_filter.ref = f"A4:O{4 + len(detail)}"

# ---------- 7. 근거 SQL ----------
ws = wb.create_sheet("7_근거SQL")
title(ws, "근거 SQL 파일")
table(ws, 3, ["파일", "용도"], [
    [r"_artifacts\_diagnostics\DMF0074_UPLOAD_180_CLOUD_RECON_20261001.sql", "[2] 배치 상태 · [3] 사유별 예상 · [4] 실제 저장 · [7] 화면 건수 대사 · [8] 부서 미등록 미조회"],
    [r"_artifacts\_diagnostics\DMF0074_UPLOAD_108_CLOUD_VERIFY_20261001.sql", "[0] 배치 목록 · [1] 버전 · [5] 부분입력 상세 · [6] PKG 판본"],
    [r"_artifacts\_deploy_packages\DMF0074_EMPTYASIS_CLOUD_20261001", "빈 금속칸 AS-IS 동일 저장 보완 (클라우드 18:22:23 적용 확인)"],
], widths=[75, 90])

wb.save(OUT)
print("saved", OUT, "| rows", len(detail), "| partial-only rows", sum(1 for d in detail if d[12] == 0 and d[13] > 0))
