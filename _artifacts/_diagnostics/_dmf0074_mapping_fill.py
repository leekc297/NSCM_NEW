#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""DMF0074 대사표 세부 축 생성·기입 (도구 --init 이 §헤더 단위로만 뽑는 한계 보완).

도구 축(E/CRUD/MODAL/COMMON/HANDLER) 뒤에 원천 md 의 **원자 항목**을 축별로 분해해 append.
  IT 이벤트·액션 21 / BT 버튼 8 / CT 조회조건 9 / PU 팝업 5 / VA Validation 10 + XP Expression /
  ED 편집제어·저장 7 / GP 격차 9 / EV Excel Upload SP Validation 49(xlsx 실측 연동)
기입값은 **FE/SP 실측 근거**만 사용. 재실행 시 기존 기입 보존(도구 parse_existing 규약과 동일 ID 키).
"""
import io
import pathlib
import re
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook  # noqa: E402

MAP = pathlib.Path(r"C:\Projects\NSCM\nscm-fe\features\demand\src\pages\metal-link-price"
                   r"\__asis-mapping__.md")
XLSX = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서"
                    r"\DMF0074_MetalLinkPrice_ExcelUpload_Validation_v1.1.xlsx")
P = "features/demand/src/pages/metal-link-price"
HDR = "| ID | 요지 (원천 md 1줄) | 구현 위치 (파일:라인) | 사유 (미구현/제외 시 필수) | 상태 |"
SEP = "|---|---|---|---|---|"

# (ID, 요지, 구현위치, 사유, 상태)
ROWS = []


def add(_id, gist, loc, why="", st="구현"):
    ROWS.append((_id, gist, loc, why, st))


S = "## 이벤트·액션 원자 축 (md §1 [1]~[21] 1:1)"
add("IT-01", "[1] ctor DMF0074_MetalLinkPrice() — 폼 생성", f"{P}/MetalLinkPrice.tsx:39 (컴포넌트 진입)")
add("IT-02", "[2] Form_Load — InitForm/InitControls/DefineProperty/CreateSpread chain",
    f"{P}/MetalLinkPrice.tsx:79~83 (mount useEffect — clearFilters + fetchStaticColumns)")
add("IT-03", "[3] InitControls — GetDMInnerTable 콤보 chain 초기 적재(7종)",
    f"{P}/components/MetalLinkPriceForm.tsx:41~43 (MyDivision/EndCustomer/SalesSite) + "
    f"TeamField:89 + PlantSiteField:253",
    "AS-IS 화면 전용 GetDMInnerTable 7종 → 공용 CMBO/카탈로그 필드로 대체(DMF0074 자체 CMBO 등록 0). "
    "동일 원천 조회이므로 등가")
add("IT-04", "[4] DefineProperty — 기본값(Period 당월~+6, Version 'Current')",
    f"{P}/useMetalLinkPriceStore.ts:47~49 (initMetalLinkPriceState)")
add("IT-05", "[5] CreateSpread/InitSpread — 동적 2단 PIVOT 헤더(6그룹×7원소)",
    f"{P}/components/MetalLinkPriceGrid.tsx:46~78 (mergedColumns — admin sClsLblNm 그룹)")
add("IT-06", "[6] clbDivisions SIC — 하위 5종 reset + Project/PlantSite 활성분기 + Version clear + InitSpread",
    f"{P}/components/MetalLinkPriceForm.tsx:69~81 (reset+version clear) + "
    f"MetalLinkPrice.tsx:88~95 (컬럼 재조회·그리드 초기화)")
add("IT-07", "[7] clbTeam SIC — EndCustomer/SalesSite reset",
    f"{P}/components/MetalLinkPriceForm.tsx:92~95")
add("IT-08", "[8] clbEndCustomer SIC — SalesSite reset",
    f"{P}/components/MetalLinkPriceForm.tsx:120~123")
add("IT-09", "[9] txtProject_KeyDown — Enter 시 Project 팝업", "—",
    "TO-BE 는 검색 아이콘 버튼으로 팝업 진입(Form:176~184). SearchArea 안에서 Enter 는 조회 트리거라 "
    "AS-IS Enter→팝업과 충돌 — 의도적 미이식(UX 표준 우선). 보완목록 R-1 로 사용자 판단 대기", "제외")
add("IT-10", "[10] txtItemId_KeyDown — Enter 시 Item 팝업", "—",
    "IT-09 와 동일 사유(버튼 진입 Form:202~207). 보완목록 R-1", "제외")
add("IT-11", "[11] FindButton override → FindData() — 검증 6 + 조회 + 전체잠금 + SmartColMerge",
    f"{P}/MetalLinkPrice.tsx:98~145 (validateFind+handleSearch) + "
    f"components/MetalLinkPriceGrid.tsx (isReadOnly + 세로 머지)")
add("IT-12", "[12] Insert/Delete/Edit/Save override 4종 — 전부 빈 몸체",
    f"{P}/MetalLinkPrice.tsx:157~164 (hotkeyConfig insert/delete/save=false)",
    "AS-IS 빈 몸체 = 기능 없음 → 버튼 미노출이 1:1 파리티")
add("IT-13", "[13] ExcelButton override — 그리드 Excel 다운로드",
    f"{P}/MetalLinkPrice.tsx:220 (excelFileName=MAIN_EXCEL_PREFIX — SubTitle 공통 EXCEL, "
    "BTN 배선 DMF0074_main_EXCEL)")
add("IT-14", "[14] btnProject_Click — SCM0002_ProjectbyDivisionPopupEx → txtProject",
    f"{P}/MetalLinkPrice.tsx:63,182 계열 ProjectByDivisionModal (Form:176~184 트리거)")
add("IT-15", "[15] btmItem_Click — DM0002_ItembyDivisionPopupEx → txtItemId",
    f"{P}/components/MetalLinkPriceForm.tsx:202~207 + ItemByDivisionModal",
    "⭐ AS-IS 는 취소 시 빈값 덮임(quirk L670) — TO-BE 는 기존값 유지로 개선(재현 금지 대상)")
add("IT-16", "[16] btnSPVer_Click — DMF0073_VersionBySalesPricePopup(SP, ±3M) → 최대 2개",
    f"{P}/components/modals/MetalLinkVersionModal.tsx:114(versionType 'SP_W'),145~157(최대2 MSG00255)",
    "AS-IS 'SP' 는 카테고리였으나 TO-BE VERSION_TYPE 정확토큰이 'SP_W' — 2026-07-10 실측 정정")
add("IT-17", "[17] btnTemplate 'Excel Form' — 전용 MetalLinkPricePopup",
    f"{P}/components/modals/MetalLinkTemplateModal.tsx (전체) + MetalLinkPrice.tsx:169~197 진입 게이트")
add("IT-18", "[18] btnExcelUpload 'Upload' — ExcelUploadForm(DIV40→formId 분기) → SP → 재조회",
    f"{P}/MetalLinkPrice.tsx:151~154(formId 분기),285~320(FcstUploadModal 2-Step onSave/onUploaded)")
add("IT-19", "[19] SetChangeControl(FILTER_PART) — Project/PlantSite 활성 분기",
    f"{P}/components/MetalLinkPriceForm.tsx:48~50(isProjectDisabled) + 253~259(PlantSiteField division)")
add("IT-20", "[20] GetSpdTitle — DMF0000_GET_COMMON_SPREAD_TITLE 로 2단 헤더 라벨 취득",
    f"{P}/components/MetalLinkPriceGrid.tsx:62~63 (admin SHEET_COLUMNS sClsLblNm)",
    "동일 원천(admin 컬럼 메타)을 직접 소비 — 별도 SPREAD_TITLE SP 불요(등가 대체)")
add("IT-21", "[21] GetDayRange — ⚠ 미호출 dead code", "—",
    "AS-IS 에서 호출되지 않는 사문 메서드 — 이관 대상 아님(파리티)", "해당없음")

S2 = "## 버튼 인벤토리 축 (md §3.1 7 + toolbar 빈 override)"
add("BT-01", "Find(toolbar) — 검증 6 + 조회 + 전체잠금 + SmartColMerge",
    f"{P}/MetalLinkPrice.tsx:158 (hotkeyConfig.search=handleSearch, hasButton('SEARCH'))")
add("BT-02", "Excel(toolbar) — SimplexelService Export",
    f"{P}/MetalLinkPrice.tsx:220 excelFileName (BTN EXCEL 배선·권한 3행 실측)")
add("BT-03", "Insert/Delete/Edit/Save(toolbar) — 빈 override, TO-BE hidden",
    f"{P}/MetalLinkPrice.tsx:159~162", "AS-IS 기능 없음 → 미노출 1:1")
add("BT-04", "btnProject '...' — Project 팝업 (81*/82 만 활성)",
    f"{P}/components/MetalLinkPriceForm.tsx:176~184 (disabled=isProjectDisabled)")
add("BT-05", "btmItem '...' — Item 팝업 (Always)",
    f"{P}/components/MetalLinkPriceForm.tsx:202~207")
add("BT-06", "btnSPVer '...' — 버전 팝업(최대 2, MSG00255)",
    f"{P}/components/MetalLinkPriceForm.tsx:234~239 + modals/MetalLinkVersionModal.tsx:145")
add("BT-07", "btnTemplate 'Excel Form' — 템플릿 전용 팝업",
    f"{P}/MetalLinkPrice.tsx:182~197 customButtons (division 미선택 시 MSG00097 차단 169~178)",
    "게이트 코드가 정본 DMF0082 의 hasButton('FORM_DATA') 와 달리 hasButton('SEARCH') — "
    "배선 실측이 template_SEARCH 라 정합(버튼코드명 차이는 결함 아님)")
add("BT-08", "btnExcelUpload 'Upload' — DIV40 formId 분기 + 완료 후 재조회",
    f"{P}/MetalLinkPrice.tsx:163 (hotkeyConfig.upload=hasButton('UPLOAD')) + 285~320")

S3 = "## 조회조건 컨트롤 축 (md §3.2 9)"
CTL = [
    ("CT-01", "clbDivisions — Division 단일선택 필수, 초기 ALL, SIC cascade",
     "components/MetalLinkPriceForm.tsx:59~88", ""),
    ("CT-02", "clbTeam — Team 멀티(ALL), UserTeamByDivision, SIC→하위 reset",
     "components/MetalLinkPriceForm.tsx:89~97 (공용 TeamField)", ""),
    ("CT-03", "clbEndCustomer — End Customer 멀티(ALL), SIC→SalesSite reset",
     "components/MetalLinkPriceForm.tsx:98~130", ""),
    ("CT-04", "clbScmSite — Sales Site 멀티(ALL), 값 수집만(핸들러 없음)",
     "components/MetalLinkPriceForm.tsx:131~162", ""),
    ("CT-05", "clbPlantSite — Plant Site 멀티, 81* 만 활성(82/기타 비활성+clear)",
     "components/MetalLinkPriceForm.tsx:253~259 (공용 PlantSiteField — division 게이트 내장)", ""),
    ("CT-06", "txtProject(+btnProject) — Project, 81*/82 만 활성",
     "components/MetalLinkPriceForm.tsx:163~190", ""),
    ("CT-07", "txtItemId(+btmItem) — Model/Nickname",
     "components/MetalLinkPriceForm.tsx:191~213", ""),
    ("CT-08", "drpPeriodMonth — Plan Date 월범위, 기본 당월~+6",
     "components/MetalLinkPriceForm.tsx:214~220 (PeriodField) + useMetalLinkPriceStore.ts:47~48", ""),
    ("CT-09", "txtVersion(+btnSPVer) — 기본 'Current', Division 변경 시 clear, 조회 필수키",
     "components/MetalLinkPriceForm.tsx:221~245 + useMetalLinkPriceStore.ts:49 + Form:79(clear)", ""),
]
for _id, g, loc, why in CTL:
    add(_id, g, f"{P}/{loc}", why)

S4 = "## 팝업 흐름 축 (md §4 A~E)"
add("PU-01", "A. SCM0002_ProjectbyDivisionPopupEx — Division/Team/EndCust/Site 필터 → ProjectCode 반환, "
    "빈값 아닐 때만 반영",
    f"{P}/MetalLinkPrice.tsx ProjectByDivisionModal 사용 (@master/components 재사용)")
add("PU-02", "B. DM0002_ItembyDivisionPopupEx — 조건 기반 Item 목록 → ItemID 반환",
    f"{P}/MetalLinkPrice.tsx ItemByDivisionModal 사용",
    "AS-IS 무조건 덮임(L670) → TO-BE 취소 시 유지 개선")
add("PU-03", "C. DMF0073_VersionBySalesPricePopup — SP 버전 ±3개월, 3개+ MSG00255 미반영",
    f"{P}/components/modals/MetalLinkVersionModal.tsx:103~157")
add("PU-04", "D. DMF0074_MetalLinkPricePopup — 템플릿 전용(자동조회 없음 → Find → 전컬럼잠금 → "
    "Excel export = 템플릿 다운로드, 0건 MSG00427)",
    f"{P}/components/modals/MetalLinkTemplateModal.tsx:68~118(Search 게이트),121~145(export),"
    "94~106(0건 안내 MSG00656)",
    "AS-IS MSG00427 은 사내 MSG_M 미등록 → 2026-07-10 MSG00656 신규 채번으로 대체(§5.1 대사 완료)")
add("PU-05", "E. ExcelUploadForm — formId 분기(40→DMF0074_40) + SaveExceptError + 완료 confirm 재조회",
    f"{P}/MetalLinkPrice.tsx:151~154,285~320 (공용 FcstUploadModal)",
    "완료 피드백은 hook #132(Alert 금지·Toast) 표준에 따라 공용 모달의 Toast(MSG00606/MSG00265)로 "
    "수행 — AS-IS confirm(Yes/No) 대신 자동 재조회. 표준 우선 의도적 편차")

S5 = "## Validation 전수 축 (md §5.1 10곳) + Expression"
VA = [
    ("VA-01", "FindData MSG00137 — 이전 작업 실행 중 가드", "MetalLinkPrice.tsx:100~103", ""),
    ("VA-02", "FindData MSG00283 — Version 공란", "MetalLinkPrice.tsx:105~108", ""),
    ("VA-03", "FindData Plan Date 필수 — AS-IS MSG00002 는 사내 실문구가 '저장 성공' 이라 사용 금지 → MSG00614",
     "MetalLinkPrice.tsx:111~117 (MSG.REQUIRED_VALUE=MSG00614)",
     "MSG ID 대사 결과 교체 — hook msg-id-asis-tobe-conversion 적용"),
    ("VA-04", "FindData MSG00049 — 날짜 형식 오류(사내 미등록)", "—",
     "TO-BE 월 피커(PeriodField)가 비정형 입력을 원천 차단 → 도달 불가. 사내 MSG_M 미등록이 정상", "해당없음"),
    ("VA-05", "FindData MSG00108 — 기간 역전(시작>종료)", "MetalLinkPrice.tsx:119~125", ""),
    ("VA-06", "FindData MSG00097 — Division 필수", "MetalLinkPrice.tsx:127~133", ""),
    ("VA-07", "btnSPVer MSG00255 — 버전 3개+ 차단", "components/modals/MetalLinkVersionModal.tsx:145~150", ""),
    ("VA-08", "btnExcelUpload MSG00265 — 업로드 완료 재조회 안내",
     "components/fcst-upload-modal/FcstUploadModal.tsx:171~175 (공용 모달 Toast)",
     "DMF0074 자체 상수 MSG.UPLOAD_DONE 은 사문(공용 모달이 담당) — 보완목록 R-2 로 정리 대상"),
    ("VA-09", "팝업D FindData MSG00097 — Division 필수",
     "components/modals/MetalLinkTemplateModal.tsx:110~117", ""),
    ("VA-10", "팝업D FindData MSG00427 — 0건(MRF:Y 버전 없음) 안내",
     "components/modals/MetalLinkTemplateModal.tsx:94~106 (MSG00656)",
     "AS-IS MSG00427 사내 미등록 → MSG00656 신규 채번 대체"),
]
for _id, g, loc, why, *st in VA:
    add(_id, g, (f"{P}/{loc}" if loc != "—" else "—"), why, st[0] if st else "구현")
add("XP-00", "Expression 0건 명기 — SetValue Auto-Sync/조건부 LockRange 없음", "—",
    "expressions.md 대조 결과 0건. 조회 후 일괄 LockColumn/SmartColMerge 는 expression 이 아닌 "
    "그리드 표준 속성(isReadOnly + 세로 머지로 재현)", "해당없음")

S6 = "## 편집제어·저장구조 축 (md §2.1 4층 + FE 판정 3)"
add("ED-01", "§2.1 그리드 잠금 — 조회 직후 전 셀 읽기전용(LockColumn L466)",
    f"{P}/components/MetalLinkPriceGrid.tsx (isReadOnly 그리드 — 셀 편집 UI 없음)")
add("ED-02", "§2.1 행 필터 — SP 가 EDIT_AUTH_V INNER JOIN 으로 무권한 division 행 통째 탈락(PKG L452/L837)",
    "SCM_MGR.PKG_DMF0074 BODY (GET L452 / TEMPLATE L837) — FE 중복 필터 없음",
    "서버측 하드필터 — FE 구현 불요가 설계 판정")
add("ED-03", "§2.1 EDIT_YN 반환(커서 마지막 컬럼)",
    "GET/TEMPLATE 커서 EDIT_YN — FE 는 표시/버튼 판단 근거로만 사용")
add("ED-04", "§2.1 EDIT_FLAG/ADMIN/시간락 — 이 화면에는 없음", "—",
    "DMF0044 류 행별 EDIT_FLAG·현재월 게이트가 AS-IS·TO-BE 모두 부재 — 만들면 오구현", "해당없음")
add("ED-05", "§2.1-FE판정1 그리드 항상 읽기전용 — EDIT_YN 으로 셀 여는 로직 금지",
    f"{P}/components/MetalLinkPriceGrid.tsx (편집 셀 타입 미사용)")
add("ED-06", "§2.1-FE판정2 EDIT_YN 실효 용도 = Upload/Excel Form 의미 판단 + 표시",
    f"{P}/MetalLinkPrice.tsx:157~164 (버튼 권한은 hasButton, 행 접근은 SP 필터)")
add("ED-07", "§2.2 저장 = 그리드 직접 SAVE 없음, Excel Upload 단일 경로(2단: 검증 P_MODE=0 → 적재 P_MODE=1)",
    f"{P}/MetalLinkPrice.tsx:285~320 (onSave=P_MODE 1) — 2026-08-16 실측 왕복 검증 완료")

S7 = "## 격차 축 (md §7 — 2026-07-10 기준 → 2026-08-18 현재 상태로 갱신)"
add("GP-01", "FE 미개발(metal-link-price 폴더 부재)",
    f"{P}/ 13파일 구현 완료(2026-07-10)", "해소됨")
add("GP-02", "메타 배선 0 (API_M/PROC_M/MNU_BTN/SHEET_MEASURE)",
    "사내 실측 2026-08-18: API_M 8 / PROC_M 24 / MNU_BTN 9 / USER_GR_BTN 24 / EXCEL_COLUMNS 2 배선 완료",
    "SHEET_MEASURE 는 이 화면이 measure 개인화를 쓰지 않아 0행이 정상(조회전용 PIVOT)")
add("GP-03", "standalone SP 5종 INVALID — 정본은 PKG_DMF0074",
    "이관 패키지 제외 명시 (_deploy_packages/DMF0074_CLOUD_20260818/_README.md §3)",
    "변환 중간 잔재 — API_M 배선은 전부 PKG 참조로 실측 확인", "제외")
add("GP-04", "MD 줄수 오기 정정(.cs 물리 760줄)", "분석 md 각주 반영 완료", "문서 사항", "해당없음")
add("GP-05", "MSG 3건 대사 함정(MSG00002/00049/00427)",
    "VA-03/VA-04/VA-10 참조 — MSG00614/도달불가/MSG00656 으로 각각 정리 완료")
add("GP-06", "PKG GTT TRUNCATE 위험(GET L153·UPLOAD L1046) — 동시세션 mutex 경합 hang",
    "미조치 — 현행 EXECUTE IMMEDIATE TRUNCATE 유지",
    "⛔ 보완목록 R-3: DELETE 전환 검토 대상(DMF0014 동형 사고). 2026-08-16/18 테스트 중 재현 없었으나 "
    "동시 다중 세션 부하 시험은 미수행", "미구현")
add("GP-07", "AS-IS 코드 quirk 4종(dtCustomer null·FP_Log2 라벨오기·dead code·btmItem 빈값덮임)",
    "TO-BE 재현 금지 — IT-15 개선 반영", "quirk 는 이관 대상 아님", "제외")
add("GP-08", "공통버튼 정책 — insert/delete/save hidden + search/excel 노출",
    f"{P}/MetalLinkPrice.tsx:157~164")
add("GP-09", "자체 버튼 3종(Excel Form/Upload/버전) 표준 배치",
    f"{P}/MetalLinkPrice.tsx:182~197(customButtons) + Form:234~239(Version 필드 내 트리거)")

# ── EV : Excel Upload SP Validation 49 (xlsx 실측 연동) ──
S8 = "## Excel Upload SP Validation 축 (PKG_DMF0074.P_DMF0074_EXCEL_UPLOAD — xlsx v1.1 실측 49판정)"
wb = load_workbook(XLSX)
ws = wb["1_Validation목록"]
n_ev = 0
for r in range(3, ws.max_row + 1):
    vid = str(ws.cell(r, 2).value or "")
    item = str(ws.cell(r, 4).value or "")
    cond = str(ws.cell(r, 6).value or "")
    msg = str(ws.cell(r, 7).value or "")
    line = ws.cell(r, 3).value
    meas = str(ws.cell(r, 10).value or "")
    note = str(ws.cell(r, 11).value or "")
    if msg in ("-", "", "None"):
        continue
    n_ev += 1
    gist = f"[{vid}] {item} — {cond} → `{msg}`"
    loc = f"PKG_DMF0074 BODY L{line} (실측 {meas})"
    why = ""
    st = "구현"
    if "DEAD" in meas or "FAIL" in meas:
        st, why = "미구현", f"{note[:120]}"
    elif "미검증" in meas or "정정됨" in meas:
        why = f"{note[:120]}" or "동일 패턴 형제 블록 — 케이스 미구성"
    add(f"EV-{n_ev:03d}", gist.replace("|", "/")[:210], loc, why.replace("|", "/"), st)

# ── 기존 파일에 append (기존 기입 보존) ──
text = MAP.read_text(encoding="utf-8")
existing = set(re.findall(r"^\|\s*([A-Z]{2,3}-\d{2,3})\s*\|", text, re.M))
out = [text.rstrip(), ""]
cur_sec = None
SECTIONS = {"IT": S, "BT": S2, "CT": S3, "PU": S4, "VA": S5, "XP": S5, "ED": S6, "GP": S7, "EV": S8}
kept = 0
for _id, gist, loc, why, st in ROWS:
    pre = _id.split("-")[0]
    sec = SECTIONS[pre]
    if sec != cur_sec:
        out += ["", sec, "", HDR, SEP]
        cur_sec = sec
    if _id in existing:
        kept += 1
        continue
    out.append(f"| {_id} | {gist} | {loc} | {why} | {st} |")
MAP.write_text("\n".join(out) + "\n", encoding="utf-8")
print(f"세부 축 {len(ROWS)}행 생성 (EV {n_ev} 포함) / 기존 보존 {kept}")
print(f"미구현 상태 행 = {[r[0] for r in ROWS if r[4] == '미구현']}")
