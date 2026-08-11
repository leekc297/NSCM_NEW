# -*- coding: utf-8 -*-
"""TOBE_UI_DB객체_전수취합 v2.2 누락 근본원인 규명

가설: 취합 = closure( API_M 에 등재된 SP/PKG ) , 그것도 **취합 대상 UI 목록에 한정**
      ⇒ 누락은 3층에서 발생한다:
         L1 취합 UI 목록 자체에서 빠진 화면
         L2 UI 는 있는데 그 기능이 API_M 에 미등재 (프레임워크 공통 SP)
         L3 SP 가 아예 없음 (BE 범용 코드가 테이블을 직접 읽음)
각 층을 수치로 확정한다.
"""
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from ora import conn_info  # noqa: E402
import oracledb  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

P = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI_DB객체_전수취합_v2.2.xlsx")
wb = load_workbook(P, read_only=True, data_only=True)
rows1 = list(wb["1_UI별_DB객체"].iter_rows(values_only=True))
hdr = [str(c) for c in rows1[0]]
i_menu, i_obj = hdr.index("MENU_ID"), hdr.index("객체명")
wb_menus, wb_objs = set(), set()
for r in rows1[1:]:
    if r and r[i_menu]:
        wb_menus.add(str(r[i_menu]).strip().upper())
    if r and r[i_obj]:
        wb_objs.add(str(r[i_obj]).strip().upper())
wb.close()
print("워크북: UI %d개 / 객체 %d개" % (len(wb_menus), len(wb_objs)))

u, p, d = conn_info()
con = oracledb.connect(user=u, password=p, dsn=d)
cur = con.cursor()

# ── L1: 취합 UI 목록 vs 실제 메뉴 마스터 ──
cur.execute("""SELECT MNU_ID FROM SCM_MGR.SCM_ADM_MNU_M WHERE NVL(USE_YN,'Y')='Y'""")
db_menus = {r[0].strip().upper() for r in cur.fetchall() if r[0]}
cur.execute("SELECT DISTINCT MNU_ID FROM SCM_MGR.SCM_ADM_MNU_BTN_MPG_M")
wired_menus = {r[0].strip().upper() for r in cur.fetchall() if r[0]}
print("\n== L1. 취합 UI 범위 ==")
print("  메뉴 마스터(USE_YN=Y) : %d" % len(db_menus))
print("  배선 보유 메뉴(BTN_MPG): %d" % len(wired_menus))
print("  워크북 UI             : %d" % len(wb_menus))
l1 = sorted(wired_menus - wb_menus)
print("  ⛔ 배선은 있는데 취합에 없는 UI: %d개" % len(l1))
print("    ", l1[:40])

# ── 취합 기점 = API_M 의 SP/PKG (워크북 UI 한정) ──
cur.execute("""SELECT DISTINCT B.MNU_ID, A.API_URL FROM SCM_MGR.SCM_ADM_MNU_BTN_MPG_M B
                JOIN SCM_MGR.SCM_ADM_API_M A ON A.API_ID = B.API_ID""")
menu_api = cur.fetchall()
def root_of(url):
    s = (url or "").strip().upper()
    if not s:
        return None
    return s.split("(")[0].split(".")[0] if "." in s else s.split("(")[0]
roots_all, roots_wb = set(), set()
for mnu, url in menu_api:
    r = root_of(url)
    if not r:
        continue
    roots_all.add(r)
    if (mnu or "").strip().upper() in wb_menus:
        roots_wb.add(r)
print("\n== 취합 기점(API_M ROOT 객체) ==")
print("  전체 배선 ROOT: %d / 워크북 UI 한정 ROOT: %d" % (len(roots_all), len(roots_wb)))

# ── L2: DB 에 있는 SP/PKG 중 API_M 에 없는 것 ──
cur.execute("""SELECT OBJECT_NAME, OBJECT_TYPE FROM ALL_OBJECTS
                WHERE OWNER='SCM_MGR' AND OBJECT_TYPE IN ('PACKAGE','PROCEDURE','FUNCTION')
                  AND STATUS='VALID'""")
db_progs = {(r[0].strip().upper(), r[1]) for r in cur.fetchall()}
unwired = sorted({n for n, t in db_progs} - roots_all)
print("\n== L2. DB 프로그램 중 API_M 미등재 ==")
print("  DB 프로그램 %d개 / API_M ROOT %d개 / 미등재 %d개" % (len(db_progs), len(roots_all), len(unwired)))
FAV = [n for n in unwired if "FAVOR" in n or "BMK" in n or "MENU" in n or "SRCH_COND" in n]
print("  그중 프레임워크 공통 계열 샘플:", FAV[:15])

# ── L3: 미포함 테이블을 참조하는 SP 가 API_M 에 있는지 ──
TARGETS = ["SCM_ADM_USER_BMK_M", "SCM_ADM_USER_MNU_SRCH_COND_M", "SCM_ADM_CMBO_SQL_M",
           "SCM_COM_SYS_LOG", "SCM_COM_SQL_LOG", "SCM_ADM_EXCEL_COLUMNS",
           "SCM_ADM_NOTICE_CONFIRM", "SCM_ADM_RECIPIENT_LIST", "SCM_ADM_SYS_VOC"]
print("\n== L3. 미포함 테이블별 '참조 SP 가 API_M 에 있는가' ==")
for t in TARGETS:
    cur.execute("""SELECT DISTINCT NAME FROM ALL_SOURCE
                    WHERE OWNER='SCM_MGR' AND TYPE IN ('PACKAGE BODY','PROCEDURE','FUNCTION')
                      AND INSTR(UPPER(TEXT), :t) > 0""", {"t": t})
    refs = {r[0].strip().upper() for r in cur.fetchall()}
    inapi = sorted(refs & roots_all)
    inwb = sorted(refs & roots_wb)
    if not refs:
        cause = "L3 참조 SP 없음 → BE 직접 접근"
    elif not inapi:
        cause = "L2 참조 SP 는 있으나 API_M 미등재"
    elif not inwb:
        cause = "L1 참조 SP 는 API_M 에 있으나 그 UI 가 취합 범위 밖"
    else:
        cause = "⚠ 기점에 있는데도 빠짐 = 폐포 누락(도구 결함)"
    print("  %-32s 참조SP %-2d | API_M내 %-2d | 워크북UI내 %-2d → %s"
          % (t, len(refs), len(inapi), len(inwb), cause))
    if refs and not inapi:
        print("      참조 SP:", sorted(refs)[:6])
con.close()
