# -*- coding: utf-8 -*-
"""워크북 v2.2 / v2.3 에서 SSP_ANL_ESS_REQ 행 조회 (읽기 전용)"""
import sys, io, json, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from openpyxl import load_workbook

BASE = r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI_DB객체_전수취합_%s.xlsx"

for ver in ("v2.2", "v2.3"):
    path = BASE % ver
    print("=" * 110)
    print("### " + path)
    print("=" * 110)
    wb = load_workbook(path, read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            hdr = next(rows)
        except StopIteration:
            continue
        hdr = [("" if h is None else str(h)) for h in hdr]
        print("--- sheet '%s' cols=%s" % (ws.title, hdr))
        n = 0
        for r in rows:
            joined = " ".join("" if v is None else str(v) for v in r)
            if "ANL_ESS_REQ" in joined.upper():
                n += 1
                print("  ROW:")
                for h, v in zip(hdr, r):
                    if v is not None and str(v).strip() != "":
                        print("     %-20s = %s" % (h, v))
        print("  (matched rows in sheet = %d)" % n)
    wb.close()
    print()
