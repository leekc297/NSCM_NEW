"""화면분석서 v1.2 생성 — 2026-08-27 라이브 대조·케이스 E2E 결과 반영 (원본 무수정, 새 파일)"""
import openpyxl, shutil, os, re
from openpyxl.styles import Font, PatternFill
D = r"C:\Projects\NSCM\Doc\화면분석서"; G = r"C:\Projects\NSCM\_artifacts\_gate_reports"
H = PatternFill("solid", fgColor="FCE4D6"); B = Font(bold=True)

def parse_report(path):
    rows = []
    for line in open(path, encoding="utf-8"):
        m = re.match(r"\| (\S+) \| (.*?) \| (.*?) \| (.*?) \| \*\*(PASS|FAIL)\*\* \|", line)
        if m: rows.append(m.groups())
    return rows

# ── DMF0014 v1.2 ──
src = os.path.join(D, "DMF0014_Excel_Upload_Validation_Rules_20260825.xlsx"); dst = os.path.join(D, "DMF0014_Excel_Upload_Validation_Rules_v1.2_20260827.xlsx")
shutil.copy2(src, dst); wb = openpyxl.load_workbook(dst); ws = wb["01_검증규칙"]
live = {"CPRJ": ("사문(주석)", "라이브 BODY L1054~1082 /* */ — 실측: 다음 규칙 CMDL 이 'Invalid Customer Model.' 로 검출"),
        "CSTM": ("사문(주석)", "라이브 BODY L1236~1262 /* */ — 실측: 'X' 가 오류 없이 SUCCESS"),
        "PRJ": ("사문(주석)", "라이브 BODY L1689~1715 /* */ — 실측 SUCCESS"),
        "MDL": ("사문(주석)", "라이브 BODY L1716~1747 /* */ — 실측 SUCCESS"),
        "SITE": ("마스킹", "BKT(3)·MUOM(4) 이 (Sales Site, Measure) 로 먼저 검사 → 실측 'Invalid Measure Upload Bucket.'"),
        "UOM": ("마스킹", "MUOM(4) 이 먼저 등록 UOM 과 비교 → 실측 'Invalid Measure Upload UOM.'"),
        "UOMC": ("마스킹", "BKT(3) 이 먼저 → 실측 'Invalid Measure Upload Bucket.'")}
c1, c2 = ws.max_column + 1, ws.max_column + 2
ws.cell(1, c1, "라이브 대조 2026-08-27").fill = H; ws.cell(1, c2, "실측 근거 (excel_upload_case_e2e · 사내 PKG_DMF0014)").fill = H
for r in range(2, ws.max_row + 1):
    cid = str(ws.cell(r, 2).value or "").strip()
    st, note = live.get(cid, ("LIVE·실측 PASS", "케이스 업로드 시 기대 메시지 그대로 검출"))
    ws.cell(r, c1, st); ws.cell(r, c2, note)
e = wb.create_sheet("05_케이스E2E_20260827")
e.append(["DMF0014 케이스별 실업로드 실측 (2026-08-27, 사내 Oracle + 로컬 BE, minyonglee1, division 81 기준행)"]); e.append([])
e.append(["케이스", "제목", "기대(문서)", "실측 MESSAGE_TEXT", "러너 판정", "해석"]); [setattr(e.cell(3, i), "font", B) for i in range(1, 7)]
interp = {"05_SITE": "순서 마스킹(BKT 선행)", "07_CPRJ": "규칙 사문 — CMDL 검출", "11_CSTM": "규칙 사문 — 통과", "15_UOM": "순서 마스킹(MUOM 선행)", "16_UOMC": "순서 마스킹(BKT 선행)", "17_PRJ": "규칙 사문 — 통과", "18_MDL": "규칙 사문 — 통과"}
for n, t, ex, g, v in parse_report(os.path.join(G, "DMF0014_excel_upload_cases_20260827_1606.md")):
    e.append([n, t, ex, g, v, interp.get(n, "기대 일치")])
e.append([]); e.append(["결론", "문서 18규칙 중 4규칙(CPRJ·CSTM·PRJ·MDL)이 라이브에서 주석 처리돼 실행되지 않음 — 복원 여부는 업무 결정. 3규칙(SITE·UOM·UOMC)은 앞 규칙에 마스킹돼 단독 메시지 관측 불가(조기 종료 구조)."])
e.append(["테스트 파일", r"Doc\단위테스트\DMF0014_ExcelUpload_ValidationCases_20260827.xlsx (19시트) / _excel_upload_cases\DMF0014\ (단일시트 업로드용)"])
for col, w in zip("ABCDEF", (12, 34, 44, 60, 10, 30)): e.column_dimensions[col].width = w
ov = wb["00_개요"]; ov.append([]); ov.append(["v1.2 (2026-08-27)", "01_검증규칙 에 '라이브 대조' 2열 추가 + 05_케이스E2E 시트 신설. 문서 규칙 18 중 4 사문·3 마스킹 실측"])
wb.save(dst); print(dst)

# ── DMF0074 v1.2 ──
src = os.path.join(D, "DMF0074_MetalLinkPrice_ExcelUpload_Validation_v1.1.xlsx"); dst = os.path.join(D, "DMF0074_MetalLinkPrice_ExcelUpload_Validation_v1.2.xlsx")
shutil.copy2(src, dst); wb = openpyxl.load_workbook(dst)
df = wb["3_발견결함"]
for r in range(1, df.max_row + 1):
    v = str(df.cell(r, 1).value or "")
    if v == "D4": df.cell(r, 7, str(df.cell(r, 7).value or "") + "\n\n[조치] ✅ 2026-08-27 사내 적용 — [E20260827-D4-NOLOADGUARD] P_MODE>0 오류잔존 → CODE '-1' NOLOAD_ERRORS / 후보 있는데 적재 0 → SAVE_VERIFY. 실증: 오류 2건 배치 [OK] 호출 → '-1' · 원장 Δ0")
    if v == "D5": df.cell(r, 7, str(df.cell(r, 7).value or "") + "\n\n[조치] ✅ 2026-08-27 사내 적용 — [E20260827-D5-ORDER] C21~C23 을 C19 앞으로 (AS-IS 도 C19 선행이었으나 사용자 요청). 실증: TC14~16 구체 메시지")
e = wb.create_sheet("8_케이스E2E_20260827")
e.append(["DMF0074 케이스별 실업로드 실측 (2026-08-27, D4/D5 배포 후, division 40 · DMF0074_40 양식, 기준행 Isuzu/MPCISZISLD1V/Isuzu_JP_TYO/USD/Mass)"])
e.append(["⚠ 금주 MRF 버전(SP_W_20260728_V002) 원장 0행이라 검증 창에서 SP_W_20260820_V002 ATTRIBUTE7 을 'Y' 로 임시 전환 후 원복(NULL) — dmf0074_mrf_window_e2e.py"]); e.append([])
e.append(["케이스", "제목", "기대(문서)", "실측 MESSAGE_TEXT", "러너 판정", "해석"]); [setattr(e.cell(4, i), "font", B) for i in range(1, 7)]
interp74 = {"01_C2": "사전분류 SKIP — Div 공란 행은 검증 대상 밖(문서 케이스 도달 불가)", "53_GATE": "D4 는 별도 P_MODE=1 호출로 PASS(dmf0074_d4_gate_verify.py) — 러너는 검증 단계 메시지만 봄"}
for n, t, ex, g, v in parse_report(os.path.join(G, "DMF0074_excel_upload_cases_20260827_1625.md")):
    e.append([n, t, ex, g, v, interp74.get(n, "기대 일치")])
e.append([]); e.append(["결론", "테스트 가능 49케이스 중 47 기대 일치 + D4 별도 PASS. N/A 6(Project Code 40 대상외 2·C3 버전없음·C17 82전용·C25 GTT·C44 MERGE)."])
e.append(["테스트 파일", r"Doc\단위테스트\DMF0074_ExcelUpload_ValidationCases_20260827.xlsx (55시트) / _excel_upload_cases\DMF0074\ (단일시트 업로드용)"])
for col, w in zip("ABCDEF", (12, 34, 44, 60, 10, 50)): e.column_dimensions[col].width = w
wb.save(dst); print(dst)
