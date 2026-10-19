"""07_build3.py — 3시트 산출물 (사용자 요청 2026-10-15: 시트 축소)
  SHEET1 사용하는 MENU 목록
  SHEET2 MENU별 API/SP/PKG 목록 + 연결된 TABLE/VIEW (진입 API → 경유 DB 객체 → TABLE/VIEW)
  SHEET3 TABLE/VIEW 목록 (중복 제거)
입력 = 03~05 산출 CSV (DB 접속 없음). 상세 근거본(10시트)은 06_build.py → 작업폴더에 별도 보관.
자가 검증: SHEET3 건수 == SHEET2 고유 TABLE/VIEW == 상세본 05 최종목록 건수
"""
import csv, json, pathlib, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = pathlib.Path(__file__).resolve().parent
OUT = pathlib.Path(r"C:\Projects\NSCM\Doc\화면분석서\TOBE_UI사용_TABLE_VIEW_SYSTEM_ADMIN_DM_AUTH_20261014.xlsx")
L = lambda p: list(csv.DictReader((D / p).open(encoding="utf-8-sig")))
mnu = {r["MNU_ID"]: r for r in L("meta/mnu_m.csv")}
entries, cm, edges, mt, ctab = L("entries.csv"), L("closure_members.csv"), L("call_edges.csv"), L("member_tables.csv"), L("combo_tables.csv")
tables = L("dict/tables.csv"); objects = L("dict/objects.csv")
sf = json.load((D / "fe_screen_files.json").open(encoding="utf-8"))
snapshot = (D / "SNAPSHOT.txt").read_text(encoding="utf-8").splitlines()[0]

SCOPE = {
    "SYSTEM": "SYS0003 SYS0007 SYS0008 SYS0001 SYS0004 SYS0016 SYS0017 SCM0016 SCM0007 SCM0019 SYS0005 SYS0006 SYS0009".split(),
    "ADMIN": "SYS0002 SYS0014 SYS0010 SYS0013".split(),
    "DM": ("DMF0004 DMF0005 DMF0006 DMF0007 DMF0020 DMF0043 DMF0043_01 DMF0008 DMF0072 "
           "DMF0003 DMF0048 DMF0014 DMF0042 DMF0044 DMF0049 DMF0074 DMF0082 DMF0021 "
           "DMF0035 DMF0040 DMF0068 DMF0053").split(),
    "AUTHORITY": "SYS0011 SYS0012".split(),
}
GROUP_OF = {s: g for g, ss in SCOPE.items() for s in ss}
GROUPS = ["SYSTEM", "ADMIN", "DM", "AUTHORITY", "COMMON"]
TEMP = {(t["OWNER"], t["TABLE_NAME"]) for t in tables if t["TEMPORARY"] == "Y"}
OPMAP = {"SELECT": "R", "INSERT": "C", "UPDATE": "U", "DELETE": "D", "MERGE": "M", "TRUNCATE": "T", "LOCK": "L"}
def opstr(ops): return "".join(v for k, v in OPMAP.items() if k in ops)
def norm_type(owner, name, t):
    if (owner, name) in TEMP: return "GTT"
    if "SYS_DICTIONARY_VIEW" in t: return "VIEW(SYS 딕셔너리)"
    if "REMOTE@" in t: return "REMOTE(DB LINK)"
    if "MATERIALIZED VIEW" in t: return "MVIEW"
    if "VIEW" in t: return "VIEW"
    if "TABLE" in t: return "TABLE"
    return t or "?"
def short(n): return n.replace("SCM_MGR.", "")

# 노드 → 테이블
node_tabs = collections.defaultdict(lambda: collections.defaultdict(set))   # node -> (owner,name,type) -> ops
for r in mt:
    if r["T_TYPE"] in ("UNRESOLVED", ""): continue
    node_tabs[r["NODE"]][(r["T_OWNER"], r["T_NAME"], norm_type(r["T_OWNER"], r["T_NAME"], r["T_TYPE"]))].add(r["OP"])
adj = collections.defaultdict(set)
for e in edges: adj[e["FROM"]].add(e["TO"])
nodes = {r["NODE"] for r in cm}
def seed_node(obj):
    o = obj.upper()
    for cand in ("SCM_MGR." + o, o):
        if cand in nodes: return cand
    hits = [n for n in nodes if n.endswith("." + o)]
    return hits[0] if hits else None
def closure(seed):
    seen, st = [], [seed]
    vis = set()
    while st:
        x = st.pop()
        if x in vis: continue
        vis.add(x); seen.append(x)
        st.extend(sorted(adj[x] - vis, reverse=True))
    return seen
combo_calls = collections.defaultdict(set); combo_tabs = collections.defaultdict(set)
for r in ctab:
    if r["TYPE"] == "CALL": combo_calls[r["COMBO"]].add(r["OBJECT"])
    elif r["TYPE"] not in ("WARN", "UNRESOLVED"): combo_tabs[r["COMBO"]].add((r["OWNER"], r["OBJECT"], norm_type(r["OWNER"], r["OBJECT"], r["TYPE"])))

# ───────── SHEET2 행 생성 ─────────
KIND_LABEL = {"CALLPROC": "버튼 API(call-proc)", "EXCEL_VALID": "Excel Upload 검증 SP", "DYNAMIC_QUERY": "Dynamic Query API",
              "COMBO": "콤보(조회조건 등)", "GRID_COMBO": "그리드 셀 콤보", "FRAMEWORK": "공통 프레임워크 SP", "DIRECT_SQL": "BE 직접 SQL"}
rows2 = []
def add_rows(e, entry_obj, chain_nodes, extra_tabs=()):
    base = [e["GROUP"], e["SCREEN"], mnu.get(e["SCREEN"], {}).get("MNU_DESC", "공통" if e["GROUP"] == "COMMON" else ""),
            KIND_LABEL.get(e["KIND"], e["KIND"]), e["VIA_MENU"], e["SHEET"], e["BTN"], e["API_ID"], entry_obj]
    emitted = False
    for n in chain_nodes:
        for (o, t, ty), ops in sorted(node_tabs.get(n, {}).items()):
            rows2.append(base + [short(n), o, t, ty, opstr(ops)]); emitted = True
    for (o, t, ty) in sorted(extra_tabs):
        rows2.append(base + ["(콤보 SQL 본문)" if e["KIND"] in ("COMBO", "GRID_COMBO") else "(BE 매퍼/JDBC)", o, t, ty, "R" if e["KIND"] != "DIRECT_SQL" else opstr({e["NOTE"]})]); emitted = True
    if not emitted:
        rows2.append(base + ["(TABLE/VIEW 없음 — DUAL 등)", "", "", "", ""])

for e in sorted(entries, key=lambda e: (GROUPS.index(e["GROUP"]), e["SCREEN"], e["KIND"], e["VIA_MENU"], e["SHEET"], e["BTN"], e["OBJECT"])):
    k = e["KIND"]
    if k in ("CALLPROC", "EXCEL_VALID", "DYNAMIC_QUERY", "FRAMEWORK"):
        sn = seed_node(e["OBJECT"])
        add_rows(e, e["OBJECT"], closure(sn) if sn else [])
    elif k in ("COMBO", "GRID_COMBO"):
        chain = []
        for c in sorted(combo_calls.get(e["OBJECT"], ())):
            sn = seed_node(c)
            if sn: chain += [x for x in closure(sn) if x not in chain]
        add_rows(e, "COMBO:" + e["OBJECT"], chain, combo_tabs.get(e["OBJECT"], ()))
    elif k == "DIRECT_SQL":
        add_rows(e, e["FE_EVIDENCE"], [], [("SCM_MGR", e["OBJECT"], norm_type("SCM_MGR", e["OBJECT"], {"TABLE"}))])

# 중복 제거(같은 메뉴·API·경유노드·테이블)
seen = set(); r2 = []
for r in rows2:
    k = tuple(r)
    if k not in seen: seen.add(k); r2.append(r)
rows2 = r2

# VIEW 내부 참조 전개 (2026-10-15 요청) — VIEW 행 바로 뒤에 그 VIEW 가 쓰는 TABLE/VIEW 를 같은 메뉴·API 로 추가
import sys as _sys; _sys.path.insert(0, str(D))
from _viewexp import expand_view
r2 = []
for r in rows2:
    r2.append(r)
    if r[12] in ("VIEW", "MVIEW"):
        for o2, n2, t2, path in expand_view(r[10], r[11])[0]:
            r2.append(r[:9] + [f"{r[9]} ▶ VIEW {path}", o2, n2, t2, "R"])
seen = set(); rows2 = []
for r in r2:
    k = tuple(r)
    if k not in seen: seen.add(k); rows2.append(r)

# ───────── SHEET3 ─────────
fin = {}
for r in rows2:
    g, scr, o, t, ty, op = r[0], r[1], r[10], r[11], r[12], r[13]
    if not t: continue
    f = fin.setdefault((o, t), dict(TYPE=ty, G=set(), OPS=set(), S=set(), API=set(), DIRECT=False, VIEWS=set()))
    f["G"].add(g); f["OPS"].update(op);
    if not scr.startswith("("): f["S"].add(scr)
    f["API"].add(r[8])
    if "▶ VIEW " in r[9]: f["VIEWS"].add(r[9].split("▶ VIEW ")[1].split(" > ")[0])
    else: f["DIRECT"] = True
TORDER = {"TABLE": 0, "GTT": 1, "VIEW": 2, "MVIEW": 3, "VIEW(SYS 딕셔너리)": 4, "REMOTE(DB LINK)": 5}
ORD = "RCUDMTL"
rows3 = []
for i, ((o, t), f) in enumerate(sorted(fin.items(), key=lambda kv: (TORDER.get(kv[1]["TYPE"], 9), kv[0][0] != "SCM_MGR", kv[0][0], kv[0][1])), 1):
    rows3.append([i, o, t, f["TYPE"]] + ["Y" if g in f["G"] else "" for g in GROUPS] +
                 ["".join(c for c in ORD if c in f["OPS"]), len(f["S"]), " ".join(sorted(f["S"])), len(f["API"]),
                  ("" if f["DIRECT"] else "VIEW 내부 참조로만 사용") + (" · 경유 VIEW: " + ", ".join(sorted(f["VIEWS"])) if f["VIEWS"] else "")])

# ───────── SHEET1 ─────────
api_cnt = collections.Counter(); tab_cnt = collections.defaultdict(set)
for r in rows2:
    if r[3] in ("버튼 API(call-proc)", "Excel Upload 검증 SP", "Dynamic Query API"): api_cnt[r[1], r[8]] += 1
    if r[11]: tab_cnt[r[1]].add((r[10], r[11]))
def path(m):
    out = []
    while m in mnu and mnu[m]["UPPR_MNU_ID"]:
        m = mnu[m]["UPPR_MNU_ID"]; out.append(m)
    return " > ".join(reversed(out))
rows1 = []
for g in ("SYSTEM", "ADMIN", "DM", "AUTHORITY"):
    for s in SCOPE[g]:
        r = mnu.get(s, {})
        es = [e for e in entries if e["SCREEN"] == s]
        pops = sorted({e["VIA_MENU"] for e in es if e["VIA_MENU"] and e["VIA_MENU"] != s})
        page = sf.get(s, {}).get("page", "")
        apis = sorted({e["OBJECT"] for e in es if e["KIND"] in ("CALLPROC", "EXCEL_VALID", "DYNAMIC_QUERY")})
        note = "" if page else "FE 페이지 없음 → UI 사용 DB 객체 없음"
        if s == "SYS0013": note = "Dynamic Query: 사용자 입력 SQL 대상 테이블은 정적 확정 불가(PKG 자체 참조분만)"
        rows1.append([g, s, r.get("MNU_DESC", ""), path(s), r.get("ROUTE_PATH", ""), page, " ".join(pops), len(apis),
                      sum(1 for e in es if e["KIND"] in ("COMBO", "GRID_COMBO")), len(tab_cnt.get(s, ())), note])
rows1.append(["COMMON", "(공통)", "로그인·권한·메뉴·그리드 컬럼·타이틀셋·조회조건·메시지·로그·Excel 공통", "", "", "BE 프레임워크 / apps/web",
              "", sum(1 for e in entries if e["GROUP"] == "COMMON" and e["KIND"] in ("FRAMEWORK", "CALLPROC")), 0, len(tab_cnt.get("(공통)", set()) | tab_cnt.get("(공통-홈 공지)", set())), "모든 화면이 공통 사용"])

# ───────── 자가 검증 ─────────
detail = D / "TOBE_UI사용_TABLE_VIEW_상세근거_20261014.xlsx"
if detail.exists():
    wb0 = openpyxl.load_workbook(detail, read_only=True)
    n05 = wb0["05_TABLE_VIEW_목록"].max_row - 1
    s05 = {(r[1], r[2]) for r in list(wb0["05_TABLE_VIEW_목록"].iter_rows(values_only=True))[1:]}
    assert n05 == len(rows3) and s05 == set(fin), ("상세본과 불일치", n05, len(rows3), sorted(s05 ^ set(fin))[:20])
    print("self-check: SHEET3 == 상세본 05 최종목록", n05)

# ───────── xlsx ─────────
wb = openpyxl.Workbook(); wb.remove(wb.active)
HF = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="305496")
THIN = Side(style="thin", color="BFBFBF"); B = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
def sheet(title, header, rows, widths, note=None):
    ws = wb.create_sheet(title); start = 1
    if note:
        ws.append([note]); ws["A1"].font = Font(italic=True, color="595959"); start = 2
    ws.append(header)
    for c in ws[start]: c.font = HF; c.fill = HFILL; c.border = B; c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows: ws.append(["" if v is None else v for v in r])
    for row in ws.iter_rows(min_row=start + 1):
        for c in row: c.border = B; c.alignment = Alignment(vertical="top")
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(header))}{ws.max_row}"
    for i, h in enumerate(header, 1): ws.column_dimensions[get_column_letter(i)].width = widths.get(h, max(10, len(h) * 2))
NOTE = f"기준: 사내오라클 {snapshot.split('FETCHED=')[1] if 'FETCHED=' in snapshot else ''} 원문 + nscm-fe/be 현행 코드 · 판정 = FE 실호출 ∩ 메타 배선 → PKG 멤버 단위 호출 폐포(미사용 멤버 제외) · 접근 R=SELECT C=INSERT U=UPDATE D=DELETE M=MERGE T=TRUNCATE"
sheet("1_사용MENU목록", ["메뉴구분", "MNU_ID", "화면명", "메뉴 경로", "ROUTE_PATH", "FE 페이지", "함께 사용하는 팝업/타메뉴", "API/SP 수", "콤보 수", "TABLE/VIEW 수", "비고"], rows1,
      {"화면명": 32, "메뉴 경로": 32, "ROUTE_PATH": 42, "FE 페이지": 48, "함께 사용하는 팝업/타메뉴": 36, "비고": 50}, NOTE)
sheet("2_MENU별_API_SP_TABLE", ["메뉴구분", "MNU_ID", "화면명", "호출유형", "VIA_MENU", "SHT_NM", "BTN", "API_ID", "진입 API(PKG.멤버/SP/콤보ID)", "경유 DB객체(PKG.멤버/SP/FN)", "OWNER", "TABLE/VIEW", "유형", "접근"], rows2,
      {"화면명": 28, "호출유형": 20, "VIA_MENU": 12, "SHT_NM": 18, "BTN": 14, "API_ID": 28, "진입 API(PKG.멤버/SP/콤보ID)": 55, "경유 DB객체(PKG.멤버/SP/FN)": 60, "OWNER": 10, "TABLE/VIEW": 40, "유형": 12, "접근": 8}, NOTE)
sheet("3_TABLE_VIEW목록", ["NO", "OWNER", "TABLE/VIEW", "유형"] + GROUPS + ["접근", "사용화면수", "사용 화면", "연결 API 수", "비고(VIEW 경유)"], rows3,
      {"TABLE/VIEW": 42, "유형": 18, "사용 화면": 90, "NO": 6, "비고(VIEW 경유)": 60}, NOTE)
wb.save(OUT)
c = collections.Counter(r[3] for r in rows3)
print("SHEET1", len(rows1), "SHEET2", len(rows2), "SHEET3", len(rows3), dict(c), {g: sum(1 for r in rows3 if r[4 + i] == "Y") for i, g in enumerate(GROUPS)})
print(OUT)
